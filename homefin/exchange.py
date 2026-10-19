"""Экспорт и импорт: CSV (Excel открывает как есть) и xlsx."""
from __future__ import annotations

import csv
from decimal import Decimal
from pathlib import Path

from .core import Finance, parse_date
from .money import fmt

KIND_RU = {"income": "Доход", "expense": "Расход", "transfer": "Перевод"}
RU_KIND = {v.lower(): k for k, v in KIND_RU.items()}
HEADER = ["Дата", "Тип", "Сумма", "Счёт", "Счёт назначения", "Категория", "Примечание"]


def _rows(fin: Finance, **filters):
    for t in reversed(fin.transactions(**filters)):
        yield [parse_date(t["date"]).strftime("%d.%m.%Y"), KIND_RU[t["kind"]], f"{t['amount_cents'] / 100:.2f}",
               t["account"], t["to_account"] or "", t["category"] or "", t["note"]]


def export_csv(fin: Finance, path, **filters) -> int:
    n = 0
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(HEADER)
        for r in _rows(fin, **filters):
            w.writerow(r)
            n += 1
    return n


def export_xlsx(fin: Finance, path, **filters) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Операции"
    ws.append(HEADER)
    for c in ws[1]:
        c.font = Font(bold=True)
    n = 0
    for r in _rows(fin, **filters):
        r[0], r[2] = parse_date(r[0]), float(r[2])
        ws.append(r)
        ws.cell(row=ws.max_row, column=1).number_format = "DD.MM.YYYY"
        ws.cell(row=ws.max_row, column=3).number_format = "#,##0.00"
        n += 1
    for i, w in enumerate((12, 10, 14, 18, 18, 22, 40), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    bills = wb.create_sheet("Коммунальные счета")
    bills.append(["Период", "Услуга", "Расход", "Ед.", "Сумма", "Срок оплаты", "Оплачен"])
    for c in bills[1]:
        c.font = Font(bold=True)
    for b in fin.bills():
        bills.append([b["period"], b["service"], float(b["consumption"]) if b["consumption"] else None, b["unit"],
                      b["amount_cents"] / 100, b["due_date"] or "", b["paid_date"] or ""])
    wb.save(path)
    return n


def import_csv(fin: Finance, path) -> tuple[int, list[str]]:
    """Импорт в формате экспорта. Счета и категории создаются при необходимости.
    Возвращает (число импортированных, список ошибок по строкам)."""
    ok, errors = 0, []
    accounts = {a["name"]: a["id"] for a in fin.accounts(True)}
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(2048)
        f.seek(0)
        delim = ";" if sample.count(";") >= sample.count(",") else ","
        rd = csv.DictReader(f, delimiter=delim)
        for i, row in enumerate(rd, 2):
            try:
                kind = RU_KIND.get((row.get("Тип") or "").strip().lower())
                if not kind:
                    raise ValueError(f"неизвестный тип {row.get('Тип')!r}")

                def acc(name):
                    name = (name or "").strip()
                    if name not in accounts:
                        accounts[name] = fin.add_account(name)
                    return accounts[name]

                cat = (row.get("Категория") or "").strip()
                fin.add_transaction(row["Дата"], kind, row["Сумма"], acc(row["Счёт"]),
                                    fin.category_id(cat, kind) if cat and kind != "transfer" else None,
                                    row.get("Примечание") or "",
                                    acc(row["Счёт назначения"]) if kind == "transfer" else None)
                ok += 1
            except (ValueError, KeyError, TypeError) as e:
                errors.append(f"строка {i}: {e}")
    return ok, errors


def report_text(fin: Finance, date_from: str, date_to: str) -> str:
    inc, exp = fin.totals(date_from, date_to)
    cur = fin.currency
    lines = [f"Отчёт {date_from} — {date_to}", f"Доходы:  {fmt(inc, cur)}", f"Расходы: {fmt(exp, cur)}",
             f"Итог:    {fmt(inc - exp, cur)}", "", "Расходы по категориям:"]
    for name, s in fin.by_category(date_from, date_to):
        share = Decimal(s) * 100 / Decimal(exp) if exp else Decimal(0)
        lines.append(f"  {name:<28}{fmt(s, cur):>18}  {share:5.1f}%")
    return "\n".join(lines)


def default_db_path() -> Path:
    base = Path.home() / ".homefin"
    base.mkdir(exist_ok=True)
    return base / "homefin.db"
