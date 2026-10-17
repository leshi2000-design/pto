"""Входящие документы: накладные, сертификаты, подписанные акты. Файл перетаскивается в программу, из текста берутся реквизиты (вид, номер, дата, сумма, контрагент),
предлагается объект для привязки; пользователь проверяет и исправляет результат до сохранения. Ничего не сохраняется само.

Текст извлекается из PDF (нужен pypdf), Word, Excel и текстовых файлов. Для сканов и фотографий нужна установленная программа Tesseract OCR (никаких Python-модулей);
без него реквизиты вводятся вручную, а программа всё равно подберёт объект по имени файла."""
import os
import re
import shutil
from datetime import date, datetime

from .config import DATA_DIR

INBOX_DIR = DATA_DIR / 'incoming'
KINDS = ['Накладная', 'Сертификат', 'Акт', 'Счёт', 'Договор', 'Другое']
MAX_CHARS = 200_000
TEXT_EXT = {'.txt', '.csv', '.md'}
IMAGE_EXT = {'.png', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp'}
MONTHS = {'января': 1, 'февраля': 2, 'марта': 3, 'апреля': 4, 'мая': 5, 'июня': 6, 'июля': 7, 'августа': 8, 'сентября': 9, 'октября': 10, 'ноября': 11, 'декабря': 12}


def db_setting(key):
    try:
        from .database import db
        return db.get_setting(key, '')
    except Exception:
        return ''


def migrate(db):
    db.execute("""CREATE TABLE IF NOT EXISTS incoming_docs(id INTEGER PRIMARY KEY, file_path TEXT NOT NULL, orig_name TEXT DEFAULT '', kind TEXT DEFAULT 'Другое', number TEXT DEFAULT '',
        doc_date TEXT DEFAULT '', amount REAL, counterparty TEXT DEFAULT '', owner_type TEXT DEFAULT '', owner_id INTEGER, note TEXT DEFAULT '', created_at TEXT DEFAULT '')""")
    db.execute('CREATE INDEX IF NOT EXISTS idx_incoming_owner ON incoming_docs(owner_type,owner_id)')


# --- ТЕКСТ ---------------------------------------------------------------------------------------

def extract_text(path):
    """(текст, замечание). Замечание непустое, если текст получить не удалось (нет модуля, скан без OCR и т. п.)."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == '.pdf':
            try:
                from pypdf import PdfReader
            except ImportError:
                return '', 'Для чтения PDF установите модуль pypdf (pip install pypdf); реквизиты введите вручную.'
            reader = PdfReader(path)
            text = '\n'.join((p.extract_text() or '') for p in reader.pages[:30])
            return text[:MAX_CHARS], '' if text.strip() else 'В PDF нет текстового слоя (скан). Реквизиты введите вручную или установите OCR (Tesseract).'
        if ext == '.docx':
            import docx
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs] + [c.text for t in d.tables for row in t.rows for c in row.cells]
            return '\n'.join(parts)[:MAX_CHARS], ''
        if ext in ('.xlsx', '.xlsm'):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets[:3]:
                for row in ws.iter_rows(max_row=300, values_only=True):
                    cells = [str(c) for c in row if c not in (None, '')]
                    if cells:
                        lines.append(' '.join(cells))
            wb.close()
            return '\n'.join(lines)[:MAX_CHARS], ''
        if ext in TEXT_EXT:
            with open(path, encoding='utf-8', errors='replace') as f:
                return f.read(MAX_CHARS), ''
        if ext in IMAGE_EXT:
            from . import invoice_parse as ip
            ok, info = ip.ocr_available(db_setting('tesseract_cmd'))
            if not ok:
                return '', info + ' Реквизиты можно ввести вручную.'
            return ip.ocr_text(ip.prepare_image(path), info, 6)[:MAX_CHARS], ''
    except Exception as e:
        return '', f'Не удалось прочитать файл: {e}'
    return '', 'Тип файла не поддерживается для распознавания; заполните реквизиты вручную.'


# --- РЕКВИЗИТЫ ----------------------------------------------------------------------------------

def _norm_date(day, month, year):
    year = int(year)
    if year < 100:
        year += 2000
    try:
        return date(year, int(month), int(day)).isoformat()
    except ValueError:
        return ''


def find_dates(text):
    out = []
    for m in re.finditer(r'(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4}|\d{2})\b', text):
        d = _norm_date(m.group(1), m.group(2), m.group(3))
        if d:
            out.append((m.start(), d))
    for m in re.finditer(r'(\d{1,2})\s+(' + '|'.join(MONTHS) + r')\s+(\d{4})', text, re.I):
        d = _norm_date(m.group(1), MONTHS[m.group(2).lower()], m.group(3))
        if d:
            out.append((m.start(), d))
    return sorted(out)


def guess_kind(text, name=''):
    low = (text[:4000] + ' ' + name).casefold()
    scores = {
        'Накладная': sum(low.count(k) for k in ('накладная', 'товарно-транспортная', 'ттн', 'товарная накладная')),
        'Сертификат': sum(low.count(k) for k in ('сертификат', 'декларация о соответствии', 'паспорт качества', 'свидетельство о соответствии')),
        'Акт': sum(low.count(k) for k in ('акт выполненных', 'акт приемки', 'акт приёмки', 'акт сдачи', 'акт №', 'акт n')),
        'Счёт': sum(low.count(k) for k in ('счет-фактура', 'счёт-фактура', 'счет на оплату', 'счёт на оплату', 'счет №', 'счёт №')),
        'Договор': sum(low.count(k) for k in ('договор №', 'договор n', 'предмет договора')),
    }
    best = max(scores, key=scores.get)
    return best if scores[best] else 'Другое'


def find_number(text, kind):
    words = {'Накладная': r'накладн\w*|ттн', 'Сертификат': r'сертификат\w*|декларац\w*', 'Акт': r'акт\w*', 'Счёт': r'сч[её]т\w*', 'Договор': r'договор\w*'}.get(kind, r'документ\w*')
    m = re.search(rf'(?:{words})[^\n№N]{{0,60}}[№N]\s*([\w\-/.]+)', text, re.I)
    if not m:
        m = re.search(r'[№N]\s*([\w\-/.]*\d[\w\-/.]*)', text)
    return m.group(1).strip('.,;') if m else ''


def find_amount(text):
    """Итоговая сумма: число после слов «итого», «всего», «к оплате»; иначе наибольшая сумма с копейками."""
    pat = r'(\d{1,3}(?:[  ]\d{3})*[,.]\d{2}|\d+[,.]\d{2})'
    cand = []
    for m in re.finditer(r'(?:итого|всего|к оплате|сумма к оплате|общая стоимость)[^\d\n]{0,50}' + pat, text, re.I):
        cand.append(m.group(1))
    if not cand:
        cand = re.findall(pat, text)
    values = []
    for c in cand:
        try:
            values.append(float(c.replace(' ', '').replace(' ', '').replace(',', '.')))
        except ValueError:
            pass
    return max(values) if values else None


def find_counterparty(text):
    m = re.search(r'((?:ООО|ЗАО|ОАО|ОДО|УП|ИП|РУП|ЧУП|ТЧУП)\s*[«"“][^»"”\n]{2,60}[»"”])', text)
    if m:
        return m.group(1)
    m = re.search(r'УНП\s*[:№]?\s*(\d{9})', text)
    return f'УНП {m.group(1)}' if m else ''


# --- ПОДБОР ОБЪЕКТА -------------------------------------------------------------------------------

def _fold(s):
    return re.sub(r'[\s_]+', ' ', str(s or '').casefold().replace('ё', 'е'))


def candidates(db, text, filename='', limit=5):
    """Объекты, которым может принадлежать документ: [(owner_type, id, подпись, баллы, [причины])]; по убыванию баллов."""
    from . import staff_domain as sd
    hay = _fold(text + ' ' + os.path.basename(filename))
    compact = re.sub(r'[\s]+', '', hay)
    result = []
    specs = {
        'gsn_projects': "SELECT id,coalesce(contract_number,''),coalesce(client_name,''),coalesce(address,''),coalesce(title,'') FROM gsn_projects",
        'contracts': "SELECT id,coalesce(contract_number,''),coalesce(nullif(client_name,''),party_name,''),coalesce(object_address,''),coalesce(object_name,'') FROM contracts",
        'gsv_projects': "SELECT id,coalesce(contract_number,''),coalesce(nullif(client_name,''),party_name,''),coalesce(address,''),coalesce(object_name,'') FROM gsv_projects",
    }
    for table, sql in specs.items():
        for rid, number, client, address, name in db.fetchall(sql):
            score, why = 0, []
            n = _fold(number).replace(' ', '')
            if len(n) >= 4 and n in compact:
                score += 10
                why.append(f'номер договора {number}')
            surname = _fold(client).split(' ')[0] if client else ''
            if len(surname) >= 4 and surname in hay:
                score += 5
                why.append(f'клиент {client.split()[0]}')
            for word in re.findall(r'[а-яёa-z]{5,}', _fold(address)):
                if word not in ('улица', 'проспект', 'переулок', 'район', 'область', 'город') and word in hay:
                    score += 2
                    why.append(f'адрес: {word}')
                    break
            if score:
                result.append((table, rid, sd.object_label(db, table, rid), score, why))
    result.sort(key=lambda r: -r[3])
    return result[:limit]


def analyze(db, path):
    """Разбор файла: {'kind','number','date','amount','counterparty','candidates','text','warning'}. Всё это только предложение для проверки."""
    text, warning = extract_text(path)
    name = os.path.basename(path)
    kind = guess_kind(text, name)
    dates = find_dates(text)
    return dict(kind=kind, number=find_number(text, kind), date=dates[0][1] if dates else '', amount=find_amount(text), counterparty=find_counterparty(text),
                candidates=candidates(db, text, name), text=text, warning=warning)


# --- СОХРАНЕНИЕ ----------------------------------------------------------------------------------

def _safe(name):
    return re.sub(r'[\\/:*?"<>|]+', '_', name).strip() or 'документ'


def _folder_for(db, owner_type, owner_id):
    try:
        if owner_type == 'gsn_projects':
            from . import gsn_domain as m
            return m.contract_folder(db, owner_id)
        if owner_type == 'contracts':
            from . import gsvm_domain as m
            return m.contract_folder(db, owner_id)
        if owner_type == 'gsv_projects':
            from . import gsv_project_domain as m
            row = db.fetchone('SELECT project_folder FROM gsv_projects WHERE id=?', (owner_id,))
            return row[0] if row and row[0] else ''
    except Exception:
        return ''
    return ''


def save(db, path, fields, copy=True, mark_signed=False):
    """Сохраняет входящий документ. fields: kind, number, doc_date (ISO), amount, counterparty, owner_type, owner_id, note.
    Файл копируется в папку объекта (если она привязана) или в папку входящих; исходный файл не удаляется. mark_signed для акта отмечает его подписанным у объекта."""
    if not os.path.isfile(path):
        raise ValueError('Файл не найден: ' + path)
    kind = fields.get('kind') or 'Другое'
    owner_type, owner_id = fields.get('owner_type') or '', fields.get('owner_id')
    target = path
    if copy:
        folder = _folder_for(db, owner_type, owner_id) if owner_type and owner_id else ''
        if not folder or not os.path.isdir(folder):
            folder = str(INBOX_DIR / date.today().strftime('%Y-%m'))
        os.makedirs(folder, exist_ok=True)
        base, ext = os.path.splitext(_safe(os.path.basename(path)))
        prefix = _safe(f'{kind}_{fields.get("number") or ""}'.strip('_'))
        target = os.path.join(folder, f'{prefix}_{base}{ext}' if prefix and prefix != 'Другое' else f'{base}{ext}')
        n = 1
        while os.path.exists(target) and os.path.abspath(target) != os.path.abspath(path):
            target = os.path.join(folder, f'{os.path.splitext(os.path.basename(target))[0].rsplit("~", 1)[0]}~{n}{ext}')
            n += 1
        if os.path.abspath(target) != os.path.abspath(path):
            shutil.copy2(path, target)
    with db.transaction():
        did = db.execute('INSERT INTO incoming_docs(file_path,orig_name,kind,number,doc_date,amount,counterparty,owner_type,owner_id,note,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                         (target, os.path.basename(path), kind, fields.get('number') or '', fields.get('doc_date') or '', fields.get('amount'), fields.get('counterparty') or '', owner_type, owner_id,
                          fields.get('note') or '', datetime.now().isoformat(timespec='seconds'))).lastrowid
        if mark_signed and kind == 'Акт' and owner_type and owner_id:
            col = 'acceptance_act_date' if owner_type == 'contracts' else 'act_date'
            day = fields.get('doc_date') or date.today().isoformat()
            db.execute(f'UPDATE {owner_type} SET {col}=?, act_signed=1 WHERE id=?', (day, owner_id))
    return did


def update(db, did, fields):
    cols = ('kind', 'number', 'doc_date', 'amount', 'counterparty', 'owner_type', 'owner_id', 'note')
    db.execute('UPDATE incoming_docs SET ' + ','.join(f'{c}=?' for c in cols) + ' WHERE id=?', (*[fields.get(c) for c in cols], did))


def delete(db, did, delete_file=False):
    row = db.fetchone('SELECT file_path FROM incoming_docs WHERE id=?', (did,))
    db.execute('DELETE FROM incoming_docs WHERE id=?', (did,))
    if delete_file and row and row[0] and os.path.isfile(row[0]):
        os.remove(row[0])


def listing(db, search='', kind='', owner_type='', owner_id=None):
    from . import staff_domain as sd
    q = f'%{search.casefold()}%'
    rows = db.fetchall("""SELECT id,created_at,kind,number,doc_date,amount,counterparty,owner_type,owner_id,file_path,orig_name,note FROM incoming_docs
        WHERE (?='' OR kind=?) AND (?='' OR (owner_type=? AND owner_id=?)) AND LOWER(coalesce(number,'')||' '||coalesce(counterparty,'')||' '||coalesce(orig_name,'')||' '||coalesce(note,'')) LIKE ?
        ORDER BY id DESC LIMIT 1000""", (kind, kind, owner_type, owner_type, owner_id, q))
    return [dict(id=r[0], created=r[1], kind=r[2], number=r[3], date=r[4], amount=r[5], counterparty=r[6], owner_type=r[7], owner_id=r[8], path=r[9], name=r[10], note=r[11],
                 object=sd.object_label(db, r[7], r[8]) if r[7] else '') for r in rows]
