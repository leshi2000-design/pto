"""Наряды: задание бригаде или работникам на объект — состав, сроки, плановые объёмы, факт при закрытии.

Наряд выдаётся на объект (договор Монтаж ГСВ / ГСН / СМР и др.), ставит работников в график занятости, предупреждает о просроченных допусках и пересечениях,
берёт плановые объёмы из вкладки «Объёмы» договора, а при закрытии принимает фактические и (по желанию) вносит их в объёмы договора."""
import os
from datetime import date, datetime

from . import staff_domain as sd
from . import volumes_domain as vd
from . import gsv_project_domain as gd
from .dates import display_date

KIND_TASK, KIND_PERMIT = 'Наряд-задание', 'Наряд-допуск'
KINDS = (KIND_TASK, KIND_PERMIT)
HAZARDS = ('Газоопасные работы', 'Огневые работы', 'Земляные работы', 'Работы на высоте', 'Работы в колодцах и замкнутых пространствах', 'Другие работы повышенной опасности')
PERMIT_COLUMNS = (('kind', "TEXT NOT NULL DEFAULT 'Наряд-задание'"), ('hazard', "TEXT DEFAULT ''"), ('place', "TEXT DEFAULT ''"), ('supervisor_id', 'INTEGER'), ('performer_id', 'INTEGER'),
                  ('admitter_id', 'INTEGER'), ('time_from', "TEXT DEFAULT ''"), ('time_to', "TEXT DEFAULT ''"), ('safety', "TEXT DEFAULT ''"), ('ppe', "TEXT DEFAULT ''"),
                  ('briefing', 'INTEGER DEFAULT 0'), ('closing_note', "TEXT DEFAULT ''"))
STATUSES = ('Выдан', 'В работе', 'Закрыт', 'Отменён')
ACTIVE = ('Выдан', 'В работе')
EXTRA_OWNERS = {'smr_contracts': 'СМР (юрлица)', 'to_contracts': 'ТО (юрлица)', 'ip_contracts': 'СМР (ИП)', 'fl_contracts': 'СМР (физлица)'}
ISO = '%Y-%m-%d'


def migrate(db):
    for sql in (
        """CREATE TABLE IF NOT EXISTS work_orders(id INTEGER PRIMARY KEY, number TEXT NOT NULL, seq_num INTEGER, year_num INTEGER, order_date TEXT NOT NULL,
             owner_type TEXT DEFAULT '', owner_id INTEGER, object_text TEXT DEFAULT '', crew_id INTEGER REFERENCES crews(id) ON DELETE SET NULL,
             leader_id INTEGER REFERENCES staff(id) ON DELETE SET NULL, start_date TEXT DEFAULT '', end_date TEXT DEFAULT '', status TEXT NOT NULL DEFAULT 'Выдан',
             task TEXT DEFAULT '', note TEXT DEFAULT '', closed_date TEXT DEFAULT '', facts_applied INTEGER NOT NULL DEFAULT 0, to_schedule INTEGER NOT NULL DEFAULT 1,
             doc_path TEXT DEFAULT '')""",
        """CREATE TABLE IF NOT EXISTS work_order_workers(order_id INTEGER NOT NULL REFERENCES work_orders(id) ON DELETE CASCADE,
             staff_id INTEGER NOT NULL REFERENCES staff(id) ON DELETE CASCADE, PRIMARY KEY(order_id,staff_id))""",
        """CREATE TABLE IF NOT EXISTS work_order_lines(id INTEGER PRIMARY KEY, order_id INTEGER NOT NULL REFERENCES work_orders(id) ON DELETE CASCADE,
             kind TEXT NOT NULL DEFAULT 'Работа', name TEXT NOT NULL, unit TEXT DEFAULT '', plan_qty REAL NOT NULL DEFAULT 0, fact_qty REAL, note TEXT DEFAULT '',
             catalog_id INTEGER, sort_order INTEGER DEFAULT 0)""",
        'CREATE INDEX IF NOT EXISTS idx_work_orders_date ON work_orders(order_date)',
        'CREATE INDEX IF NOT EXISTS idx_work_order_lines_order ON work_order_lines(order_id)',
    ):
        db.execute(sql)
    have = {r[1] for r in db.fetchall('PRAGMA table_info(work_orders)')}
    for name, ddl in PERMIT_COLUMNS:
        if name not in have:
            db.execute(f'ALTER TABLE work_orders ADD COLUMN {name} {ddl}')
    if 'order_id' not in {r[1] for r in db.fetchall('PRAGMA table_info(staff_assign)')}:
        db.execute('ALTER TABLE staff_assign ADD COLUMN order_id INTEGER')
    db.execute('CREATE TRIGGER IF NOT EXISTS work_orders_cleanup AFTER DELETE ON work_orders BEGIN DELETE FROM staff_assign WHERE order_id=old.id; END')


def _day(value):
    return sd._day(value)


# --- ОБЪЕКТЫ -------------------------------------------------------------------------------------

def object_label(db, owner_type, owner_id, text=''):
    if owner_type in EXTRA_OWNERS and owner_id:
        from . import board_domain as board
        return board.link_label(db, owner_type, owner_id) or '(объект удалён)'
    return sd.object_label(db, owner_type, owner_id, text)


def objects(db, search=''):
    """Объекты для выбора: [(owner_type, id, подпись)]. Те же, что в назначениях, плюс договоры СМР / ТО."""
    out = sd.objects(db, search)
    q = search.casefold().strip()
    for table in EXTRA_OWNERS:
        for (rid,) in db.fetchall(f'SELECT id FROM {table} ORDER BY id DESC LIMIT 300'):
            label = object_label(db, table, rid)
            if not q or q in label.casefold():
                out.append((table, rid, label))
    return out


# --- НОМЕР -----------------------------------------------------------------------------------------

def next_number(db, order_date=None, kind=KIND_TASK):
    """Номер по виду документа: наряды-задания «01/26», наряды-допуски «Д-01/26»; счётчик свой для каждого года и вида."""
    day = _day(order_date) or date.today()
    yy = day.year
    row = db.fetchone('SELECT coalesce(max(seq_num),0) FROM work_orders WHERE year_num=? AND kind=?', (yy, kind))
    seq = (row[0] or 0) + 1
    return seq, yy, ('Д-' if kind == KIND_PERMIT else '') + f'{seq:02d}/{yy % 100:02d}'


# --- ЗАПИСЬ ----------------------------------------------------------------------------------------

def _clean_lines(lines):
    out = []
    for r in lines:
        name = str(r.get('name') or '').strip()
        plan = float(str(r.get('plan_qty') or 0).replace(',', '.') or 0)
        if not name and not plan:
            continue
        if not name:
            raise ValueError('Укажите наименование работы или материала')
        if plan < 0:
            raise ValueError(f'«{name}»: количество не может быть отрицательным')
        fact = r.get('fact_qty')
        fact = None if fact in (None, '') else float(str(fact).replace(',', '.'))
        if fact is not None and fact < 0:
            raise ValueError(f'«{name}»: фактическое количество не может быть отрицательным')
        out.append((r.get('kind') if r.get('kind') in vd.KINDS else 'Работа', name, str(r.get('unit') or '').strip(), plan, fact, str(r.get('note') or '').strip(), r.get('catalog_id') or None))
    return out


def save_order(db, values, workers, lines, oid=None):
    """Сохраняет наряд. workers — id сотрудников (если не заданы, берутся члены бригады). Возвращает id."""
    status = values.get('status') or 'Выдан'
    if status not in STATUSES:
        raise ValueError('Неизвестный статус наряда')
    if not _day(values.get('order_date')):
        raise ValueError('Укажите дату наряда')
    start, end = values.get('start_date') or '', values.get('end_date') or ''
    if not _day(start) or not _day(end):
        raise ValueError('Укажите сроки работ (с … по …)')
    if _day(end) < _day(start):
        raise ValueError('Окончание работ раньше начала')
    owner_type, owner_id = values.get('owner_type') or '', values.get('owner_id') or None
    object_text = str(values.get('object_text') or '').strip()
    if not (owner_type and owner_id) and not object_text:
        raise ValueError('Выберите объект или впишите его название')
    crew_id = values.get('crew_id') or None
    workers = list(dict.fromkeys(workers or (sd.crew_members(db, crew_id) if crew_id else [])))
    if not workers:
        raise ValueError('Выберите бригаду или хотя бы одного работника')
    leader_id = values.get('leader_id') or None
    if not leader_id and crew_id:
        row = db.fetchone('SELECT leader_id FROM crews WHERE id=?', (crew_id,))
        leader_id = row[0] if row else None
    clean = _clean_lines(lines)
    kind = values.get('kind') or KIND_TASK
    if kind not in KINDS:
        raise ValueError('Неизвестный вид наряда')
    permit = (str(values.get('hazard') or ''), str(values.get('place') or '').strip(), values.get('supervisor_id') or None, values.get('performer_id') or None,
              values.get('admitter_id') or None, str(values.get('time_from') or '').strip(), str(values.get('time_to') or '').strip(), str(values.get('safety') or '').strip(),
              str(values.get('ppe') or '').strip(), int(bool(values.get('briefing'))), str(values.get('closing_note') or '').strip())
    if kind == KIND_PERMIT:
        if not permit[0]:
            raise ValueError('Выберите вид работ повышенной опасности')
        if not permit[3]:
            raise ValueError('Назначьте ответственного исполнителя работ')
        if permit[3] not in workers:
            raise ValueError('Ответственный исполнитель должен входить в состав работников')
        if not permit[7]:
            raise ValueError('Укажите меры безопасности')
    number = str(values.get('number') or '').strip()
    with db.transaction():
        if not number:
            seq, yy, number = next_number(db, values['order_date'], kind)
        else:
            seq, yy = None, _day(values['order_date']).year
        dup = db.fetchone('SELECT id FROM work_orders WHERE number=? AND year_num=? AND kind=? AND id<>?', (number, yy, kind, oid or 0))
        if dup:
            raise ValueError(f'Наряд № {number} за {yy} год уже есть')
        row = (number, seq, yy, values['order_date'], owner_type, owner_id, object_text, crew_id, leader_id, start, end, status,
               str(values.get('task') or '').strip(), str(values.get('note') or '').strip(), int(bool(values.get('to_schedule', True))))
        if oid:
            old = db.fetchone('SELECT seq_num FROM work_orders WHERE id=?', (oid,))
            if seq is None and old:
                row = (row[0], old[0], *row[2:])
            db.execute('UPDATE work_orders SET number=?,seq_num=?,year_num=?,order_date=?,owner_type=?,owner_id=?,object_text=?,crew_id=?,leader_id=?,start_date=?,end_date=?,'
                       'status=?,task=?,note=?,to_schedule=? WHERE id=?', (*row, oid))
            db.execute('UPDATE work_orders SET kind=?,hazard=?,place=?,supervisor_id=?,performer_id=?,admitter_id=?,time_from=?,time_to=?,safety=?,ppe=?,briefing=?,closing_note=? WHERE id=?', (kind, *permit, oid))
        else:
            oid = db.execute('INSERT INTO work_orders(number,seq_num,year_num,order_date,owner_type,owner_id,object_text,crew_id,leader_id,start_date,end_date,status,task,note,to_schedule) '
                             'VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)', row).lastrowid
            db.execute('UPDATE work_orders SET kind=?,hazard=?,place=?,supervisor_id=?,performer_id=?,admitter_id=?,time_from=?,time_to=?,safety=?,ppe=?,briefing=?,closing_note=? WHERE id=?', (kind, *permit, oid))
        db.execute('DELETE FROM work_order_workers WHERE order_id=?', (oid,))
        for sid in workers:
            db.execute('INSERT INTO work_order_workers(order_id,staff_id) VALUES(?,?)', (oid, sid))
        old_facts = {}
        db.execute('DELETE FROM work_order_lines WHERE order_id=?', (oid,))
        for i, (kind, name, unit, plan, fact, note, cat) in enumerate(clean):
            db.execute('INSERT INTO work_order_lines(order_id,kind,name,unit,plan_qty,fact_qty,note,catalog_id,sort_order) VALUES(?,?,?,?,?,?,?,?,?)', (oid, kind, name, unit, plan, fact, note, cat, i))
        sync_assignments(db, oid)
    return oid


def sync_assignments(db, oid):
    """Ставит работников наряда в график занятости (отменённый и закрытый наряд график не занимает)."""
    db.execute('DELETE FROM staff_assign WHERE order_id=?', (oid,))
    o = db.fetchone('SELECT number,owner_type,owner_id,object_text,start_date,end_date,status,to_schedule FROM work_orders WHERE id=?', (oid,))
    if not o or o[6] not in ACTIVE or not o[7]:
        return 0
    n = 0
    for (sid,) in db.fetchall('SELECT staff_id FROM work_order_workers WHERE order_id=?', (oid,)):
        db.execute('INSERT INTO staff_assign(staff_id,owner_type,owner_id,object_text,start_date,end_date,note,order_id) VALUES(?,?,?,?,?,?,?,?)',
                   (sid, o[1] or '', o[2], o[3] or '', o[4], o[5], f'Наряд № {o[0]}', oid))
        n += 1
    return n


def delete_order(db, oid):
    db.execute('DELETE FROM work_orders WHERE id=?', (oid,))


# --- ЧТЕНИЕ ----------------------------------------------------------------------------------------

def get_order(db, oid):
    cur = db.execute('SELECT * FROM work_orders WHERE id=?', (oid,))
    row = cur.fetchone()
    if not row:
        return None
    o = dict(zip([c[0] for c in cur.description], row))
    o['workers'] = [r[0] for r in db.fetchall('SELECT w.staff_id FROM work_order_workers w JOIN staff s ON s.id=w.staff_id WHERE w.order_id=? ORDER BY s.name', (oid,))]
    cur = db.execute('SELECT id,kind,name,unit,plan_qty,fact_qty,note,catalog_id FROM work_order_lines WHERE order_id=? ORDER BY sort_order,id', (oid,))
    o['lines'] = [dict(zip([c[0] for c in cur.description], r)) for r in cur.fetchall()]
    o['object'] = object_label(db, o['owner_type'], o['owner_id'], o['object_text'])
    return o


def list_orders(db, search='', status='', limit=500, kind=''):
    rows = db.fetchall("""SELECT o.id,o.number,o.order_date,o.owner_type,o.owner_id,o.object_text,coalesce(c.name,''),o.start_date,o.end_date,o.status,
        (SELECT count(*) FROM work_order_workers w WHERE w.order_id=o.id),(SELECT count(*) FROM work_order_lines l WHERE l.order_id=o.id),o.kind
        FROM work_orders o LEFT JOIN crews c ON c.id=o.crew_id WHERE (?='' OR o.status=?) AND (?='' OR o.kind=?)
        ORDER BY coalesce(o.year_num,0) DESC, o.kind, coalesce(o.seq_num,0) DESC, o.id DESC LIMIT ?""", (status, status, kind, kind, limit))
    out = [dict(id=r[0], number=r[1], date=r[2], object=object_label(db, r[3], r[4], r[5]), crew=r[6], start=r[7], end=r[8], status=r[9], workers=r[10], lines=r[11], kind=r[12]) for r in rows]
    if search:
        needle = search.casefold()
        out = [r for r in out if needle in f'{r["number"]} {r["object"]} {r["crew"]} {r["status"]}'.casefold()]
    return out


def worker_names(db, oid):
    return [r[0] for r in db.fetchall('SELECT s.name FROM work_order_workers w JOIN staff s ON s.id=w.staff_id WHERE w.order_id=? ORDER BY s.name', (oid,))]


# --- ПРОВЕРКИ --------------------------------------------------------------------------------------

def warnings(db, workers, start, end, oid=None):
    """Предупреждения при выдаче: просроченные допуски и пересечения с другими объектами. Не блокируют сохранение."""
    s, e = _day(start), _day(end)
    if not s or not e:
        return []
    out = []
    for sid in dict.fromkeys(workers):
        name = (db.fetchone('SELECT name FROM staff WHERE id=?', (sid,)) or ('?',))[0]
        expired = sd.expired_for(db, sid, s)
        if expired:
            out.append(f'{name}: просрочен допуск — ' + ', '.join(expired))
        for aid, osid, ocid, otype, ooid, text, a, b, note in db.fetchall(
                'SELECT id,staff_id,crew_id,owner_type,owner_id,object_text,start_date,end_date,note FROM staff_assign WHERE start_date<=? AND end_date>=? AND coalesce(order_id,0)<>?',
                (e.strftime(ISO), s.strftime(ISO), oid or 0)):
            if sid in sd.members_of(db, osid, ocid):
                out.append(f'{name}: уже назначен на «{sd.object_label(db, otype, ooid, text)}» с {display_date(a)} по {display_date(b)}')
    return out


# --- ОБЪЁМЫ ----------------------------------------------------------------------------------------

def lines_from_volumes(db, owner_type, owner_id):
    """Плановые строки наряда из вкладки «Объёмы» договора."""
    if owner_type not in vd.OWNERS or not owner_id:
        return []
    return [dict(kind=v['kind'], name=v['name'], unit=v['unit'], plan_qty=v['quantity'], catalog_id=v['catalog_id'], note='') for v in vd.items(db, owner_type, owner_id)]


def close_order(db, oid, facts, closed_date=None):
    """Закрывает наряд. facts: {id строки: фактическое количество}; строки без факта принимаются по плану."""
    o = get_order(db, oid)
    if not o:
        raise ValueError('Наряд не найден')
    if o['status'] == 'Отменён':
        raise ValueError('Отменённый наряд закрыть нельзя')
    with db.transaction():
        for line in o['lines']:
            fact = facts.get(line['id'])
            if fact in (None, ''):
                fact = line['fact_qty'] if line['fact_qty'] is not None else line['plan_qty']
            fact = float(str(fact).replace(',', '.'))
            if fact < 0:
                raise ValueError(f'«{line["name"]}»: фактическое количество не может быть отрицательным')
            db.execute('UPDATE work_order_lines SET fact_qty=? WHERE id=?', (fact, line['id']))
        db.execute("UPDATE work_orders SET status='Закрыт',closed_date=? WHERE id=?", (closed_date or date.today().isoformat(), oid))
        sync_assignments(db, oid)


def apply_facts_to_volumes(db, oid):
    """Факт закрытого наряда добавляется к объёмам договора (одинаковые вид и наименование суммируются). Возвращает число строк."""
    o = get_order(db, oid)
    if not o or o['status'] != 'Закрыт':
        raise ValueError('Вносить в объёмы можно только закрытый наряд')
    if o['facts_applied']:
        raise ValueError('Факт этого наряда уже внесён в объёмы')
    if o['owner_type'] not in vd.OWNERS or not o['owner_id']:
        raise ValueError('У объекта наряда нет вкладки «Объёмы»')
    rows = vd.items(db, o['owner_type'], o['owner_id'])
    index = {(r['kind'], r['name'].casefold()): r for r in rows}
    n = 0
    with db.transaction():
        for line in o['lines']:
            fact = line['fact_qty']
            if not fact:
                continue
            hit = index.get((line['kind'], line['name'].casefold()))
            if hit:
                hit['quantity'] = float(hit['quantity']) + float(fact)
            else:
                new = dict(kind=line['kind'], name=line['name'], unit=line['unit'], quantity=fact, catalog_id=line['catalog_id'], note=f'Наряд № {o["number"]}')
                rows.append(new)
                index[(new['kind'], new['name'].casefold())] = new
            n += 1
        vd.save_items(db, o['owner_type'], o['owner_id'], rows)
        db.execute('UPDATE work_orders SET facts_applied=1 WHERE id=?', (oid,))
    return n


# --- ДОКУМЕНТ ----------------------------------------------------------------------------------------

def _num(v):
    if v in (None, ''):
        return ''
    return ('%g' % float(v))


def _staff_name(db, sid):
    return (db.fetchone('SELECT name FROM staff WHERE id=?', (sid,)) or ('',))[0] if sid else ''


def tags(db, oid):
    o = get_order(db, oid)
    names = worker_names(db, oid)
    crew = (db.fetchone('SELECT name FROM crews WHERE id=?', (o['crew_id'],)) or ('',))[0] if o['crew_id'] else ''
    leader = (db.fetchone('SELECT name FROM staff WHERE id=?', (o['leader_id'],)) or ('',))[0] if o['leader_id'] else ''
    plan = [f'{i}. {l["name"]} — {_num(l["plan_qty"])} {l["unit"]}'.rstrip() for i, l in enumerate(o['lines'], 1)]
    fact = [f'{i}. {l["name"]} — {_num(l["fact_qty"] if l["fact_qty"] is not None else "")} {l["unit"]}'.rstrip() for i, l in enumerate(o['lines'], 1)]
    return {'НОМЕР_НАРЯДА': o['number'], 'ДАТА_НАРЯДА': display_date(o['order_date']), 'ОБЪЕКТ': o['object'], 'БРИГАДА': crew, 'БРИГАДИР': leader,
            'РАБОТНИКИ': '\n'.join(f'{i}. {n}' for i, n in enumerate(names, 1)), 'СРОК_С': display_date(o['start_date']), 'СРОК_ПО': display_date(o['end_date']),
            'ЗАДАНИЕ': o['task'], 'РАБОТЫ': '\n'.join(plan), 'ФАКТ': '\n'.join(fact), 'ПРИМЕЧАНИЕ': o['note'], 'СТАТУС': o['status'], 'ДАТА_ЗАКРЫТИЯ': display_date(o['closed_date']),
            'ВИД_ДОКУМЕНТА': o['kind'], 'ВИД_РАБОТ': o['hazard'], 'МЕСТО_РАБОТ': o['place'] or o['object'], 'ОТВ_РУКОВОДИТЕЛЬ': _staff_name(db, o['supervisor_id']),
            'ОТВ_ИСПОЛНИТЕЛЬ': _staff_name(db, o['performer_id']), 'ДОПУСКАЮЩИЙ': _staff_name(db, o['admitter_id']), 'ВРЕМЯ_С': o['time_from'], 'ВРЕМЯ_ПО': o['time_to'],
            'МЕРЫ_БЕЗОПАСНОСТИ': o['safety'], 'СИЗ': o['ppe'], 'ИНСТРУКТАЖ': 'проведён' if o['briefing'] else 'не проведён', 'ОТМЕТКА_О_ЗАКРЫТИИ': o['closing_note']}


TAG_HELP = [('НОМЕР_НАРЯДА', 'номер наряда'), ('ДАТА_НАРЯДА', 'дата выдачи'), ('ОБЪЕКТ', 'объект'), ('БРИГАДА', 'бригада'), ('БРИГАДИР', 'бригадир'), ('РАБОТНИКИ', 'список работников'),
            ('СРОК_С', 'начало работ'), ('СРОК_ПО', 'окончание работ'), ('ЗАДАНИЕ', 'задание'), ('РАБОТЫ', 'плановые работы и материалы'), ('ФАКТ', 'фактически выполнено'),
            ('ПРИМЕЧАНИЕ', 'примечание'), ('СТАТУС', 'статус'), ('ДАТА_ЗАКРЫТИЯ', 'дата закрытия'),
            ('ВИД_ДОКУМЕНТА', 'наряд-задание / наряд-допуск'), ('ВИД_РАБОТ', 'вид работ повышенной опасности (допуск)'), ('МЕСТО_РАБОТ', 'место проведения работ'),
            ('ОТВ_РУКОВОДИТЕЛЬ', 'ответственный руководитель работ'), ('ОТВ_ИСПОЛНИТЕЛЬ', 'ответственный исполнитель'), ('ДОПУСКАЮЩИЙ', 'допускающий'),
            ('ВРЕМЯ_С', 'время начала'), ('ВРЕМЯ_ПО', 'время окончания'), ('МЕРЫ_БЕЗОПАСНОСТИ', 'меры безопасности'), ('СИЗ', 'средства индивидуальной защиты'),
            ('ИНСТРУКТАЖ', 'целевой инструктаж проведён'), ('ОТМЕТКА_О_ЗАКРЫТИИ', 'отметка о закрытии допуска')]
TEMPLATE_SETTING = 'orders_template'
PERMIT_TEMPLATE_SETTING = 'orders_template_permit'


def template_setting(kind=KIND_TASK):
    return PERMIT_TEMPLATE_SETTING if kind == KIND_PERMIT else TEMPLATE_SETTING


def create_default_template(kind=KIND_TASK):
    folder = gd.TEMPLATES_DIR / 'наряды'
    folder.mkdir(parents=True, exist_ok=True)
    permit = kind == KIND_PERMIT
    path = folder / ('Шаблон_наряда_допуска.docx' if permit else 'Шаблон_наряда.docx')
    if path.exists():
        return path
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = docx.Document()
    if permit:
        doc.add_heading('НАРЯД-ДОПУСК № {НОМЕР_НАРЯДА}', 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph('на производство работ повышенной опасности: {ВИД_РАБОТ}').alignment = WD_ALIGN_PARAGRAPH.CENTER
        lines = ('Дата выдачи: {ДАТА_НАРЯДА}', 'Объект: {ОБЪЕКТ}', 'Место проведения работ: {МЕСТО_РАБОТ}', 'Срок действия: с {СРОК_С} {ВРЕМЯ_С} по {СРОК_ПО} {ВРЕМЯ_ПО}',
                 'Содержание работ: {ЗАДАНИЕ}', 'Ответственный руководитель работ: {ОТВ_РУКОВОДИТЕЛЬ}', 'Ответственный исполнитель: {ОТВ_ИСПОЛНИТЕЛЬ}', 'Допускающий: {ДОПУСКАЮЩИЙ}',
                 'Состав бригады:', '{РАБОТНИКИ}', 'Меры безопасности:', '{МЕРЫ_БЕЗОПАСНОСТИ}', 'СИЗ: {СИЗ}', 'Целевой инструктаж: {ИНСТРУКТАЖ}',
                 'Объём работ:', '{РАБОТЫ}', 'Отметка о закрытии наряда-допуска ({ДАТА_ЗАКРЫТИЯ}): {ОТМЕТКА_О_ЗАКРЫТИИ}',
                 'Наряд выдал ____________________   Допускающий ____________________   Ответственный исполнитель ____________________')
    else:
        doc.add_heading('НАРЯД-ЗАДАНИЕ № {НОМЕР_НАРЯДА}', 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
        lines = ('Дата выдачи: {ДАТА_НАРЯДА}', 'Объект: {ОБЪЕКТ}', 'Бригада: {БРИГАДА}', 'Бригадир: {БРИГАДИР}', 'Срок выполнения: с {СРОК_С} по {СРОК_ПО}',
                 'Задание: {ЗАДАНИЕ}', 'Состав работников:', '{РАБОТНИКИ}', 'Объём работ и материалы:', '{РАБОТЫ}', 'Выполнено фактически:', '{ФАКТ}', 'Примечание: {ПРИМЕЧАНИЕ}',
                 'Наряд выдал ____________________        Наряд принял ____________________')
    for line in lines:
        doc.add_paragraph(line)
    doc.save(path)
    return path


def template_path(db, kind=KIND_TASK):
    path = db.get_setting(template_setting(kind), '')
    return path if path and os.path.exists(path) else str(create_default_template(kind))


def generate(db, oid, out_dir):
    """Формирует документ наряда (формат Word или Excel — как у шаблона) в папке out_dir; возвращает путь."""
    o = get_order(db, oid)
    if not o:
        raise ValueError('Наряд не найден')
    os.makedirs(out_dir, exist_ok=True)
    prefix = 'Наряд-допуск' if o['kind'] == KIND_PERMIT else 'Наряд'
    base = os.path.join(out_dir, gd.safe_name(f'{prefix}_{o["number"].replace("/", "-")}').replace(' ', '_'))
    out = gd.render_file(template_path(db, o['kind']), base + '.docx', tags(db, oid))
    db.execute('UPDATE work_orders SET doc_path=? WHERE id=?', (out, oid))
    return out
