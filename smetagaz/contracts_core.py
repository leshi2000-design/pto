"""Общее ядро модулей «Юрлица» и «СМР»: договоры с контрагентами, акты, реквизиты, НДС, теги и документы.

Оба модуля устроены одинаково (договор → акты → справки о стоимости, оплаты, папка), но независимы друг от друга и от остальных разделов:
у каждого свои таблицы (`le_*` и `smr_*`), свои теги и шаблоны. Различия:
  * «Юрлица» — контрагент только организация из справочника юрлиц; смет нет;
  * «СМР» — контрагент физическое лицо (из «Клиентов») или юрлицо; есть смета; исполнительной документации и сертификатов нет.
Модуль не зависит от Qt.
"""
import hashlib
import json
import os
import re
from datetime import date, datetime
from pathlib import Path

from . import gsv_project_domain as gd
from . import gsvm_domain as md
from .agenda_domain import short_name
from .gsvm_docs import FORMATS, apply_format, money, norm_tag, tag_text

MODS = {
    'le': dict(title='Юрлица', contracts='le_contracts', acts='le_acts', tags='le_tags', docs='le_docs', section='Юрлица', code='04', party='legal'),
    'smr': dict(title='СМР', contracts='smr_contracts', acts='smr_acts', tags='smr_tags', docs='smr_docs', section='СМР', code='05', party='both'),
}
DOC_KINDS = {
    'contract': ('Договор', 'contract', 'Договор'),
    'act': ('Акт выполненных работ', 'act', 'Акт'),
    'statement': ('Справка о стоимости выполненных работ', 'act', 'Справка'),
}
STATUSES = ('Действует', 'Завершён', 'Расторгнут')
LEGAL_FIELDS = ('full_name', 'legal_address', 'postal_address', 'bank', 'bic', 'account', 'okpo', 'head_position', 'head_name', 'head_position_gen', 'head_name_gen', 'basis')


def cfg(mod):
    return MODS[mod]


# --- МИГРАЦИЯ --------------------------------------------------------------------------------

def _cols(db, table):
    return {r[1] for r in db.fetchall(f'PRAGMA table_info({table})')}


def _add(db, table, items):
    have = _cols(db, table)
    for name, ddl in items:
        if name not in have:
            db.execute(f'ALTER TABLE {table} ADD COLUMN {name} {ddl}')


def migrate(db):
    _add(db, 'le_clients', [('full_name', "TEXT DEFAULT ''"), ('legal_address', "TEXT DEFAULT ''"), ('postal_address', "TEXT DEFAULT ''"), ('bank', "TEXT DEFAULT ''"),
                            ('bic', "TEXT DEFAULT ''"), ('account', "TEXT DEFAULT ''"), ('okpo', "TEXT DEFAULT ''"), ('head_position', "TEXT DEFAULT 'Директор'"),
                            ('head_name', "TEXT DEFAULT ''"), ('head_position_gen', "TEXT DEFAULT 'директора'"), ('head_name_gen', "TEXT DEFAULT ''"),
                            ('basis', "TEXT DEFAULT 'Устава'")])
    db.execute("""CREATE TABLE IF NOT EXISTS le_contacts(id INTEGER PRIMARY KEY, client_id INTEGER NOT NULL REFERENCES le_clients(id) ON DELETE CASCADE,
        name TEXT DEFAULT '', position TEXT DEFAULT '', phone TEXT DEFAULT '', email TEXT DEFAULT '')""")
    common = [('subject', "TEXT DEFAULT ''"), ('object_address', "TEXT DEFAULT ''"), ('start_date', "TEXT DEFAULT ''"), ('end_date', "TEXT DEFAULT ''"),
              ('vat_rate', 'REAL DEFAULT 0'), ('signed', 'INTEGER DEFAULT 0'), ('folder', "TEXT DEFAULT ''"), ('seq_num', 'INTEGER'), ('year_num', 'INTEGER')]
    _add(db, 'le_contracts', common + [('source_type', "TEXT DEFAULT ''"), ('source_id', 'INTEGER')])
    _add(db, 'le_acts', [('signed', 'INTEGER DEFAULT 0')])
    db.execute("""CREATE TABLE IF NOT EXISTS smr_contracts(id INTEGER PRIMARY KEY, contract_number TEXT DEFAULT '', contract_date TEXT DEFAULT '',
        party_type TEXT NOT NULL DEFAULT 'person' CHECK(party_type IN ('person','legal')), person_id INTEGER, legal_id INTEGER REFERENCES le_clients(id) ON DELETE RESTRICT,
        object_name TEXT DEFAULT '', object_address TEXT DEFAULT '', subject TEXT DEFAULT '', start_date TEXT DEFAULT '', end_date TEXT DEFAULT '',
        amount REAL DEFAULT 0, vat_included INTEGER DEFAULT 0, vat_rate REAL DEFAULT 0, signed INTEGER DEFAULT 0, status TEXT DEFAULT 'Действует', note TEXT DEFAULT '',
        folder TEXT DEFAULT '', seq_num INTEGER, year_num INTEGER, estimate_id INTEGER UNIQUE REFERENCES estimates(id) ON DELETE SET NULL,
        est_materials REAL DEFAULT 0, est_works REAL DEFAULT 0, est_rev_synced INTEGER DEFAULT 0, est_synced_at TEXT DEFAULT '')""")
    db.execute("""CREATE TABLE IF NOT EXISTS smr_acts(id INTEGER PRIMARY KEY, contract_id INTEGER NOT NULL REFERENCES smr_contracts(id) ON DELETE CASCADE,
        act_number TEXT DEFAULT '', act_date TEXT DEFAULT '', amount REAL DEFAULT 0, vat_included INTEGER DEFAULT 0, description TEXT DEFAULT '', note TEXT DEFAULT '',
        signed INTEGER DEFAULT 0)""")
    for mod, c in MODS.items():
        db.execute(f"""CREATE TABLE IF NOT EXISTS {c['tags']}(id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE, source TEXT NOT NULL, fmt TEXT NOT NULL DEFAULT 'text',
            empty_text TEXT DEFAULT '', auto_key TEXT UNIQUE, enabled INTEGER NOT NULL DEFAULT 1, note TEXT DEFAULT '')""")
        db.execute(f"""CREATE TABLE IF NOT EXISTS {c['docs']}(scope TEXT NOT NULL, ref_id INTEGER NOT NULL, kind TEXT NOT NULL, file_path TEXT NOT NULL, data_hash TEXT NOT NULL DEFAULT '',
            generated_at TEXT NOT NULL DEFAULT '', linked INTEGER NOT NULL DEFAULT 0, PRIMARY KEY(scope,ref_id,kind))""")
        db.execute(f"CREATE INDEX IF NOT EXISTS idx_{c['acts']}_contract ON {c['acts']}(contract_id)")
        # оплаты не должны исчезать вместе с договором
        db.execute(f"DROP TRIGGER IF EXISTS protect_paid_{c['contracts']}")
        db.execute(f"""CREATE TRIGGER protect_paid_{c['contracts']} BEFORE DELETE ON {c['contracts']}
            WHEN EXISTS(SELECT 1 FROM payments WHERE owner_type='{c['contracts']}' AND owner_id=old.id)
            BEGIN SELECT RAISE(ABORT,'У договора есть оплаты. Удалите оплаты или сохраните договор для истории расчётов.'); END""")
        db.execute(f"""CREATE TRIGGER IF NOT EXISTS cleanup_docs_{c['contracts']} AFTER DELETE ON {c['contracts']} BEGIN
            DELETE FROM {c['docs']} WHERE scope='contract' AND ref_id=old.id; END""")
        db.execute(f"""CREATE TRIGGER IF NOT EXISTS cleanup_act_docs_{c['acts']} AFTER DELETE ON {c['acts']} BEGIN
            DELETE FROM {c['docs']} WHERE scope='act' AND ref_id=old.id; END""")
    # заказчик-юрлицо в карточках проектов и монтажа ГСВ
    for table in ('gsv_projects', 'contracts'):
        _add(db, table, [('le_client_id', 'INTEGER'), ('party_name', "TEXT DEFAULT ''")])
    if db.get_setting('le_signed_default', '0') != '1':
        db.execute("UPDATE le_contracts SET signed=1 WHERE coalesce(contract_date,'')<>''")
        db.execute("UPDATE le_acts SET signed=1 WHERE coalesce(act_date,'')<>''")
        db.set_setting('le_signed_default', '1')
    db.execute("CREATE INDEX IF NOT EXISTS idx_le_contacts_client ON le_contacts(client_id)")


# --- ПРОСТЫЕ ДАННЫЕ --------------------------------------------------------------------------

def _row(db, sql, params):
    cur = db.execute(sql, params)
    row = cur.fetchone()
    return dict(zip([c[0] for c in cur.description], row)) if row else None


def contract(db, mod, cid):
    row = _row(db, f"SELECT * FROM {cfg(mod)['contracts']} WHERE id=?", (cid,))
    if not row:
        raise ValueError('Сначала сохраните договор')
    return row


def act(db, mod, aid):
    row = _row(db, f"SELECT * FROM {cfg(mod)['acts']} WHERE id=?", (aid,))
    if not row:
        raise ValueError('Акт не найден')
    return row


def acts(db, mod, cid):
    cur = db.execute(f"SELECT * FROM {cfg(mod)['acts']} WHERE contract_id=? ORDER BY act_date,id", (cid,))
    return [dict(zip([c[0] for c in cur.description], r)) for r in cur.fetchall()]


def legal(db, lid):
    return _row(db, 'SELECT * FROM le_clients WHERE id=?', (lid,)) if lid else None


def contacts(db, lid):
    cur = db.execute('SELECT id,name,position,phone,email FROM le_contacts WHERE client_id=? ORDER BY id', (lid,))
    return [dict(zip([c[0] for c in cur.description], r)) for r in cur.fetchall()]


def save_legal(db, values, lid=None, contact_rows=()):
    """Сохраняет юрлицо и его контактных лиц. Название — краткое имя в справочнике; полное — для документов."""
    name = str(values.get('name') or '').strip()
    if not name:
        raise ValueError('Укажите наименование организации')
    unp = str(values.get('unp') or '').strip()
    if unp and not re.fullmatch(r'[0-9A-Za-zА-Яа-я]{9,12}', unp):
        raise ValueError('УНП: 9 цифр (допускаются буквы для ИНН других стран, 9–12 знаков)')
    keys = ['name', 'unp', 'address', 'contact_person', 'phone', 'email', 'note'] + list(LEGAL_FIELDS)
    data = {k: str(values.get(k) or '').strip() for k in keys}
    data['address'] = data['address'] or data['legal_address']
    with db.transaction():
        if lid:
            db.execute('UPDATE le_clients SET ' + ','.join(f'{k}=?' for k in keys) + ' WHERE id=?', (*[data[k] for k in keys], lid))
        else:
            lid = db.execute('INSERT INTO le_clients(' + ','.join(keys) + ') VALUES(' + ','.join('?' for _ in keys) + ')', [data[k] for k in keys]).lastrowid
        db.execute('DELETE FROM le_contacts WHERE client_id=?', (lid,))
        for c in contact_rows:
            if any(str(c.get(k) or '').strip() for k in ('name', 'phone', 'email')):
                db.execute('INSERT INTO le_contacts(client_id,name,position,phone,email) VALUES(?,?,?,?,?)',
                           (lid, *[str(c.get(k) or '').strip() for k in ('name', 'position', 'phone', 'email')]))
    return lid


def party(db, mod, c):
    """Контрагент договора в едином виде (юрлицо или физлицо)."""
    kind = 'legal' if mod == 'le' or c.get('party_type') == 'legal' else 'person'
    p = dict(kind=kind, name='', full_name='', unp='', okpo='', legal_address='', postal_address='', bank='', bic='', account='', head_position='', head_name='',
             head_position_gen='', head_name_gen='', basis='', head_short='', contact='', phone='', email='', passport='', passport_issuer='', passport_date='', address='')
    if kind == 'legal':
        l = legal(db, c.get('client_id') if mod == 'le' else c.get('legal_id'))
        if l:
            p.update({k: l.get(k) or '' for k in ('name', 'unp', 'okpo', 'legal_address', 'postal_address', 'bank', 'bic', 'account', 'head_position', 'head_name',
                                                    'head_position_gen', 'head_name_gen', 'basis', 'phone', 'email')})
            p['full_name'] = l.get('full_name') or l.get('name') or ''
            p['address'] = l.get('legal_address') or l.get('address') or ''
            p['postal_address'] = l.get('postal_address') or p['address']
            p['head_short'] = short_name(l.get('head_name') or '')
            contacts_list = contacts(db, l['id'])
            p['contact'] = (contacts_list[0]['name'] if contacts_list else l.get('contact_person')) or ''
            if contacts_list and contacts_list[0]['phone']:
                p['phone'] = contacts_list[0]['phone']
            if contacts_list and contacts_list[0]['email']:
                p['email'] = contacts_list[0]['email']
    else:
        from .gsv_domain import get_client
        client = get_client(db, c.get('person_id')) if c.get('person_id') else None
        if client:
            p.update(name=client['name'] or '', full_name=client['name'] or '', phone=client.get('phone') or '', passport=client.get('passport') or '',
                     passport_issuer=client.get('passport_issuer') or '', passport_date=client.get('passport_date') or '', address=client.get('address') or '')
            p['head_short'] = short_name(client['name'] or '')
    return p


def party_label(db, mod, c):
    p = party(db, mod, c)
    return p['name'] or '(контрагент не выбран)'


def vat_parts(amount, included, rate):
    """(без НДС, НДС, всего). Сумма договора — итоговая; при «в том числе НДС» налог выделяется из неё."""
    amount = float(amount or 0)
    rate = float(rate or 0)
    if not included or rate <= 0:
        return round(amount, 2), 0.0, round(amount, 2)
    vat = round(amount * rate / (100 + rate), 2)
    return round(amount - vat, 2), vat, round(amount, 2)


# --- НОМЕР, ПАПКА, ОПЛАТЫ --------------------------------------------------------------------

def next_number(db, mod, contract_date=None):
    """(порядковый, год, 'NN-<код раздела>/YY'): для юрлиц раздел 04, для СМР раздел 05."""
    c = cfg(mod)
    day = gd._parse(contract_date) or date.today()
    yy = day.year % 100
    top = db.fetchone(f"SELECT coalesce(max(seq_num),0) FROM {c['contracts']} WHERE year_num=?", (yy,))[0]
    pattern = re.compile(rf"^(\d+)-{c['code']}/(\d{{2}})$")
    for (number,) in db.fetchall(f"SELECT contract_number FROM {c['contracts']} WHERE contract_number LIKE '%-{c['code']}/%'"):
        m = pattern.match(str(number or '').strip())
        if m and int(m.group(2)) == yy:
            top = max(top, int(m.group(1)))
    return top + 1, yy, f'{top + 1:02}-{c["code"]}/{yy:02}'


def folder_name(number, address='', party_name=''):
    """«01-04.26, адрес объекта (контрагент)»: в номере «/» заменяется на «.»."""
    name = str(number or '').strip().replace('/', '.')
    if str(address or '').strip():
        name += f', {" ".join(str(address).split())[:90]}'
    if str(party_name or '').strip():
        name += f' ({" ".join(str(party_name).split())})'
    return gd.safe_name(name)


def suggested_folder_name(db, mod, cid):
    c = contract(db, mod, cid)
    p = party(db, mod, c)
    label = p['name'] if p['kind'] == 'legal' else short_name(p['name'])
    return folder_name(c['contract_number'], c.get('object_address') or c.get('object_name'), label)


def contract_folder(db, mod, cid):
    row = db.fetchone(f"SELECT folder FROM {cfg(mod)['contracts']} WHERE id=?", (cid,))
    return row[0] if row and row[0] and os.path.isdir(row[0]) else ''


def link_folder(db, mod, cid, path, create=False):
    path = str(path)
    if create:
        os.makedirs(path, exist_ok=True)
    if not os.path.isdir(path):
        raise ValueError('Папка не найдена: ' + path)
    db.execute(f"UPDATE {cfg(mod)['contracts']} SET folder=? WHERE id=?", (path, cid))
    return path


def section_root(mod):
    return gd.PROJECTS_DIR / cfg(mod)['section']


def payments(db, mod, cid):
    from . import payments_domain
    return sorted(((d, float(a), n or '') for _i, d, a, n in payments_domain.history(db, cfg(mod)['contracts'], cid)), key=lambda p: p[0])


# --- СМЕТА (только СМР) ----------------------------------------------------------------------

def estimate_state(db, cid):
    row = db.fetchone('SELECT estimate_id,est_rev_synced,est_synced_at FROM smr_contracts WHERE id=?', (cid,))
    if not row or not row[0]:
        return 'none'
    if not row[2]:
        return 'unsynced'
    return 'stale' if md.last_change(db, row[0]) > (row[1] or 0) else 'fresh'


def sync_estimate(db, cid):
    row = db.fetchone('SELECT estimate_id FROM smr_contracts WHERE id=?', (cid,))
    if not row or not row[0]:
        raise ValueError('К договору не привязана смета')
    s = md.estimate_summary(db, row[0])
    if not s:
        raise ValueError('Смета не найдена')
    db.execute('UPDATE smr_contracts SET amount=?, est_materials=?, est_works=?, est_rev_synced=?, est_synced_at=? WHERE id=?',
               (round(s['total'], 2), round(s['materials'], 2), round(s['works'], 2), md.last_change(db, row[0]), datetime.now().isoformat(timespec='seconds'), cid))
    return s


def link_estimate(db, cid, eid):
    from . import payments_domain
    for table in ('contracts', 'smr_contracts'):
        other = db.fetchone(f'SELECT id FROM {table} WHERE estimate_id=?' + (' AND id<>?' if table == 'smr_contracts' else ''), (eid, cid) if table == 'smr_contracts' else (eid,))
        if other:
            raise ValueError('Эта смета уже привязана к другому договору')
    payments_domain.link(db, 'smr_contracts', cid, eid)
    db.execute("UPDATE smr_contracts SET est_synced_at='', est_rev_synced=0 WHERE id=?", (cid,))


# --- ТЕГИ -------------------------------------------------------------------------------------

STATIC_SOURCES = [
    ('contract.number', 'Договор: номер'), ('contract.date', 'Договор: дата'), ('contract.subject', 'Договор: предмет'), ('contract.object', 'Договор: объект'),
    ('contract.object_address', 'Договор: адрес объекта'), ('contract.start', 'Договор: начало работ'), ('contract.end', 'Договор: окончание работ'),
    ('contract.amount', 'Договор: сумма (итого)'), ('contract.amount_net', 'Договор: сумма без НДС'), ('contract.vat', 'Договор: сумма НДС'),
    ('contract.vat_rate', 'Договор: ставка НДС, %'), ('contract.vat_note', 'Договор: пояснение про НДС'), ('payments.paid', 'Оплачено всего'), ('payments.rest', 'Остаток к оплате'),
    ('payments.dates', 'Оплаты: даты (по строкам)'), ('payments.amounts', 'Оплаты: суммы (по строкам)'),
    ('party.name', 'Контрагент: краткое наименование / ФИО'), ('party.full_name', 'Контрагент: полное наименование / ФИО'), ('party.unp', 'Контрагент: УНП'),
    ('party.okpo', 'Контрагент: ОКПО'), ('party.legal_address', 'Контрагент: юридический адрес'), ('party.postal_address', 'Контрагент: почтовый адрес'),
    ('party.bank', 'Контрагент: банк'), ('party.bic', 'Контрагент: БИК'), ('party.account', 'Контрагент: расчётный счёт'),
    ('party.head_position', 'Контрагент: должность руководителя'), ('party.head_name', 'Контрагент: ФИО руководителя'),
    ('party.head_position_gen', 'Контрагент: должность руководителя (в родительном падеже)'), ('party.head_name_gen', 'Контрагент: ФИО руководителя (в родительном падеже)'),
    ('party.head_short', 'Контрагент: инициалы и фамилия (подпись)'), ('party.basis', 'Контрагент: действует на основании'), ('party.contact', 'Контрагент: контактное лицо'),
    ('party.phone', 'Контрагент: телефон'), ('party.email', 'Контрагент: e-mail'), ('party.passport', 'Физлицо: паспорт'), ('party.passport_issuer', 'Физлицо: кем выдан'),
    ('party.passport_date', 'Физлицо: дата выдачи'), ('party.address', 'Контрагент: адрес'),
    ('act.number', 'Акт: номер'), ('act.date', 'Акт: дата'), ('act.amount', 'Акт: сумма (итого)'), ('act.amount_net', 'Акт: сумма без НДС'), ('act.vat', 'Акт: сумма НДС'),
    ('act.description', 'Акт: описание работ'),
    ('estimate.materials', 'Смета: стоимость материалов'), ('estimate.works', 'Смета: стоимость работ'),
]
BUILTIN_TAGS = [
    ('НОМЕР_ДОГОВОРА', 'contract.number', 'text'), ('ДАТА_ДОГОВОРА', 'contract.date', 'date_long'), ('ДАТА_ДОГОВОРА_С', 'contract.date', 'date_short'),
    ('ПРЕДМЕТ', 'contract.subject', 'text'), ('ОБЪЕКТ', 'contract.object', 'text'), ('АДРЕС_ОБЪЕКТА', 'contract.object_address', 'text'),
    ('НАЧАЛО_РАБОТ', 'contract.start', 'date_long'), ('ОКОНЧАНИЕ_РАБОТ', 'contract.end', 'date_long'),
    ('СУММА', 'contract.amount', 'money'), ('СУММА_ПРОПИСЬЮ', 'contract.amount', 'money_words'), ('СУММА_БЕЗ_НДС', 'contract.amount_net', 'money'),
    ('СУММА_НДС', 'contract.vat', 'money'), ('СТАВКА_НДС', 'contract.vat_rate', 'text'), ('НДС_ПОЯСНЕНИЕ', 'contract.vat_note', 'text'),
    ('ОПЛАЧЕНО', 'payments.paid', 'money'), ('ОСТАТОК', 'payments.rest', 'money'),
    ('КОНТРАГЕНТ', 'party.full_name', 'text'), ('КОНТРАГЕНТ_КРАТКО', 'party.name', 'text'), ('УНП', 'party.unp', 'text'), ('ОКПО', 'party.okpo', 'text'),
    ('АДРЕС_ЮРИДИЧЕСКИЙ', 'party.legal_address', 'text'), ('АДРЕС_ПОЧТОВЫЙ', 'party.postal_address', 'text'), ('БАНК', 'party.bank', 'text'), ('БИК', 'party.bic', 'text'),
    ('РАСЧЕТНЫЙ_СЧЕТ', 'party.account', 'text'), ('ДОЛЖНОСТЬ_РУКОВОДИТЕЛЯ', 'party.head_position', 'text'), ('ФИО_РУКОВОДИТЕЛЯ', 'party.head_name', 'text'),
    ('ДОЛЖНОСТЬ_РУКОВОДИТЕЛЯ_РОД', 'party.head_position_gen', 'text'), ('ФИО_РУКОВОДИТЕЛЯ_РОД', 'party.head_name_gen', 'text'),
    ('ФИО_РУКОВОДИТЕЛЯ_СОКР', 'party.head_short', 'text'), ('ОСНОВАНИЕ', 'party.basis', 'text'), ('КОНТАКТНОЕ_ЛИЦО', 'party.contact', 'text'),
    ('ТЕЛЕФОН', 'party.phone', 'text'), ('EMAIL', 'party.email', 'text'), ('ПАСПОРТ', 'party.passport', 'text'), ('КЕМ_ВЫДАН', 'party.passport_issuer', 'text'),
    ('ДАТА_ВЫДАЧИ', 'party.passport_date', 'date_short'), ('АДРЕС_КОНТРАГЕНТА', 'party.address', 'text'),
    ('НОМЕР_АКТА', 'act.number', 'text'), ('ДАТА_АКТА', 'act.date', 'date_long'), ('ДАТА_АКТА_С', 'act.date', 'date_short'),
    ('СУММА_АКТА', 'act.amount', 'money'), ('СУММА_АКТА_ПРОПИСЬЮ', 'act.amount', 'money_words'), ('СУММА_АКТА_БЕЗ_НДС', 'act.amount_net', 'money'),
    ('СУММА_АКТА_НДС', 'act.vat', 'money'), ('ОПИСАНИЕ_РАБОТ', 'act.description', 'text'), ('МЕСЯЦ_РАБОТ', 'act.date', 'month_name'), ('МЕСЯЦ_РАБОТ_ГОД', 'act.date', 'month_year'),
    ('СТОИМОСТЬ_МАТЕРИАЛОВ', 'estimate.materials', 'money'), ('СТОИМОСТЬ_РАБОТ', 'estimate.works', 'money'),
]


def available_sources(db, mod):
    return [s for s in STATIC_SOURCES if mod == 'smr' or not s[0].startswith('estimate.')]


def sync_tags(db, mod):
    table = cfg(mod)['tags']
    for name, source, fmt in BUILTIN_TAGS:
        if source.startswith('estimate.') and mod != 'smr':
            continue
        if not db.fetchone(f'SELECT 1 FROM {table} WHERE auto_key=?', (f'b:{name}',)):
            base, n, nm = name, 1, name
            while db.fetchone(f'SELECT 1 FROM {table} WHERE name=?', (nm,)):
                n += 1
                nm = f'{base}_{n}'
            db.execute(f'INSERT INTO {table}(name,source,fmt,auto_key) VALUES(?,?,?,?)', (nm, source, fmt, f'b:{name}'))


def list_tags(db, mod):
    sync_tags(db, mod)
    return db.fetchall(f"SELECT id,name,source,fmt,empty_text,coalesce(auto_key,''),enabled,note FROM {cfg(mod)['tags']} ORDER BY (auto_key IS NULL), id")


def save_tag(db, mod, tag_id, name, source, fmt, empty_text='', enabled=True, note=''):
    table = cfg(mod)['tags']
    name = norm_tag(name)
    if not re.fullmatch(r'[0-9A-ZА-ЯЁ_]+', name):
        raise ValueError('Имя тега: буквы, цифры и подчёркивание (пробелы заменяются на _)')
    if fmt not in FORMATS:
        raise ValueError('Неизвестный формат')
    if not source:
        raise ValueError('Выберите источник данных')
    if db.fetchone(f'SELECT id FROM {table} WHERE name=? AND id IS NOT ?', (name, tag_id)):
        raise ValueError(f'Тег {{{name}}} уже существует')
    if tag_id:
        db.execute(f'UPDATE {table} SET name=?,source=?,fmt=?,empty_text=?,enabled=?,note=? WHERE id=?', (name, source, fmt, empty_text, int(bool(enabled)), note, tag_id))
        return tag_id
    return db.execute(f'INSERT INTO {table}(name,source,fmt,empty_text,enabled,note) VALUES(?,?,?,?,?,?)', (name, source, fmt, empty_text, int(bool(enabled)), note)).lastrowid


def delete_tag(db, mod, tag_id):
    table = cfg(mod)['tags']
    row = db.fetchone(f'SELECT auto_key FROM {table} WHERE id=?', (tag_id,))
    if not row:
        return 'missing'
    if row[0]:
        db.execute(f'UPDATE {table} SET enabled=0 WHERE id=?', (tag_id,))
        return 'disabled'
    db.execute(f'DELETE FROM {table} WHERE id=?', (tag_id,))
    return 'deleted'


def context(db, mod, cid, act_id=None):
    c = contract(db, mod, cid)
    p = party(db, mod, c)
    net, vat, gross = vat_parts(c.get('amount'), c.get('vat_included'), c.get('vat_rate'))
    pays = payments(db, mod, cid)
    paid = sum(a for _d, a, _n in pays)
    ctx = {
        'contract.number': c.get('contract_number') or '', 'contract.date': c.get('contract_date') or '', 'contract.subject': c.get('subject') or '',
        'contract.object': c.get('object_name') or '', 'contract.object_address': c.get('object_address') or '', 'contract.start': c.get('start_date') or c.get('contract_date') or '',
        'contract.end': c.get('end_date') or '', 'contract.amount': gross, 'contract.amount_net': net, 'contract.vat': vat,
        'contract.vat_rate': f'{float(c.get("vat_rate") or 0):g}' if vat else '', 'contract.vat_note': (f'в том числе НДС {float(c["vat_rate"]):g}% — {money(vat)}' if vat else 'НДС не облагается'),
        'payments.paid': paid, 'payments.rest': gross - paid, 'payments.dates': '\n'.join(gd.date_short(d).rstrip('г.') for d, _a, _n in pays),
        'payments.amounts': '\n'.join(money(a) for _d, a, _n in pays),
        'estimate.materials': c.get('est_materials') or 0, 'estimate.works': c.get('est_works') or 0,
    }
    ctx.update({f'party.{k}': v for k, v in p.items() if k != 'kind'})
    if act_id:
        a = act(db, mod, act_id)
        a_net, a_vat, a_gross = vat_parts(a.get('amount'), c.get('vat_included') if a.get('vat_included') is None else a.get('vat_included'), c.get('vat_rate'))
        ctx.update({'act.number': a.get('act_number') or '', 'act.date': a.get('act_date') or '', 'act.amount': a_gross, 'act.amount_net': a_net, 'act.vat': a_vat,
                    'act.description': a.get('description') or ''})
    else:
        ctx.update({'act.number': '', 'act.date': '', 'act.amount': '', 'act.amount_net': '', 'act.vat': '', 'act.description': ''})
    return ctx


def tag_map(db, mod, cid, act_id=None):
    sync_tags(db, mod)
    ctx = context(db, mod, cid, act_id)
    result = {}
    for _id, name, source, fmt, empty_text, _auto, enabled, _note in db.fetchall(f"SELECT id,name,source,fmt,empty_text,coalesce(auto_key,''),enabled,note FROM {cfg(mod)['tags']}"):
        if not enabled:
            continue
        text = apply_format(ctx.get(source, ''), fmt)
        result[name] = text if text != '' else (empty_text or '')
    return result


# --- ДОКУМЕНТЫ --------------------------------------------------------------------------------

DEFAULT_BODIES = {
    'contract': ['ДОГОВОР № {НОМЕР_ДОГОВОРА}', '{ПРЕДМЕТ}', '{ДАТА_ДОГОВОРА}',
                 'Подрядчик и {КОНТРАГЕНТ}, в лице {ДОЛЖНОСТЬ_РУКОВОДИТЕЛЯ_РОД} {ФИО_РУКОВОДИТЕЛЯ_РОД}, действующего на основании {ОСНОВАНИЕ}, заключили настоящий договор.',
                 'Объект: {ОБЪЕКТ}, {АДРЕС_ОБЪЕКТА}.', 'Сроки: с {НАЧАЛО_РАБОТ} по {ОКОНЧАНИЕ_РАБОТ}.',
                 'Стоимость работ: {СУММА} BYN ({СУММА_ПРОПИСЬЮ}), {НДС_ПОЯСНЕНИЕ}.',
                 'Реквизиты заказчика: {КОНТРАГЕНТ}, УНП {УНП}, юридический адрес: {АДРЕС_ЮРИДИЧЕСКИЙ}, р/с {РАСЧЕТНЫЙ_СЧЕТ} в {БАНК}, БИК {БИК}. Тел.: {ТЕЛЕФОН}, e-mail: {EMAIL}.',
                 'Заказчик: ______________ {ФИО_РУКОВОДИТЕЛЯ_СОКР}'],
    'act': ['АКТ № {НОМЕР_АКТА} выполненных работ', 'к договору № {НОМЕР_ДОГОВОРА} от {ДАТА_ДОГОВОРА}', '{ДАТА_АКТА}',
            'Подрядчик выполнил, а Заказчик {КОНТРАГЕНТ} принял работы: {ОПИСАНИЕ_РАБОТ} на объекте: {ОБЪЕКТ}, {АДРЕС_ОБЪЕКТА}.',
            'Стоимость выполненных работ: {СУММА_АКТА} BYN ({СУММА_АКТА_ПРОПИСЬЮ}), в том числе НДС {СУММА_АКТА_НДС} BYN.',
            'Заказчик: ______________ {ФИО_РУКОВОДИТЕЛЯ_СОКР}'],
    'statement': ['СПРАВКА о стоимости выполненных работ и затрат', 'по акту № {НОМЕР_АКТА} к договору № {НОМЕР_ДОГОВОРА}', '{ДАТА_АКТА}',
                  'Объект: {ОБЪЕКТ}, {АДРЕС_ОБЪЕКТА}. Заказчик: {КОНТРАГЕНТ}.', 'Работы: {ОПИСАНИЕ_РАБОТ}.',
                  'Всего: {СУММА_АКТА} BYN, в том числе без НДС {СУММА_АКТА_БЕЗ_НДС} BYN, НДС {СУММА_АКТА_НДС} BYN.', 'Заказчик: ______________ {ФИО_РУКОВОДИТЕЛЯ_СОКР}'],
}


def templates_dir(mod):
    return gd.TEMPLATES_DIR / cfg(mod)['section'].lower()


def template_setting(mod, kind):
    return f'{mod}_tpl_{kind}'


def create_default_template(mod, kind):
    folder = templates_dir(mod)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f'Шаблон_{DOC_KINDS[kind][2]}.docx'
    if path.exists():
        return path
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = docx.Document()
    lines = DEFAULT_BODIES[kind]
    doc.add_heading(lines[0], 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in lines[1:]:
        doc.add_paragraph(line)
    doc.save(path)
    return path


def template_path(db, mod, kind):
    path = db.get_setting(template_setting(mod, kind), '')
    return path if path and os.path.exists(path) else str(create_default_template(mod, kind))


def _scope(kind):
    return DOC_KINDS[kind][1]


def data_hash(db, mod, cid, act_id=None):
    return hashlib.sha256(json.dumps(tag_map(db, mod, cid, act_id), ensure_ascii=False, sort_keys=True).encode('utf-8')).hexdigest()


def doc_row(db, mod, kind, ref_id):
    return db.fetchone(f"SELECT file_path,data_hash,linked FROM {cfg(mod)['docs']} WHERE scope=? AND ref_id=? AND kind=?", (_scope(kind), ref_id, kind))


def _ids(db, mod, kind, ref_id):
    """(id договора, id акта) для вида документа."""
    if _scope(kind) == 'contract':
        return ref_id, None
    a = act(db, mod, ref_id)
    return a['contract_id'], ref_id


def doc_state(db, mod, kind, ref_id):
    row = doc_row(db, mod, kind, ref_id)
    if not row:
        return 'none'
    if not os.path.exists(row[0]):
        return 'missing'
    if row[2]:
        return 'linked'
    cid, aid = _ids(db, mod, kind, ref_id)
    return 'fresh' if row[1] == data_hash(db, mod, cid, aid) else 'stale'


def doc_path(db, mod, kind, ref_id):
    row = doc_row(db, mod, kind, ref_id)
    return row[0] if row else ''


def generate(db, mod, kind, ref_id):
    if kind not in DOC_KINDS:
        raise ValueError('Неизвестный документ')
    cid, aid = _ids(db, mod, kind, ref_id)
    folder = contract_folder(db, mod, cid)
    if not folder:
        raise ValueError('Сначала привяжите или создайте папку договора кнопкой «Папка договора»')
    tags = tag_map(db, mod, cid, aid)
    title, scope, prefix = DOC_KINDS[kind]
    number = tags.get('НОМЕР_ДОГОВОРА') or str(cid)
    suffix = f'_{tags.get("НОМЕР_АКТА") or aid}' if aid else ''
    out = os.path.join(folder, gd.safe_name(f'{prefix}{suffix}_{number.replace("/", "-")}').replace(' ', '_') + '.docx')
    gd.render_docx(template_path(db, mod, kind), out, tags, normalize=norm_tag)
    db.execute(f"INSERT OR REPLACE INTO {cfg(mod)['docs']}(scope,ref_id,kind,file_path,data_hash,generated_at,linked) VALUES(?,?,?,?,?,?,0)",
               (scope, ref_id, kind, out, data_hash(db, mod, cid, aid), datetime.now().isoformat(timespec='seconds')))
    return out


def link_doc(db, mod, kind, ref_id, path):
    if not os.path.isfile(path):
        raise ValueError('Файл не найден: ' + path)
    db.execute(f"INSERT OR REPLACE INTO {cfg(mod)['docs']}(scope,ref_id,kind,file_path,data_hash,generated_at,linked) VALUES(?,?,?,?,?,?,1)",
               (_scope(kind), ref_id, kind, path, '', datetime.now().isoformat(timespec='seconds')))


def remove_doc(db, mod, kind, ref_id, delete_file=False):
    row = doc_row(db, mod, kind, ref_id)
    if row and delete_file and not row[2] and os.path.isfile(row[0]):
        os.remove(row[0])
    db.execute(f"DELETE FROM {cfg(mod)['docs']} WHERE scope=? AND ref_id=? AND kind=?", (_scope(kind), ref_id, kind))


# --- ЗАПИСЬ ДОГОВОРА И АКТОВ ------------------------------------------------------------------

def save_contract(db, mod, values, cid=None):
    """Сохраняет договор. Номер, не заданный вручную, назначается автоматически. Возвращает id."""
    c = cfg(mod)
    fields = ['contract_number', 'contract_date', 'object_name', 'object_address', 'subject', 'start_date', 'end_date', 'amount', 'vat_included', 'vat_rate', 'signed', 'status', 'note']
    fields += ['direction', 'client_id', 'source_type', 'source_id'] if mod == 'le' else ['party_type', 'person_id', 'legal_id', 'estimate_id']
    data = {k: values.get(k) for k in fields if k in values}
    if mod == 'le' and not data.get('client_id'):
        raise ValueError('Выберите юрлицо')
    if mod == 'smr':
        kind = data.get('party_type') or 'person'
        if kind == 'legal' and not data.get('legal_id'):
            raise ValueError('Выберите юрлицо')
        if kind == 'person' and not data.get('person_id'):
            raise ValueError('Выберите клиента (физическое лицо)')
    if float(data.get('vat_rate') or 0) < 0 or float(data.get('amount') or 0) < 0:
        raise ValueError('Сумма и ставка НДС не могут быть отрицательными')
    with db.transaction():
        if cid:
            db.execute(f"UPDATE {c['contracts']} SET " + ','.join(f'{k}=?' for k in data) + ' WHERE id=?', (*data.values(), cid))
        else:
            number = str(data.get('contract_number') or '').strip()
            seq = year = None
            if not number or number.upper().startswith('XX'):
                seq, year, number = next_number(db, mod, data.get('contract_date'))
            else:
                m = re.match(rf"^(\d+)-{c['code']}/(\d{{2}})$", number)
                if m:
                    seq, year = int(m.group(1)), int(m.group(2))
            data.update(contract_number=number, seq_num=seq, year_num=year)
            cid = db.execute(f"INSERT INTO {c['contracts']}(" + ','.join(data) + ') VALUES(' + ','.join('?' for _ in data) + ')', tuple(data.values())).lastrowid
    return cid


def save_act(db, mod, values, aid=None):
    c = cfg(mod)
    cid = values.get('contract_id')
    if not cid:
        raise ValueError('Выберите договор')
    day = str(values.get('act_date') or '')
    if day:
        date.fromisoformat(day)
    data = dict(contract_id=cid, act_number=str(values.get('act_number') or '').strip(), act_date=day, amount=float(values.get('amount') or 0),
                description=str(values.get('description') or '').strip(), note=str(values.get('note') or '').strip(), signed=int(bool(values.get('signed'))))
    if data['amount'] < 0:
        raise ValueError('Сумма акта не может быть отрицательной')
    if not data['act_number']:
        n = db.fetchone(f"SELECT count(*) FROM {c['acts']} WHERE contract_id=?", (cid,))[0] + 1
        data['act_number'] = str(n)
    if data['signed'] and not day:
        raise ValueError('Для подписанного акта укажите дату')
    if aid:
        db.execute(f"UPDATE {c['acts']} SET " + ','.join(f'{k}=?' for k in data) + ' WHERE id=?', (*data.values(), aid))
        return aid
    return db.execute(f"INSERT INTO {c['acts']}(" + ','.join(data) + ') VALUES(' + ','.join('?' for _ in data) + ')', tuple(data.values())).lastrowid


def contract_from_source(db, source_type, source_id):
    """Договор юрлица на основе карточки проекта / монтажа ГСВ (заказчик — юрлицо). Повторный вызов возвращает уже созданный договор."""
    found = db.fetchone("SELECT id FROM le_contracts WHERE source_type=? AND source_id=?", (source_type, source_id))
    if found:
        return found[0], False
    if source_type == 'gsv_projects':
        row = db.fetchone('SELECT le_client_id,contract_number,contract_date,object_name,address,cost,due_date,act_date,contract_signed FROM gsv_projects WHERE id=?', (source_id,))
        direction = 'Проектирование ГСВ'
        subject = 'Разработка проектной документации газоснабжения'
        values = dict(client_id=row[0], contract_number=row[1], contract_date=row[2], object_name=row[3], object_address=row[4], amount=row[5], end_date=row[6], signed=row[8])
    elif source_type == 'contracts':
        row = db.fetchone('SELECT le_client_id,contract_number,contract_date,object_name,object_address,contract_amount,work_end_date,acceptance_act_date,contract_signed FROM contracts WHERE id=?', (source_id,))
        direction = 'Монтажные работы'
        subject = 'Монтаж внутреннего газопровода и газоиспользующего оборудования'
        values = dict(client_id=row[0], contract_number=row[1], contract_date=row[2], object_name=row[3], object_address=row[4], amount=row[5], end_date=row[6], signed=row[8])
    else:
        raise ValueError('Неизвестный источник')
    if not row or not row[0]:
        raise ValueError('В карточке не выбрано юрлицо-заказчик')
    values.update(direction=direction, subject=subject, start_date=row[2] or '', status='Действует', vat_included=0, vat_rate=0, source_type=source_type, source_id=source_id)
    cid = save_contract(db, 'le', values)
    if row[7]:
        save_act(db, 'le', dict(contract_id=cid, act_date=row[7], amount=row[5], description=subject, signed=1))
    return cid, True


# --- СВЕРКА ПЕРЕД ФОРМИРОВАНИЕМ ----------------------------------------------------------------

def check(db, mod, kind, ref_id):
    """Замечания перед формированием документа: [{'level': 'error'|'warn', 'text': …}]."""
    cid, aid = _ids(db, mod, kind, ref_id)
    c = contract(db, mod, cid)
    p = party(db, mod, c)
    out = []

    def add(level, text):
        out.append(dict(level=level, text=text))
    if not (c.get('contract_number') or '').strip():
        add('error', 'Не указан номер договора')
    if not c.get('contract_date'):
        add('error', 'Не указана дата договора')
    if not p['name']:
        add('error', 'Не выбран контрагент')
    if p['kind'] == 'legal':
        for key, label in (('unp', 'УНП'), ('legal_address', 'юридический адрес'), ('bank', 'банк'), ('account', 'расчётный счёт'), ('head_name', 'ФИО руководителя'),
                           ('head_name_gen', 'ФИО руководителя в родительном падеже'), ('basis', 'основание (устав/доверенность)')):
            if not p[key]:
                add('warn', f'У юрлица не заполнено: {label}')
    else:
        if not p['passport']:
            add('warn', 'У клиента не заполнены паспортные данные')
        if not p['address']:
            add('warn', 'У клиента не указан адрес')
    if not float(c.get('amount') or 0):
        add('warn', 'Сумма договора равна нулю')
    if not (c.get('subject') or '').strip() and kind == 'contract':
        add('warn', 'Не указан предмет договора')
    if aid:
        a = act(db, mod, aid)
        if not a.get('act_date'):
            add('error', 'У акта не указана дата')
        if not float(a.get('amount') or 0):
            add('error', 'У акта нулевая сумма')
        if not (a.get('description') or '').strip():
            add('warn', 'Не заполнено описание работ в акте')
    if mod == 'smr' and estimate_state(db, cid) == 'stale':
        add('error', 'Смета изменена после получения данных — обновите стоимость в карточке договора')
    return out
