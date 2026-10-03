"""
Отдельный от обычных бэкапов снимок реестров договоров (ГСВ и ГСН) в один
Excel-файл. Файл не копится — при каждом обновлении он перезаписывается,
так что по этому пути всегда лежит последняя версия таблицы.
"""
import os
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

EXPORT_FILENAME = "Реестры_договоров.xlsx"


def _write_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)


def export_registries(db, destination):
    """Overwrite `destination` with a fresh two-sheet snapshot (ГСВ, ГСН)."""
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    gsv_rows = db.fetchall(
        """SELECT c.contract_number, c.contract_date, c.object_name, c.client_name, c.client_phone,
                  c.contract_amount, e.title
           FROM contracts c LEFT JOIN estimates e ON c.estimate_id = e.id ORDER BY c.id DESC"""
    )
    _write_sheet(
        wb, "ГСВ",
        ["№ Договора", "Дата", "Объект / Адрес", "Клиент", "Телефон", "Сумма (руб)", "Привязанная смета"],
        [[num or "Б/Н", date or "", obj or "", client or "", phone or "", float(amount or 0), title or "— нет сметы —"]
         for num, date, obj, client, phone, amount, title in gsv_rows],
    )

    gsn_rows = db.fetchall(
        "SELECT contract_number, contract_date, title, client_name, phone, address FROM gsn_projects ORDER BY id DESC"
    )
    _write_sheet(
        wb, "ГСН",
        ["Договор", "Дата", "Объект строительства", "Клиент", "Телефон", "Адрес объекта"],
        [[num or "", date or "", title or "", client or "", phone or "", address or ""]
         for num, date, title, client, phone, address in gsn_rows],
    )

    # Never let a plain string be reinterpreted as a formula by Excel (same guard as elsewhere in the app).
    for sheet in wb.worksheets:
        for row_cells in sheet:
            for cell in row_cells:
                if isinstance(cell.value, str): cell.data_type = "s"

    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=destination.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, destination)
    except Exception:
        Path(tmp).unlink(missing_ok=True)
        raise
    return len(gsv_rows), len(gsn_rows)
