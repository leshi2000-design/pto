"""Раздел «Монтаж ГСВ»: теги, редактор тегов и формирование документов Word / Excel.

Теги этого раздела хранятся в таблице `gsvm_tags` и не пересекаются с тегами «Проектов ГСВ» и других модулей.
Каждый тег — это имя, источник данных (поле карточки, оборудование, трубопровод), формат и текст «если пусто».
"""
import hashlib
import json
import os
import re
from datetime import datetime

from . import gsv_project_domain as gd
from . import gsvm_domain as md
from .gsv_project_domain import MONTHS_GENITIVE, get_initials_first, num_to_words_byn

TEMPLATES_DIR = gd.TEMPLATES_DIR / 'gsvm'
MONTHS_NOMINATIVE = ['', 'январь', 'февраль', 'март', 'апрель', 'май', 'июнь', 'июль', 'август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь']

# вид документа -> (название, расширение, группа: 'main' — вкладка «Договор и клиент», 'id' — исполнительная документация, префикс файла)
DOC_KINDS = {
    'contract': ('Договор на монтаж', 'docx', 'main', 'Договор'),
    'act': ('Акт выполненных работ и справка о стоимости', 'docx', 'main', 'Акт_и_справка'),
    'card': ('Карточка клиента с оплатами', 'xlsx', 'main', 'Карточка_клиента'),
    'registry': ('Реестр ГСВ', 'xlsx', 'id', 'Реестр_ГСВ'),
    'warranty': ('Гарантийный паспорт', 'docx', 'id', 'Гарантийный_паспорт'),
    'accept': ('Акт приемки объекта в эксплуатацию', 'docx', 'id', 'Акт_приемки'),
    'building_passport': ('Строительный паспорт объекта', 'docx', 'id', 'Строительный_паспорт'),
    'hidden_works': ('Акты скрытых работ', 'docx', 'id', 'Акты_скрытых_работ'),
    'device_revision': ('Акт ревизии тех.устройств', 'docx', 'id', 'Акт_ревизии_тех_устройств'),
    'equipment_revision': ('Акт о проведении ревизии оборудования', 'docx', 'id', 'Акт_ревизии_оборудования'),
}
ID_KINDS = [k for k, v in DOC_KINDS.items() if v[2] == 'id']
MAIN_KINDS = [k for k, v in DOC_KINDS.items() if v[2] == 'main']
FORMATS = {
    'text': 'Как есть', 'upper': 'ЗАГЛАВНЫМИ', 'date_long': 'Дата: 12 ноября 2026', 'date_quoted': 'Дата: «12» ноября 2026',
    'date_short': 'Дата: 12.11.2026', 'month_name': 'Месяц: ноябрь', 'month_year': 'Месяц и год: ноябрь 2026',
    'money': 'Деньги: 1 250,00', 'money_words': 'Сумма прописью', 'number2': 'Число с двумя знаками: 12,50',
}
PAYMENT_TAGS = gd.PAYMENT_TAGS


def norm_tag(raw):
    return re.sub(r'\s+', '_', str(raw).strip().upper())


def tag_text(text):
    return re.sub(r'[^0-9A-ZА-ЯЁ]+', '_', str(text).upper()).strip('_')


# --- ИСТОЧНИКИ И ТЕГИ ------------------------------------------------------------------------

STATIC_SOURCES = [
    ('contract.number', 'Номер договора'), ('contract.date', 'Дата договора'), ('client.name', 'Клиент: ФИО полностью'),
    ('client.short', 'Клиент: инициалы и фамилия'), ('client.passport', 'Клиент: паспорт (серия, номер)'),
    ('client.passport_date', 'Клиент: дата выдачи паспорта'), ('client.passport_issuer', 'Клиент: кем выдан паспорт'),
    ('client.phone', 'Клиент: телефоны'), ('client.address', 'Клиент: адрес'), ('object.name', 'Наименование объекта газоснабжения'),
    ('object.address', 'Адрес объекта'), ('notes', 'Примечания по объекту'), ('work.start', 'Начало работ'), ('work.end', 'Окончание работ'),
    ('work.cost', 'Стоимость работ'), ('work.act_date', 'Дата акта'), ('project.code', 'Шифр проекта (номер проектной документации)'),
    ('project.designer', 'Проектировщик'), ('project.date', 'Дата проекта (месяц и год)'), ('estimate.materials', 'Смета: стоимость материалов'),
    ('estimate.works', 'Смета: стоимость работ'), ('payments.dates', 'Оплаты: даты (по строкам)'), ('payments.amounts', 'Оплаты: суммы (по строкам)'),
    ('payments.notes', 'Оплаты: примечания (по строкам)'), ('payments.paid', 'Оплачено всего'), ('payments.rest', 'Остаток к оплате'),
    ('joints.total', 'Количество стыков всего'),
]
BUILTIN_TAGS = [
    ('НОМЕР_ДОГОВОРА', 'contract.number', 'text'), ('ДАТА_ДОГОВОРА', 'contract.date', 'date_long'), ('ДАТА_ДОГОВОРА_С', 'contract.date', 'date_short'),
    ('КЛИЕНТ', 'client.name', 'text'), ('КЛИЕНТ_СОКР', 'client.short', 'text'), ('ПАСПОРТ', 'client.passport', 'text'),
    ('ДАТА_ВЫДАЧИ', 'client.passport_date', 'date_short'), ('КЕМ_ВЫДАН', 'client.passport_issuer', 'text'), ('ТЕЛЕФОН', 'client.phone', 'text'),
    ('АДРЕС_КЛИЕНТА', 'client.address', 'text'), ('ОБЪЕКТ', 'object.name', 'text'), ('АДРЕС', 'object.address', 'text'),
    ('ПРИМЕЧАНИЕ', 'notes', 'text'), ('НАЧАЛО_РАБОТ', 'work.start', 'date_long'), ('ОКОНЧАНИЕ_РАБОТ', 'work.end', 'date_long'),
    ('СТОИМОСТЬ', 'work.cost', 'money'), ('СТОИМОСТЬ_ПРОПИСЬЮ', 'work.cost', 'money_words'), ('МЕСЯЦ_РАБОТ', 'work.act_date', 'month_name'),
    ('МЕСЯЦ_РАБОТ_ГОД', 'work.act_date', 'month_year'), ('ДАТА_АКТА', 'work.act_date', 'date_quoted'), ('ДАТА_АКТА_С', 'work.act_date', 'date_short'),
    ('НОМЕР_ПРОЕКТА', 'project.code', 'text'), ('ПРОЕКТИРОВЩИК', 'project.designer', 'text'), ('ДАТА_ПРОЕКТА', 'project.date', 'month_year'),
    ('СТОИМОСТЬ_МАТЕРИАЛОВ', 'estimate.materials', 'money'), ('СТОИМОСТЬ_РАБОТ_СМЕТА', 'estimate.works', 'money'),
    ('ОПЛАЧЕНО', 'payments.paid', 'money'), ('ОСТАТОК', 'payments.rest', 'money'),
    ('ДАТА_ОПЛАТЫ', 'payments.dates', 'text'), ('СУММА_ОПЛАТЫ', 'payments.amounts', 'text'), ('ПРИМЕЧАНИЕ_ОПЛАТЫ', 'payments.notes', 'text'),
    ('СТЫКОВ_ВСЕГО', 'joints.total', 'text'),
]


def available_sources(db):
    """[(ключ источника, подпись)] — поля карточки, оборудование и трубопроводы справочника ГСВ."""
    result = list(STATIC_SOURCES)
    for kind in equipment_kinds(db):
        label = md.EQUIPMENT_LABELS.get(kind, kind)
        result += [(f'equipment:{kind}:model', f'Оборудование «{label}»: тип / модель'), (f'equipment:{kind}:serial', f'Оборудование «{label}»: серийный номер')]
    for pid, name in db.fetchall('SELECT id,name FROM gsv_pipelines ORDER BY name'):
        result += [(f'pipeline:{pid}:name', f'Трубопровод «{name}»: наименование'), (f'pipeline:{pid}:name_cert', f'Трубопровод «{name}»: наименование и сертификат'),
                   (f'pipeline:{pid}:cert', f'Трубопровод «{name}»: номер сертификата'), (f'pipeline:{pid}:diameter', f'Трубопровод «{name}»: диаметр'),
                   (f'pipeline:{pid}:qty', f'Трубопровод «{name}»: количество, м'), (f'pipeline:{pid}:joints', f'Трубопровод «{name}»: количество стыков')]
    return result


def equipment_kinds(db):
    kinds = list(md.EQUIPMENT_KINDS)
    for (k,) in db.fetchall("SELECT DISTINCT kind FROM gsvm_equipment WHERE kind<>''"):
        if k not in kinds:
            kinds.append(k)
    return kinds


def sync_tags(db):
    """Создаёт недостающие стандартные теги, теги оборудования и трубопроводов. Переименованные пользователем теги не пересоздаются."""
    def add(auto_key, name, source, fmt, note=''):
        if db.fetchone('SELECT 1 FROM gsvm_tags WHERE auto_key=?', (auto_key,)):
            return
        base, n = name, 1
        while db.fetchone('SELECT 1 FROM gsvm_tags WHERE name=?', (name,)):
            n += 1
            name = f'{base}_{n}'
        db.execute('INSERT INTO gsvm_tags(name,source,fmt,auto_key,note) VALUES(?,?,?,?,?)', (name, source, fmt, auto_key, note))

    for name, source, fmt in BUILTIN_TAGS:
        add(f'b:{name}', name, source, fmt)
    for kind in equipment_kinds(db):
        base = tag_text(kind)
        add(f'e:{kind}:model', f'{base}_МОДЕЛЬ', f'equipment:{kind}:model', 'text', 'Пусто, если такого оборудования в объекте нет')
        add(f'e:{kind}:serial', f'{base}_СЕРИЙНЫЙ_НОМЕР', f'equipment:{kind}:serial', 'text', 'Пусто, если такого оборудования в объекте нет')
    for pid, pname in db.fetchall('SELECT id,name FROM gsv_pipelines ORDER BY id'):
        d = md.diameter_of(pname)
        short = tag_text(d) if d else str(pid)
        add(f'p:{pid}:full', f'ТРУБОПРОВОД_{tag_text(pname)}', f'pipeline:{pid}:name_cert', 'text', 'Наименование и номер сертификата; пусто, если трубопровод не используется')
        add(f'p:{pid}:short', f'ТРУБА_{short}', f'pipeline:{pid}:name_cert', 'text', 'Сокращённый тег: только диаметр')
        add(f'p:{pid}:qty', f'ДЛИНА_{short}', f'pipeline:{pid}:qty', 'number2', 'Количество, м')
        add(f'p:{pid}:joints', f'СТЫКИ_{short}', f'pipeline:{pid}:joints', 'text', 'Количество стыков, шт.')


def list_tags(db):
    sync_tags(db)
    return db.fetchall('SELECT id,name,source,fmt,empty_text,coalesce(auto_key,\'\'),enabled,note FROM gsvm_tags ORDER BY (auto_key IS NULL), id')


def save_tag(db, tag_id, name, source, fmt, empty_text='', enabled=True, note=''):
    name = norm_tag(name)
    if not re.fullmatch(r'[0-9A-ZА-ЯЁ_]+', name):
        raise ValueError('Имя тега: буквы, цифры и подчёркивание (пробелы заменяются на _)')
    if fmt not in FORMATS:
        raise ValueError('Неизвестный формат')
    if not source:
        raise ValueError('Выберите источник данных')
    clash = db.fetchone('SELECT id FROM gsvm_tags WHERE name=? AND id IS NOT ?', (name, tag_id))
    if clash:
        raise ValueError(f'Тег {{{name}}} уже существует')
    if tag_id:
        db.execute('UPDATE gsvm_tags SET name=?,source=?,fmt=?,empty_text=?,enabled=?,note=? WHERE id=?', (name, source, fmt, empty_text, int(bool(enabled)), note, tag_id))
        return tag_id
    return db.execute('INSERT INTO gsvm_tags(name,source,fmt,empty_text,enabled,note) VALUES(?,?,?,?,?,?)', (name, source, fmt, empty_text, int(bool(enabled)), note)).lastrowid


def delete_tag(db, tag_id):
    """Собственные теги удаляются; стандартные и автоматические отключаются (чтобы не появиться заново)."""
    row = db.fetchone('SELECT auto_key FROM gsvm_tags WHERE id=?', (tag_id,))
    if not row:
        return 'missing'
    if row[0]:
        db.execute('UPDATE gsvm_tags SET enabled=0 WHERE id=?', (tag_id,))
        return 'disabled'
    db.execute('DELETE FROM gsvm_tags WHERE id=?', (tag_id,))
    return 'deleted'


# --- ЗНАЧЕНИЯ --------------------------------------------------------------------------------

def money(value):
    try:
        return f'{float(value):,.2f}'.replace(',', ' ').replace('.', ',')
    except (TypeError, ValueError):
        return ''


def apply_format(value, fmt):
    value = '' if value is None else str(value)
    if not value:
        return ''
    if fmt == 'upper':
        return value.upper()
    if fmt in ('date_long', 'date_quoted', 'date_short', 'month_name', 'month_year'):
        m = re.match(r'^(\d{4})-(\d{2})(?:-(\d{2}))?', value)
        if not m:
            return value
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3) or 1)
        if fmt == 'month_name':
            return MONTHS_NOMINATIVE[mo]
        if fmt == 'month_year':
            return f'{MONTHS_NOMINATIVE[mo]} {y}'
        if not m.group(3):
            return value
        return {'date_long': f'{d} {MONTHS_GENITIVE[mo]} {y}', 'date_quoted': f'«{d:02}» {MONTHS_GENITIVE[mo]} {y}', 'date_short': f'{d:02}.{mo:02}.{y}'}[fmt]
    if fmt == 'money':
        return money(value)
    if fmt == 'money_words':
        try:
            return num_to_words_byn(float(value))
        except ValueError:
            return value
    if fmt == 'number2':
        try:
            return f'{float(value):.2f}'.replace('.', ',')
        except ValueError:
            return value
    return value


def context(db, cid):
    """Исходные значения источников по договору (даты — ГГГГ-ММ-ДД, суммы — числа)."""
    cur = db.execute('SELECT * FROM contracts WHERE id=?', (cid,))
    row = cur.fetchone()
    if not row:
        raise ValueError('Сначала сохраните договор')
    c = dict(zip([x[0] for x in cur.description], row))
    client = {}
    if c.get('client_id'):
        from .gsv_domain import get_client
        client = get_client(db, c['client_id']) or {}
    name = client.get('name') or c.get('client_name') or c.get('party_name') or ''
    is_org = bool(c.get('le_client_id'))
    from . import payments_domain
    pays = sorted(((d, float(a), n or '') for _i, d, a, n in payments_domain.history(db, 'contracts', cid)), key=lambda p: p[0])
    paid = sum(a for _d, a, _n in pays)
    cost = float(c.get('contract_amount') or 0)
    ctx = {
        'contract.number': c.get('contract_number') or '', 'contract.date': c.get('contract_date') or '', 'client.name': name,
        'client.short': (name if is_org else get_initials_first(name)) if name else '', 'client.passport': client.get('passport') or c.get('passport_series_number') or '',
        'client.passport_date': client.get('passport_date') or c.get('passport_issue_date') or '',
        'client.passport_issuer': client.get('passport_issuer') or c.get('passport_issued_by') or '',
        'client.phone': client.get('phone') or c.get('client_phone') or '', 'client.address': client.get('address') or c.get('client_address') or '',
        'object.name': c.get('object_name') or '', 'object.address': c.get('object_address') or '', 'notes': c.get('notes') or '',
        'work.start': c.get('work_start_date') or c.get('contract_date') or '', 'work.end': c.get('work_end_date') or '', 'work.cost': cost,
        'work.act_date': c.get('acceptance_act_date') or '', 'project.code': c.get('designer_code') or '', 'project.designer': c.get('designer') or '',
        'project.date': c.get('project_month') or '', 'estimate.materials': c.get('est_materials') or 0, 'estimate.works': c.get('est_works') or 0,
        'payments.dates': '\n'.join(gd.date_short(d).rstrip('г.') for d, _a, _n in pays), 'payments.amounts': '\n'.join(money(a) for _d, a, _n in pays),
        'payments.notes': '\n'.join(n for _d, _a, n in pays), 'payments.paid': paid, 'payments.rest': cost - paid, 'joints.total': md.joints_total(db, cid),
    }
    eq = md.equipment(db, cid)
    for kind in equipment_kinds(db):
        mine = [e for e in eq if md._norm(e['kind']) == md._norm(kind)]
        ctx[f'equipment:{kind}:model'] = '; '.join(e['model'] for e in mine if e['model'])
        ctx[f'equipment:{kind}:serial'] = '; '.join(e['serial'] for e in mine if e['serial'])
    joint_by_pipeline = {}
    for j in md.joints(db, cid):
        if j['pipeline_id']:
            joint_by_pipeline[j['pipeline_id']] = joint_by_pipeline.get(j['pipeline_id'], 0) + j['count']
    used = {}
    for p in md.pipelines(db, cid):
        used.setdefault(p['pipeline_id'], dict(p, quantity=0.0))['quantity'] += p['quantity']
    for pid, p in used.items():
        cert = p['cert_number'] or ''
        ctx[f'pipeline:{pid}:name'] = p['name']
        ctx[f'pipeline:{pid}:cert'] = cert
        ctx[f'pipeline:{pid}:name_cert'] = p['name'] + (f', сертификат № {cert}' if cert else '')
        ctx[f'pipeline:{pid}:diameter'] = md.diameter_of(p['name'])
        ctx[f'pipeline:{pid}:qty'] = f'{p["quantity"]:.2f}'
        ctx[f'pipeline:{pid}:joints'] = str(joint_by_pipeline.get(pid, 0))
    return ctx


def tag_map(db, cid):
    """{ИМЯ_ТЕГА: значение} — только включённые теги; неиспользуемое оборудование и трубопроводы дают пустую строку."""
    sync_tags(db)
    ctx = context(db, cid)
    result = {}
    for _id, name, source, fmt, empty_text, _auto, enabled, _note in db.fetchall('SELECT id,name,source,fmt,empty_text,coalesce(auto_key,\'\'),enabled,note FROM gsvm_tags'):
        if not enabled:
            continue
        raw = ctx.get(source, '')
        text = apply_format(raw, fmt)
        result[name] = text if text != '' else (empty_text or '')
    return result


def payments(db, cid):
    from . import payments_domain
    return sorted(((d, float(a), n or '') for _i, d, a, n in payments_domain.history(db, 'contracts', cid)), key=lambda p: p[0])


def data_hash(db, cid):
    return hashlib.sha256(json.dumps(tag_map(db, cid), ensure_ascii=False, sort_keys=True).encode('utf-8')).hexdigest()


# --- ШАБЛОНЫ И ДОКУМЕНТЫ ---------------------------------------------------------------------

def template_setting(kind):
    return f'gsvm_tpl_{kind}'


def template_path(db, kind):
    path = db.get_setting(template_setting(kind), '')
    return path if path and os.path.exists(path) else str(create_default_template(kind))


DEFAULT_BODIES = {
    'contract': ['ДОГОВОР № {НОМЕР_ДОГОВОРА}', 'на монтаж внутреннего газопровода и газоиспользующего оборудования', '{ДАТА_ДОГОВОРА}',
                 'Заказчик: {КЛИЕНТ}, паспорт {ПАСПОРТ}, выдан {КЕМ_ВЫДАН} {ДАТА_ВЫДАЧИ}, проживающий: {АДРЕС_КЛИЕНТА}, тел.: {ТЕЛЕФОН}.',
                 'Объект газоснабжения: {ОБЪЕКТ}, адрес: {АДРЕС}.', 'Проектная документация: {НОМЕР_ПРОЕКТА}, проектировщик {ПРОЕКТИРОВЩИК}, {ДАТА_ПРОЕКТА}.',
                 'Начало работ: {НАЧАЛО_РАБОТ}. Окончание работ: {ОКОНЧАНИЕ_РАБОТ}.', 'Стоимость работ: {СТОИМОСТЬ} BYN ({СТОИМОСТЬ_ПРОПИСЬЮ}).',
                 'Примечание: {ПРИМЕЧАНИЕ}', 'Заказчик: ______________ {КЛИЕНТ_СОКР}'],
    'act': ['АКТ выполненных работ', 'к договору № {НОМЕР_ДОГОВОРА} от {ДАТА_ДОГОВОРА}', '{ДАТА_АКТА}',
            'Подрядчик выполнил, а Заказчик {КЛИЕНТ} принял работы по монтажу газопровода на объекте: {ОБЪЕКТ}, {АДРЕС}.',
            'Справка о стоимости выполненных работ: {СТОИМОСТЬ} BYN ({СТОИМОСТЬ_ПРОПИСЬЮ}). Оплачено: {ОПЛАЧЕНО} BYN, остаток: {ОСТАТОК} BYN.',
            'Работы выполнены в {МЕСЯЦ_РАБОТ_ГОД}.', 'Заказчик: ______________ {КЛИЕНТ_СОКР}'],
    'warranty': ['ГАРАНТИЙНЫЙ ПАСПОРТ', 'Объект: {ОБЪЕКТ}, {АДРЕС}', 'Заказчик: {КЛИЕНТ}', 'Договор № {НОМЕР_ДОГОВОРА} от {ДАТА_ДОГОВОРА}',
                 'Работы выполнены: {НАЧАЛО_РАБОТ} — {ОКОНЧАНИЕ_РАБОТ}', 'Дата акта: {ДАТА_АКТА}'],
    'accept': ['АКТ приемки объекта в эксплуатацию', 'Объект: {ОБЪЕКТ}, {АДРЕС}', 'Заказчик: {КЛИЕНТ}', 'Договор № {НОМЕР_ДОГОВОРА}', '{ДАТА_АКТА}'],
    'building_passport': ['СТРОИТЕЛЬНЫЙ ПАСПОРТ ОБЪЕКТА', 'Объект: {ОБЪЕКТ}, {АДРЕС}', 'Проект: {НОМЕР_ПРОЕКТА}, {ПРОЕКТИРОВЩИК}, {ДАТА_ПРОЕКТА}',
                          'Заказчик: {КЛИЕНТ}', 'Стыков всего: {СТЫКОВ_ВСЕГО}'],
    'hidden_works': ['АКТ освидетельствования скрытых работ', 'Объект: {ОБЪЕКТ}, {АДРЕС}', 'Договор № {НОМЕР_ДОГОВОРА}', '{ДАТА_АКТА}', 'Стыков всего: {СТЫКОВ_ВСЕГО}'],
    'device_revision': ['АКТ ревизии технических устройств', 'Объект: {ОБЪЕКТ}, {АДРЕС}', 'Заказчик: {КЛИЕНТ}', '{ДАТА_АКТА}'],
    'equipment_revision': ['АКТ о проведении ревизии оборудования', 'Объект: {ОБЪЕКТ}, {АДРЕС}', 'Заказчик: {КЛИЕНТ}', '{ДАТА_АКТА}'],
}


def create_default_template(kind):
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    ext = DOC_KINDS[kind][1]
    path = TEMPLATES_DIR / f'Шаблон_{DOC_KINDS[kind][3]}.{ext}'
    if path.exists():
        return path
    if ext == 'xlsx':
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = DOC_KINDS[kind][0][:30]
        ws['A1'] = DOC_KINDS[kind][0] + ' — {НОМЕР_ДОГОВОРА}'
        ws['A1'].font = Font(bold=True, size=14)
        rows = [('Клиент', '{КЛИЕНТ}'), ('Телефон', '{ТЕЛЕФОН}'), ('Паспорт', '{ПАСПОРТ}'), ('Кем выдан', '{КЕМ_ВЫДАН}'), ('Дата выдачи', '{ДАТА_ВЫДАЧИ}'),
                ('Адрес клиента', '{АДРЕС_КЛИЕНТА}'), ('Объект', '{ОБЪЕКТ}'), ('Адрес объекта', '{АДРЕС}'), ('Дата договора', '{ДАТА_ДОГОВОРА_С}'),
                ('Начало работ', '{НАЧАЛО_РАБОТ}'), ('Окончание работ', '{ОКОНЧАНИЕ_РАБОТ}'), ('Дата акта', '{ДАТА_АКТА_С}'), ('Стоимость, BYN', '{СТОИМОСТЬ}'),
                ('Оплачено, BYN', '{ОПЛАЧЕНО}'), ('Остаток, BYN', '{ОСТАТОК}'), ('Стыков всего', '{СТЫКОВ_ВСЕГО}')]
        for r, (a, b) in enumerate(rows, 3):
            ws.cell(r, 1, a).font = Font(bold=True)
            ws.cell(r, 2, b)
        start = len(rows) + 5
        if kind == 'card':
            ws.cell(start - 1, 1, 'Оплаты (на каждую оплату добавляются свои столбцы)').font = Font(bold=True)
            for c, (a, b) in enumerate((('Дата оплаты', '{ДАТА_ОПЛАТЫ}'), ('Сумма оплаты', '{СУММА_ОПЛАТЫ}'), ('Примечание', '{ПРИМЕЧАНИЕ_ОПЛАТЫ}')), 1):
                ws.cell(start, c, a).font = Font(bold=True)
                ws.cell(start + 1, c, b)
        ws.column_dimensions['A'].width = 24
        ws.column_dimensions['B'].width = 44
        wb.save(path)
        return path
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = docx.Document()
    lines = DEFAULT_BODIES[kind]
    heading = doc.add_heading(lines[0], 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in lines[1:]:
        doc.add_paragraph(line)
    doc.save(path)
    return path


def doc_row(db, cid, kind):
    return db.fetchone('SELECT file_path,data_hash,linked FROM gsvm_docs WHERE contract_id=? AND kind=?', (cid, kind))


def doc_state(db, cid, kind):
    """'none' — не создан, 'missing' — файл пропал, 'linked' — привязан готовый файл, 'stale' — данные изменились, 'fresh' — актуален."""
    row = doc_row(db, cid, kind)
    if not row:
        return 'none'
    if not os.path.exists(row[0]):
        return 'missing'
    if row[2]:
        return 'linked'
    return 'fresh' if row[1] == data_hash(db, cid) else 'stale'


def generate(db, cid, kind):
    """Формирует документ в папке договора (старый файл перезаписывается); возвращает путь."""
    if kind not in DOC_KINDS:
        raise ValueError('Неизвестный документ')
    folder = md.contract_folder(db, cid)
    if not folder:
        raise ValueError('Сначала привяжите или создайте папку договора кнопкой «Папка договора»')
    tags = tag_map(db, cid)
    title, ext, _group, prefix = DOC_KINDS[kind]
    number = tags.get('НОМЕР_ДОГОВОРА') or str(cid)
    out = os.path.join(folder, gd.safe_name(f'{prefix}_{number.replace("/", "-")}').replace(' ', '_') + '.' + ext)
    template = template_path(db, kind)
    if ext == 'xlsx':
        gd.render_xlsx(template, out, tags, payments(db, cid), normalize=norm_tag)
    else:
        gd.render_docx(template, out, tags, normalize=norm_tag)
    db.execute('INSERT OR REPLACE INTO gsvm_docs(contract_id,kind,file_path,data_hash,generated_at,linked) VALUES(?,?,?,?,?,0)',
               (cid, kind, out, data_hash(db, cid), datetime.now().isoformat(timespec='seconds')))
    return out


def link_doc(db, cid, kind, path):
    if not os.path.isfile(path):
        raise ValueError('Файл не найден: ' + path)
    db.execute('INSERT OR REPLACE INTO gsvm_docs(contract_id,kind,file_path,data_hash,generated_at,linked) VALUES(?,?,?,?,?,1)',
               (cid, kind, path, '', datetime.now().isoformat(timespec='seconds')))


def remove_doc(db, cid, kind, delete_file=False):
    """Убирает документ из списка; файл удаляется только если создан программой (не привязан) и delete_file=True."""
    row = doc_row(db, cid, kind)
    if row and delete_file and not row[2] and os.path.isfile(row[0]):
        os.remove(row[0])
    db.execute('DELETE FROM gsvm_docs WHERE contract_id=? AND kind=?', (cid, kind))


def doc_path(db, cid, kind):
    row = doc_row(db, cid, kind)
    return row[0] if row else ''
