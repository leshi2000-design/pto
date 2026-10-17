"""Накладные: окно проверки распознанной накладной (фото, скан, PDF, Excel) и вкладка «Накладные» со списком и товарами."""
import os
from concurrent.futures import ThreadPoolExecutor

from PyQt6.QtCore import Qt, QTimer
from PyQt6.QtGui import QColor
from PyQt6.QtWidgets import (QAbstractItemView, QCheckBox, QComboBox, QDialog, QDoubleSpinBox, QFileDialog, QFormLayout, QGroupBox, QHBoxLayout, QHeaderView, QLabel, QLineEdit,
                             QMessageBox, QPushButton, QSplitter, QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget)

from . import incoming_domain as inc
from . import invoice_parse as ip
from . import invoices_domain as inv
from .database import db
from .date_widgets import DateEdit

COLUMNS = ['Наименование', 'Ед.', 'Кол-во', 'Цена', 'Сумма', 'НДС %', 'Сумма НДС', 'Всего с НДС', 'Проверка / справочник']
NUMERIC = (2, 3, 4, 5, 6, 7)
RED, AMBER, GREEN = QColor('#FDE2E1'), QColor('#FEF3C7'), QColor('#E3F4EA')


def tesseract_setting():
    return db.get_setting('tesseract_cmd', '')


class InvoiceReviewDialog(QDialog):
    """Распознанная накладная: реквизиты и товары можно править; сохраняется только по кнопке. data — результат invoice_parse (или None: тогда разбирается path)."""
    def __init__(self, path=None, parent=None, data=None, iid=None, sync=False):
        super().__init__(parent)
        self.path, self.iid, self.saved, self.saved_id = path, iid, False, None
        self.setWindowTitle('Накладная · проверка')
        self.resize(1180, 820)
        layout = QVBoxLayout(self)
        self.title = QLabel(f'<b>{os.path.basename(path)}</b>' if path else '<b>Накладная</b>')
        layout.addWidget(self.title)
        self.warning = QLabel()
        self.warning.setWordWrap(True)
        self.warning.setStyleSheet('color: #B45309;')
        layout.addWidget(self.warning)

        top = QHBoxLayout()
        head = QGroupBox('Накладная')
        hf = QFormLayout(head)
        self.kind = QComboBox()
        self.kind.setEditable(True)
        self.kind.addItems(['ТН-2', 'ТТН', 'Накладная', 'Счёт-фактура', 'Акт приёма-передачи'])
        self.number = QLineEdit()
        from PyQt6.QtCore import QDate
        self.date = DateEdit(QDate.currentDate())
        self.date.setCalendarPopup(True)
        self.currency = QLineEdit('BYN')
        self.note = QLineEdit()
        for label, w in (('Вид', self.kind), ('Номер', self.number), ('Дата', self.date), ('Валюта', self.currency), ('Примечание', self.note)):
            hf.addRow(label, w)
        top.addWidget(head, 1)
        sup = QGroupBox('Поставщик (грузоотправитель)')
        sf = QFormLayout(sup)
        self.s_name, self.s_unp, self.s_address, self.s_account, self.s_bank, self.s_bic = (QLineEdit() for _ in range(6))
        for label, w in (('Название', self.s_name), ('УНП', self.s_unp), ('Адрес', self.s_address), ('Расчётный счёт', self.s_account), ('Банк', self.s_bank), ('БИК', self.s_bic)):
            sf.addRow(label, w)
        top.addWidget(sup, 2)
        buy = QGroupBox('Получатель (покупатель)')
        bf = QFormLayout(buy)
        self.b_name, self.b_unp, self.b_address = (QLineEdit() for _ in range(3))
        for label, w in (('Название', self.b_name), ('УНП', self.b_unp), ('Адрес', self.b_address)):
            bf.addRow(label, w)
        top.addWidget(buy, 1)
        layout.addLayout(top)

        self.table = QTableWidget(0, len(COLUMNS))
        self.table.setHorizontalHeaderLabels(COLUMNS)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(8, QHeaderView.ResizeMode.Stretch)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.verticalHeader().setVisible(True)
        layout.addWidget(self.table, 1)
        bar = QHBoxLayout()
        for text, fn in (('Добавить строку', self.add_row), ('Удалить строку', self.remove_row), ('Пересчитать суммы', self.recalc_all), ('Подставить названия из справочника', self.use_catalog_names)):
            b = QPushButton(text)
            b.clicked.connect(fn)
            bar.addWidget(b)
        bar.addStretch()
        self.totals = QLabel()
        bar.addWidget(self.totals)
        layout.addLayout(bar)

        after = QGroupBox('После сохранения')
        af = QVBoxLayout(after)
        row = QHBoxLayout()
        self.opt_supplier = QCheckBox('Добавить поставщика в «Справочник юрлиц» (если там нет такого УНП)')
        self.opt_supplier.setChecked(True)
        row.addWidget(self.opt_supplier)
        af.addLayout(row)
        row = QHBoxLayout()
        self.opt_catalog = QCheckBox('Обновить закупочные цены в справочнике материалов и добавить новые позиции. Наценка, %:')
        self.markup = QDoubleSpinBox()
        self.markup.setRange(0, 500)
        self.markup.setDecimals(1)
        self.markup.setValue(float(db.get_setting('invoice_markup', '0') or 0))
        row.addWidget(self.opt_catalog)
        row.addWidget(self.markup)
        row.addStretch()
        af.addLayout(row)
        row = QHBoxLayout()
        self.opt_stock = QCheckBox('Оприходовать товары на склад («Списание → материалы») на дату накладной')
        row.addWidget(self.opt_stock)
        af.addLayout(row)
        self.opt_copy = QCheckBox('Сохранить файл накладной в папке «Входящие документы»')
        self.opt_copy.setChecked(bool(path))
        self.opt_copy.setVisible(bool(path))
        af.addWidget(self.opt_copy)
        layout.addWidget(after)
        if iid:
            after.setVisible(False)

        buttons = QHBoxLayout()
        buttons.addStretch()
        skip = QPushButton('Пропустить')
        skip.clicked.connect(self.reject)
        self.ok = QPushButton('Сохранить накладную')
        self.ok.setProperty('type', 'primary')
        self.ok.clicked.connect(self.save)
        buttons.addWidget(skip)
        buttons.addWidget(self.ok)
        layout.addLayout(buttons)
        self.table.itemChanged.connect(self.cell_changed)
        self._busy = False
        self._pool = None
        self.result = None
        if data is not None:
            self.fill(data)
        elif path:
            self.start_analysis(sync)

    # --- распознавание в фоне (OCR занимает секунды)
    def start_analysis(self, sync=False):
        self.ok.setEnabled(False)
        self.warning.setText('Распознаю документ… (для фотографий это занимает 5–20 секунд)')
        if sync:
            self.finish_analysis(self._work())
            return
        self._pool = ThreadPoolExecutor(max_workers=1)
        self._future = self._pool.submit(self._work)
        self._timer = QTimer(self)
        self._timer.timeout.connect(self.poll)
        self._timer.start(150)

    def _work(self):
        try:
            return ip.parse_file(self.path, tesseract_setting())
        except Exception as e:                                   # результат — исключение: показываем пользователю, а не роняем окно
            return e

    def poll(self):
        if not self._future.done():
            return
        self._timer.stop()
        self.finish_analysis(self._future.result())

    def finish_analysis(self, result):
        self.ok.setEnabled(True)
        if isinstance(result, Exception):
            self.fill(dict(kind='Накладная', number='', date='', currency='BYN', supplier={}, buyer={}, items=[], totals={}, warnings=[], source='', text=''))
            self.warning.setText(f'Не удалось распознать: {result}\nМожно ввести накладную вручную: добавьте строки и заполните реквизиты.')
            return
        self.fill(result)

    def done(self, code):
        if self._pool:
            self._pool.shutdown(wait=False, cancel_futures=True)
        super().done(code)

    # --- заполнение и расчёт
    def fill(self, r):
        self.result = r
        s, b = r.get('supplier') or {}, r.get('buyer') or {}
        self.kind.setCurrentText(r.get('kind') or 'Накладная')
        self.number.setText(r.get('number') or '')
        self._set_date(r.get('date') or '')
        self.currency.setText(r.get('currency') or 'BYN')
        for w, key in ((self.s_name, 'name'), (self.s_unp, 'unp'), (self.s_address, 'address'), (self.s_account, 'account'), (self.s_bank, 'bank'), (self.s_bic, 'bic')):
            w.setText(s.get(key, ''))
        for w, key in ((self.b_name, 'name'), (self.b_unp, 'unp'), (self.b_address, 'address')):
            w.setText(b.get(key, ''))
        warnings = list(r.get('warnings') or [])
        self.warning.setText('\n'.join(warnings))
        self.doc_totals = r.get('totals') or {}
        self.table.blockSignals(True)
        self.table.setRowCount(0)
        for it in r.get('items') or []:
            self._append(it)
        self.table.blockSignals(False)
        self.refresh_rows()

    def _set_date(self, iso):
        from PyQt6.QtCore import QDate
        d = QDate.fromString(iso, 'yyyy-MM-dd')
        self.date.setDate(d if d.isValid() else QDate.currentDate())

    def _append(self, it):
        r = self.table.rowCount()
        self.table.insertRow(r)
        values = [it.get('name') or '', it.get('unit') or '', it.get('qty'), it.get('price'), it.get('sum'), it.get('vat_rate'), it.get('vat_sum'), it.get('total')]
        for c, v in enumerate(values):
            text = v if c < 2 else ('' if v is None else (f'{v:g}' if c in (2, 5) else f'{float(v):.2f}'))
            item = QTableWidgetItem(str(text))
            if c in NUMERIC:
                item.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            self.table.setItem(r, c, item)
        info = QTableWidgetItem('')
        info.setFlags(info.flags() & ~Qt.ItemFlag.ItemIsEditable)
        self.table.setItem(r, 8, info)
        self.table.item(r, 0).setData(Qt.ItemDataRole.UserRole, it.get('material_id'))

    def row_item(self, r):
        def num(c):
            t = self.table.item(r, c).text() if self.table.item(r, c) else ''
            return ip.to_number(t)
        return dict(name=self.table.item(r, 0).text().strip(), unit=self.table.item(r, 1).text().strip(), qty=num(2), price=num(3), sum=num(4), vat_rate=num(5), vat_sum=num(6), total=num(7),
                    material_id=self.table.item(r, 0).data(Qt.ItemDataRole.UserRole))

    def items(self):
        return [self.row_item(r) for r in range(self.table.rowCount())]

    def cell_changed(self, item):
        if item.column() in (2, 3, 4, 5, 6, 7):
            self.table.blockSignals(True)
            self.recalc_row(item.row(), item.column())
            self.table.blockSignals(False)
        self.refresh_rows()

    def recalc_row(self, r, changed):
        it = self.row_item(r)
        q, p = it['qty'], it['price']
        if changed in (2, 3) and q is not None and p is not None:
            it['sum'] = round(q * p, 2)
        if changed in (2, 3, 4, 5) and it['sum'] is not None and it['vat_rate'] is not None:
            it['vat_sum'] = round(it['sum'] * it['vat_rate'] / 100, 2)
        if changed in (2, 3, 4, 5, 6) and it['sum'] is not None:
            it['total'] = round(it['sum'] + (it['vat_sum'] or 0), 2)
        for c, key in ((4, 'sum'), (6, 'vat_sum'), (7, 'total')):
            self.table.item(r, c).setText('' if it[key] is None else f'{it[key]:.2f}')

    def recalc_all(self):
        self.table.blockSignals(True)
        for r in range(self.table.rowCount()):
            self.recalc_row(r, 3)
        self.table.blockSignals(False)
        self.refresh_rows()

    def refresh_rows(self):
        """Подсветка: красный — не сходится арифметика, жёлтый — исправлено автоматически, зелёный — найдено в справочнике; справа пояснение."""
        fixed = {i: (it.get('fixed') or '') for i, it in enumerate((self.result or {}).get('items') or [])}
        self.table.blockSignals(True)
        for r in range(self.table.rowCount()):
            it = self.row_item(r)
            problem = ip._check_row({k: it[k] for k in ('qty', 'price', 'sum', 'vat_sum', 'total')})
            matches = self.cached_match(it['name'], it['unit'])
            note = problem
            color = RED if problem else None
            if not problem and fixed.get(r):
                note, color = 'исправлено по арифметике: ' + fixed[r], AMBER
            if matches:
                m = matches[0]
                same = inv._norm(it['name']) == inv._norm(m[1])
                note = (note + ' · ' if note else '') + (f'в справочнике: {m[1]}' + ('' if same else f' ({int(m[4] * 100)}%)'))
                if color is None:
                    color = GREEN
            for c in range(self.table.columnCount()):
                cell = self.table.item(r, c)
                if cell:
                    cell.setBackground(color if color else QColor(Qt.GlobalColor.transparent))
            self.table.item(r, 8).setText(note)
        self.table.blockSignals(False)
        s, v, t = inv.totals_of(self.items())
        doc = getattr(self, 'doc_totals', {}) or {}
        text = f'Строк: {self.table.rowCount()} · без НДС {s:,.2f} · НДС {v:,.2f} · всего {t:,.2f}'.replace(',', ' ')
        if doc.get('total') is not None and abs(doc['total'] - t) > 0.05 and abs(doc['total'] - s) > 0.05:
            text += f'  ⚠ в документе итог {doc["total"]:,.2f}'.replace(',', ' ')
        self.totals.setText(text)

    def cached_match(self, name, unit):
        if not name:
            return []
        if not hasattr(self, '_index'):
            self._index, self._matches = inv.catalog_index(db), {}
        key = (name, unit)
        if key not in self._matches:
            self._matches[key] = inv.match_catalog(db, name, unit, 1, self._index)
        return self._matches[key]

    def use_catalog_names(self):
        self.table.blockSignals(True)
        count = 0
        for r in range(self.table.rowCount()):
            it = self.row_item(r)
            m = (self.cached_match(it['name'], it['unit']) or [None])[0]
            if m and m[4] >= 0.7 and m[1] != it['name']:
                self.table.item(r, 0).setText(m[1])
                self.table.item(r, 0).setData(Qt.ItemDataRole.UserRole, m[0])
                count += 1
        self.table.blockSignals(False)
        self.refresh_rows()
        QMessageBox.information(self, 'Справочник', f'Заменено названий: {count}. Проверьте, что позиции действительно те же (единицы, размеры).')

    def add_row(self):
        self.table.blockSignals(True)
        self._append(dict(name='', unit='шт', qty=1.0))
        self.table.blockSignals(False)
        self.table.setCurrentCell(self.table.rowCount() - 1, 0)
        self.refresh_rows()

    def remove_row(self):
        rows = sorted({i.row() for i in self.table.selectedItems()}, reverse=True)
        for r in rows:
            self.table.removeRow(r)
        self.refresh_rows()

    # --- сохранение
    def data(self):
        s, v, t = inv.totals_of(self.items())
        doc = getattr(self, 'doc_totals', {}) or {}
        return dict(kind=self.kind.currentText().strip() or 'Накладная', number=self.number.text().strip(), doc_date=self.date.date().toString('yyyy-MM-dd'), currency=self.currency.text().strip() or 'BYN',
                    source=(self.result or {}).get('source', ''), note=self.note.text().strip(), file_path=self.path or '',
                    supplier_name=self.s_name.text().strip(), supplier_unp=self.s_unp.text().strip(), supplier_address=self.s_address.text().strip(), supplier_account=self.s_account.text().strip().upper(),
                    supplier_bank=self.s_bank.text().strip(), supplier_bic=self.s_bic.text().strip().upper(),
                    buyer_name=self.b_name.text().strip(), buyer_unp=self.b_unp.text().strip(), buyer_address=self.b_address.text().strip(),
                    sum_total=s, vat_total=v, total_with_vat=t, items=self.items())

    def save(self):
        data = self.data()
        try:
            inv.validate(data)
        except ValueError as e:
            QMessageBox.warning(self, 'Накладная', str(e))
            return
        bad = [i + 1 for i, it in enumerate(data['items']) if ip._check_row({k: it[k] for k in ('qty', 'price', 'sum', 'vat_sum', 'total')})]
        if bad and QMessageBox.question(self, 'Проверка', f'В строках {", ".join(map(str, bad[:10]))} не сходится арифметика. Всё равно сохранить?') != QMessageBox.StandardButton.Yes:
            return
        try:
            with db.transaction():
                if self.path and self.opt_copy.isChecked() and not self.iid:
                    did = inc.save(db, self.path, dict(kind='Накладная', number=data['number'], doc_date=data['doc_date'], amount=data['total_with_vat'], counterparty=data['supplier_name'],
                                                       note=f'Товаров: {len(data["items"])}'), copy=True)
                    data['incoming_id'] = did
                    data['file_path'] = db.fetchone('SELECT file_path FROM incoming_docs WHERE id=?', (did,))[0]
                iid = inv.save(db, data, self.iid)
                if not self.iid:
                    if self.opt_supplier.isChecked() and (data['supplier_unp'] or data['supplier_name']):
                        lid, _created = inv.ensure_supplier(db, data)
                        db.execute('UPDATE invoices SET legal_id=? WHERE id=?', (lid, iid))
                    if self.opt_catalog.isChecked():
                        inv.apply_to_catalog(db, iid, self.markup.value())
                        db.set_setting('invoice_markup', str(self.markup.value()))
                    if self.opt_stock.isChecked():
                        inv.receive_to_stock(db, iid, data['doc_date'] or None)
        except (ValueError, OSError) as e:
            QMessageBox.warning(self, 'Накладная не сохранена', str(e))
            return
        self.saved, self.saved_id = True, iid
        self.accept()


class InvoicesTab(QWidget):
    """Список накладных и товаров выбранной накладной; поиск по номеру, поставщику, УНП и названию товара."""
    def __init__(self, parent=None):
        super().__init__(parent)
        layout = QVBoxLayout(self)
        bar = QHBoxLayout()
        self.search = QLineEdit()
        self.search.setPlaceholderText('Номер накладной, поставщик, УНП или название товара…')
        bar.addWidget(self.search, 1)
        for text, fn in (('Накладная (фото / Excel / PDF)…', self.add), ('Открыть и править', self.open), ('В склад', self.to_stock), ('В справочник цен', self.to_catalog),
                         ('Экспорт в Excel', self.export), ('Удалить', self.delete)):
            b = QPushButton(text)
            b.clicked.connect(fn)
            if text.startswith('Накладная'):
                b.setProperty('type', 'primary')
            if text == 'Удалить':
                b.setProperty('type', 'danger')
            bar.addWidget(b)
        layout.addLayout(bar)
        split = QSplitter(Qt.Orientation.Vertical)
        self.table = QTableWidget(0, 8)
        self.table.setHorizontalHeaderLabels(['Дата', '№ накладной', 'Поставщик', 'УНП', 'Сумма с НДС', 'Позиций', 'На складе', 'Источник'])
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        self.table.doubleClicked.connect(lambda *_: self.open())
        self.table.itemSelectionChanged.connect(self.show_items)
        split.addWidget(self.table)
        self.lines = QTableWidget(0, 7)
        self.lines.setHorizontalHeaderLabels(['Наименование', 'Ед.', 'Кол-во', 'Цена', 'Сумма', 'НДС %', 'Всего с НДС'])
        self.lines.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.lines.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.lines.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        split.addWidget(self.lines)
        layout.addWidget(split, 1)
        self.ids = []
        self.search.textChanged.connect(self.load_data)
        self.load_data()

    def load_data(self, *_):
        rows = inv.listing(db, self.search.text())
        self.ids = [r[0] for r in rows]
        self.table.setRowCount(len(rows))
        for r, (_id, day, number, name, unp, total, count, stocked, source) in enumerate(rows):
            for c, v in enumerate((day or '', number, name, unp, f'{total:,.2f}'.replace(',', ' ') if total is not None else '', count, 'да' if stocked else '', {'excel': 'Excel', 'ocr': 'фото/скан', 'text': 'PDF'}.get(source, source))):
                self.table.setItem(r, c, QTableWidgetItem(str(v)))
        self.show_items()

    def current_id(self):
        r = self.table.currentRow()
        return self.ids[r] if 0 <= r < len(self.ids) else None

    def show_items(self):
        iid = self.current_id()
        data = inv.get(db, iid) if iid else None
        items = data['items'] if data else []
        self.lines.setRowCount(len(items))
        for r, it in enumerate(items):
            for c, v in enumerate((it['name'], it['unit'], it['qty'], it['price'], it['sum'], it['vat_rate'], it['total'])):
                self.lines.setItem(r, c, QTableWidgetItem('' if v is None else (f'{v:g}' if isinstance(v, float) and c == 2 else str(v))))

    def add(self):
        paths, _ = QFileDialog.getOpenFileNames(self, 'Накладная', '', 'Накладные (*.xlsx *.xlsm *.pdf *.jpg *.jpeg *.png *.tif *.tiff *.bmp);;Все файлы (*)')
        if paths:
            process_invoice_files(self, paths)
            self.load_data()

    def open(self):
        iid = self.current_id()
        if not iid:
            return
        data = inv.get(db, iid)
        result = dict(kind=data['kind'], number=data['number'], date=data['doc_date'], currency=data['currency'], source=data['source'], warnings=[],
                      supplier=dict(name=data['supplier_name'], unp=data['supplier_unp'], address=data['supplier_address'], account=data['supplier_account'], bank=data['supplier_bank'], bic=data['supplier_bic']),
                      buyer=dict(name=data['buyer_name'], unp=data['buyer_unp'], address=data['buyer_address']),
                      items=data['items'], totals=dict(sum=data['sum_total'], vat=data['vat_total'], total=data['total_with_vat']))
        d = InvoiceReviewDialog(data['file_path'] or None, self, data=result, iid=iid)
        d.note.setText(data['note'])
        d.exec()
        self.load_data()

    def to_stock(self):
        iid = self.current_id()
        if not iid:
            return
        if QMessageBox.question(self, 'Склад', 'Оприходовать все товары этой накладной на склад?') != QMessageBox.StandardButton.Yes:
            return
        try:
            n = inv.receive_to_stock(db, iid)
        except ValueError as e:
            QMessageBox.warning(self, 'Склад', str(e))
            return
        QMessageBox.information(self, 'Склад', f'Оприходовано позиций: {n}')
        self.load_data()

    def to_catalog(self):
        iid = self.current_id()
        if not iid:
            return
        pct, ok = QInputDialogGetDouble(self)
        if not ok:
            return
        r = inv.apply_to_catalog(db, iid, pct)
        QMessageBox.information(self, 'Справочник', f'Обновлено позиций: {r["updated"]}, добавлено новых: {r["added"]}, пропущено: {r["skipped"]}')

    def export(self):
        iid = self.current_id()
        if not iid:
            return
        data = inv.get(db, iid)
        path, _ = QFileDialog.getSaveFileName(self, 'Экспорт накладной', f'Накладная_{data["number"]}.xlsx'.replace('/', '-'), 'Excel (*.xlsx)')
        if path:
            export_excel(data, path)

    def delete(self):
        iid = self.current_id()
        if iid and QMessageBox.question(self, 'Удалить', 'Удалить накладную из программы? Складские движения и справочник не изменятся.') == QMessageBox.StandardButton.Yes:
            inv.delete(db, iid)
            self.load_data()


def QInputDialogGetDouble(parent):
    from PyQt6.QtWidgets import QInputDialog
    value, ok = QInputDialog.getDouble(parent, 'Справочник цен', 'Наценка на закупочную цену для новых позиций, %:', float(db.get_setting('invoice_markup', '0') or 0), 0, 500, 1)
    if ok:
        db.set_setting('invoice_markup', str(value))
    return value, ok


def export_excel(data, path):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Накладная'
    ws.append([f'{data["kind"]} № {data["number"]} от {data["doc_date"]}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append(['Поставщик', data['supplier_name'], 'УНП', data['supplier_unp']])
    ws.append(['Счёт', data['supplier_account'], 'Банк', data['supplier_bank'], 'БИК', data['supplier_bic']])
    ws.append(['Получатель', data['buyer_name'], 'УНП', data['buyer_unp']])
    ws.append([])
    ws.append(['№', 'Наименование', 'Ед.', 'Количество', 'Цена', 'Стоимость', 'НДС, %', 'Сумма НДС', 'Всего с НДС'])
    for c in ws[6]:
        c.font = Font(bold=True)
    for n, it in enumerate(data['items'], 1):
        ws.append([n, it['name'], it['unit'], it['qty'], it['price'], it['sum'], it['vat_rate'], it['vat_sum'], it['total']])
    ws.append(['', 'Итого', '', '', '', data['sum_total'], '', data['vat_total'], data['total_with_vat']])
    for col, width in zip('ABCDEFGHI', (5, 60, 8, 12, 12, 14, 8, 12, 14)):
        ws.column_dimensions[col].width = width
    wb.save(path)
    return path


def process_invoice_files(parent, paths):
    """Окно проверки для каждой накладной подряд. -> число сохранённых."""
    saved = 0
    for path in paths:
        if os.path.isfile(path):
            d = InvoiceReviewDialog(path, parent)
            d.exec()
            saved += int(d.saved)
    return saved


class OcrSettingsBox(QGroupBox):
    """Настройки → «Распознавание документов»: путь к Tesseract и проверка, что фото накладных можно читать."""
    def __init__(self, parent=None):
        super().__init__('Распознавание накладных по фотографии', parent)
        layout = QVBoxLayout(self)
        layout.addWidget(QLabel('Excel и PDF с текстом читаются сразу. Для фотографий и сканов нужна бесплатная программа Tesseract OCR (Windows: сборка UB Mannheim; при установке '
                                'отметьте язык Russian). Если она стоит не в стандартной папке, укажите путь к tesseract.exe.'))
        row = QHBoxLayout()
        self.path = QLineEdit(tesseract_setting())
        self.path.setPlaceholderText('например C:\\Program Files\\Tesseract-OCR\\tesseract.exe (пусто — искать автоматически)')
        browse = QPushButton('Выбрать…')
        browse.clicked.connect(self.browse)
        check = QPushButton('Проверить')
        check.clicked.connect(self.check)
        row.addWidget(self.path, 1)
        row.addWidget(browse)
        row.addWidget(check)
        layout.addLayout(row)
        self.status = QLabel('')
        self.status.setWordWrap(True)
        layout.addWidget(self.status)
        self.refresh_status()

    def browse(self):
        path, _ = QFileDialog.getOpenFileName(self, 'Tesseract', '', 'tesseract (tesseract.exe tesseract);;Все файлы (*)')
        if path:
            self.path.setText(path)
            self.check()

    def check(self):
        db.set_setting('tesseract_cmd', self.path.text().strip())
        self.refresh_status()

    def refresh_status(self):
        ok, info = ip.ocr_available(self.path.text().strip())
        self.status.setText(f'<span style="color:#15803D">✔ Распознавание фотографий доступно ({info})</span>' if ok else f'<span style="color:#B45309">⚠ {info}</span>')
