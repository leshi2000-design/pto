"""
Редактор сметы: таблица позиций, скидки/наценки, клиент и договор, окно «Расшифровка и маржа», экспорт. Начислений на работы в смете нет.
"""
import re
import os
import logging

import pandas as pd
import openpyxl
from openpyxl.styles import Font, Border, Side
from openpyxl.utils import get_column_letter

import docx

from PyQt6.QtWidgets import (QVBoxLayout, QHBoxLayout, QPushButton, QLabel,
                              QTableWidgetItem, QInputDialog, QMessageBox, QFileDialog,
                              QDialog, QLineEdit, QComboBox, QMenu, QCheckBox, QGroupBox,
                              QDoubleSpinBox, QGridLayout)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QAction, QTextDocument, QShortcut, QKeySequence
from PyQt6.QtPrintSupport import QPrinter

from .database import db
from .workspace_view import choose_client
from .platform_utils import open_local
from .workspace_view import ExportDialog
from .widgets import ReorderTableWidget, SmartTableManager
from .dialogs_common import PaymentDialog, CustomItemDialog, WorkSelectionDialog, MaterialSelectionDialog
from .contract_card import ContractCardDialog
from . import estimates_domain as ed


class EstimateEditorDialog(QDialog):
    def __init__(self, estimate_id, title, parent=None):
        super().__init__(parent)
        self.estimate_id = estimate_id
        self.title = title
        self.setWindowTitle(f"Смета: {self.title}")
        self.resize(1200, 850)
        self._updating = False

        self.base_dir_total = 0.0
        self.base_work_total = 0.0
        self.base_mat_total = 0.0

        self.adj_mat_total = 0.0
        self.adj_work_total = 0.0
        self.adj_subtotal = 0.0

        self.social_val = 0.0
        self.overhead_val = 0.0
        self.profit_val = 0.0
        self.vat_val = 0.0
        self.grand_total = 0.0

        layout = QVBoxLayout(self)

        QShortcut(QKeySequence("Ctrl+S"), self).activated.connect(self.save_meta)
        QShortcut(QKeySequence("Insert"), self).activated.connect(self.add_custom)
        QShortcut(QKeySequence("Delete"), self).activated.connect(self.del_item)

        top_bar = QHBoxLayout()
        btn_client=QPushButton("Выбрать клиента из базы")
        btn_client.clicked.connect(lambda: choose_client(self,"estimates",self.estimate_id))
        top_bar.addWidget(btn_client)
        btn_unified=QPushButton("Документ по шаблону")
        btn_unified.clicked.connect(self.export_template)
        top_bar.addWidget(btn_unified)
        self._export_table="estimates"
        btn_rename = QPushButton("Переименовать")
        btn_rename.clicked.connect(self.rename_estimate)
        top_bar.addWidget(btn_rename)

        btn_pay = QPushButton("История оплат")
        btn_pay.setProperty("type", "primary")
        btn_pay.clicked.connect(self.open_payments)
        top_bar.addWidget(btn_pay)

        self.btn_contract = QPushButton("Договор ▼")
        self.contract_menu = QMenu(self)
        self.btn_contract.setMenu(self.contract_menu)
        self.contract_menu.aboutToShow.connect(self.fill_contract_menu)
        top_bar.addWidget(self.btn_contract)

        btn_breakdown = QPushButton("Расшифровка и маржа")
        btn_breakdown.setToolTip("Работы по статьям, материалы с маржой и итоги в одном окне; экспорт по шаблону")
        btn_breakdown.clicked.connect(self.show_work_breakdown)
        top_bar.addWidget(btn_breakdown)

        btn_export = QPushButton("Экспорт ▼")
        btn_export.setProperty("type", "primary")
        export_menu = QMenu(self)

        act_excel_norm = QAction("Смета (Excel - стандартная)", self)
        act_excel_norm.triggered.connect(lambda: self.export_excel(False))
        export_menu.addAction(act_excel_norm)

        act_excel_tpl = QAction("Смета (Excel - по красивому шаблону)", self)
        act_excel_tpl.triggered.connect(lambda: self.export_excel(True))
        export_menu.addAction(act_excel_tpl)

        act_pdf_full = QAction("Вся смета в PDF", self)
        act_pdf_full.triggered.connect(self.export_estimate_pdf)
        export_menu.addAction(act_pdf_full)

        export_menu.addSeparator()
        act_dos_pdf = QAction("Карточка клиента (PDF)", self)
        act_dos_pdf.triggered.connect(self.export_dossier_pdf)
        export_menu.addAction(act_dos_pdf)
        act_dos_doc = QAction("Карточка клиента (Word / Бланк)", self)
        act_dos_doc.triggered.connect(self.export_dossier_word)
        export_menu.addAction(act_dos_doc)
        act_dos_xls = QAction("Карточка клиента (Excel)", self)
        act_dos_xls.triggered.connect(self.export_dossier_excel)
        export_menu.addAction(act_dos_xls)

        btn_export.setMenu(export_menu)
        top_bar.addStretch()
        top_bar.addWidget(btn_export)
        layout.addLayout(top_bar)

        meta_group = QGroupBox("Данные клиента, статусы и прикрепленные файлы")
        meta_layout = QVBoxLayout(meta_group)

        row1 = QHBoxLayout()
        row1.addWidget(QLabel("Клиент:"))
        self.inp_client = QLineEdit()
        row1.addWidget(self.inp_client)
        row1.addWidget(QLabel("Телефон:"))
        self.inp_phone = QLineEdit()
        row1.addWidget(self.inp_phone)
        btn_legal = QPushButton("Юрлицо…")
        btn_legal.setToolTip("Выбрать организацию из справочника «Юрлица»")
        btn_legal.clicked.connect(self.choose_legal)
        row1.addWidget(btn_legal)
        btn_from_contract = QPushButton("Клиент из договора…")
        btn_from_contract.setToolTip("Взять клиента из карточки договора (сам договор к смете не привязывается)")
        btn_from_contract.clicked.connect(self.client_from_contract)
        row1.addWidget(btn_from_contract)
        meta_layout.addLayout(row1)
        self.lbl_contract = QLabel("")
        self.lbl_contract.setWordWrap(True)
        meta_layout.addWidget(self.lbl_contract)

        row2 = QHBoxLayout()
        row2.addWidget(QLabel("Текущие статусы:"))
        self.cb_status1 = QCheckBox("Предварительная смета")
        self.cb_status2 = QCheckBox("Передано в работу")
        self.cb_status3 = QCheckBox("Ожидает оплаты")
        self.cb_status4 = QCheckBox("Оплачено")
        for cb in [self.cb_status1, self.cb_status2, self.cb_status3, self.cb_status4]:
            row2.addWidget(cb)
            cb.clicked.connect(self.save_meta)
        self.inp_client.editingFinished.connect(self.save_meta)
        self.inp_phone.editingFinished.connect(self.save_meta)
        row2.addStretch()
        meta_layout.addLayout(row2)

        file_layout = QHBoxLayout()
        file_layout.addWidget(QLabel("Файлы:"))
        self.combo_files = QComboBox()
        self.combo_files.setMinimumWidth(250)
        btn_attach = QPushButton("Прикрепить")
        btn_open_file = QPushButton("Открыть")
        btn_del_file = QPushButton("Удалить")

        btn_attach.clicked.connect(self.attach_file)
        btn_open_file.clicked.connect(self.open_attached_file)
        btn_del_file.clicked.connect(self.delete_attached_file)

        file_layout.addWidget(self.combo_files)
        file_layout.addWidget(btn_attach)
        file_layout.addWidget(btn_open_file)
        file_layout.addWidget(btn_del_file)
        file_layout.addStretch()
        meta_layout.addLayout(file_layout)

        layout.addWidget(meta_group)

        main_h = QHBoxLayout()

        self.table = ReorderTableWidget()
        self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels(["Наименование позиции", "Тип", "Ед.", "Кол-во", "Цена (руб)", "Сумма (руб)"])
        self.table_manager = SmartTableManager(self.table, main_col=0)
        self.table.order_changed.connect(self.save_rows_order)
        main_h.addWidget(self.table, stretch=7)

        right_panel = QVBoxLayout()

        adj_group = QGroupBox("Скидки (–) и наценки (+) в %")
        adj_layout = QGridLayout(adj_group)

        self.spin_mat_adj = QDoubleSpinBox()
        self.spin_mat_adj.setRange(-100.0, 1000.0)
        self.spin_mat_adj.valueChanged.connect(self.save_meta)

        self.spin_work_adj = QDoubleSpinBox()
        self.spin_work_adj.setRange(-100.0, 1000.0)
        self.spin_work_adj.valueChanged.connect(self.save_meta)

        self.spin_total_adj = QDoubleSpinBox()
        self.spin_total_adj.setRange(-100.0, 1000.0)
        self.spin_total_adj.valueChanged.connect(self.save_meta)

        adj_layout.addWidget(QLabel("На материалы:"), 0, 0)
        adj_layout.addWidget(self.spin_mat_adj, 0, 1)
        adj_layout.addWidget(QLabel("На работы:"), 1, 0)
        adj_layout.addWidget(self.spin_work_adj, 1, 1)
        adj_layout.addWidget(QLabel("На итог:"), 2, 0)
        adj_layout.addWidget(self.spin_total_adj, 2, 1)
        right_panel.addWidget(adj_group)

        self.totals_lbl = QLabel()
        self.totals_lbl.setStyleSheet("font-size: 11pt; line-height: 1.5;")
        self.totals_lbl.setAlignment(Qt.AlignmentFlag.AlignRight)
        right_panel.addWidget(self.totals_lbl)

        self.lbl_grand_total = QLabel("ВСЕГО: 0.00")
        accent = db.get_setting("accent_color", "#0284C7")
        self.lbl_grand_total.setStyleSheet(f"font-size: 16pt; font-weight: bold; color: {accent};")
        self.lbl_grand_total.setAlignment(Qt.AlignmentFlag.AlignRight)
        right_panel.addWidget(self.lbl_grand_total)

        right_panel.addStretch()
        main_h.addLayout(right_panel, stretch=2)
        layout.addLayout(main_h)

        bottom_bar = QHBoxLayout()
        btn_add_mat = QPushButton("Из справочника")
        btn_add_mat.setProperty("type", "primary")
        btn_add_work = QPushButton("Добавить работу")
        btn_add_custom = QPushButton("Произвольная позиция")
        btn_add_section = QPushButton("Добавить раздел")
        self.btn_del = QPushButton("Удалить (Del)")
        self.btn_del.setProperty("type", "danger")

        bottom_bar.addWidget(btn_add_mat)
        bottom_bar.addWidget(btn_add_work)
        bottom_bar.addWidget(btn_add_custom)
        bottom_bar.addWidget(btn_add_section)
        bottom_bar.addWidget(self.btn_del)
        bottom_bar.addStretch()

        layout.addLayout(bottom_bar)

        btn_add_mat.clicked.connect(self.add_material)
        btn_add_work.clicked.connect(self.add_work)
        btn_add_custom.clicked.connect(self.add_custom)
        btn_add_section.clicked.connect(self.add_section)
        self.btn_del.clicked.connect(self.del_item)
        self.table.itemChanged.connect(self.on_cell_changed)

        self.load_meta()
        self.load_files_combo()
        self.load_items()

    def show_work_breakdown(self):
        from .work_breakdown_view import WorkBreakdownDialog
        WorkBreakdownDialog(self.estimate_id, self).exec()

    def load_files_combo(self):
        self.combo_files.clear()
        files = db.fetchall("SELECT id, file_name, file_path FROM attachments WHERE estimate_id=?", (self.estimate_id,))
        if files:
            for fid, fname, fpath in files:
                self.combo_files.addItem(fname, fpath)
        else:
            self.combo_files.addItem("Нет прикрепленных файлов", None)

    def attach_file(self):
        from .file_jobs import attach_files
        attach_files(self)

    def open_attached_file(self):
        fpath = self.combo_files.currentData()
        if fpath and os.path.exists(fpath):
            open_local(os.path.abspath(fpath))
        else:
            QMessageBox.warning(self, "Внимание", "Файл не выбран или не найден на диске.")

    def delete_attached_file(self):
        fpath = self.combo_files.currentData()
        if fpath:
            row = db.fetchone("SELECT id FROM attachments WHERE estimate_id=? AND file_path=?", (self.estimate_id, fpath))
            if row:
                db.execute("DELETE FROM attachments WHERE id=?", (row[0],))
                try:
                    if os.path.exists(fpath): os.remove(fpath)
                except OSError as e:
                    logging.warning(f"Не удалось удалить файл {fpath}: {e}")
                logging.info(f"Удален прикрепленный файл ID {row[0]}")
                self.load_files_combo()
                QMessageBox.information(self, "Успех", "Файл удален.")

    def save_rows_order(self):
        for row in range(self.table.rowCount()):
            item_id = self.table.item(row, 0).data(Qt.ItemDataRole.UserRole)
            if item_id:
                db.execute("UPDATE estimate_items SET sort_order=? WHERE id=?", (row, item_id))

    def get_max_sort_order(self):
        res = db.fetchone("SELECT MAX(sort_order) FROM estimate_items WHERE estimate_id=?", (self.estimate_id,))
        return (res[0] or 0) + 1

    def rename_estimate(self):
        new_title, ok = QInputDialog.getText(self, "Переименование", "Новое наименование объекта:", text=self.title)
        if ok and new_title.strip():
            self.title = new_title.strip()
            self.setWindowTitle(f"Смета: {self.title}")
            db.execute("UPDATE estimates SET title=? WHERE id=?", (self.title, self.estimate_id))
            if self.parent() and hasattr(self.parent(), 'load_data'):
                self.parent().load_data()

    def open_payments(self):
        res = db.fetchone("SELECT total FROM estimates WHERE id=?", (self.estimate_id,))
        total = res[0] if res and res[0] else 0.0
        dlg = PaymentDialog(self.estimate_id, total, self)
        dlg.exec()
        self.load_items()

    def open_contract_card(self):
        """Создание (или открытие) договора «Монтаж ГСВ» на основе этой сметы — как и прежде."""
        dlg = ContractCardDialog(self.estimate_id, self.title, self)
        dlg.exec()
        self.refresh_contract_info()

    # --- клиент и договор ---
    def refresh_contract_info(self):
        info = ed.contract_of(db, self.estimate_id)
        legal = db.fetchone("SELECT le_client_id,party_name FROM estimates WHERE id=?", (self.estimate_id,))
        is_legal = bool(legal and legal[0])
        self.inp_client.setReadOnly(is_legal)
        self.inp_phone.setReadOnly(is_legal)
        if is_legal:
            self.inp_client.setText(legal[1] or "")
        if info:
            self.lbl_contract.setText(f"📎 Договор: {info['label']}  (смета привязана; клиент и оплаты ведутся вместе с договором)")
            self.lbl_contract.setStyleSheet("color: #2563EB;")
        else:
            self.lbl_contract.setText("Смета без договора. Договор можно привязать позже — меню «Договор».")
            self.lbl_contract.setStyleSheet("color: #65758b;")

    def fill_contract_menu(self):
        self.contract_menu.clear()
        info = ed.contract_of(db, self.estimate_id)
        if info:
            self.contract_menu.addAction(f"Открыть: {info['label']}").triggered.connect(self.open_linked_contract)
            self.contract_menu.addAction("Отвязать договор").triggered.connect(self.unlink_contract)
        else:
            self.contract_menu.addAction("Привязать существующий договор…").triggered.connect(self.link_contract)
            self.contract_menu.addSeparator()
            self.contract_menu.addAction("Создать договор «Монтаж ГСВ» из сметы…").triggered.connect(self.open_contract_card)

    def pick_contract(self, with_estimate, title):
        from .estimates_ui import pick_contract
        return pick_contract(self, with_estimate, title)

    def link_contract(self):
        pick = self.pick_contract(False, "Привязать договор к смете")
        if not pick:
            return
        try:
            ed.link_contract(db, self.estimate_id, *pick)
        except ValueError as e:
            QMessageBox.warning(self, "Договор", str(e))
            return
        self.load_meta()
        self.refresh_contract_info()
        if self.parent() and hasattr(self.parent(), 'load_data'):
            self.parent().load_data()

    def unlink_contract(self):
        if QMessageBox.question(self, "Договор", "Отвязать договор от сметы? Смета и договор останутся, связь между ними исчезнет.") != QMessageBox.StandardButton.Yes:
            return
        try:
            ed.unlink_contract(db, self.estimate_id)
        except ValueError as e:
            QMessageBox.warning(self, "Договор", str(e))
            return
        self.refresh_contract_info()
        if self.parent() and hasattr(self.parent(), 'load_data'):
            self.parent().load_data()

    def open_linked_contract(self):
        info = ed.contract_of(db, self.estimate_id)
        if info:
            from .workspace_view import open_record
            open_record(info['table'], info['id'], self)
            self.refresh_contract_info()

    def choose_legal(self):
        d = QDialog(self)
        d.setWindowTitle("Юрлицо — клиент сметы")
        d.resize(480, 420)
        lay = QVBoxLayout(d)
        from PyQt6.QtWidgets import QListWidget, QListWidgetItem
        lst = QListWidget()
        for lid, name in db.fetchall("SELECT id,name FROM le_clients ORDER BY name"):
            it = QListWidgetItem(name)
            it.setData(Qt.ItemDataRole.UserRole, lid)
            lst.addItem(it)
        lay.addWidget(lst, 1)
        ok = QPushButton("Выбрать")
        ok.setProperty("type", "primary")
        lay.addWidget(ok)
        ok.clicked.connect(lambda: d.accept() if lst.currentItem() else None)
        lst.itemDoubleClicked.connect(lambda *_: d.accept())
        if d.exec() and lst.currentItem():
            try:
                ed.set_client(db, self.estimate_id, "legal", lst.currentItem().data(Qt.ItemDataRole.UserRole))
            except ValueError as e:
                QMessageBox.warning(self, "Клиент", str(e))
                return
            self.load_meta()
            self.refresh_contract_info()
            if self.parent() and hasattr(self.parent(), 'load_data'):
                self.parent().load_data()

    def client_from_contract(self):
        pick = self.pick_contract(True, "Клиент из карточки договора")
        if not pick:
            return
        if ed.contract_of(db, self.estimate_id):
            QMessageBox.information(self, "Клиент", "Смета привязана к договору — клиент задаётся в карточке договора.")
            return
        if not ed.apply_contract_client(db, self.estimate_id, *pick):
            QMessageBox.warning(self, "Клиент", "В выбранном договоре клиент не указан.")
            return
        self.load_meta()
        self.refresh_contract_info()
        if self.parent() and hasattr(self.parent(), 'load_data'):
            self.parent().load_data()

    def load_meta(self):
        self._updating = True
        row = db.fetchone("""SELECT coalesce(nullif(client_name,''),party_name), client_phone, statuses, mat_adj_pct, work_adj_pct, total_adj_pct
                             FROM estimates WHERE id=?""", (self.estimate_id,))
        self.meta_cache = {'mat_adj': 0.0, 'work_adj': 0.0, 'tot_adj': 0.0}
        if row:
            c_name, c_phone, statuses, mat_adj, work_adj, tot_adj = row
            self.inp_client.setText(c_name or "")
            self.inp_phone.setText(c_phone or "")
            s_list = statuses or ""
            self.cb_status1.setChecked("Предварительная смета" in s_list)
            self.cb_status2.setChecked("Передано в работу" in s_list)
            self.cb_status3.setChecked("Ожидает оплаты" in s_list)
            self.cb_status4.setChecked("Оплачено" in s_list)

            self.spin_mat_adj.setValue(float(mat_adj or 0.0))
            self.spin_work_adj.setValue(float(work_adj or 0.0))
            self.spin_total_adj.setValue(float(tot_adj or 0.0))

            self.meta_cache = {
                'mat_adj': float(mat_adj or 0.0), 'work_adj': float(work_adj or 0.0), 'tot_adj': float(tot_adj or 0.0)
            }
        self._updating = False
        self.refresh_contract_info()

    def save_meta(self):
        if self._updating: return
        active = []
        if self.cb_status1.isChecked(): active.append("Предварительная смета")
        if self.cb_status2.isChecked(): active.append("Передано в работу")
        if self.cb_status3.isChecked(): active.append("Ожидает оплаты")
        if self.cb_status4.isChecked(): active.append("Оплачено")
        s_str = ", ".join(active)

        legal = db.fetchone("SELECT le_client_id FROM estimates WHERE id=?", (self.estimate_id,))
        if legal and legal[0]:      # клиент-юрлицо задаётся выбором из справочника, а не вводом текста
            db.execute("UPDATE estimates SET statuses=?, mat_adj_pct=?, work_adj_pct=?, total_adj_pct=? WHERE id=?",
                       (s_str, self.spin_mat_adj.value(), self.spin_work_adj.value(), self.spin_total_adj.value(), self.estimate_id))
        else:
            db.execute("""UPDATE estimates SET client_name=?, client_phone=?, statuses=?,
                          mat_adj_pct=?, work_adj_pct=?, total_adj_pct=? WHERE id=?""",
                       (self.inp_client.text(), self.inp_phone.text(), s_str,
                        self.spin_mat_adj.value(), self.spin_work_adj.value(), self.spin_total_adj.value(),
                        self.estimate_id))

        self.meta_cache['mat_adj'] = self.spin_mat_adj.value()
        self.meta_cache['work_adj'] = self.spin_work_adj.value()
        self.meta_cache['tot_adj'] = self.spin_total_adj.value()
        self.sync_total()
        if self.parent() and hasattr(self.parent(), 'load_data'):
            self.parent().load_data()

    def load_items(self):
        self._updating = True
        self.table.setRowCount(0)
        self.base_dir_total = 0.0
        self.base_work_total = 0.0
        self.base_mat_total = 0.0

        rows = db.fetchall("SELECT id, name, item_type, unit, quantity, price, sum FROM estimate_items WHERE estimate_id=? ORDER BY sort_order ASC, id ASC", (self.estimate_id,))
        self.table.setRowCount(len(rows))
        for r, (item_id, name, itype, unit, qty, price, sm) in enumerate(rows):
            item_name = QTableWidgetItem(name or "")
            item_name.setData(Qt.ItemDataRole.UserRole, item_id)
            self.table.setItem(r, 0, item_name)

            i_type = QTableWidgetItem(itype or "Материал")
            i_type.setFlags(i_type.flags() & ~Qt.ItemFlag.ItemIsEditable)

            if itype == "Раздел":
                for col_idx in range(6):
                    bg_col = QTableWidgetItem()
                    if col_idx == 0: bg_col = item_name
                    elif col_idx == 1: bg_col = i_type
                    font = bg_col.font()
                    font.setBold(True)
                    bg_col.setFont(font)
                    if col_idx not in [0, 1]:
                        bg_col.setFlags(bg_col.flags() & ~Qt.ItemFlag.ItemIsEditable)
                        self.table.setItem(r, col_idx, bg_col)
                self.table.setItem(r, 1, i_type)
            else:
                self.table.setItem(r, 1, i_type)

                i_u = QTableWidgetItem(unit or "")
                i_u.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.table.setItem(r, 2, i_u)

                q_val = float(qty) if qty is not None else 0.0
                i_qty = QTableWidgetItem(f"{q_val:.3f}")
                i_qty.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(r, 3, i_qty)

                p_val = float(price) if price is not None else 0.0
                i_prc = QTableWidgetItem(f"{p_val:,.2f}".replace(",", " "))
                i_prc.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.table.setItem(r, 4, i_prc)

                sm_val = float(sm) if sm is not None else 0.0
                i_sm = QTableWidgetItem(f"{sm_val:,.2f}".replace(",", " "))
                i_sm.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                i_sm.setFlags(i_sm.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.table.setItem(r, 5, i_sm)

                self.base_dir_total += sm_val
                if itype == "Работа": self.base_work_total += sm_val
                else: self.base_mat_total += sm_val

        self._updating = False
        self.sync_total()

    def add_material(self):
        dlg = MaterialSelectionDialog(self)
        if dlg.exec() == QDialog.DialogCode.Accepted and dlg.selected_material:
            mat = dlg.selected_material
            qty, ok = QInputDialog.getDouble(self, "Количество", f"Укажите объем ({mat['unit']}):", 1.0, 0.001, 1000000.0, 3)
            if ok:
                sm = qty * mat['price']
                so = self.get_max_sort_order()
                db.execute("INSERT INTO estimate_items (estimate_id, item_type, name, unit, quantity, price, sum, sort_order, purchase_price) VALUES (?, 'Материал', ?, ?, ?, ?, ?, ?, ?)",
                           (self.estimate_id, mat['name'], mat['unit'], qty, mat['price'], sm, so, mat['purchase_price']))
                self.load_items()

    def add_work(self):
        dlg = WorkSelectionDialog(self)
        if dlg.exec() == QDialog.DialogCode.Accepted and dlg.selected_work:
            work = dlg.selected_work
            qty, ok = QInputDialog.getDouble(self, "Объем работ", f"Укажите объем ({work['unit']}):", 1.0, 0.001, 1000000.0, 3)
            if ok:
                sm = qty * work['price']
                so = self.get_max_sort_order()
                # The cost breakdown is snapshotted here (like price/purchase_price) so a later
                # edit to the catalog formula never changes an estimate that's already been issued.
                db.execute("""INSERT INTO estimate_items (estimate_id, item_type, name, unit, quantity, price, sum, sort_order, purchase_price,
                              labor_hours, hourly_rate, overhead_pct, profit_pct, other_costs) VALUES (?, 'Работа', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                           (self.estimate_id, work['name'], work['unit'], qty, work['price'], sm, so, work['purchase_price'],
                            work['labor_hours'], work['hourly_rate'], work['overhead_pct'], work['profit_pct'], work['other_costs']))
                self.load_items()

    def add_custom(self):
        dlg = CustomItemDialog(self)
        if dlg.exec() == QDialog.DialogCode.Accepted and dlg.result_data:
            d = dlg.result_data
            qty, ok = QInputDialog.getDouble(self, "Количество", f"Укажите объем ({d['unit']}):", 1.0, 0.001, 1000000.0, 3)
            if ok:
                sm = qty * d['price']
                so = self.get_max_sort_order()
                db.execute("INSERT INTO estimate_items (estimate_id, item_type, name, unit, quantity, price, sum, sort_order, purchase_price) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                           (self.estimate_id, d['type'], d['name'], d['unit'], qty, d['price'], sm, so, d['purchase_price']))
                self.load_items()

    def add_section(self):
        name, ok = QInputDialog.getText(self, "Новый раздел", "Название раздела:")
        if ok and name.strip():
            so = self.get_max_sort_order()
            db.execute("INSERT INTO estimate_items (estimate_id, item_type, name, unit, quantity, price, sum, sort_order, purchase_price) VALUES (?, 'Раздел', ?, '', 0, 0, 0, ?, 0.0)",
                       (self.estimate_id, name.strip(), so))
            self.load_items()

    def on_cell_changed(self, item):
        if self._updating: return
        row = item.row()
        col = item.column()
        if col in [0, 2, 3, 4]:
            try:
                item_id = self.table.item(row, 0).data(Qt.ItemDataRole.UserRole)
                itype = self.table.item(row, 1).text()
                if itype == "Раздел":
                    name = self.table.item(row, 0).text()
                    db.execute("UPDATE estimate_items SET name=? WHERE id=?", (name, item_id))
                else:
                    name = self.table.item(row, 0).text()
                    unit = self.table.item(row, 2).text()
                    qty_str = self.table.item(row, 3).text().replace(' ', '').replace(',', '.')
                    price_str = self.table.item(row, 4).text().replace(' ', '').replace(',', '.')

                    try:
                        qty = float(qty_str)
                        price = float(price_str)
                    except ValueError:
                        QMessageBox.warning(self, "Ошибка ввода", "Количество и Цена должны быть числами.")
                        self.load_items()
                        return

                    new_sum = qty * price
                    db.execute("UPDATE estimate_items SET name=?, unit=?, quantity=?, price=?, sum=? WHERE id=?", (name, unit, qty, price, new_sum, item_id))
                self.load_items()
            except Exception as e:
                logging.error(f"Ошибка при изменении ячейки: {e}")

    def del_item(self):
        curr = self.table.currentRow()
        if curr >= 0:
            item_id = self.table.item(curr, 0).data(Qt.ItemDataRole.UserRole)
            db.execute("DELETE FROM estimate_items WHERE id=?", (item_id,))
            self.load_items()

    def export_template(self):
        self.save_meta();self.sync_total()
        ExportDialog('estimates',self.estimate_id,self).exec()

    def sync_total(self):
        t = ed.totals(db, self.estimate_id)
        mat_adj, work_adj, tot_adj = float(t['mat_adj']), float(t['work_adj']), float(t['tot_adj'])
        self.base_mat_total, self.base_work_total, self.base_dir_total = float(t['base_mat']), float(t['base_work']), float(t['base'])
        self.adj_mat_total, self.adj_work_total = float(t['adj_mat']), float(t['adj_work'])
        self.adj_subtotal = float(t['subtotal'])
        self.grand_total = ed.recalc(db, self.estimate_id)
        if self.parent() and hasattr(self.parent(), 'load_data'):
            self.parent().load_data()

        t_text = f"Прямые затраты (базовые): {self.base_dir_total:,.2f} руб.<br>"
        m_str = f" (с учетом {mat_adj}%)" if mat_adj != 0 else ""
        t_text += f"<span style='color: #64748B;'>— Материалы{m_str}: {self.adj_mat_total:,.2f} руб.</span><br>"
        w_str = f" (с учетом {work_adj}%)" if work_adj != 0 else ""
        t_text += f"<span style='color: #64748B;'>— Работы{w_str}: {self.adj_work_total:,.2f} руб.</span><br>"
        if tot_adj != 0:
            t_text += f"<br>Промежуточный итог: {self.adj_subtotal:,.2f} руб.<br>"
            t_text += f"Скидка/Надбавка ({tot_adj}%): {(self.grand_total - self.adj_subtotal):,.2f} руб.<br>"

        self.totals_lbl.setText(t_text.replace(",", " "))
        self.lbl_grand_total.setText(f"ВСЕГО: {self.grand_total:,.2f} руб.".replace(",", " "))

    def export_estimate_pdf(self):
        safe_title = re.sub(r'[\\/*?:"<>|]', "_", self.title)
        path, _ = QFileDialog.getSaveFileName(self, "Экспорт сметы в PDF", f"Смета_{safe_title}.pdf", "PDF (*.pdf)")
        if path:
            meta = db.fetchone("SELECT date, coalesce(nullif(client_name,''),party_name), client_phone FROM estimates WHERE id=?", (self.estimate_id,))
            items = db.fetchall("SELECT item_type, name, unit, quantity, price, sum FROM estimate_items WHERE estimate_id=? ORDER BY sort_order ASC", (self.estimate_id,))

            accent = db.get_setting("accent_color", "#0284C7")
            font = db.get_setting("export_font", "Segoe UI")
            f_size = db.get_setting("export_font_size", "13")
            company = db.get_setting("export_company_name", 'ООО "ГазМонтаж"')

            mat_adj = self.meta_cache.get('mat_adj', 0.0)
            work_adj = self.meta_cache.get('work_adj', 0.0)
            tot_adj = self.meta_cache.get('tot_adj', 0.0)

            rows_html = ""
            for itype, name, unit, qty, price, sm in items:
                if itype == "Раздел":
                    rows_html += f"<tr><td colspan='6' style='background-color: #e2e8f0; color: black; font-weight: bold; padding: 8px;'>{name}</td></tr>"
                else:
                    q_v = float(qty) if qty else 0.0
                    p_v = float(price) if price else 0.0
                    s_v = float(sm) if sm else 0.0
                    rows_html += f"<tr><td style='color: black;'>{name}</td><td style='color: black;'>{itype}</td><td align='center' style='color: black;'>{unit}</td><td align='right' style='color: black;'>{q_v:.3f}</td><td align='right' style='color: black;'>{p_v:,.2f}</td><td align='right' style='color: black;'>{s_v:,.2f}</td></tr>"

            rows_html += f"<tr style='font-weight: bold;'><td colspan='5' align='right'>Прямые затраты (базовые):</td><td align='right'>{self.base_dir_total:,.2f}</td></tr>"

            if mat_adj != 0:
                rows_html += f"<tr><td colspan='5' align='right' style='color: #64748b;'>В т.ч. материалы (база):</td><td align='right' style='color: #64748b;'>{self.base_mat_total:,.2f}</td></tr>"
                rows_html += f"<tr><td colspan='5' align='right'>Скидка/Надбавка на материалы ({mat_adj}%):</td><td align='right'>{(self.adj_mat_total - self.base_mat_total):,.2f}</td></tr>"

            if work_adj != 0:
                rows_html += f"<tr><td colspan='5' align='right' style='color: #64748b;'>В т.ч. работы (база):</td><td align='right' style='color: #64748b;'>{self.base_work_total:,.2f}</td></tr>"
                rows_html += f"<tr><td colspan='5' align='right'>Скидка/Надбавка на работы ({work_adj}%):</td><td align='right'>{(self.adj_work_total - self.base_work_total):,.2f}</td></tr>"

            if tot_adj != 0:
                rows_html += f"<tr style='font-weight: bold;'><td colspan='5' align='right'>Промежуточный итог:</td><td align='right'>{self.adj_subtotal:,.2f}</td></tr>"
                rows_html += f"<tr><td colspan='5' align='right'>Итоговая скидка/надбавка ({tot_adj}%):</td><td align='right'>{(self.grand_total - self.adj_subtotal):,.2f}</td></tr>"

            rows_html += f"<tr style='font-weight: bold; font-size: {int(f_size)+1}px; background-color: #e2e8f0;'><td colspan='5' align='right'>ВСЕГО ПО СМЕТЕ:</td><td align='right'>{self.grand_total:,.2f}</td></tr>"

            html = f"""
            <html><head><meta charset="utf-8"></head>
            <body style="font-family: '{font}', Arial, sans-serif; font-size: {f_size}px; padding: 20px; color: black;">
                <div style="text-align: right; color: #64748b; font-size: {int(f_size)-2}px; margin-bottom: 10px;">{company}</div>
                <h1 style="color: {accent}; border-bottom: 2px solid {accent}; padding-bottom: 5px;">СМЕТА НА ОБЪЕКТ: {self.title}</h1>
                <p><b>Дата:</b> {meta[0]} | <b>Заказчик:</b> {meta[1] or '—'} | <b>Телефон:</b> {meta[2] or '—'}</p>
                <table border="1" cellpadding="6" cellspacing="0" style="border-collapse: collapse; width: 100%; border: 1px solid #cbd5e1; font-size: {f_size}px; margin-top: 15px;">
                    <tr style="background-color: {accent}; color: white;">
                        <th>Наименование</th><th>Тип</th><th>Ед.</th><th>Кол-во</th><th>Цена (руб)</th><th>Сумма (руб)</th>
                    </tr>
                    {rows_html}
                </table>
            </body></html>
            """
            printer = QPrinter(QPrinter.PrinterMode.HighResolution)
            printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
            printer.setOutputFileName(path)
            doc_p = QTextDocument()
            doc_p.setHtml(html)
            doc_p.print(printer)
            QMessageBox.information(self, "Успех", "Смета успешно сохранена в PDF.")

    def export_excel(self, use_template=False):
        safe_title = re.sub(r'[\\/*?:"<>|]', "_", self.title)
        if not use_template:
            save_path, _ = QFileDialog.getSaveFileName(self, "Экспорт сметы", f"Смета_{safe_title}.xlsx", "Excel (*.xlsx)")
            if not save_path: return
            try:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Смета"

                font_bold = Font(name="Segoe UI", size=11, bold=True)
                border_thin = Border(
                    left=Side(style='thin', color='CBD5E1'),
                    right=Side(style='thin', color='CBD5E1'),
                    top=Side(style='thin', color='CBD5E1'),
                    bottom=Side(style='thin', color='CBD5E1')
                )

                ws.append([f"СМЕТА НА ОБЪЕКТ: {self.title}"])
                ws.append([])
                headers = ["Наименование позиции", "Тип", "Ед.", "Кол-во", "Цена (руб)", "Сумма (руб)"]
                ws.append(headers)

                items = db.fetchall("SELECT item_type, name, unit, quantity, price, sum FROM estimate_items WHERE estimate_id=? ORDER BY sort_order ASC", (self.estimate_id,))
                for itype, name, unit, qty, price, sm in items:
                    if itype == "Раздел":
                        ws.append([name, "Раздел", "", "", "", ""])
                    else:
                        ws.append([name, itype, unit, qty, price, sm])

                ws.append([])
                ws.append(["", "", "", "", "Прямые затраты (базовые):", self.base_dir_total])

                mat_adj = self.meta_cache.get('mat_adj', 0.0)
                if mat_adj != 0:
                    ws.append(["", "", "", "", "В т.ч. материалы (база):", self.base_mat_total])
                    ws.append(["", "", "", "", f"Скидка/Надбавка на материалы ({mat_adj}%):", self.adj_mat_total - self.base_mat_total])

                work_adj = self.meta_cache.get('work_adj', 0.0)
                if work_adj != 0:
                    ws.append(["", "", "", "", "В т.ч. работы (база):", self.base_work_total])
                    ws.append(["", "", "", "", f"Скидка/Надбавка на работы ({work_adj}%):", self.adj_work_total - self.base_work_total])

                tot_adj = self.meta_cache.get('tot_adj', 0.0)
                if tot_adj != 0:
                    ws.append(["", "", "", "", "Промежуточный итог:", self.adj_subtotal])
                    ws.append(["", "", "", "", f"Итоговая скидка/надбавка ({tot_adj}%):", self.grand_total - self.adj_subtotal])

                ws.append(["", "", "", "", "ВСЕГО ПО СМЕТЕ:", self.grand_total])
                ws.cell(row=ws.max_row, column=5).font = font_bold
                ws.cell(row=ws.max_row, column=6).font = font_bold

                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

                for sheet in wb.worksheets:

                    for row_cells in sheet:

                        for cell in row_cells:

                            if isinstance(cell.value, str): cell.data_type = "s"

                wb.save(save_path)
                QMessageBox.information(self, "Успех", "Смета выгружена в Excel.")
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить файл:\n{e}")
        else:
            QMessageBox.warning(self, "Внимание", "Экспорт по шаблону: кнопка «Документ по шаблону» (теги сметы настраиваются в разделе «Сметы → Шаблоны и теги»).")

    def build_dossier_html(self):
        meta = db.fetchone("SELECT client_name, client_phone, statuses, total, paid FROM estimates WHERE id=?", (self.estimate_id,))
        payments = db.fetchall("SELECT date, amount FROM payments WHERE estimate_id=? ORDER BY date DESC", (self.estimate_id,))
        c_name, c_phone, statuses, total, paid = meta
        total = total or 0.0
        paid = paid or 0.0
        debt = total - paid
        return c_name, c_phone, statuses, total, paid, debt, payments

    def export_dossier_pdf(self):
        safe_title = re.sub(r'[\\/*?:"<>|]', "_", self.title)
        path, _ = QFileDialog.getSaveFileName(self, "Экспорт карточки", f"Карточка_Клиента_{safe_title}.pdf", "PDF (*.pdf)")
        if path:
            c_name, c_phone, statuses, total, paid, debt, payments = self.build_dossier_html()
            accent = db.get_setting("accent_color", "#0284C7")
            font = db.get_setting("export_font", "Segoe UI")
            f_size = db.get_setting("export_font_size", "13")
            company = db.get_setting("export_company_name", 'ООО "ГазМонтаж"')

            html = f"""
            <html><head><meta charset="utf-8"></head>
            <body style="font-family: '{font}', Arial, sans-serif; font-size: {f_size}px; padding: 30px; color: black;">
                <div style="text-align: right; color: #64748b; font-size: {int(f_size)-2}px; margin-bottom: 5px;">{company}</div>
                <div style="border-bottom: 3px solid {accent}; padding-bottom: 10px; margin-bottom: 20px;">
                    <h1 style="color: {accent}; margin: 0;">КАРТОЧКА КЛИЕНТА</h1>
                    <p style="color: #64748b; margin: 5px 0 0 0; font-size: {f_size}px;">Объект / Шифр: <b>{self.title}</b></p>
                </div>
                
                <table width="100%" style="font-size: {f_size}px; margin-bottom: 25px; background: #f8fafc; padding: 15px; border-radius: 8px;">
                    <tr><td width="25%"><b>Заказчик:</b></td><td>{c_name or '—'}</td></tr>
                    <tr><td style="padding-top: 8px;"><b>Телефон:</b></td><td style="padding-top: 8px;">{c_phone or '—'}</td></tr>
                    <tr><td style="padding-top: 8px;"><b>Статусы:</b></td><td style="padding-top: 8px;">{statuses or '—'}</td></tr>
                </table>
                
                <h3 style="color: #0f172a; border-left: 4px solid {accent}; padding-left: 10px;">Финансовая сводка</h3>
                <table width="100%" style="font-size: {int(f_size)+1}px; margin-bottom: 30px; background: #f1f5f9; padding: 15px; border-radius: 8px;">
                    <tr><td width="30%"><b>Сумма сметы:</b></td><td><b>{total:,.2f} руб.</b></td></tr>
                    <tr><td style="padding-top: 8px;"><b>Оплачено:</b></td><td style="color: #16a34a; padding-top: 8px;"><b>{paid:,.2f} руб.</b></td></tr>
                    <tr><td style="padding-top: 8px;"><b>Остаток долга:</b></td><td style="color: #dc2626; padding-top: 8px;"><b>{debt:,.2f} руб.</b></td></tr>
                </table>

                <h3 style="color: #0f172a; border-left: 4px solid {accent}; padding-left: 10px;">История поступления платежей</h3>
                <table border="1" cellpadding="10" cellspacing="0" style="border-collapse: collapse; width: 100%; border: 1px solid #cbd5e1; font-size: {f_size}px;">
                    <tr style="background-color: {accent}; color: white; text-align: left;">
                        <th width="40%">Дата платежа</th><th>Внесенная сумма (руб)</th>
                    </tr>
            """
            if payments:
                for dt, amt in payments:
                    v = float(amt) if amt is not None else 0.0
                    html += f"<tr><td>{dt}</td><td>{v:,.2f} руб.</td></tr>"
            else:
                html += "<tr><td colspan='2' style='text-align: center; color: #64748b; padding: 15px;'>Платежи по данному объекту пока не вносились</td></tr>"

            html += """
                </table>
            </body></html>
            """
            printer = QPrinter(QPrinter.PrinterMode.HighResolution)
            printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
            printer.setOutputFileName(path)
            doc_pdf = QTextDocument()
            doc_pdf.setHtml(html)
            doc_pdf.print(printer)
            QMessageBox.information(self, "Успех", "Карточка клиента сохранена в PDF.")

    def export_dossier_word(self):
        safe_title = re.sub(r'[\\/*?:"<>|]', "_", self.title)
        path, _ = QFileDialog.getSaveFileName(self, "Экспорт карточки", f"Карточка_Клиента_{safe_title}.docx", "Word Document (*.docx)")
        if path:
            c_name, c_phone, statuses, total, paid, debt, payments = self.build_dossier_html()
            doc = docx.Document()

            title = doc.add_heading("КАРТОЧКА КЛИЕНТА", 0)
            doc.add_paragraph(f"Объект / Шифр: {self.title}")

            table_info = doc.add_table(rows=3, cols=2)
            table_info.style = 'Table Grid'
            table_info.rows[0].cells[0].text = "Заказчик:"
            table_info.rows[0].cells[1].text = str(c_name or '—')
            table_info.rows[1].cells[0].text = "Телефон:"
            table_info.rows[1].cells[1].text = str(c_phone or '—')
            table_info.rows[2].cells[0].text = "Статусы:"
            table_info.rows[2].cells[1].text = str(statuses or '—')

            doc.add_heading("Финансовая сводка", level=2)
            table_fin = doc.add_table(rows=3, cols=2)
            table_fin.style = 'Table Grid'
            table_fin.rows[0].cells[0].text = "Сумма сметы:"
            table_fin.rows[0].cells[1].text = f"{total:,.2f} руб."
            table_fin.rows[1].cells[0].text = "Оплачено:"
            table_fin.rows[1].cells[1].text = f"{paid:,.2f} руб."
            table_fin.rows[2].cells[0].text = "Остаток долга:"
            table_fin.rows[2].cells[1].text = f"{debt:,.2f} руб."

            doc.add_heading("История поступления платежей", level=2)
            if payments:
                table_pay = doc.add_table(rows=1, cols=2)
                table_pay.style = 'Table Grid'
                hdr_cells = table_pay.rows[0].cells
                hdr_cells[0].text = 'Дата платежа'
                hdr_cells[1].text = 'Внесенная сумма (руб)'
                for dt, amt in payments:
                    row_cells = table_pay.add_row().cells
                    row_cells[0].text = str(dt)
                    row_cells[1].text = f"{float(amt or 0):,.2f} руб."
            else:
                doc.add_paragraph("Платежи по данному объекту пока не вносились.")

            doc.save(path)
            QMessageBox.information(self, "Успех", "Карточка клиента сохранена в Word.")

    def export_dossier_excel(self):
        safe_title = re.sub(r'[\\/*?:"<>|]', "_", self.title)
        path, _ = QFileDialog.getSaveFileName(self, "Экспорт карточки", f"Карточка_Клиента_{safe_title}.xlsx", "Excel (*.xlsx)")
        if path:
            c_name, c_phone, statuses, total, paid, debt, payments = self.build_dossier_html()
            df_info = pd.DataFrame([
                ["Шифр / Объект", self.title], ["Клиент", c_name], ["Телефон", c_phone],
                ["Статусы", statuses], ["Общая сумма", total], ["Оплачено", paid], ["Долг", debt]
            ], columns=["Показатель", "Значение"])
            df_pay = pd.DataFrame(payments, columns=["Дата", "Сумма (руб)"])
            with pd.ExcelWriter(path) as writer:
                df_info.to_excel(writer, sheet_name="Карточка клиента", index=False)
                df_pay.to_excel(writer, sheet_name="История платежей", index=False)
            QMessageBox.information(self, "Успех", "Карточка сохранена в Excel.")
