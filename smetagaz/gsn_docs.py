"""Раздел «Монтаж ГСН»: теги, редактор тегов и формирование 14 документов Word.

Теги этого раздела хранятся в таблице `gsn_tags` и не пересекаются с тегами других модулей. Каждый тег — имя, источник данных (поле карточки, трубопровод,
материал, сертификаты, аттестаты), формат и текст «если пусто». У каждой вводимой в карточке величины есть свой тег."""
import hashlib
import json
import os
import re
from datetime import datetime

from . import gsv_project_domain as gd
from . import gsn_domain as gn
from .gsv_project_domain import get_initials_first
from .gsvm_docs import apply_format, money, norm_tag, tag_text, FORMATS  # noqa: F401  (общий разбор формата и имени тега)

TEMPLATES_DIR = gd.TEMPLATES_DIR / 'gsn'

# вид документа -> (название, расширение, группа: 'main' — вкладка «Договор и клиент», 'id' — исполнительная документация, префикс файла)
DOC_KINDS = {
    'contract': ('Договор ГСН', 'docx', 'main', 'Договор_ГСН'),
    'title': ('Титульник ГСН', 'docx', 'id', 'Титульник_ГСН'),
    'registry': ('Реестр исполнительной документации', 'docx', 'id', 'Реестр_ИД'),
    'warranty': ('Гарантийный паспорт объекта', 'docx', 'id', 'Гарантийный_паспорт'),
    'payment_cert': ('Справка об оплате выполненных работ', 'docx', 'main', 'Справка_об_оплате'),
    'accept': ('Акт приемки объекта в эксплуатацию', 'docx', 'id', 'Акт_приемки'),
    'building_passport': ('Строительный паспорт объекта', 'docx', 'id', 'Строительный_паспорт'),
    'hidden_works': ('Акт на скрытые работы', 'docx', 'id', 'Акт_скрытых_работ'),
    'device_revision': ('Акт ревизии технических устройств', 'docx', 'id', 'Акт_ревизии_устройств'),
    'vkk_input': ('Протокол ВКК (ввод)', 'docx', 'id', 'Протокол_ВКК_ввод'),
    'vkk_coupling': ('Протокол ВКК (муфта)', 'docx', 'id', 'Протокол_ВКК_муфта'),
    'vkk_pipe': ('Протокол ВКК (труба)', 'docx', 'id', 'Протокол_ВКК_труба'),
    'vkk_joint': ('Протокол ВКК (неразъемное соединение)', 'docx', 'id', 'Протокол_ВКК_соединение'),
    'materials_transfer': ('Акт передачи материалов', 'docx', 'id', 'Акт_передачи_материалов'),
}
ID_KINDS = [k for k, v in DOC_KINDS.items() if v[2] == 'id']
MAIN_KINDS = [k for k, v in DOC_KINDS.items() if v[2] == 'main']


# --- ИСТОЧНИКИ И ТЕГИ ------------------------------------------------------------------------

STATIC_SOURCES = [
    ('contract.number', 'Номер договора'), ('contract.date', 'Дата заключения договора'), ('contract.signed', 'Договор подписан (да / нет)'),
    ('client.name', 'Клиент: ФИО полностью'), ('client.short', 'Клиент: инициалы и фамилия'), ('client.phone', 'Клиент: телефоны'), ('client.address', 'Клиент: адрес'),
    ('client.passport', 'Клиент: паспорт (серия, номер)'), ('client.passport_issuer', 'Клиент: кем выдан паспорт'), ('client.passport_date', 'Клиент: дата выдачи паспорта'),
    ('object.name', 'Наименование объекта (по проекту и ТУ)'), ('object.address', 'Адрес объекта'), ('notes', 'Примечания по объекту'),
    ('project.code', 'Шифр проекта'), ('project.designer', 'Проектировщик'), ('project.date', 'Дата проекта (месяц и год)'), ('project.tu', 'Технические условия (ТУ)'),
    ('work.start', 'Начало работ'), ('work.end', 'Окончание работ'), ('work.act_date', 'Дата акта'), ('work.act_signed', 'Акт подписан (да / нет)'),
    ('work.cost', 'Стоимость работ'), ('estimate.materials', 'Смета: стоимость материалов'), ('estimate.works', 'Смета: стоимость работ'),
    ('payments.dates', 'Оплаты: даты (по строкам)'), ('payments.amounts', 'Оплаты: суммы (по строкам)'), ('payments.notes', 'Оплаты: примечания (по строкам)'),
    ('payments.paid', 'Оплачено всего'), ('payments.rest', 'Остаток к оплате'), ('payments.last_date', 'Дата последней оплаты'),
    ('gas.pe_length', 'Протяжённость ПЭ трубопровода, м'), ('gas.steel_length', 'Протяжённость стального трубопровода, м'), ('gas.total_length', 'Общая протяжённость, м'),
    ('gas.pressure', 'Тип газопровода (низкое / среднее давление)'), ('gas.pressure_text', 'Газопровод низкого / среднего давления (в родительном падеже)'),
    ('gas.tie_in', 'Тип врезки (в сталь / в ПЭ)'), ('gas.tie_in_text', 'Врезка: «в стальной газопровод» / «в полиэтиленовый газопровод»'),
    ('gas.couplings', 'Количество муфт'), ('gas.pipes_text', 'Трубопроводы: ПЭ и сталь с протяжённостью'),
    ('materials.list', 'Материалы объекта (по строкам: название и количество)'), ('materials.names', 'Материалы объекта: названия (по строкам)'),
    ('certs.list', 'Сертификаты (по строкам: название, номер)'), ('certs.numbers', 'Сертификаты: номера (по строкам)'), ('certs.names', 'Сертификаты: названия (по строкам)'),
    ('attest.list', 'Аттестаты сварщиков (по строкам: документ, сварщик, номер)'), ('attest.welders', 'Сварщики (по строкам)'), ('attest.numbers', 'Аттестаты: номера (по строкам)'),
]
BUILTIN_TAGS = [
    ('НОМЕР_ДОГОВОРА', 'contract.number', 'text'), ('ДАТА_ДОГОВОРА', 'contract.date', 'date_long'), ('ДАТА_ДОГОВОРА_С', 'contract.date', 'date_short'),
    ('ДАТА_ДОГОВОРА_КАВ', 'contract.date', 'date_quoted'), ('ДОГОВОР_ПОДПИСАН', 'contract.signed', 'text'),
    ('КЛИЕНТ', 'client.name', 'text'), ('КЛИЕНТ_СОКР', 'client.short', 'text'), ('ТЕЛЕФОН', 'client.phone', 'text'), ('АДРЕС_КЛИЕНТА', 'client.address', 'text'),
    ('ПАСПОРТ', 'client.passport', 'text'), ('КЕМ_ВЫДАН', 'client.passport_issuer', 'text'), ('ДАТА_ВЫДАЧИ', 'client.passport_date', 'date_short'),
    ('ОБЪЕКТ', 'object.name', 'text'), ('АДРЕС', 'object.address', 'text'), ('ПРИМЕЧАНИЕ', 'notes', 'text'),
    ('ШИФР_ПРОЕКТА', 'project.code', 'text'), ('НОМЕР_ПРОЕКТА', 'project.code', 'text'), ('ПРОЕКТИРОВЩИК', 'project.designer', 'text'),
    ('ДАТА_ПРОЕКТА', 'project.date', 'month_year'), ('ТУ', 'project.tu', 'text'),
    ('НАЧАЛО_РАБОТ', 'work.start', 'date_long'), ('НАЧАЛО_РАБОТ_С', 'work.start', 'date_short'), ('ОКОНЧАНИЕ_РАБОТ', 'work.end', 'date_long'), ('ОКОНЧАНИЕ_РАБОТ_С', 'work.end', 'date_short'),
    ('ДАТА_АКТА', 'work.act_date', 'date_quoted'), ('ДАТА_АКТА_С', 'work.act_date', 'date_short'), ('ДАТА_АКТА_ДЛИННАЯ', 'work.act_date', 'date_long'),
    ('МЕСЯЦ_РАБОТ', 'work.act_date', 'month_name'), ('МЕСЯЦ_РАБОТ_ГОД', 'work.act_date', 'month_year'), ('АКТ_ПОДПИСАН', 'work.act_signed', 'text'),
    ('СТОИМОСТЬ', 'work.cost', 'money'), ('СТОИМОСТЬ_ПРОПИСЬЮ', 'work.cost', 'money_words'),
    ('СТОИМОСТЬ_МАТЕРИАЛОВ', 'estimate.materials', 'money'), ('СТОИМОСТЬ_РАБОТ_СМЕТА', 'estimate.works', 'money'),
    ('ОПЛАЧЕНО', 'payments.paid', 'money'), ('ОПЛАЧЕНО_ПРОПИСЬЮ', 'payments.paid', 'money_words'), ('ОСТАТОК', 'payments.rest', 'money'),
    ('ДАТА_ОПЛАТЫ', 'payments.dates', 'text'), ('СУММА_ОПЛАТЫ', 'payments.amounts', 'text'), ('ПРИМЕЧАНИЕ_ОПЛАТЫ', 'payments.notes', 'text'),
    ('ДАТА_ПОСЛЕДНЕЙ_ОПЛАТЫ', 'payments.last_date', 'date_short'),
    ('ДЛИНА_ПЭ', 'gas.pe_length', 'number2'), ('ДЛИНА_СТАЛЬ', 'gas.steel_length', 'number2'), ('ДЛИНА_ВСЕГО', 'gas.total_length', 'number2'),
    ('ДАВЛЕНИЕ', 'gas.pressure', 'text'), ('ДАВЛЕНИЕ_ТЕКСТ', 'gas.pressure_text', 'text'), ('ВРЕЗКА', 'gas.tie_in', 'text'), ('ВРЕЗКА_ТЕКСТ', 'gas.tie_in_text', 'text'),
    ('МУФТЫ_КОЛ', 'gas.couplings', 'text'), ('ТРУБОПРОВОДЫ', 'gas.pipes_text', 'text'),
    ('МАТЕРИАЛЫ', 'materials.list', 'text'), ('МАТЕРИАЛЫ_НАЗВАНИЯ', 'materials.names', 'text'),
    ('СЕРТИФИКАТЫ', 'certs.list', 'text'), ('СЕРТИФИКАТЫ_НОМЕРА', 'certs.numbers', 'text'), ('СЕРТИФИКАТЫ_НАЗВАНИЯ', 'certs.names', 'text'),
    ('АТТЕСТАТЫ', 'attest.list', 'text'), ('СВАРЩИКИ', 'attest.welders', 'text'), ('АТТЕСТАТЫ_НОМЕРА', 'attest.numbers', 'text'),
]
ROLE_TAGS = {'pe_pipe': 'ПЭ_ТРУБА', 'steel_pipe': 'СТАЛЬ_ТРУБА', 'transition': 'НЕРАЗЪЕМНОЕ_СОЕДИНЕНИЕ', 'plug': 'ЗАГЛУШКА', 'coupling': 'МУФТА'}


def available_sources(db):
    result = list(STATIC_SOURCES)
    for role, label in gn.ROLES.items():
        result += [(f'role:{role}:name', f'«{label}»: позиция справочника'), (f'role:{role}:qty', f'«{label}»: количество'), (f'role:{role}:unit', f'«{label}»: единица'),
                   (f'role:{role}:cert', f'«{label}»: номера сертификатов'), (f'role:{role}:cert_names', f'«{label}»: названия и номера сертификатов')]
    return result


def sync_tags(db):
    """Создаёт недостающие стандартные теги и теги ролей материалов. Переименованные и отключённые пользователем теги не затрагиваются."""
    def add(auto_key, name, source, fmt, note=''):
        if db.fetchone('SELECT 1 FROM gsn_tags WHERE auto_key=?', (auto_key,)):
            return
        base, n = name, 1
        while db.fetchone('SELECT 1 FROM gsn_tags WHERE name=?', (name,)):
            n += 1
            name = f'{base}_{n}'
        db.execute('INSERT INTO gsn_tags(name,source,fmt,auto_key,note) VALUES(?,?,?,?,?)', (name, source, fmt, auto_key, note))
    for name, source, fmt in BUILTIN_TAGS:
        add(f'b:{name}', name, source, fmt)
    for role, base in ROLE_TAGS.items():
        add(f'r:{role}:name', base, f'role:{role}:name', 'text', 'Пусто, если этого материала в объекте нет')
        add(f'r:{role}:qty', f'{base}_КОЛ', f'role:{role}:qty', 'text', 'Количество; пусто, если материала в объекте нет')
        add(f'r:{role}:unit', f'{base}_ЕД', f'role:{role}:unit', 'text')
        add(f'r:{role}:cert', f'{base}_СЕРТ', f'role:{role}:cert', 'text', 'Номера сертификатов позиции')
        add(f'r:{role}:cert_names', f'{base}_СЕРТ_ПОЛНО', f'role:{role}:cert_names', 'text', 'Название и номер сертификатов позиции')


def list_tags(db):
    sync_tags(db)
    return db.fetchall("SELECT id,name,source,fmt,empty_text,coalesce(auto_key,''),enabled,note FROM gsn_tags ORDER BY (auto_key IS NULL), id")


def save_tag(db, tag_id, name, source, fmt, empty_text='', enabled=True, note=''):
    name = norm_tag(name)
    if not re.fullmatch(r'[0-9A-ZА-ЯЁ_]+', name):
        raise ValueError('Имя тега: буквы, цифры и подчёркивание (пробелы заменяются на _)')
    if fmt not in FORMATS:
        raise ValueError('Неизвестный формат')
    if not source:
        raise ValueError('Выберите источник данных')
    if db.fetchone('SELECT id FROM gsn_tags WHERE name=? AND id IS NOT ?', (name, tag_id)):
        raise ValueError(f'Тег {{{name}}} уже существует')
    if tag_id:
        db.execute('UPDATE gsn_tags SET name=?,source=?,fmt=?,empty_text=?,enabled=?,note=? WHERE id=?', (name, source, fmt, empty_text, int(bool(enabled)), note, tag_id))
        return tag_id
    return db.execute('INSERT INTO gsn_tags(name,source,fmt,empty_text,enabled,note) VALUES(?,?,?,?,?,?)', (name, source, fmt, empty_text, int(bool(enabled)), note)).lastrowid


def delete_tag(db, tag_id):
    """Собственные теги удаляются; стандартные отключаются (чтобы не появиться заново)."""
    row = db.fetchone('SELECT auto_key FROM gsn_tags WHERE id=?', (tag_id,))
    if not row:
        return 'missing'
    if row[0]:
        db.execute('UPDATE gsn_tags SET enabled=0 WHERE id=?', (tag_id,))
        return 'disabled'
    db.execute('DELETE FROM gsn_tags WHERE id=?', (tag_id,))
    return 'deleted'


# --- ЗНАЧЕНИЯ --------------------------------------------------------------------------------

def _num(value):
    value = float(value or 0)
    return f'{value:g}'.replace('.', ',')


def payments(db, cid):
    from . import payments_domain
    return sorted(((d, float(a), n or '') for _i, d, a, n in payments_domain.history(db, 'gsn_projects', cid)), key=lambda p: p[0])


def context(db, cid):
    """Исходные значения источников по договору (даты — ГГГГ-ММ-ДД, суммы — числа)."""
    c = gn.get(db, cid)
    client = {}
    if c.get('client_id'):
        from .gsv_domain import get_client
        client = get_client(db, c['client_id']) or {}
    name = client.get('name') or c.get('client_name') or ''
    pays = payments(db, cid)
    paid = sum(a for _d, a, _n in pays)
    cost = float(c.get('contract_amount') or 0)
    pe, steel = float(c.get('pe_length') or 0), float(c.get('steel_length') or 0)
    mats = gn.materials(db, cid, c)
    certs = gn.certificates(db, cid, c)
    atts = gn.attestations(db, cid, c)
    pipes = []
    if pe:
        pipes.append(f'полиэтиленовый {gn.PE_DIAMETER} — {_num(pe)} м')
    if steel:
        pipes.append(f'стальной {gn.STEEL_DIAMETER} — {_num(steel)} м')
    ctx = {
        'contract.number': c.get('contract_number') or '', 'contract.date': c.get('contract_date') or '', 'contract.signed': 'да' if c.get('contract_signed') else 'нет',
        'client.name': name, 'client.short': get_initials_first(name) if name else '', 'client.phone': client.get('phone') or c.get('phone') or '',
        'client.address': client.get('address') or c.get('client_address') or '', 'client.passport': client.get('passport') or c.get('passport') or '',
        'client.passport_issuer': client.get('passport_issuer') or c.get('passport_issuer') or '', 'client.passport_date': client.get('passport_date') or c.get('passport_issue_date') or '',
        'object.name': c.get('title') or '', 'object.address': c.get('address') or '', 'notes': c.get('notes') or '',
        'project.code': c.get('designer_code') or '', 'project.designer': c.get('designer') or '', 'project.date': c.get('project_month') or '', 'project.tu': c.get('tu_text') or '',
        'work.start': c.get('work_start_date') or c.get('contract_date') or '', 'work.end': c.get('work_end_date') or '', 'work.act_date': c.get('act_date') or '',
        'work.act_signed': 'да' if c.get('act_signed') else 'нет', 'work.cost': cost, 'estimate.materials': c.get('est_materials') or 0, 'estimate.works': c.get('est_works') or 0,
        'payments.dates': '\n'.join(gd.date_short(d).rstrip('г.') for d, _a, _n in pays), 'payments.amounts': '\n'.join(money(a) for _d, a, _n in pays),
        'payments.notes': '\n'.join(n for _d, _a, n in pays), 'payments.paid': paid, 'payments.rest': cost - paid, 'payments.last_date': pays[-1][0] if pays else '',
        'gas.pe_length': pe or '', 'gas.steel_length': steel or '', 'gas.total_length': (pe + steel) or '', 'gas.pressure': gn.PRESSURES.get(c.get('pressure_type'), ''),
        'gas.pressure_text': gn.PRESSURE_TEXT.get(c.get('pressure_type'), ''), 'gas.tie_in': gn.TIE_INS.get(c.get('tie_in'), '').lower(), 'gas.tie_in_text': gn.TIE_TEXT.get(c.get('tie_in'), ''),
        'gas.couplings': str(int(c.get('couplings') or 0)) if c.get('tie_in') else '', 'gas.pipes_text': '; '.join(pipes),
        'materials.list': '\n'.join(f'{m["name"]} — {_num(m["quantity"])} {m["unit"]}'.strip() for m in mats), 'materials.names': '\n'.join(m['name'] for m in mats),
        'certs.list': '\n'.join(f'{x["name"]}' + (f', № {x["number"]}' if x['number'] else '') for x in certs), 'certs.numbers': '\n'.join(x['number'] for x in certs if x['number']),
        'certs.names': '\n'.join(x['name'] for x in certs),
        'attest.list': '\n'.join(' · '.join(p for p in (a['title'], a['welder'], ('№ ' + a['number']) if a['number'] else '') if p) for a in atts),
        'attest.welders': '\n'.join(dict.fromkeys(a['welder'] for a in atts if a['welder'])), 'attest.numbers': '\n'.join(a['number'] for a in atts if a['number']),
    }
    by_role = {m['role']: m for m in mats}
    for role in gn.ROLES:
        m = by_role.get(role)
        item_certs = []
        if m and m['item_id']:
            item_certs = db.fetchall('SELECT c.name,c.number FROM gsn_item_certs l JOIN gsn_certificates c ON c.id=l.cert_id WHERE l.item_id=? ORDER BY c.id', (m['item_id'],))
        ctx[f'role:{role}:name'] = m['name'] if m else ''
        ctx[f'role:{role}:qty'] = _num(m['quantity']) if m else ''
        ctx[f'role:{role}:unit'] = m['unit'] if m else ''
        ctx[f'role:{role}:cert'] = '; '.join(n for _t, n in item_certs if n)
        ctx[f'role:{role}:cert_names'] = '; '.join(t + (f', № {n}' if n else '') for t, n in item_certs)
    return ctx


def tag_map(db, cid):
    """{ИМЯ_ТЕГА: значение} — только включённые теги; пустые значения заменяются текстом «если пусто»."""
    sync_tags(db)
    ctx = context(db, cid)
    result = {}
    for _id, name, source, fmt, empty_text, _auto, enabled, _note in db.fetchall("SELECT id,name,source,fmt,empty_text,coalesce(auto_key,''),enabled,note FROM gsn_tags"):
        if not enabled:
            continue
        text = apply_format(ctx.get(source, ''), fmt)
        result[name] = text if text != '' else (empty_text or '')
    return result


def data_hash(db, cid):
    return hashlib.sha256(json.dumps(tag_map(db, cid), ensure_ascii=False, sort_keys=True).encode('utf-8')).hexdigest()


# --- ШАБЛОНЫ И ДОКУМЕНТЫ ---------------------------------------------------------------------

def template_setting(kind):
    return f'gsn_tpl_{kind}'


def template_path(db, kind):
    path = db.get_setting(template_setting(kind), '')
    return path if path and os.path.exists(path) else str(create_default_template(kind))


_HEAD = ['Объект: {ОБЪЕКТ}', 'Адрес объекта: {АДРЕС}', 'Заказчик: {КЛИЕНТ}', 'Договор № {НОМЕР_ДОГОВОРА} от {ДАТА_ДОГОВОРА_С}']
_GAS = ['Газопровод {ДАВЛЕНИЕ_ТЕКСТ}. Трубопроводы: {ТРУБОПРОВОДЫ}. Врезка: {ВРЕЗКА_ТЕКСТ}.']
DEFAULT_BODIES = {
    'contract': ['ДОГОВОР № {НОМЕР_ДОГОВОРА}', 'на монтаж газопровода-ввода', '{ДАТА_ДОГОВОРА}',
                 'Заказчик: {КЛИЕНТ}, паспорт {ПАСПОРТ}, выдан {КЕМ_ВЫДАН} {ДАТА_ВЫДАЧИ}, адрес: {АДРЕС_КЛИЕНТА}, тел.: {ТЕЛЕФОН}.',
                 'Объект: {ОБЪЕКТ}, адрес: {АДРЕС}.', 'Проект: {ШИФР_ПРОЕКТА}, проектировщик {ПРОЕКТИРОВЩИК}, {ДАТА_ПРОЕКТА}. Технические условия: {ТУ}.',
                 *_GAS, 'Начало работ: {НАЧАЛО_РАБОТ}. Окончание работ: {ОКОНЧАНИЕ_РАБОТ}.', 'Стоимость работ: {СТОИМОСТЬ} BYN ({СТОИМОСТЬ_ПРОПИСЬЮ}).',
                 'Примечание: {ПРИМЕЧАНИЕ}', 'Заказчик: ______________ {КЛИЕНТ_СОКР}'],
    'title': ['ИСПОЛНИТЕЛЬНАЯ ДОКУМЕНТАЦИЯ', 'на монтаж газопровода-ввода', *_HEAD, 'Проект: {ШИФР_ПРОЕКТА}, {ПРОЕКТИРОВЩИК}, {ДАТА_ПРОЕКТА}', 'Технические условия: {ТУ}', *_GAS],
    'registry': ['РЕЕСТР ИСПОЛНИТЕЛЬНОЙ ДОКУМЕНТАЦИИ', *_HEAD, 'Документы: договор, титульный лист, гарантийный паспорт, акт приемки, строительный паспорт, акт на скрытые работы, акт ревизии, '
                 'протоколы ВКК, акт передачи материалов.', 'Сертификаты:', '{СЕРТИФИКАТЫ}', 'Аттестаты сварщиков:', '{АТТЕСТАТЫ}'],
    'warranty': ['ГАРАНТИЙНЫЙ ПАСПОРТ ОБЪЕКТА', *_HEAD, *_GAS, 'Работы выполнены: {НАЧАЛО_РАБОТ} — {ОКОНЧАНИЕ_РАБОТ}', 'Дата акта: {ДАТА_АКТА}'],
    'payment_cert': ['СПРАВКА об оплате выполненных работ', *_HEAD, 'Стоимость работ: {СТОИМОСТЬ} BYN ({СТОИМОСТЬ_ПРОПИСЬЮ}).', 'Оплачено: {ОПЛАЧЕНО} BYN, остаток: {ОСТАТОК} BYN.',
                     'Даты оплат: {ДАТА_ОПЛАТЫ}', 'Суммы: {СУММА_ОПЛАТЫ}'],
    'accept': ['АКТ приемки объекта в эксплуатацию', *_HEAD, *_GAS, 'Материалы:', '{МАТЕРИАЛЫ}', '{ДАТА_АКТА}'],
    'building_passport': ['СТРОИТЕЛЬНЫЙ ПАСПОРТ ОБЪЕКТА', *_HEAD, 'Проект: {ШИФР_ПРОЕКТА}, {ПРОЕКТИРОВЩИК}, {ДАТА_ПРОЕКТА}', 'ТУ: {ТУ}', *_GAS, 'Материалы:', '{МАТЕРИАЛЫ}'],
    'hidden_works': ['АКТ освидетельствования скрытых работ', *_HEAD, *_GAS, '{ДАТА_АКТА}'],
    'device_revision': ['АКТ ревизии технических устройств', *_HEAD, 'Устройства: {НЕРАЗЪЕМНОЕ_СОЕДИНЕНИЕ}{ЗАГЛУШКА}, {МУФТА} — {МУФТА_КОЛ_ИЛИ}', '{ДАТА_АКТА}'],
    'vkk_input': ['ПРОТОКОЛ ВИЗУАЛЬНО-ИЗМЕРИТЕЛЬНОГО КОНТРОЛЯ (ввод)', *_HEAD, *_GAS, 'Сварщики: {СВАРЩИКИ}', '{ДАТА_АКТА}'],
    'vkk_coupling': ['ПРОТОКОЛ ВИЗУАЛЬНО-ИЗМЕРИТЕЛЬНОГО КОНТРОЛЯ (муфта)', *_HEAD, 'Муфты: {МУФТА}, количество {МУФТА_КОЛ}. Сертификаты: {МУФТА_СЕРТ}', 'Сварщики: {СВАРЩИКИ}', '{ДАТА_АКТА}'],
    'vkk_pipe': ['ПРОТОКОЛ ВИЗУАЛЬНО-ИЗМЕРИТЕЛЬНОГО КОНТРОЛЯ (труба)', *_HEAD, 'ПЭ труба: {ПЭ_ТРУБА}, {ДЛИНА_ПЭ} м. Сертификаты: {ПЭ_ТРУБА_СЕРТ}',
                 'Стальная труба: {СТАЛЬ_ТРУБА}, {ДЛИНА_СТАЛЬ} м. Сертификаты: {СТАЛЬ_ТРУБА_СЕРТ}', 'Сварщики: {СВАРЩИКИ}', '{ДАТА_АКТА}'],
    'vkk_joint': ['ПРОТОКОЛ ВИЗУАЛЬНО-ИЗМЕРИТЕЛЬНОГО КОНТРОЛЯ (неразъёмное соединение)', *_HEAD, 'Соединение ПЭ/сталь: {НЕРАЗЪЕМНОЕ_СОЕДИНЕНИЕ}, {НЕРАЗЪЕМНОЕ_СОЕДИНЕНИЕ_КОЛ} шт. '
                  'Сертификаты: {НЕРАЗЪЕМНОЕ_СОЕДИНЕНИЕ_СЕРТ}', 'Сварщики: {СВАРЩИКИ}', '{ДАТА_АКТА}'],
    'materials_transfer': ['АКТ передачи материалов', *_HEAD, 'Переданы материалы:', '{МАТЕРИАЛЫ}', 'Сертификаты:', '{СЕРТИФИКАТЫ_НОМЕРА}', '{ДАТА_АКТА}'],
}
DEFAULT_BODIES['device_revision'] = ['АКТ ревизии технических устройств', *_HEAD, 'Устройства и материалы:', '{МАТЕРИАЛЫ}', 'Сертификаты:', '{СЕРТИФИКАТЫ}', '{ДАТА_АКТА}']


def create_default_template(kind):
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    path = TEMPLATES_DIR / f'Шаблон_{DOC_KINDS[kind][3]}.docx'
    if path.exists():
        return path
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = docx.Document()
    lines = DEFAULT_BODIES[kind]
    doc.add_heading(lines[0], 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in lines[1:]:
        doc.add_paragraph(line)
    doc.save(path)
    return path


def doc_row(db, cid, kind):
    return db.fetchone('SELECT file_path,data_hash,linked FROM gsn_docs WHERE contract_id=? AND kind=?', (cid, kind))


def doc_state(db, cid, kind):
    """'none' — не создан, 'missing' — файл пропал, 'linked' — привязан готовый файл, 'stale' — данные изменились, 'fresh' — актуален."""
    row = doc_row(db, cid, kind)
    if not row:
        return 'none'
    if not os.path.exists(row[0]):
        return 'missing'
    if row[2]:
        return 'linked'
    return 'fresh' if row[1] == data_hash(db, cid) else 'stale'


def generate(db, cid, kind):
    """Формирует документ в папке договора (старый файл перезаписывается); возвращает путь."""
    if kind not in DOC_KINDS:
        raise ValueError('Неизвестный документ')
    folder = gn.contract_folder(db, cid)
    if not folder:
        raise ValueError('Сначала привяжите или создайте папку договора кнопкой «Папка договора»')
    tags = tag_map(db, cid)
    title, ext, _group, prefix = DOC_KINDS[kind]
    number = tags.get('НОМЕР_ДОГОВОРА') or str(cid)
    out = os.path.join(folder, gd.safe_name(f'{prefix}_{number.replace("/", "-")}').replace(' ', '_') + '.' + ext)
    gd.render_docx(template_path(db, kind), out, tags, normalize=norm_tag)
    db.execute('INSERT OR REPLACE INTO gsn_docs(contract_id,kind,file_path,data_hash,generated_at,linked) VALUES(?,?,?,?,?,0)',
               (cid, kind, out, data_hash(db, cid), datetime.now().isoformat(timespec='seconds')))
    return out


def link_doc(db, cid, kind, path):
    if not os.path.isfile(path):
        raise ValueError('Файл не найден: ' + path)
    db.execute('INSERT OR REPLACE INTO gsn_docs(contract_id,kind,file_path,data_hash,generated_at,linked) VALUES(?,?,?,?,?,1)',
               (cid, kind, path, '', datetime.now().isoformat(timespec='seconds')))


def remove_doc(db, cid, kind, delete_file=False):
    row = doc_row(db, cid, kind)
    if row and delete_file and not row[2] and os.path.isfile(row[0]):
        os.remove(row[0])
    db.execute('DELETE FROM gsn_docs WHERE contract_id=? AND kind=?', (cid, kind))


def doc_path(db, cid, kind):
    row = doc_row(db, cid, kind)
    return row[0] if row else ''


# --- СВЕРКА ПЕРЕД ФОРМИРОВАНИЕМ ---------------------------------------------------------------

def check(db, cid, kind):
    """Замечания перед формированием: [{'level': 'error'|'warn', 'text'}] — пустые поля, устаревшая смета, отсутствие материалов, неизвестные теги шаблона."""
    from .preflight import _need, _issue, _template_issues
    tags = tag_map(db, cid)
    c = gn.get(db, cid)
    issues = []
    _need(tags, 'КЛИЕНТ', 'ФИО клиента', issues)
    _need(tags, 'ОБЪЕКТ', 'наименование объекта', issues)
    _need(tags, 'АДРЕС', 'адрес объекта', issues)
    _need(tags, 'ДАТА_ДОГОВОРА_С', 'дата заключения договора', issues)
    _need(tags, 'НОМЕР_ДОГОВОРА', 'номер договора', issues, 'warn')
    if kind in ('contract', 'payment_cert'):
        _need(tags, 'ПАСПОРТ', 'паспорт (серия, номер)', issues)
        _need(tags, 'КЕМ_ВЫДАН', 'кем выдан паспорт', issues, 'warn')
        _need(tags, 'ДАТА_ВЫДАЧИ', 'дата выдачи паспорта', issues, 'warn')
        _need(tags, 'АДРЕС_КЛИЕНТА', 'адрес клиента', issues)
        if tags.get('СТОИМОСТЬ') in ('', '0,00'):
            issues.append(_issue('error', 'Не указана стоимость работ'))
    if kind == 'contract':
        _need(tags, 'ОКОНЧАНИЕ_РАБОТ', 'окончание работ', issues)
    if kind in ID_KINDS:
        _need(tags, 'ДАТА_АКТА_С', 'дата акта', issues, 'warn')
        _need(tags, 'ШИФР_ПРОЕКТА', 'шифр проекта', issues, 'warn')
        if not (c.get('pe_length') or c.get('steel_length')):
            issues.append(_issue('warn', 'Не указана протяжённость трубопроводов'))
        if not c.get('tie_in'):
            issues.append(_issue('warn', 'Не выбран тип врезки'))
        if not c.get('pressure_type'):
            issues.append(_issue('warn', 'Не выбран тип газопровода (низкое / среднее давление)'))
        if not gn.certificates(db, cid, c):
            issues.append(_issue('warn', 'Нет сертификатов: отметьте «по умолчанию» или привяжите сертификаты к позициям в Справочнике ГСН'))
    needs = {'vkk_pipe': ('pe_pipe', 'steel_pipe'), 'vkk_coupling': ('coupling',), 'vkk_joint': ('transition',), 'vkk_input': ()}
    have = {m['role'] for m in gn.materials(db, cid, c)}
    if kind in needs and needs[kind] and not (set(needs[kind]) & have):
        issues.append(_issue('warn', 'В объекте нет материалов для этого протокола: ' + ', '.join(gn.ROLES[r] for r in needs[kind])))
    state = estimate_state_for(db, cid)
    if state == 'stale':
        issues.append(_issue('error', 'Смета была изменена — обновите данные («Получить стоимость из сметы»), иначе стоимость в документе будет устаревшей'))
    elif state == 'unsynced':
        issues.append(_issue('warn', 'Данные из привязанной сметы ещё не получены — стоимость введена вручную'))
    ignore = {name for (name,) in db.fetchall("SELECT name FROM gsn_tags WHERE coalesce(auto_key,'') LIKE 'r:%'")}
    issues += _template_issues(template_path(db, kind), tags, norm_tag, ignore)
    seen, unique = set(), []
    for i in issues:
        if i['text'] not in seen:
            seen.add(i['text'])
            unique.append(i)
    return unique


def estimate_state_for(db, cid):
    return gn.estimate_state(db, cid)
