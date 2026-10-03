"""Per-record export profiles. Output is atomically replaced only after success."""
import json
import os
import tempfile
from pathlib import Path

DEFAULTS={'title':'','company':'','font':'Arial','font_size':10,'landscape':False,'margin_mm':15,'include_private':False,'columns':[]}
PRIVATE={'passport','passport_series_number','passport_issued_by','passport_issue_date','passport_issuer','passport_date','client_address','address','phone','client_phone'}

def record_data(db,table,rid):
    from .data_services import SOURCES
    if table not in SOURCES and table!='crm.clients':raise ValueError('Неизвестный реестр')
    with db._lock:
        cur=db.execute(f'SELECT * FROM {table} WHERE id=?',(rid,))
        row=cur.fetchone()
    if row is None:raise ValueError('Запись удалена')
    data=dict(zip([c[0] for c in cur.description],row))
    if data.get('client_id'):
        client=db.fetchone('SELECT name,phone,address FROM crm.clients WHERE id=?',(data['client_id'],))
        if client:data.update(client_name=client[0],client_phone=client[1],client_address=client[2])
    sections=[]
    if table=='estimates':
        rows=db.fetchall('SELECT name,item_type,unit,quantity,price,sum FROM estimate_items WHERE estimate_id=? ORDER BY sort_order,id',(rid,))
        sections.append(('Позиции сметы',['Наименование','Тип','Ед.','Кол-во','Цена','Сумма'],rows))
        sections.append(('Оплаты',['Дата','Сумма'],db.fetchall('SELECT date,amount FROM payments WHERE estimate_id=? ORDER BY date,id',(rid,))))
    if table in ('contracts','gsv_projects'):
        equipment,key=('contract_equipment','contract_id') if table=='contracts' else ('gsv_equipment','project_id')
        rows=db.fetchall(f'SELECT equipment_kind,equipment_model,equipment_name,certificate_number FROM {equipment} WHERE {key}=?',(rid,))
        sections.append(('Оборудование',['Вид','Тип / модель','№ сертификата / паспорта'],[(r[0],r[1] or r[2],r[3]) for r in rows]))
        rows=db.fetchall('SELECT p.name,op.quantity,p.unit,c.cert_number FROM object_pipelines op JOIN gsv_pipelines p ON p.id=op.pipeline_id LEFT JOIN certificates c ON c.id=p.certificate_id WHERE op.owner_type=? AND op.owner_id=? ORDER BY op.id',(table,rid))
        sections.append(('Трубопроводы',['Вид','Количество','Ед.','Сертификат'],rows))
        from .gsv_domain import dossier
        docs=dossier(db,table,rid)
        sections.append(('Исполнительная документация',['Источник','Название','Номер'],[(r['source'],r['title'],r['number']) for r in docs]))
    return data,sections

LABELS={'id':'ID','title':'Название','name':'Наименование','client_name':'Клиент','client_phone':'Телефон клиента','client_address':'Адрес клиента','date':'Дата','total':'Итого','paid':'Оплачено','statuses':'Статусы','contract_number':'Номер договора','contract_date':'Дата договора','object_name':'Объект','contract_amount':'Сумма договора','pd_number':'Шифр проекта','notes':'Примечания','phone':'Телефон','address':'Адрес','passport':'Паспорт','description':'Описание','cost':'Стоимость','work_status':'Статус работ','client_status':'Статус клиента','quantity':'Количество','price':'Цена','unit':'Единица','note':'Примечание','file_name':'Файл','file_path':'Путь','certificate':'Аттестация','cert_number':'Номер сертификата','valid_to':'Действителен до','event_date':'Дата события','event_time':'Время','client_id':'ID клиента','estimate_id':'ID сметы','due_date':'Срок','act_date':'Дата акта','status':'Статус','urgency':'Срочность','created_at':'Создано','birth_date':'Дата рождения','welding_type':'Вид сварки','stamp':'Клеймо','grade':'Разряд','valid_until':'Аттестация действительна до','document_type':'Тип документа','document_date':'Дата документа'}

def get_profile(db,key):
    row=db.fetchone('SELECT options FROM export_profiles WHERE record_key=?',(key,))
    try: opts=json.loads(row[0]) if row else {}
    except (ValueError,TypeError):opts={}
    return {**DEFAULTS,'company':db.get_setting('export_company_name',''),'font':db.get_setting('export_font','Arial'),'font_size':int(db.get_setting('export_font_size','10')),**opts}


def export_record(db,table,rid,path,options=None):
    if table in ('stock_acts','defect_acts'):
        from .stock_exports import export_document
        return export_document(db,table,rid,path,options)
    opts={**DEFAULTS,**(options or {})}
    data,sections=record_data(db,table,rid)
    selected=opts['columns'] or list(data)
    rows=[(LABELS.get(k,k),data[k]) for k in selected if k in data and (opts['include_private'] or k not in PRIVATE)]
    title=opts['title'] or str(data.get('title') or data.get('name') or data.get('pd_number') or data.get('contract_number') or 'Документ')
    path=Path(path);suffix=path.suffix.lower()
    path.parent.mkdir(parents=True,exist_ok=True)
    fd,temp=tempfile.mkstemp(dir=path.parent,suffix=suffix);os.close(fd)
    size=max(6,min(24,int(opts['font_size'])));margin=max(5,min(40,int(opts['margin_mm'])))
    try:
        if suffix=='.xlsx':
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.styles import Font
            wb=Workbook(write_only=True);ws=wb.create_sheet('Документ')
            ws.page_setup.orientation='landscape' if opts['landscape'] else 'portrait'
            def append(values):
                cells=[]
                for value in values:
                    cell=WriteOnlyCell(ws,value=value if isinstance(value,(int,float)) else str(value or ''))
                    if isinstance(value,str):cell.data_type='s' # no spreadsheet formula injection
                    cell.font=Font(name=opts['font'],size=size);cells.append(cell)
                ws.append(cells)
            append([title]);append([opts['company']])
            for row in rows:append(row)
            for heading,headers,items in sections:
                append([]);append([heading]);append(headers)
                for row in items:append(row)
            wb.save(temp)
        elif suffix=='.docx':
            from docx import Document
            from docx.shared import Pt,Mm
            from docx.enum.section import WD_ORIENT
            doc=Document();section=doc.sections[0]
            if opts['landscape']:section.orientation=WD_ORIENT.LANDSCAPE;section.page_width,section.page_height=section.page_height,section.page_width
            section.left_margin=section.right_margin=section.top_margin=section.bottom_margin=Mm(margin)
            normal=doc.styles['Normal'];normal.font.name=opts['font'];normal.font.size=Pt(size)
            doc.add_heading(title,0)
            if opts['company']:doc.add_paragraph(opts['company'])
            t=doc.add_table(rows=0,cols=2);t.style='Table Grid'
            for a,b in rows:
                cells=t.add_row().cells;cells[0].text=str(a);cells[1].text=str(b if b is not None else '')
            for heading,headers,items in sections:
                doc.add_heading(heading,1);t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
                for cell,value in zip(t.rows[0].cells,headers):cell.text=value
                for row in items:
                    for cell,value in zip(t.add_row().cells,row):cell.text=str(value if value is not None else '')
            doc.save(temp)
        elif suffix=='.pdf':
            from reportlab.pdfgen.canvas import Canvas
            from reportlab.lib.pagesizes import A4,landscape
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.lib.utils import simpleSplit
            font_path=Path(__file__).parent/'assets'/'DejaVuSans.ttf'
            if not font_path.exists():raise RuntimeError('Не найден шрифт PDF assets/DejaVuSans.ttf')
            pdfmetrics.registerFont(TTFont('SmetaSans',str(font_path)))
            page=landscape(A4) if opts['landscape'] else A4
            c=Canvas(temp,pagesize=page);c.setTitle(title)
            x=margin*72/25.4;y=page[1]-x;number=1
            def line(value,bold=False):
                nonlocal y,number
                for text in simpleSplit(str(value),'SmetaSans',size,page[0]-2*x):
                    if y<x+size*2:
                        c.setFont('SmetaSans',8);c.drawString(x,x/2,str(number));c.showPage();number+=1;y=page[1]-x
                    c.setFont('SmetaSans',size);c.drawString(x,y,text);y-=size*1.5
            line(title);line(opts['company'])
            for a,b in rows:line(f'{a}: {b if b is not None else ""}')
            for heading,headers,items in sections:
                line('');line(heading);line(' | '.join(headers))
                for row in items:line(' | '.join(str(v if v is not None else '') for v in row))
            c.setFont('SmetaSans',8);c.drawString(x,x/2,str(number));c.save()
        else:raise ValueError('Поддерживаются .docx, .xlsx и .pdf')
        os.replace(temp,path)
    finally:
        if os.path.exists(temp):os.unlink(temp)
    return str(path)
