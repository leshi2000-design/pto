"""Editable document exports; draft C-29/C-1 layouts are explicitly identified."""
from pathlib import Path
import os
import tempfile
import json
from xml.sax.saxutils import escape
from .stock_domain import fmt

FORM_NOTICE='Рабочий проект для заполнения формы. Макет не подтверждён как точное воспроизведение действующего бланка РБ.'

def document_data(db,source,rid,mode='act'):
    if source not in ('stock_acts','defect_acts'):raise ValueError('Неверный вид документа')
    with db.transaction():
        cur=db.execute(f'SELECT * FROM {source} WHERE id=?',(rid,));row=cur.fetchone()
        if not row:raise ValueError('Документ не найден')
        d=dict(zip([c[0] for c in cur.description],row));sections=[]
        if source=='defect_acts':
            title='Дефектный акт';notice='Проект для формы С-1. '+FORM_NOTICE
            rows=db.fetchall('SELECT defect,work,unit,quantity,notes FROM defect_lines WHERE act_id=? ORDER BY id',(rid,))
            sections.append(('Дефекты и необходимые работы',['№','Место и описание дефекта','Необходимые работы','Ед.','Объём','Примечание'],[(i,*r) for i,r in enumerate(rows,1)]))
        else:
            title='Отчёт о расходе материалов в сопоставлении с нормами' if mode=='c29' else 'Акт списания материалов'
            notice='Проект для формы С-29. '+FORM_NOTICE if mode=='c29' else ''
            status={'draft':'ЧЕРНОВИК — списание не проведено','posted':'ПРОВЕДЁН','reversed':'ОТМЕНЁН — материалы возвращены на остаток'}[d['status']]
            d['document_status']=status
            rows=db.fetchall('SELECT material_name,unit,norm_qty,actual_qty,note FROM stock_act_lines WHERE act_id=? ORDER BY id',(rid,))
            sections.append(('Расход материалов',['№','Наименование материала','Ед.','По норме','Фактически','Отклонение (+/−)','Примечание'],[(i,name,unit,fmt(norm) if norm is not None else 'Не задано',fmt(actual),fmt(actual-norm) if norm is not None else '',note) for i,(name,unit,norm,actual,note) in enumerate(rows,1)]))
            snapshots=json.loads(d['calculation'] or '[]')
            if snapshots:
                details=[]
                for v in snapshots:
                    for mat in v['items']:details.append((v['name'],v['basis'],v['volume'],mat['name'],mat['rate'],mat['unit'],mat['quantity']))
                sections.append(('Расчёт по сохранённым нормам',['Работа / диаметр','База','Объём','Материал','На единицу','Ед.','Расчётный расход'],details))
            if d.get('defect_id'):
                defect=db.fetchone('SELECT number,act_date FROM defect_acts WHERE id=?',(d['defect_id'],));d['defect_reference']='№ '+defect[0]+' от '+defect[1]
                rows=db.fetchall('SELECT defect,work,unit,quantity FROM defect_lines WHERE act_id=? ORDER BY id',(d['defect_id'],))
                sections.append(('Работы по дефектному акту',['Дефект','Работа','Ед.','Объём'],rows))
    return title,notice,d,sections

HEADER_LABELS={'organization':'Организация','object_name':'Объект строительства','act_date':'Дата','number':'Номер','document_status':'Состояние','basis':'Основание','reason':'Основание обследования','conditions':'Условия выполнения работ','defect_reference':'Дефектный акт','notes':'Примечание'}

def export_document(db,source,rid,path,options=None):
    opts={'mode':'act','title':'','font_size':10,'font':'Arial','margin_mm':15,'landscape':True,**(options or {})}
    title,notice,d,sections=document_data(db,source,rid,opts['mode']);title=opts['title'] or title
    headers=[(label,str(d.get(key) or '')) for key,label in HEADER_LABELS.items() if d.get(key)]
    tail=['УТВЕРЖДАЮ: '+(d.get('approved_by') or '________________________'), 'Ответственные / комиссия: '+(d.get('commission') or '________________________'),'Подписи: ________________________     Дата: ________________________']
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);ext=path.suffix.lower();fd,temp=tempfile.mkstemp(dir=path.parent,suffix=ext);os.close(fd)
    size=max(7,min(16,int(opts['font_size'])));margin=max(8,min(30,int(opts['margin_mm'])))
    try:
        if ext=='.pdf':
            from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
            from reportlab.lib.pagesizes import A4,landscape
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.lib import colors
            from reportlab.lib.units import mm
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            font='SGExport';pdfmetrics.registerFont(TTFont(font,str(Path(__file__).parent/'assets/DejaVuSans.ttf')))
            pagesize=landscape(A4) if opts['landscape'] else A4;usable=pagesize[0]-2*margin*mm
            style=ParagraphStyle('body',fontName=font,fontSize=size,leading=size*1.35,spaceAfter=5)
            title_style=ParagraphStyle('title',parent=style,fontSize=size+5,leading=(size+5)*1.3,spaceAfter=12)
            def p(s):return Paragraph(escape(str(s if s is not None else '')).replace('\n','<br/>'),style)
            story=[Paragraph(escape(title),title_style)]
            if notice:story.extend([p(notice),Spacer(1,6)])
            for label,value in headers:story.append(p(label+': '+value))
            for idx,(heading,cols,rows) in enumerate(sections):
                if idx:story.append(PageBreak())
                story.extend([Spacer(1,10),p(heading)])
                weights=[0.5 if h=='№' else 0.8 if h in ('Ед.','База') else 2.8 if h in ('Наименование материала','Место и описание дефекта','Необходимые работы','Материал','Работа / диаметр','Дефект','Работа') else 1.3 for h in cols]
                widths=[usable*w/sum(weights) for w in weights];cells=[[p(c) for c in cols]]+[[p(c) for c in row] for row in rows]
                t=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.4,colors.grey),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8edf2')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]));story.append(t)
            story.append(Spacer(1,16));story.extend(p(v) for v in tail)
            def footer(canvas,doc):canvas.setFont(font,8);canvas.drawRightString(pagesize[0]-margin*mm,8*mm,str(doc.page))
            SimpleDocTemplate(temp,pagesize=pagesize,leftMargin=margin*mm,rightMargin=margin*mm,topMargin=margin*mm,bottomMargin=margin*mm).build(story,onFirstPage=footer,onLaterPages=footer)
        elif ext=='.docx':
            from docx import Document
            from docx.shared import Mm,Pt,RGBColor
            from docx.enum.section import WD_ORIENT
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            doc=Document();sec=doc.sections[0];sec.page_width=Mm(297 if opts['landscape'] else 210);sec.page_height=Mm(210 if opts['landscape'] else 297);sec.orientation=WD_ORIENT.LANDSCAPE if opts['landscape'] else WD_ORIENT.PORTRAIT
            sec.left_margin=sec.right_margin=sec.top_margin=sec.bottom_margin=Mm(margin);normal=doc.styles['Normal'];normal.font.name=opts['font'];normal.font.size=Pt(size);normal.paragraph_format.space_after=Pt(5)
            for name in ('Title','Heading 1'):doc.styles[name].font.color.rgb=RGBColor(0,0,0);doc.styles[name].font.name=opts['font']
            for style in doc.styles:
                for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
            doc.styles['Title'].font.size=Pt(18)
            doc.add_paragraph(title,'Title')
            if notice:doc.add_paragraph(notice)
            for label,value in headers:doc.add_paragraph(label+': '+value)
            for idx,(heading,cols,rows) in enumerate(sections):
                if idx:doc.add_page_break()
                doc.add_heading(heading,1);t=doc.add_table(rows=1,cols=len(cols));t.style='Table Grid';t.autofit=False
                weights=[0.5 if h=='№' else 0.8 if h in ('Ед.','База') else 2.8 if h in ('Наименование материала','Место и описание дефекта','Необходимые работы','Материал','Работа / диаметр','Дефект','Работа') else 1.3 for h in cols]
                widths=[Mm(((297 if opts['landscape'] else 210)-2*margin)*w/sum(weights)) for w in weights]
                for col,width in zip(t.columns,widths):col.width=width
                repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
                for c,h in zip(t.rows[0].cells,cols):c.text=h
                for row in rows:
                    for c,v in zip(t.add_row().cells,row):c.text=str(v if v is not None else '')
                for row in t.rows:
                    for cell,width in zip(row.cells,widths):cell.width=width
            for v in tail:doc.add_paragraph(v)
            doc.save(temp)
        elif ext=='.xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font,Alignment,PatternFill,Border,Side
            from openpyxl.utils import get_column_letter
            wb=Workbook();wb.remove(wb.active)
            for idx,(heading,cols,rows) in enumerate(sections):
                ws=wb.create_sheet(('Акт' if idx==0 else 'Приложение '+str(idx))[:31]);n=len(cols)
                def append(values,merge=False,bold=False):
                    ws.append([str(v) if v is not None else '' for v in values]);r=ws.max_row
                    if merge:ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=n)
                    for c in ws[r]:
                        if c.value is not None:c.data_type='s'
                        c.font=Font(name=opts['font'],size=size,bold=bold);c.alignment=Alignment(vertical='top',wrap_text=True)
                    ws.row_dimensions[r].height=30
                append([title],True,True)
                if notice:append([notice],True)
                for label,value in headers:append([label+': '+value],True)
                append([heading],True,True);append(cols,bold=True);headerrow=ws.max_row
                for row in rows:append(row)
                last=ws.max_row;ws.auto_filter.ref=f'A{headerrow}:{get_column_letter(n)}{last}';ws.freeze_panes=f'A{headerrow+1}';ws.print_title_rows=f'1:{headerrow}'
                for row in ws.iter_rows(min_row=headerrow,max_row=last):
                    for c in row:c.border=Border(bottom=Side(style='thin',color='AAAAAA'))
                for c in ws[headerrow]:c.fill=PatternFill('solid',fgColor='E8EDF2')
                for c,h in enumerate(cols,1):ws.column_dimensions[get_column_letter(c)].width=8 if h in ('№','Ед.','База') else 38 if h in ('Наименование материала','Место и описание дефекта','Необходимые работы','Материал','Работа / диаметр','Дефект','Работа') else 20
                for v in tail:append([v],True)
                ws.page_setup.orientation='landscape' if opts['landscape'] else 'portrait';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
                ws.print_options.horizontalCentered=True
            wb.save(temp)
        else:raise ValueError('Выберите .docx, .xlsx или .pdf')
        os.replace(temp,path);return str(path)
    finally:
        if os.path.exists(temp):os.unlink(temp)

from PyQt6.QtWidgets import QDialog,QVBoxLayout,QFormLayout,QComboBox,QLineEdit,QSpinBox,QCheckBox,QLabel,QPushButton,QFileDialog
from PyQt6.QtCore import QTimer
from concurrent.futures import ThreadPoolExecutor
from .database import db

class LegacyStockExportDialog(QDialog):
    def __init__(self,source,rid,parent=None):
        super().__init__(parent);self.source=source;self.rid=rid;self.key=f'stock_export:{source}:{rid}';self.future=None;self.pool=ThreadPoolExecutor(max_workers=1);self.setWindowTitle('Экспорт акта');self.resize(650,450)
        row=db.fetchone('SELECT options FROM export_profiles WHERE record_key=?',(self.key,));opts=json.loads(row[0]) if row else {};l=QVBoxLayout(self);f=QFormLayout();l.addLayout(f)
        self.mode=QComboBox();self.mode.addItem('Акт списания','act');self.mode.addItem('Сопоставление расхода — проект для С-29','c29');self.mode.setCurrentIndex(max(0,self.mode.findData(opts.get('mode'))));self.mode.setVisible(source=='stock_acts');f.addRow('Документ',self.mode)
        self.title=QLineEdit(opts.get('title',''));self.title.setPlaceholderText('Стандартное название');f.addRow('Заголовок',self.title);self.font=QLineEdit(opts.get('font','Arial'));f.addRow('Шрифт Word / Excel',self.font);self.size=QSpinBox();self.size.setRange(7,16);self.size.setValue(opts.get('font_size',10));f.addRow('Размер',self.size);self.margin=QSpinBox();self.margin.setRange(8,30);self.margin.setValue(opts.get('margin_mm',15));f.addRow('Поля, мм',self.margin);self.landscape=QCheckBox('Альбомная ориентация');self.landscape.setChecked(opts.get('landscape',True));l.addWidget(self.landscape)
        self.format=QComboBox();self.format.addItems(['docx','xlsx','pdf']);f.addRow('Формат',self.format);label=QLabel(FORM_NOTICE+' Настройки сохраняются отдельно для этого документа.');label.setWordWrap(True);l.addWidget(label);self.info=QLabel();self.info.setWordWrap(True);l.addWidget(self.info);self.button=QPushButton('Сохранить файл');self.button.clicked.connect(self.run);l.addWidget(self.button);self.timer=QTimer(self);self.timer.timeout.connect(self.poll)
    def run(self):
        opts=dict(mode=self.mode.currentData(),title=self.title.text(),font=self.font.text() or 'Arial',font_size=self.size.value(),margin_mm=self.margin.value(),landscape=self.landscape.isChecked());ext=self.format.currentText();path,_=QFileDialog.getSaveFileName(self,'Сохранить акт',f'act_{self.rid}.{ext}',f'{ext.upper()} (*.{ext})')
        if not path:return
        if not path.lower().endswith('.'+ext):path+='.'+ext
        db.execute('INSERT INTO export_profiles VALUES(?,?) ON CONFLICT(record_key) DO UPDATE SET options=excluded.options',(self.key,json.dumps(opts,ensure_ascii=False)));self.button.setEnabled(False);self.info.setText('Формирование документа…');self.future=self.pool.submit(export_document,db,self.source,self.rid,path,opts);self.timer.start(100)
    def poll(self):
        if self.future and self.future.done():
            self.timer.stop();self.button.setEnabled(True)
            try:self.info.setText('Сохранено: '+self.future.result())
            except Exception as e:self.info.setText('Ошибка: '+str(e))
    def reject(self):
        if self.future and not self.future.done():return
        self.pool.shutdown(wait=False);super().reject()


def StockExportDialog(source,rid,parent=None):
    from .report_dialog import ReportTemplateDialog
    return ReportTemplateDialog(source,rid,parent)
