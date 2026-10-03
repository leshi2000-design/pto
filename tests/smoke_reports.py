import os,sys,tempfile,time
from pathlib import Path
root=tempfile.TemporaryDirectory();os.environ['SMETAGAZ_DATA_DIR']=root.name;os.environ.setdefault('QT_QPA_PLATFORM','offscreen');sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from PyQt6.QtWidgets import QApplication,QMessageBox
from PyQt6.QtCore import Qt
app=QApplication([])
def warning(*a,**kw):raise AssertionError(str(a[1:]))
QMessageBox.warning=warning
from smetagaz.database import db
db.init_db()
for i in range(12):db.execute('INSERT INTO estimates(title,date,client_name,client_phone,total,paid,statuses) VALUES(?,?,?,?,?,?,?)',(f'Газоснабжение жилого дома · участок {i+1}','2026-09-19',f'Заказчик {i+1}','+375 29 123-45-67',1200+i*155,500,'В работе'))
rid=db.fetchone('SELECT id FROM estimates LIMIT 1')[0];db.execute("INSERT INTO estimate_items(estimate_id,name,item_type,unit,quantity,price,sum) VALUES(?,'Труба ПЭ 25','Материал','м',10,12,120)",(rid,))
from smetagaz.main_window import MainWindow
main=MainWindow();main.resize(1400,900);main.show()
for i,b in enumerate(main.tabs_buttons):
    if 'смет' in b.text().lower():main.switch_tab(i,b);break
app.processEvents()
folder=Path(os.environ.get('SMETAGAZ_SCREENSHOTS',root.name));folder.mkdir(parents=True,exist_ok=True);main.grab().save(str(folder/'interface-light.png'));print('Main window size:',main.width(),main.height())
from smetagaz.report_dialog import ReportTemplateDialog
from docx import Document
source=Path(root.name)/'estimate.docx';doc=Document();doc.add_paragraph('{{title}}');t=doc.add_table(rows=1,cols=2);t.cell(0,0).text='{{items.name}}';t.cell(0,1).text='{{items.quantity}}';doc.save(source);tid=db.execute('INSERT INTO report_templates(kind,name,file_path) VALUES(?,?,?)',('estimates','Смета · основной шаблон',str(source))).lastrowid
report=ReportTemplateDialog('estimates',rid,main);report.author.setText('Иванов И.И.');report.show();report.validate();assert 'пройдена' in report.status.text();assert report.tags.rowCount()>20;report.search.setText('items.name');assert sum(not report.tags.isRowHidden(r) for r in range(report.tags.rowCount()))==1;report.search.clear();app.processEvents();report.grab().save(str(folder/'template-dialog.png'));report.reject()
from smetagaz.stock_view import Registry
for i in range(420):db.execute('INSERT INTO stock_materials(name,category,unit) VALUES(?,?,?)',(f'Материал {i}','Тестовая категория','шт'))
view=Registry('balances');view.category.setCurrentIndex(view.category.findData('Тестовая категория'));assert view.table.rowCount()==200
from smetagaz.report_templates import context
query,params=view.pager.filters.apply(view.report_query,view.report_params);ctx,tables=context(db,'balances',filters={'query':query,'params':params});assert len(tables['items'])==420
from smetagaz.contract_card import ContractCardDialog
card=ContractCardDialog();assert card.executive.templates_table.rowCount()==8;assert card.executive.generate_button.isEnabled();card.close()
main.toggle_theme();app.processEvents();main.grab().save(str(folder/'interface-dark.png'));main.toggle_theme();app.processEvents()
main.backup_service.close();main.close();app.processEvents();print('REPORT GUI OK: templates, tag search, row expansion, balances across pages, GSV separation, light/dark theme')
