import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from openpyxl import load_workbook
import pytest
from smetagaz.database import DatabaseManager
from smetagaz.contracts_excel_export import export_registries


@pytest.fixture
def database(tmp_path):
    d = DatabaseManager(tmp_path / 'smetagaz.db')
    d.init_db()
    yield d
    d.close()


def test_export_registries_writes_both_sheets_and_overwrites(database, tmp_path):
    db = database
    est = db.execute("INSERT INTO estimates(title) VALUES('Смета 1')").lastrowid
    db.execute(
        "INSERT INTO contracts(estimate_id,contract_number,contract_date,object_name,client_name,client_phone,contract_amount) "
        "VALUES(?, '1/26','2026-01-05','Дом, ул. Садовая','Иванов И.И.','+375291234567',1500.0)", (est,)
    )
    db.execute(
        "INSERT INTO gsn_projects(contract_number,contract_date,title,client_name,phone,address) "
        "VALUES('2/26','2026-01-06','Газоснабжение дома','Петров П.П.','+375297654321','Минск')"
    )

    path = tmp_path / 'reports' / 'registries.xlsx'
    gsv_count, gsn_count = export_registries(db, path)
    assert (gsv_count, gsn_count) == (1, 1)
    assert path.exists()

    wb = load_workbook(path)
    assert wb.sheetnames == ['ГСВ', 'ГСН', 'Юрлица', 'СМР']
    gsv = wb['ГСВ']
    assert gsv.cell(1, 1).value == '№ Договора'
    assert gsv.cell(2, 1).value == '1/26'
    assert gsv.cell(2, 4).value == 'Иванов И.И.'
    assert gsv.cell(2, 7).value == 'Смета 1'
    gsn = wb['ГСН']
    assert gsn.cell(2, 1).value == '2/26'
    assert gsn.cell(2, 4).value == 'Петров П.П.'

    # A second export must overwrite the same file, not add to it.
    db.execute(
        "INSERT INTO contracts(contract_number,contract_date,object_name) VALUES('3/26','2026-01-07','Ещё один дом')"
    )
    gsv_count2, gsn_count2 = export_registries(db, path)
    assert gsv_count2 == 2
    wb2 = load_workbook(path)
    assert wb2['ГСВ'].max_row == 3  # header + 2 rows, not accumulated from the first export
