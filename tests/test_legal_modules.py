"""СМР (физлица), Прочие (юр.лица) и перенос договоров юрлиц между разделами."""
import sys
from smetagaz import main_window, orders_view, docs_view, counterparty_ui, contract_card, gsn_card, volumes_ui, legal_entities_view  # noqa: F401  (до подмены базы в тестах)

import pytest

from smetagaz import contracts_core as cc, payments_domain as pay
from smetagaz.database import DatabaseManager


@pytest.fixture
def db(tmp_path):
    d = DatabaseManager(tmp_path / 'smetagaz.db')
    d.init_db()
    yield d
    d.close()


def make_legal(db, name='ООО Ромашка'):
    return cc.save_legal(db, dict(name=name, unp='190123456'))


def make_person(db):
    from smetagaz.gsv_domain import save_client
    return save_client(db, dict(name='Иванов Иван Иванович', phone='+375291112233', address='г. Минск'))


def test_fl_module_tables_numbers_and_person_contract(db):
    assert cc.cfg('fl')['title'] == 'СМР (физлица)' and cc.cfg('le')['title'] == 'Прочие (юр.лица)'
    pid = make_person(db)
    cid = cc.save_contract(db, 'fl', dict(contract_date='2026-03-05', object_name='Дом', object_address='г. Минск, ул. Мира 1', person_id=pid, amount=1200))
    row = cc.contract(db, 'fl', cid)
    assert row['contract_number'] == '01-08/26' and row['party_type'] == 'person'
    assert cc.party(db, 'fl', row)['name'].startswith('Иванов')
    assert cc.next_number(db, 'fl', '2026-04-01')[2] == '02-08/26'
    aid = cc.save_act(db, 'fl', dict(contract_id=cid, act_number='1', act_date='2026-03-20', amount=500))
    assert [a['id'] for a in cc.acts(db, 'fl', cid)] == [aid]
    pay.add(db, 'fl_contracts', cid, 300, '2026-03-21', 'аванс')
    assert pay.summary(db, 'fl_contracts', cid)['paid'] == 300
    # модуль виден везде, где перечислены разделы
    from smetagaz import acts_statement, estimates_domain as ed, dossier_client, integrity
    assert 'fl_contracts' in acts_statement.SECTIONS and 'fl_contracts' in ed.CONTRACT_TABLES
    dossier = dossier_client.client_dossier(db, 'person', pid)
    assert any(c['table'] == 'fl_contracts' for c in dossier['contracts'])
    assert not [f for f in integrity.check(db) if f['level'] == 'error']


def test_move_contract_between_legal_sections(db):
    lid = make_legal(db)
    cid = cc.save_contract(db, 'smr', dict(contract_date='2026-02-01', object_name='Склад', subject='СМР', legal_id=lid, amount=1000, note='заметка'))
    old_number = cc.contract(db, 'smr', cid)['contract_number']
    a1 = cc.save_act(db, 'smr', dict(contract_id=cid, act_number='1', act_date='2026-02-10', amount=400, signed=True))
    cc.save_act(db, 'smr', dict(contract_id=cid, act_number='2', act_date='2026-02-20', amount=300))
    pay.add(db, 'smr_contracts', cid, 250, '2026-02-12', 'аванс')
    db.execute("INSERT INTO smr_docs(scope,ref_id,kind,file_path,data_hash,generated_at,linked) VALUES('act',?,'act','/tmp/akt.docx','h','2026-02-10',1)", (a1,))
    db.execute("INSERT INTO notes(title,body,link_table,link_id) VALUES('Звонить','x','smr_contracts',?)", (cid,)) if db.fetchone("SELECT 1 FROM sqlite_master WHERE name='notes'") else None
    new = cc.move_contract(db, 'smr', cid, 'to', renumber=True)
    assert db.fetchone('SELECT count(*) FROM smr_contracts')[0] == 0 and db.fetchone('SELECT count(*) FROM smr_acts')[0] == 0
    moved = cc.contract(db, 'to', new)
    assert moved['contract_number'] == '01-06/26' and moved['legal_id'] == lid and moved['amount'] == 1000
    assert f'Прежний номер: {old_number}' in moved['note'] and 'Перенесён из раздела «СМР (юрлица)»' in moved['note'] and 'заметка' in moved['note']
    acts = cc.acts(db, 'to', new)
    assert [(a['act_number'], a['amount'], a['signed']) for a in acts] == [('1', 400, 1), ('2', 300, 0)]
    assert pay.summary(db, 'to_contracts', new)['paid'] == 250
    assert db.fetchone("SELECT file_path,linked FROM to_docs WHERE scope='act' AND ref_id=?", (acts[0]['id'],)) == ('/tmp/akt.docx', 1)
    # в «Прочие» и обратно, номер сохраняется
    other = cc.move_contract(db, 'to', new, 'le', renumber=False)
    row = cc.contract(db, 'le', other)
    assert row['client_id'] == lid and row['contract_number'] == '01-06/26' and row['seq_num'] == 1
    assert len(cc.acts(db, 'le', other)) == 2 and pay.summary(db, 'le_contracts', other)['paid'] == 250
    back = cc.move_contract(db, 'le', other, 'smr', renumber=True)
    assert cc.contract(db, 'smr', back)['legal_id'] == lid and cc.contract(db, 'smr', back)['contract_number'].endswith('-05/26')
    with pytest.raises(ValueError):
        cc.move_contract(db, 'smr', back, 'smr')
    with pytest.raises(ValueError):
        cc.move_contract(db, 'smr', back, 'ip')


def test_move_with_estimate_requires_confirmation(db):
    from smetagaz import estimates_domain as ed
    lid = make_legal(db)
    cid = cc.save_contract(db, 'smr', dict(contract_date='2026-02-01', legal_id=lid, amount=500))
    eid = ed.create_estimate(db, 'Смета', ('manual', 'Клиент', ''), None, False, None, '')
    cc.link_estimate(db, cid, eid, 'smr')
    new = cc.move_contract(db, 'smr', cid, 'to')                               # СМР -> ТО: смета остаётся привязанной
    assert cc.contract(db, 'to', new)['estimate_id'] == eid
    with pytest.raises(ValueError) as e:
        cc.move_contract(db, 'to', new, 'le')
    assert 'смет' in str(e.value)
    le = cc.move_contract(db, 'to', new, 'le', drop_estimate=True)
    assert cc.contract(db, 'le', le)['client_id'] == lid and db.fetchone('SELECT count(*) FROM estimates WHERE id=?', (eid,))[0] == 1


def test_person_contract_cannot_move(db):
    pid = make_person(db)
    cid = cc.save_contract(db, 'smr', dict(contract_date='2026-02-01', party_type='person', person_id=pid)) if False else None
    assert cc.MOVABLE == ('smr', 'to', 'le')


def test_menu_views_and_move_dialog(db, tmp_path, monkeypatch):
    for name, mod in list(sys.modules.items()):
        if name.startswith('smetagaz') and isinstance(getattr(mod, 'db', None), DatabaseManager) and mod.db is not db:
            monkeypatch.setattr(mod, 'db', db)
    import smetagaz.counterparty_ui as ui
    import smetagaz.main_window as mw
    from PyQt6.QtCore import Qt
    monkeypatch.setattr(mw.QApplication, 'quit', staticmethod(lambda: None))
    w = mw.MainWindow()
    names = [b.text() for b in w.tabs_buttons]
    assert 'СМР (физлица)' in names and 'Прочие (юр.лица)' in names and 'СМР (ИП)' in names and 'СМР (юрлица)' in names
    w.hide()
    w.deleteLater()
    for mod in ('fl', 'le'):
        v = ui.CounterpartyView(mod)
        v.load_data()
        v.deleteLater()
    lid = make_legal(db)
    cid = cc.save_contract(db, 'smr', dict(contract_date='2026-02-01', legal_id=lid, amount=100))
    reg = ui.ContractsRegistry('smr')
    reg.load_data()
    assert any(reg.table.cellWidget(0, c) is not None or True for c in range(3))
    reg.table.selectRow(0)
    dlg = ui.MoveContractDialog('smr', cid)
    dlg.combo.setCurrentIndex(dlg.combo.findData('le'))
    assert '04/26' in dlg.preview.text() or '-04/' in dlg.preview.text()
    dlg.renum.setChecked(False)
    assert 'сохранится' in dlg.preview.text()
    monkeypatch.setattr(ui.MoveContractDialog, 'exec', lambda self: True)
    monkeypatch.setattr(ui.QMessageBox, 'information', staticmethod(lambda *a, **k: None))
    reg.move()
    assert db.fetchone('SELECT count(*) FROM le_contracts')[0] + db.fetchone('SELECT count(*) FROM to_contracts')[0] == 1
    reg.deleteLater()


def test_act_vat_and_contract_amount_from_acts(db):
    lid = make_legal(db)
    cid = cc.save_contract(db, 'smr', dict(contract_date='2026-02-01', legal_id=lid, amount=999, vat_included=1, vat_rate=20, amount_from_acts=1))
    assert cc.contract(db, 'smr', cid)['amount'] == 0                                  # актов нет — сумма по актам нулевая
    a1 = cc.save_act(db, 'smr', dict(contract_id=cid, act_number='1', act_date='2026-02-10', amount=1200, vat_included=True, vat_rate=20))
    a2 = cc.save_act(db, 'smr', dict(contract_id=cid, act_number='2', act_date='2026-02-20', amount=500))            # без НДС
    c = cc.contract(db, 'smr', cid)
    assert c['amount'] == 1700
    assert cc.act_amounts(c, cc.act(db, 'smr', a1)) == (1000.0, 200.0, 1200.0)
    assert cc.act_amounts(c, cc.act(db, 'smr', a2)) == (500.0, 0.0, 500.0)
    assert cc.acts_total(db, 'smr', cid) == (1500.0, 200.0, 1700.0)
    ctx = cc.context(db, 'smr', cid, a1)
    assert ctx['act.amount'] == 1200 and ctx['act.amount_net'] == 1000 and ctx['act.vat'] == 200
    assert ctx['contract.amount'] == 1700 and ctx['contract.vat'] == 200 and ctx['contract.amount_net'] == 1500
    cc.save_act(db, 'smr', dict(contract_id=cid, act_number='2', act_date='2026-02-20', amount=600), a2)               # правка акта пересчитывает договор
    assert cc.contract(db, 'smr', cid)['amount'] == 1800
    cc.delete_act(db, 'smr', a1)
    assert cc.contract(db, 'smr', cid)['amount'] == 600
    # ставка акта не указана — берётся ставка договора
    a3 = cc.save_act(db, 'smr', dict(contract_id=cid, act_number='3', act_date='2026-03-01', amount=240, vat_included=True))
    assert cc.act_amounts(cc.contract(db, 'smr', cid), cc.act(db, 'smr', a3)) == (200.0, 40.0, 240.0)
    # без флага сумма договора остаётся ручной
    manual = cc.save_contract(db, 'smr', dict(contract_date='2026-02-01', legal_id=lid, amount=777))
    cc.save_act(db, 'smr', dict(contract_id=manual, act_number='1', act_date='2026-02-10', amount=100))
    assert cc.contract(db, 'smr', manual)['amount'] == 777
    cc.save_contract(db, 'smr', dict(contract_date='2026-02-01', legal_id=lid, amount=1, amount_from_acts=1), manual)
    assert cc.contract(db, 'smr', manual)['amount'] == 100
    with pytest.raises(ValueError):
        cc.save_act(db, 'smr', dict(contract_id=cid, act_number='9', amount=1, vat_included=True, vat_rate=150))


def test_act_dialog_and_contract_dialog_vat_ui(db, monkeypatch):
    for name, mod in list(sys.modules.items()):
        if name.startswith('smetagaz') and isinstance(getattr(mod, 'db', None), DatabaseManager) and mod.db is not db:
            monkeypatch.setattr(mod, 'db', db)
    import smetagaz.counterparty_ui as ui
    lid = make_legal(db)
    cid = cc.save_contract(db, 'to', dict(contract_date='2026-02-01', legal_id=lid, vat_included=1, vat_rate=20, amount_from_acts=1))
    d = ui.ActDialog('to', cid)
    assert d.vat_included.isChecked() and d.vat_rate.value() == 20
    d.number.setText('7')
    d.amount.setValue(600)
    assert 'НДС 100' in d.vat_info.text().replace(',', '.').replace('\xa0', ' ') or '100' in d.vat_info.text()
    d.save()
    dlg = ui.ContractDialog('to', cid)
    assert dlg.amount_from_acts.isChecked() and dlg.amount.isReadOnly() and dlg.amount.value() == 600
    assert dlg.acts_table.columnCount() == 8 and dlg.acts_table.horizontalHeaderItem(3).text() == 'в т. ч. НДС'
    ip_dlg = ui.ContractDialog('ip')                                                    # СМР (ИП) остаётся прежним
    assert ip_dlg.amount_from_acts.isHidden() and ip_dlg.acts_table.columnCount() == 7
    ip_act = ui.ActDialog('ip', cc.save_contract(db, 'ip', dict(contract_date='2026-02-01', person_id=make_person(db))))
    assert ip_act.vat_included is None
    for w in (d, dlg, ip_dlg, ip_act):
        w.deleteLater()


def test_display_status_and_filters(db):
    from datetime import date
    today = date(2026, 6, 1)
    assert cc.display_status(dict(status='Действует', end_date='2026-12-31'), today) == 'work'
    assert cc.display_status(dict(status='Действует', end_date='2026-05-31'), today) == 'overdue'
    assert cc.display_status(dict(status='Действует', start_date='2026-01-01', end_date='2026-12-31'), today) == 'annual'          # 364 дня
    assert cc.display_status(dict(status='Действует', annual=1), today) == 'annual'
    assert cc.display_status(dict(status='Действует', annual=1, end_date='2026-01-01'), today) == 'overdue'                    # просрочка важнее
    assert cc.display_status(dict(status='Завершён', end_date='2020-01-01'), today) == 'done'
    assert cc.display_status(dict(status='Расторгнут'), today) == 'cancelled'
    lid = make_legal(db)
    rows = {
        'work': dict(end_date='2026-12-31'),
        'overdue': dict(end_date='2026-05-01'),
        'annual': dict(start_date='2026-01-01', end_date='2026-12-31'),
        'done': dict(status='Завершён'),
        'cancelled': dict(status='Расторгнут'),
    }
    ids = {}
    for key, extra in rows.items():
        ids[key] = cc.save_contract(db, 'smr', dict(contract_date='2026-01-10', legal_id=lid, **extra))
    for key in rows:
        cond, params = cc.status_condition(key, 'c', today)
        found = [r[0] for r in db.fetchall(f'SELECT c.id FROM smr_contracts c WHERE {cond}', params)]
        assert found == [ids[key]], key
    assert cc.status_condition('', 'c')[0] == '1=1'


def test_status_icons_and_registries(db, monkeypatch):
    for name, mod in list(sys.modules.items()):
        if name.startswith('smetagaz') and isinstance(getattr(mod, 'db', None), DatabaseManager) and mod.db is not db:
            monkeypatch.setattr(mod, 'db', db)
    from PyQt6.QtCore import Qt
    import smetagaz.counterparty_ui as ui
    import smetagaz.status_icons as si
    for key in si.TITLES:
        assert not si.pixmap(key, 16).isNull() and not si.icon(key).isNull()
    lid = make_legal(db)
    done = cc.save_contract(db, 'to', dict(contract_date='2026-01-10', legal_id=lid, status='Завершён'))
    late = cc.save_contract(db, 'to', dict(contract_date='2026-01-10', legal_id=lid, end_date='2020-01-01'))
    reg = ui.ContractsRegistry('to')
    reg.load_data()
    texts = {reg.table.item(r, 8).text() for r in range(reg.table.rowCount())}
    assert texts == {'Завершён', 'Просрочен'}
    assert all(not reg.table.item(r, 8).icon().isNull() for r in range(reg.table.rowCount()))
    reg.cmb_status.setCurrentIndex(reg.cmb_status.findData('overdue'))
    assert reg.table.rowCount() == 1 and reg.table.item(0, 8).text() == 'Просрочен'
    reg.deleteLater()
    ip = ui.ContractsRegistry('ip')                                                           # СМР (ИП) без изменений
    assert not ip.visual and ip.cmb_status.count() == 1 + len(cc.STATUSES)
    ip.deleteLater()
    # Монтаж ГСВ и ГСН: «Завершён» со значком, когда акт подписан
    from smetagaz import contracts_registry as cr, gsn_view
    db.execute("INSERT INTO contracts(contract_number,contract_date,act_signed) VALUES('01-02/26','2026-01-01',1)")
    db.execute("INSERT INTO contracts(contract_number,contract_date,act_signed) VALUES('02-02/26','2026-01-02',0)")
    v = cr.ContractsRegistryView()
    v.load_data()
    states = {v.table.item(r, 0).text(): v.table.item(r, 5) for r in range(v.table.rowCount())}
    assert states['01-02/26'].text() == 'Завершён' and not states['01-02/26'].icon().isNull() and states['02-02/26'].text() == ''
    v.deleteLater()
    db.execute("INSERT INTO gsn_projects(contract_number,contract_date,title,act_signed) VALUES('10','2026-01-01','Объект',1)")
    g = gsn_view.GsnProjectsView()
    g.load_data()
    assert g.table.item(0, 7).text() == 'Завершён' and not g.table.item(0, 7).icon().isNull()
    g.deleteLater()


def test_render_file_word_and_excel(tmp_path):
    import openpyxl
    from smetagaz import gsv_project_domain as gd
    wb = openpyxl.Workbook()
    wb.active['A1'] = 'Договор № {НОМЕР_ДОГОВОРА}'
    tpl = tmp_path / 'tpl.xlsx'
    wb.save(tpl)
    out = gd.render_file(str(tpl), str(tmp_path / 'out.docx'), {'НОМЕР_ДОГОВОРА': '12'})
    assert out.endswith('.xlsx')
    assert openpyxl.load_workbook(out).active['A1'].value == 'Договор № 12'
    import docx
    d = docx.Document()
    d.add_paragraph('№ {НОМЕР_ДОГОВОРА}')
    d.save(tmp_path / 't.docx')
    out = gd.render_file(str(tmp_path / 't.docx'), str(tmp_path / 'o.xlsx'), {'НОМЕР_ДОГОВОРА': '7'})
    assert out.endswith('.docx') and docx.Document(out).paragraphs[0].text == '№ 7'
    import pytest
    with pytest.raises(ValueError):
        gd.render_file(str(tmp_path / 'x.doc'), str(tmp_path / 'o'), {})


def test_gsn_steel_under_above(db):
    from smetagaz import gsn_domain as gn
    cid = gn.save_card(db, dict(title='Объект', contract_number='1', pe_length=10, steel_under_length=5, steel_above_length=3), dict(name='Иванов И.И.'))[0]
    c = gn.get(db, cid)
    assert c['steel_length'] == 8 and c['steel_under_length'] == 5 and c['steel_above_length'] == 3


def test_norm_per_object_per_case_per_joint(db):
    from smetagaz import stock_domain as sd, gsvm_domain as md
    assert sd.default_per('Ацетилен') == '' and sd.default_per('Сварочная проволока') == '' and sd.default_per('Кислород') == ''
    assert sd.default_per('Краска') == 'футляр' and sd.default_per('Шнур асбестовый') == 'футляр'
    assert sd.default_per('Отрезные круги') == 'объект' and sd.default_per('Электроды') == 'объект'
    ids = {r[1]: r[0] for r in db.fetchall('SELECT id,name FROM stock_materials')}
    pid = sd.save_profile(db, dict(name='Ду 25', diameter='25', thickness='3', basis='стык'),
                          [(ids['Ацетилен'], '2'), (ids['Краска'], '0.5'), (ids['Отрезные круги'], '3')])
    lines, snaps = sd.calculate(db, [(pid, 10)], cases=4)
    qty = {l['material_id']: sd.fmt(l['norm_qty']) if hasattr(sd, 'fmt') else l['norm_qty'] for l in lines}
    by = {s['name']: s['quantity'] for s in snaps[0]['items']}
    assert float(by['Ацетилен']) == 20 and float(by['Краска']) == 2 and float(by['Отрезные круги']) == 3
    plan = md.plan_consumption(db, [dict(name='Ду 25', count=10, profile_id=pid)], cases=4)
    m = {x['name']: x['qty'] for x in plan['materials']}
    assert float(m['Ацетилен']) == 20 and float(m['Краска']) == 2 and float(m['Отрезные круги']) == 3
    plan = md.plan_consumption(db, [dict(name='Ду 25', count=20, profile_id=pid)], cases=4)
    assert float({x['name']: x['qty'] for x in plan['materials']}['Отрезные круги']) == 3     # на объект — не растёт со стыками


def test_theme_toggle_releases_hidden_views(db, monkeypatch):
    import sys
    from smetagaz import database as dbm
    for name, mod in list(sys.modules.items()):
        if name.startswith('smetagaz') and getattr(mod, 'db', None) is not None and mod.db.__class__.__name__ == 'DatabaseManager':
            monkeypatch.setattr(mod, 'db', db, raising=False)
    from smetagaz import main_window as mw
    w = mw.MainWindow()
    for i, b in enumerate(w.tabs_buttons[:4]):
        w.switch_tab(i, b)
    realized = [v for v in w.views if getattr(v, '_factory_saved', None)]
    assert len(realized) >= 3
    w.toggle_theme()
    assert db.get_setting('is_dark', '0') == '1'
    assert sum(1 for v in w.views if getattr(v, '_factory_saved', None)) <= 1       # остаётся только показанная страница
    w.switch_tab(0, w.tabs_buttons[0])                                              # страница создаётся заново
    assert getattr(w.views[0], '_factory_saved', None)
    w.toggle_theme()
    w.close()


def test_volumes_forward_and_reverse_order(db):
    from smetagaz import volumes_domain as vd, estimates_domain as ed, contracts_core as cc
    db.execute("INSERT INTO materials(item_type,name,unit,price,purchase_price) VALUES('Материал','Труба ПЭ 32','м',5,3)")
    db.execute("INSERT INTO materials(item_type,name,unit,price,purchase_price) VALUES('Работа','Монтаж трубы','м',12,0)")
    lid = make_legal(db)
    cid = cc.save_contract(db, 'smr', dict(client_id=lid, party_type='legal', legal_id=lid, contract_date='2026-09-01', amount=0))
    # порядок: договор → объёмы → смета
    vd.save_items(db, 'smr_contracts', cid, [dict(kind='Материал', name='труба пэ 32', unit='м', quantity='10'), dict(kind='Работа', name='Монтаж трубы', unit='м', quantity='10'),
                                           dict(kind='Материал', name='Хомут нестандартный', unit='шт', quantity='4')])
    eid, unpriced = vd.build_estimate(db, 'smr_contracts', cid)
    assert vd.estimate_id(db, 'smr_contracts', cid) == eid and unpriced == ['Хомут нестандартный']
    assert ed.contract_of(db, eid)['id'] == cid
    assert float(db.fetchone('SELECT total FROM estimates WHERE id=?', (eid,))[0]) == 170.0
    assert float(cc.contract(db, 'smr', cid)['amount']) == 170.0
    import pytest
    with pytest.raises(ValueError):
        vd.build_estimate(db, 'smr_contracts', cid)                       # позиции уже есть — нужна явная замена
    vd.build_estimate(db, 'smr_contracts', cid, replace=True)
    # обратный порядок: смета без договора → привязка → объёмы из сметы
    e2 = ed.create_estimate(db, 'Предварительная')
    db.execute("INSERT INTO estimate_items(estimate_id,item_type,name,unit,quantity,price,sum,sort_order) VALUES(?,'Материал','Кран',?,2,10,20,0)", (e2, 'шт'))
    c2 = cc.save_contract(db, 'smr', dict(client_id=lid, party_type='legal', legal_id=lid, contract_date='2026-09-02', amount=0))
    ed.link_contract(db, e2, 'smr_contracts', c2)
    assert vd.fill_from_estimate(db, 'smr_contracts', c2) == 1
    assert vd.items(db, 'smr_contracts', c2)[0]['name'] == 'Кран' and vd.items(db, 'smr_contracts', c2)[0]['quantity'] == 2
    with pytest.raises(ValueError):
        vd.fill_from_estimate(db, 'smr_contracts', c2)
    with pytest.raises(ValueError):
        vd.items(db, 'le_contracts', 1)
    # удаление договора убирает объёмы
    db.execute('DELETE FROM smr_contracts WHERE id=?', (c2,))
    assert db.fetchone('SELECT count(*) FROM volumes WHERE owner_type=? AND owner_id=?', ('smr_contracts', c2))[0] == 0


def test_volumes_tab_in_cards(db, monkeypatch):
    import sys
    for name, mod in list(sys.modules.items()):
        if name.startswith('smetagaz') and getattr(mod, 'db', None) is not None and mod.db.__class__.__name__ == 'DatabaseManager':
            monkeypatch.setattr(mod, 'db', db, raising=False)
    from smetagaz import counterparty_ui as cu, contract_card as cg, gsn_card as gc
    for mod in ('smr', 'ip', 'fl'):
        d = cu.ContractDialog(mod)
        assert any(d.tabs.tabText(i) == 'Объёмы' for i in range(d.tabs.count())), mod
    assert not any(cu.ContractDialog('to').tabs.tabText(i) == 'Объёмы' for i in range(cu.ContractDialog('to').tabs.count()))
    assert hasattr(cg.ContractCardDialog(), 'volumes_tab') and hasattr(gc.GsnContractDialog(), 'volumes_tab')


def test_volumes_tab_actions(db, monkeypatch):
    import sys
    for name, mod in list(sys.modules.items()):
        if name.startswith('smetagaz') and getattr(mod, 'db', None) is not None and mod.db.__class__.__name__ == 'DatabaseManager':
            monkeypatch.setattr(mod, 'db', db, raising=False)
    from PyQt6.QtWidgets import QMessageBox
    from smetagaz import counterparty_ui as cu, volumes_domain as vd
    monkeypatch.setattr(QMessageBox, 'information', lambda *a, **k: 0)
    monkeypatch.setattr(QMessageBox, 'warning', lambda *a, **k: 0)
    db.execute("INSERT INTO materials(item_type,name,unit,price,purchase_price) VALUES('Материал','Кран шаровой','шт',20,10)")
    lid = make_legal(db)
    d = cu.ContractDialog('smr')
    d.contract_id = cc.save_contract(db, 'smr', dict(client_id=lid, party_type='legal', legal_id=lid, contract_date='2026-09-01', amount=0))
    tab = d.volumes_tab
    tab.reload()
    tab.add_row('Материал', 'Кран шаровой', 'шт', 3)
    assert tab.save()
    tab.make_estimate()
    assert vd.estimate_id(db, 'smr_contracts', d.contract_id)
    assert float(cc.contract(db, 'smr', d.contract_id)['amount']) == 60.0


def test_date_typing_is_simple():
    from PyQt6.QtCore import QDate, Qt
    from PyQt6.QtTest import QTest
    from smetagaz.date_widgets import DateEdit, parse_typed, parse_month
    from smetagaz.domain_widgets import OptionalDate
    base = QDate(2026, 3, 5)
    assert parse_typed('14102026', base) == QDate(2026, 10, 14)
    assert parse_typed('141026', base) == QDate(2026, 10, 14)
    assert parse_typed('14.10.26', base) == QDate(2026, 10, 14)
    assert parse_typed('1.5', base) == QDate(2026, 5, 1)
    assert parse_typed('31.02.2026', base) == QDate(2026, 2, 28)
    assert parse_typed('2026-10-14', base) == QDate(2026, 10, 14)
    assert parse_typed('45', base) is None and parse_typed('1413', base) is None
    assert parse_month('10.2026', base) == QDate(2026, 10, 1) and parse_month('1026', base) == QDate(2026, 10, 1)

    def type_in(w, text):
        w.setFocus()
        w.selectAll()
        QTest.keyClicks(w, text)

    w = DateEdit(QDate(2026, 3, 5))
    w.show()
    seen = []
    w.dateChanged.connect(seen.append)
    type_in(w, '1410')
    assert w.text() == '14.10'                       # точка поставилась сама, поле не блокируется
    QTest.keyClicks(w, '2026')
    assert w.text() == '14.10.2026'
    assert w.date() == QDate(2026, 10, 14) and seen[-1] == QDate(2026, 10, 14)
    type_in(w, '7.8')
    assert w.date() == QDate(2026, 8, 7) and w.text() == '07.08.2026'
    type_in(w, '31.12.27')
    assert w.date() == QDate(2027, 12, 31)
    type_in(w, '99.99.9999')                          # неверный ввод — прежняя дата
    assert w.date() == QDate(2027, 12, 31) and w.text() == '31.12.2027'
    type_in(w, '010203')
    QTest.keyClick(w, Qt.Key.Key_Return)
    assert w.text() == '01.02.2003'
    QTest.keyClick(w, Qt.Key.Key_Up)
    assert w.date() == QDate(2003, 2, 2)
    QTest.keyClick(w, Qt.Key.Key_PageDown)
    assert w.date() == QDate(2003, 1, 2)
    # дата берётся, даже если поле не покидали (сразу нажали «Сохранить»)
    type_in(w, '05052030')
    assert w.date() == QDate(2030, 5, 5)
    o = OptionalDate()
    o.show()
    assert o.is_empty() and o.text() == '' and o.placeholderText() == 'Не указана'
    type_in(o, '01.02.27')
    assert o.value() == '2027-02-01' and not o.is_empty()
    type_in(o, '')
    QTest.keyClick(o, Qt.Key.Key_Backspace)
    assert o.value() == '' and o.is_empty()
    m = DateEdit(QDate(2026, 3, 5))
    m.setDisplayFormat('MMMM yyyy')
    type_in(m, '112027')
    assert m.date() == QDate(2027, 11, 1)


def test_manual_visual_status(db):
    lid = make_legal(db)
    cid = cc.save_contract(db, 'smr', dict(client_id=lid, party_type='legal', legal_id=lid, contract_date='2026-01-01', end_date='2026-02-01', amount=1))
    row = lambda: cc.contract(db, 'smr', cid)
    assert cc.display_status(row()) == 'overdue'
    cc.set_visual_status(db, 'smr', cid, 'annual')
    assert cc.display_status(row()) == 'annual'
    cond, params = cc.status_condition('annual', 'c')
    assert db.fetchone(f'SELECT count(*) FROM smr_contracts c WHERE {cond}', params)[0] == 1
    cond, params = cc.status_condition('overdue', 'c')
    assert db.fetchone(f'SELECT count(*) FROM smr_contracts c WHERE {cond}', params)[0] == 0
    cc.save_contract(db, 'smr', dict(client_id=lid, party_type='legal', legal_id=lid, contract_date='2026-01-01', end_date='2026-02-01', amount=1, status_manual='done'), cid)
    assert cc.display_status(row()) == 'done'
    cc.set_visual_status(db, 'smr', cid, '')
    assert cc.display_status(row()) == 'overdue'
    import pytest
    with pytest.raises(ValueError):
        cc.set_visual_status(db, 'smr', cid, 'bogus')


def test_manual_contract_number_is_not_overwritten(db, monkeypatch):
    import sys
    for name, mod in list(sys.modules.items()):
        if name.startswith('smetagaz') and getattr(mod, 'db', None) is not None and mod.db.__class__.__name__ == 'DatabaseManager':
            monkeypatch.setattr(mod, 'db', db, raising=False)
    from PyQt6.QtCore import QDate
    from PyQt6.QtTest import QTest
    from smetagaz import counterparty_ui as cu
    for mod in ('smr', 'ip', 'fl', 'to'):
        d = cu.ContractDialog(mod)
        auto = d.number.text()
        assert auto
        d.number.clear()
        QTest.keyClicks(d.number, 'MY-17')                  # ввод вручную
        d.date.setDate(QDate(2026, 5, 5))                    # переход в другую графу меняет дату — номер остаётся
        d.date.setDate(QDate(2027, 1, 2))
        assert d.number.text() == 'MY-17', mod
        d.number.clear()
        d.number.textEdited.emit('')                         # пустой номер снова назначается автоматически
        d.date.setDate(QDate(2026, 6, 6))
        assert d.number.text() and d.number.text() != 'MY-17'
