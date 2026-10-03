import os
from datetime import date
import pytest
import docx
import openpyxl
from smetagaz.database import DatabaseManager
from smetagaz import gsv_project_domain as g
from smetagaz import payments_domain


@pytest.fixture
def d(tmp_path, monkeypatch):
    monkeypatch.setattr(g, 'PROJECTS_DIR', tmp_path / 'projects')
    monkeypatch.setattr(g, 'TEMPLATES_DIR', tmp_path / 'templates')
    d = DatabaseManager(tmp_path / 'smetagaz.db')
    d.init_db()
    yield d
    d.close()


def make_project(d, **extra):
    cid = d.execute("INSERT INTO crm.clients(name,phone,address,passport,passport_issuer,passport_date) VALUES('Иванов Иван Иванович','375291112233 (основной); 375447778899 (жена)','г. Минск, ул. Лесная 5','МР 1234567','Фрунзенский РУВД','2015-03-04')").lastrowid
    cols = dict(pd_number='01-26 ГСВ', seq_num=1, year_num=26, object_name='Жилой дом', address='д. Ключи', client_name='Иванов Иван Иванович', client_id=cid,
                contract_number='01-03/26', contract_date='2026-11-12', due_date='2026-12-12', act_date='2026-12-10', cost=250, notes='Срочно',
                contract_signed=1, act_signed=1)
    cols.update(extra)
    return d.execute(f"INSERT INTO gsv_projects({','.join(cols)}) VALUES({','.join('?' for _ in cols)})", tuple(cols.values())).lastrowid


def test_defaults_statuses_and_migration(d):
    assert [r[1] for r in g.catalog(d, 'work')] == ['Не сделано', 'В работе', 'Не хватает данных', 'Сделано']
    assert [r[1] for r in g.catalog(d, 'client')] == ['Договор не подписан', 'Договор подписан', 'Акт подписан', 'Передано клиенту', 'Не оплачено', 'Оплачено']


def test_multiple_statuses_per_project_and_text_mirror(d):
    pid = make_project(d)
    ids = {r[1]: r[0] for kind in ('work', 'client') for r in g.catalog(d, kind)}
    g.set_project_statuses(d, pid, [ids['В работе'], ids['Не хватает данных'], ids['Договор подписан'], ids['Не оплачено']])
    assert g.project_statuses(d, pid, 'work') == ['В работе', 'Не хватает данных']
    assert d.fetchone('SELECT work_status,client_status FROM gsv_projects WHERE id=?', (pid,)) == ('В работе, Не хватает данных', 'Договор подписан, Не оплачено')
    with pytest.raises(Exception):
        d.execute('DELETE FROM gsv_status_catalog WHERE id=?', (ids['В работе'],))   # статус используется — удалить нельзя


def test_phones_roundtrip():
    text = '375291112233 (основной); 375447778899 (жена)'
    assert g.parse_phones(text) == [('375291112233', 'основной'), ('375447778899', 'жена')]
    assert g.join_phones(g.parse_phones(text)) == text
    assert g.parse_phones('') == [] and g.parse_phones('123') == [('123', '')]


def test_tags_formats(d):
    pid = make_project(d)
    payments_domain.add(d, 'gsv_projects', pid, 100, '2026-11-15', 'аванс')
    payments_domain.add(d, 'gsv_projects', pid, 50.5, '2026-12-01', 'доплата')
    t = g.project_tags(d, pid)
    assert t['ДАТА_ЗАКЛЮЧЕНИЯ_П'] == '12 ноября 2026г.' and t['ДАТА_ЗАКЛЮЧЕНИЯ_С'] == '12.11.2026г.'
    assert t['СРОК_ИСПОЛНЕНИЯ_П'] == '12 декабря 2026г.' and t['ДАТА_АКТА_С'] == '10.12.2026г.'
    assert t['КЛИЕНТ'] == 'Иванов Иван Иванович' and t['КЛИЕНТ_СОКР'] == 'И.И. Иванов'
    assert t['СТОИМОСТЬ'] == '250,00' and t['СУММА_ПРОПИСЬЮ'].lower().startswith('двести пятьдесят')
    assert t['ПАСПОРТ'] == 'МР 1234567' and t['КЕМ_ВЫДАН'] == 'Фрунзенский РУВД' and t['ДАТА_ВЫДАЧИ'] == '04.03.2015'
    assert t['ОПЛАЧЕНО'] == '150,50' and t['ОСТАТОК'] == '99,50'
    assert t['ДАТА_ОПЛАТЫ'].splitlines() == ['15.11.2026', '01.12.2026'] and t['ПРИМЕЧАНИЕ_ОПЛАТЫ'].splitlines() == ['аванс', 'доплата']
    assert t['ДОГОВОР_ПОДПИСАН'] == 'Да'


def test_default_due_is_30_days():
    assert g.default_due('2026-11-12') == date(2026, 12, 12)
    assert g.default_due(date(2026, 1, 31)) == date(2026, 3, 2)


def test_word_documents_and_regeneration(d):
    pid = make_project(d)
    assert g.doc_state(d, pid, 'contract') == 'none'
    path = g.generate(d, pid, 'contract')
    text = '\n'.join(p.text for p in docx.Document(path).paragraphs)
    assert '01-03/26' in text and '12 ноября 2026г.' in text and 'Иванов Иван Иванович' in text and '{' not in text
    assert os.path.dirname(path).endswith('01-26 ГСВ Иванов Иван Иванович') or 'Проекты ГСВ' in path
    assert g.doc_state(d, pid, 'contract') == 'fresh'
    d.execute("UPDATE gsv_projects SET cost=300 WHERE id=?", (pid,))
    assert g.doc_state(d, pid, 'contract') == 'stale'             # данные изменились — нужна кнопка «переформировать»
    g.generate(d, pid, 'contract')
    assert g.doc_state(d, pid, 'contract') == 'fresh'
    payments_domain.add(d, 'gsv_projects', pid, 10, '2026-11-20', '')
    assert g.doc_state(d, pid, 'contract') == 'stale'             # оплата тоже меняет документ
    os.remove(path)
    assert g.doc_state(d, pid, 'contract') == 'missing'


def test_act_requires_date(d):
    pid = make_project(d, act_date=None, act_signed=0)
    with pytest.raises(ValueError, match='дату акта'):
        g.generate(d, pid, 'act')
    d.execute("UPDATE gsv_projects SET act_date='2026-12-10' WHERE id=?", (pid,))
    text = '\n'.join(p.text for p in docx.Document(g.generate(d, pid, 'act')).paragraphs)
    assert '10 декабря 2026г.' in text


def test_custom_template_tags_with_spaces_and_case(d, tmp_path):
    pid = make_project(d)
    tpl = tmp_path / 'my.docx'
    doc = docx.Document()
    doc.add_paragraph('Выдан {КЕМ ВЫДАН} {дата выдачи}; адрес {Адрес клиента}; неизвестный {НЕТ_ТАКОГО}')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'Статус: {СТАТУС КЛИЕНТА}'
    doc.save(tpl)
    d.set_setting('gsvp_tpl_contract', str(tpl))
    ids = {r[1]: r[0] for r in g.catalog(d, 'client')}
    g.set_project_statuses(d, pid, [ids['Договор подписан'], ids['Оплачено']])
    out = docx.Document(g.generate(d, pid, 'contract'))
    assert out.paragraphs[0].text == 'Выдан Фрунзенский РУВД 04.03.2015; адрес г. Минск, ул. Лесная 5; неизвестный {НЕТ_ТАКОГО}'
    assert out.tables[0].cell(0, 0).text == 'Статус: Договор подписан, Оплачено'


def test_excel_card_has_column_per_payment(d):
    pid = make_project(d)
    for day, amount, note in (('2026-11-15', 100, 'аванс'), ('2026-12-01', 50.5, 'доплата'), ('2027-01-10', 99.5, 'остаток')):
        payments_domain.add(d, 'gsv_projects', pid, amount, day, note)
    ws = openpyxl.load_workbook(g.generate(d, pid, 'card')).active
    # шаблон по умолчанию: блок из трёх столбцов повторяется для каждой оплаты
    header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'Дата оплаты')
    values = [ws.cell(header_row + 1, c).value for c in range(1, 10)]
    assert values == ['15.11.2026', 100.0, 'аванс', '01.12.2026', 50.5, 'доплата', '10.01.2027', 99.5, 'остаток']
    found = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert found['Клиент'] == 'Иванов Иван Иванович' and found['Объект'] == 'Жилой дом'


def test_excel_card_without_payments(d):
    pid = make_project(d)
    ws = openpyxl.load_workbook(g.generate(d, pid, 'card')).active
    header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'Дата оплаты')
    assert [ws.cell(header_row + 1, c).value for c in (1, 2, 3)] == [None, None, None] or ws.cell(header_row + 1, 1).value in ('', None)


def test_export_all_cards_overwrites_file(d, tmp_path):
    pid = make_project(d)
    payments_domain.add(d, 'gsv_projects', pid, 100, '2026-11-15', 'аванс')
    make_project(d, pd_number='02-26 ГСВ', seq_num=2, contract_number='02-03/26', client_name='Петров')
    path = tmp_path / 'out' / 'cards.xlsx'
    assert g.export_cards(d, path) == 2
    assert g.export_cards(d, path) == 2
    ws = openpyxl.load_workbook(path).active
    assert [c.value for c in ws[1]] == g.EXPORT_HEADERS and ws.max_row == 3
    row = {h: ws.cell(3, i + 1).value for i, h in enumerate(g.EXPORT_HEADERS)}
    assert row['№ ПД'] == '01-26 ГСВ' and row['Оплачено'] == 100.0 and row['Остаток'] == 150.0 and 'аванс' in row['Оплаты']


def test_calendar_marks_projects_with_client_name(d):
    from smetagaz import agenda_domain as a
    make_project(d, contract_signed=0)
    events = a.events_between(d, date(2026, 11, 12), date(2026, 11, 12), date(2026, 11, 1))
    assert [(e['kind'], e['title']) for e in events] == [('contract', 'Заключение договора Иванов И.И. (не подписан)')]
    acts = a.events_between(d, date(2026, 12, 10), date(2026, 12, 10), date(2026, 11, 1))
    assert acts[0]['title'] == 'Подписан акт Иванов И.И.'
