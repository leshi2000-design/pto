import docx
from docx.shared import Pt, RGBColor
from smetagaz import gsv_project_domain as gd


def test_template_formatting_is_preserved(tmp_path):
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run('Договор № ').bold = True
    r = p.add_run('{НОМЕР_')
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(255, 0, 0)
    p.add_run('ДОГОВОРА} от ').italic = True
    p.add_run('{ДАТА}, ').underline = True
    p.add_run('{{ДВОЙНОЙ}} и {НЕИЗВЕСТНЫЙ}')
    t = d.add_table(1, 1)
    t.cell(0, 0).paragraphs[0].add_run('в таблице {КЛИЕНТ}').bold = True
    d.sections[0].header.paragraphs[0].text = 'Шапка {КЛИЕНТ}'
    src, out = tmp_path / 't.docx', tmp_path / 'o.docx'
    d.save(src)
    gd.render_docx(str(src), str(out), {'НОМЕР_ДОГОВОРА': '7/26', 'ДАТА': '01.01.2026', 'КЛИЕНТ': 'Иванов', 'ДВОЙНОЙ': 'ок'})
    o = docx.Document(str(out))
    runs = o.paragraphs[0].runs
    assert [r.text for r in runs] == ['Договор № ', '7/26', ' от ', '01.01.2026, ', 'ок и {НЕИЗВЕСТНЫЙ}']
    assert runs[0].bold and runs[1].font.size == Pt(20) and runs[1].font.color.rgb == RGBColor(255, 0, 0) and runs[2].italic and runs[3].underline
    assert not runs[1].bold and not runs[2].bold
    cell_run = o.tables[0].cell(0, 0).paragraphs[0].runs[0]
    assert cell_run.text == 'в таблице Иванов' and cell_run.bold
    assert o.sections[0].header.paragraphs[0].text == 'Шапка Иванов'
