from datetime import date
import openpyxl
import pytest
from smetagaz.database import DatabaseManager
from smetagaz import acts_statement as st, payments_domain


@pytest.fixture
def d(tmp_path):
    d = DatabaseManager(tmp_path / 'smetagaz.db')
    d.init_db()
    yield d
    d.close()


def test_default_period():
    assert st.default_period(date(2026, 10, 7)) == (2026, 9)          # до 10-го — за прошлый месяц
    assert st.default_period(date(2026, 10, 10)) == (2026, 9)
    assert st.default_period(date(2026, 10, 11)) == (2026, 10)
    assert st.default_period(date(2027, 1, 5)) == (2026, 12)


def data(d, tmp_path):
    act = tmp_path / 'акт.docx'
    act.write_bytes(b'x')
    p1 = d.execute("INSERT INTO gsv_projects(pd_number,contract_number,contract_date,client_name,object_name,address,cost,act_date,act_signed) VALUES('01-26 ГСВ','01-03/26','2026-08-01','Иванов','Дом','Минск',250,'2026-09-05',1)").lastrowid
    d.execute("INSERT INTO gsv_projects(pd_number,contract_number,client_name,cost,act_date,act_signed) VALUES('02-26 ГСВ','02-03/26','Петров',300,'2026-09-30',1)")
    d.execute("INSERT INTO gsv_projects(pd_number,contract_number,client_name,cost,act_date,act_signed) VALUES('03-26 ГСВ','03-03/26','Сидоров',100,'2026-09-12',0)")
    d.execute("INSERT INTO gsv_projects(pd_number,contract_number,client_name,cost,act_date,act_signed) VALUES('04-26 ГСВ','04-03/26','Другой месяц',100,'2026-10-01',1)")
    d.execute("INSERT INTO gsv_projects(pd_number,contract_number,client_name,cost,act_date,act_signed) VALUES('05-26 ГСВ','05-03/26','Без акта',100,NULL,0)")
    d.execute("INSERT INTO gsv_project_docs(project_id,kind,file_path,data_hash,generated_at) VALUES(?,?,?,?,?)", (p1, 'act', str(act), '', ''))
    payments_domain.add(d, 'gsv_projects', p1, 100, '2026-09-06', '')
    c = d.execute("INSERT INTO contracts(contract_number,contract_date,client_name,object_name,object_address,contract_amount,acceptance_act_date,act_signed) VALUES('01-02/26','2026-08-20','Козлов','Баня','Брест',1000,'2026-09-15',1)").lastrowid
    d.execute("INSERT INTO contracts(contract_number,client_name,contract_amount,acceptance_act_date,act_signed) VALUES('02-02/26','Не подписал',5,'15.09.2026',0)")
    return act


def test_statement_for_both_sections(d, tmp_path):
    data(d, tmp_path)
    g = st.statement(d, 'gsv_projects', 2026, 9)
    assert [r['client'] for r in g['signed']] == ['Иванов', 'Петров'] and [r['client'] for r in g['unsigned']] == ['Сидоров']
    assert g['signed'][0]['paid'] == 100 and g['signed'][0]['rest'] == 150 and g['signed'][0]['file'] and not g['signed'][1]['file']
    m = st.statement(d, 'contracts', 2026, 9)
    assert [r['client'] for r in m['signed']] == ['Козлов'] and [r['client'] for r in m['unsigned']] == ['Не подписал']     # дата ДД.ММ.ГГГГ тоже распознаётся
    assert st.statement(d, 'gsv_projects', 2026, 10)['signed'][0]['client'] == 'Другой месяц'
    with pytest.raises(ValueError):
        st.statement(d, 'gsn_projects', 2026, 9)


def test_export_and_collect(d, tmp_path):
    data(d, tmp_path)
    path = tmp_path / 'out' / 'vedomost.xlsx'
    st.export_statement(d, 'gsv_projects', 2026, 9, path)
    wb = openpyxl.load_workbook(path)
    ws = wb['Подписанные акты']
    assert 'сентябрь 2026' in ws['A1'].value and [c.value for c in ws[4]][:5] == ['№', 'Дата акта', '№ договора', 'Дата договора', 'Клиент']
    assert ws['E5'].value == 'Иванов' and ws['B5'].value == '05.09.2026' and ws['H7'].value == 550 and ws['F7'].value == 'актов: 2'
    assert wb['Не подписаны']['E5'].value == 'Сидоров'
    result = st.collect_acts(d, 'gsv_projects', 2026, 9, tmp_path / 'бухгалтерия')
    assert result['total'] == 2 and len(result['copied']) == 1 and [r['client'] for r in result['missing']] == ['Петров']
    files = sorted(p.name for p in __import__('pathlib').Path(result['folder']).iterdir())
    assert len(files) == 2 and any(f.endswith('.xlsx') for f in files) and any(f.startswith('2026-09-05') for f in files)
