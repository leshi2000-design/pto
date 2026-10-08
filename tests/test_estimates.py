import pytest
from decimal import Decimal
from smetagaz import estimates_domain as ed, contracts_core as cc, payments_domain, report_templates as reports
from smetagaz.database import DatabaseManager


@pytest.fixture
def db(tmp_path):
    d = DatabaseManager(tmp_path / 'smetagaz.db')
    d.init_db()
    yield d
    d.close()


def add_items(db, eid):
    db.execute("INSERT INTO estimate_items(estimate_id,name,item_type,unit,quantity,price,sum,purchase_price) VALUES(?,?,?,?,?,?,?,?)", (eid, 'Труба', 'Материал', 'м', 10, 10, 100, 6))
    db.execute("INSERT INTO estimate_items(estimate_id,name,item_type,unit,quantity,price,sum,purchase_price,labor_hours,hourly_rate,overhead_pct,profit_pct,other_costs) "
               "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)", (eid, 'Монтаж', 'Работа', 'шт', 2, 50, 100, 30, 1, 10, 15, 10, 0))
    db.execute("INSERT INTO estimate_items(estimate_id,name,item_type,quantity,price,sum) VALUES(?,?,?,?,?,?)", (eid, 'Раздел 1', 'Раздел', 0, 0, 0))


def test_totals_without_surcharges_and_migration(db):
    eid = ed.create_estimate(db, 'Котельная')
    add_items(db, eid)
    db.execute('UPDATE estimates SET mat_adj_pct=10, work_adj_pct=-10, total_adj_pct=5 WHERE id=?', (eid,))
    t = ed.totals(db, eid)
    assert t['adj_mat'] == 110 and t['adj_work'] == 90 and t['total'] == Decimal('210')
    assert ed.recalc(db, eid) == 210.0
    # старая смета с начислениями: флаги отключаются, итог пересчитывается один раз
    db.execute('UPDATE estimates SET has_vat=1,has_social=1,has_overhead=1,has_profit=1,total=999 WHERE id=?', (eid,))
    db.set_setting('est_no_surcharges', '0')
    ed.migrate(db)
    assert db.fetchone('SELECT has_vat,has_social,has_overhead,has_profit,total FROM estimates WHERE id=?', (eid,)) == (0, 0, 0, 0, 210.0)


def test_margin_and_breakdown(db):
    eid = ed.create_estimate(db, 'Смета')
    add_items(db, eid)
    m = ed.margin_summary(db, eid)
    assert m['cost_materials'] == 60 and m['cost_works'] == 60 and m['sale_total'] == 200 and m['margin_total'] == 80
    assert m['margin_materials'] == 40 and m['margin_works'] == 40 and m['margin_pct'] == 40
    works, mats, info = ed.breakdown(db, eid)
    assert works[0]['has_breakdown'] == 'Да' and works[0]['margin_sum'] == 40 and mats[0]['margin_sum'] == 40
    assert info['materials_margin_total'] == 40 and info['grand_total'] == 200


def test_client_and_contract_choice(db):
    pid = db.execute("INSERT INTO crm.clients(name,phone) VALUES('Петров Пётр','+375291112233')").lastrowid
    lid = cc.save_legal(db, dict(name='ООО Ромашка'))
    # по умолчанию — без договора
    e1 = ed.create_estimate(db, 'Для Петрова', client=('person', pid))
    assert ed.contract_of(db, e1) is None and ed.client_label(db, e1) == 'Петров Пётр'
    assert db.fetchone('SELECT client_phone,client_id FROM estimates WHERE id=?', (e1,)) == ('+375291112233', pid)
    e2 = ed.create_estimate(db, '', client=('legal', lid))
    assert ed.client_label(db, e2) == 'ООО Ромашка' and db.fetchone('SELECT title,client_id,client_name FROM estimates WHERE id=?', (e2,)) == ('Смета · ООО Ромашка', None, '')
    # клиент из карточки договора, договор не привязывается
    cid = cc.save_contract(db, 'smr', dict(party_type='person', person_id=pid, contract_date='2026-05-01', amount=1))
    e3 = ed.create_estimate(db, 'Под договор', contract=('smr_contracts', cid))
    assert ed.client_label(db, e3) == 'Петров Пётр' and ed.contract_of(db, e3) is None
    # привязка позже
    ed.link_contract(db, e3, 'smr_contracts', cid)
    assert ed.contract_of(db, e3)['section'] == 'СМР'
    assert any(c[1] == cid for c in ed.contract_candidates(db, with_estimate=True))
    assert not any(c[1] == cid for c in ed.contract_candidates(db))        # у договора уже есть смета
    with pytest.raises(ValueError):
        ed.link_contract(db, e1, 'smr_contracts', cid)
    with pytest.raises(ValueError):
        ed.set_client(db, e3, 'legal', lid)                                # клиент привязанной сметы задаётся в договоре
    # отвязка; с оплатами — запрещена
    assert ed.unlink_contract(db, e3)['table'] == 'smr_contracts' and ed.contract_of(db, e3) is None
    ed.link_contract(db, e3, 'smr_contracts', cid)
    payments_domain.add(db, 'estimates', e3, 10, '2026-05-02')
    with pytest.raises(ValueError):
        ed.unlink_contract(db, e3)
    # создание сразу с привязкой и договор монтажа ГСВ
    mid = db.execute("INSERT INTO contracts(contract_number,contract_date,client_id) VALUES('01-02/26','2026-01-01',?)", (pid,)).lastrowid
    e4 = ed.create_estimate(db, 'С привязкой', contract=('contracts', mid), link=True)
    assert ed.contract_of(db, e4)['table'] == 'contracts' and ed.contract_client(db, 'contracts', mid) == ('person', pid)
    with pytest.raises(ValueError):
        ed.create_estimate(db, 'x', contract=('le_contracts', 1), link=True)


def test_tags_context_and_template_export(db, tmp_path):
    from docx import Document
    pid = db.execute("INSERT INTO crm.clients(name) VALUES('Сидоров')").lastrowid
    eid = ed.create_estimate(db, 'Баня', client=('person', pid))
    add_items(db, eid)
    ed.recalc(db, eid)
    ctx, tables = reports.context(db, 'estimates', eid)
    assert ctx['клиент'] == 'Сидоров' and ctx['итого'] == '200,00' and ctx['маржа'] == '80,00' and ctx['есть_договор'] == 'Нет'
    assert ctx['vat_total'] == 0 and len(tables['works']) == 1 and len(tables['materials']) == 1 and len(tables['items']) == 3
    ctx2, tables2 = reports.context(db, 'estimate_breakdown', eid)
    assert ctx2['маржа_материалы'] == '40,00' and len(tables2['items']) == 3
    # свой тег и отключение встроенного
    tid = ed.save_tag(db, None, 'Сумма Работ', 'works_total', 'money_words')
    row = db.fetchone('SELECT name FROM est_tags WHERE id=?', (tid,))
    assert row[0] == 'сумма_работ'
    ctx3, _ = reports.context(db, 'estimates', eid)
    assert ctx3['сумма_работ'].startswith('Сто')
    auto = db.fetchone("SELECT id FROM est_tags WHERE auto_key='b:маржа'")[0]
    assert ed.delete_tag(db, auto) == 'disabled' and 'маржа' not in reports.context(db, 'estimates', eid)[0]
    with pytest.raises(ValueError):
        ed.save_tag(db, None, 'сумма работ', 'works_total', 'money')
    # экспорт по шаблону с русскими тегами и таблицей
    src = tmp_path / 't.docx'
    doc = Document()
    doc.add_paragraph('{{клиент}}: {{итого}} ({{итого_прописью}})')
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = 'Работа'
    t.rows[1].cells[0].text = '{{works.name}}'
    t.rows[1].cells[1].text = '{{works.margin_sum}}'
    doc.save(src)
    db.execute('UPDATE estimates SET prepared_by=? WHERE id=?', ('Иванов', eid))
    tpl = db.execute('INSERT INTO report_templates(kind,name,file_path) VALUES(?,?,?)', ('estimates', 'Смета', str(src))).lastrowid
    out = tmp_path / 'out.docx'
    reports.export(db, 'estimates', eid, tpl, out)
    result = Document(out)
    assert result.paragraphs[0].text.startswith('Сидоров: 200,00 (Двести')
    assert result.tables[0].cell(1, 0).text == 'Монтаж'
