import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import sqlite3
import zipfile
from pathlib import Path
import pytest
from smetagaz.database import DatabaseManager
from smetagaz.data_services import search,link_client
from smetagaz.backup import create_backup,restore_backup
from smetagaz.exports import export_record
from smetagaz.security import password_hash,verify,encrypt_backup,decrypt_backup

@pytest.fixture
def database(tmp_path):
    d=DatabaseManager(tmp_path/'smetagaz.db');d.init_db();yield d;d.close()

def test_transactions_and_isolation(database,tmp_path):
    with database.transaction() as cur:
        cur.execute('INSERT INTO estimates(title) VALUES(?)',('Один',));rid=cur.lastrowid
        with pytest.raises(ValueError):
            with database.transaction():
                database.execute('INSERT INTO estimates(title) VALUES("Ошибка")');raise ValueError()
    assert database.fetchone('SELECT count(*) FROM estimates')[0]==1
    assert rid>0
    other=DatabaseManager(tmp_path/'other.db');assert other is not database;other.close()

def test_clients_and_search(database):
    a=database.execute('INSERT INTO estimates(title,client_name,client_phone) VALUES(?,?,?)',('Газоснабжение','ИВАНОВ','+43 555')).lastrowid
    cid=database.fetchone('SELECT client_id FROM estimates WHERE id=?',(a,))[0]
    b=database.execute('INSERT INTO gsv_projects(pd_number) VALUES("ПД-25")').lastrowid
    link_client(database,'gsv_projects',b,cid)
    database.execute('UPDATE crm.clients SET name="Петров",phone="12345" WHERE id=?',(cid,))
    assert database.fetchone('SELECT client_name,phone FROM gsv_projects WHERE id=?',(b,))==('Петров','12345')
    assert database.fetchone('SELECT client_name FROM estimates WHERE id=?',(a,))[0]=='Петров'
    assert len(search(database,'ПЕТРОВ'))==3
    assert search(database,'газоснаб')
    assert not search(database,'" OR *')
    database.execute('DELETE FROM estimates WHERE id=?',(a,))
    assert not search(database,'газоснаб')

def test_backup_restore_and_traversal(database,tmp_path):
    a=database.execute('INSERT INTO estimates(title,client_name) VALUES("Тест","Клиент")').lastrowid
    source=tmp_path/'example.txt';source.write_text('данные',encoding='utf8')
    database.execute('INSERT INTO attachments(estimate_id,file_name,file_path) VALUES(?,?,?)',(a,source.name,str(source)))
    archive=tmp_path/'backup.zip';manifest=create_backup(database,archive)
    assert not manifest['missing']
    dest=tmp_path/'restored';restore_backup(archive,dest)
    restored=DatabaseManager(dest/'smetagaz.db');restored.init_db()
    assert restored.fetchone('SELECT name FROM crm.clients')[0]=='Клиент'
    f=restored.fetchone('SELECT file_path FROM attachments')[0]
    assert (dest/f).read_text(encoding='utf8')=='данные';restored.close()
    with pytest.raises(ValueError):restore_backup(archive,dest)
    bad=tmp_path/'bad.zip'
    with zipfile.ZipFile(bad,'w') as z:z.writestr('../escape','bad')
    with pytest.raises(ValueError):restore_backup(bad,tmp_path/'badrestore')
    assert not (tmp_path/'escape').exists()

def test_exports(database,tmp_path):
    rid=database.execute('INSERT INTO estimates(title,client_name,client_phone,total) VALUES(?,?,?,?)',('=1+1','Петров','SECRET',125)).lastrowid
    database.execute('INSERT INTO estimate_items(estimate_id,name,unit,quantity,price,sum) VALUES(?,?,?,?,?,?)',(rid,'Труба <ГОСТ>','м',5,25,125))
    for ext in ['docx','xlsx','pdf']:
        path=tmp_path/('test.'+ext);export_record(database,'estimates',rid,path)
        assert path.stat().st_size>1000
    from openpyxl import load_workbook
    book=load_workbook(tmp_path/'test.xlsx');assert book.active['A1'].data_type=='s'
    assert 'SECRET' not in str(list(book.active.values))
    from docx import Document
    doc=Document(tmp_path/'test.docx');assert any('Труба' in c.text for t in doc.tables for r in t.rows for c in r.cells)

def test_password_and_encrypted_backup(tmp_path):
    stored=password_hash('test-password-123');assert verify('test-password-123',stored);assert not verify('wrong-password',stored)
    src=tmp_path/'plain';src.write_bytes(os.urandom(2200000));enc=tmp_path/'copy.sgb';out=tmp_path/'out'
    encrypt_backup(src,enc,'test-password-123');decrypt_backup(enc,out,'test-password-123');assert src.read_bytes()==out.read_bytes()
    with pytest.raises(ValueError):decrypt_backup(enc,out,'wrong-password')
    assert not out.exists()

def test_bulk_attachments(database,tmp_path):
    from smetagaz.file_jobs import import_files
    rid=database.execute('INSERT INTO estimates(title) VALUES("Files")').lastrowid
    a=tmp_path/'a';b=tmp_path/'b';a.mkdir();b.mkdir();(a/'same.txt').write_text('1');(b/'same.txt').write_text('2')
    result=import_files(database,rid,[a/'same.txt',b/'same.txt',a/'missing'])
    assert result['imported']==2 and len(result['failed'])==1
    files=database.fetchall('SELECT file_path FROM attachments');assert len({r[0] for r in files})==2

def test_legacy_columns_survive(tmp_path):
    path=tmp_path/'smetagaz.db'
    with sqlite3.connect(path) as c:
        c.execute('CREATE TABLE estimates(id INTEGER PRIMARY KEY,title TEXT,date TEXT,total REAL)');c.execute('INSERT INTO estimates VALUES(1,"old","2025-01-01",100)')
        c.execute('CREATE TABLE estimate_items(id INTEGER PRIMARY KEY,estimate_id INTEGER,name TEXT,unit TEXT,quantity REAL,price REAL,sum REAL,purchase_price REAL,custom_field TEXT)')
        c.execute('INSERT INTO estimate_items VALUES(1,1,"old","m",1,100,100,60,"keep")')
    d=DatabaseManager(path);d.init_db();assert d.fetchone('SELECT purchase_price,custom_field FROM estimate_items')==(60,'keep');d.close()

def test_backup_corruption_rejected(database,tmp_path):
    archive=tmp_path/'good.zip';create_backup(database,archive)
    bad=tmp_path/'bad.zip'
    with zipfile.ZipFile(archive) as src,zipfile.ZipFile(bad,'w') as dst:
        for name in src.namelist():dst.writestr(name,src.read(name) if name!='clients.db' else b'broken')
    with pytest.raises(ValueError):restore_backup(bad,tmp_path/'output')
    assert not (tmp_path/'output').exists()

def test_reference_delete_and_rollback(database):
    rid=database.execute('INSERT INTO estimates(title,client_name) VALUES("Test","Client")').lastrowid
    database.execute('INSERT INTO estimate_items(estimate_id,name) VALUES(?,"Item")',(rid,))
    with pytest.raises(RuntimeError):
        with database.transaction():
            database.execute('UPDATE crm.clients SET name="Wrong"');raise RuntimeError()
    assert database.fetchone('SELECT client_name FROM estimates')[0]=='Client'
    database.execute('DELETE FROM estimates WHERE id=?',(rid,))
    assert database.fetchone('SELECT count(*) FROM estimate_items')[0]==0


def test_search_paging_and_reopen(database):
    with database.transaction():database.executemany('INSERT INTO materials(name) VALUES(?)',((f'Газовая труба {i}',) for i in range(215)))
    first=search(database,'ГАЗОВАЯ',100);second=search(database,'ГАЗОВАЯ',100,100)
    assert len(first)==len(second)==100
    assert not {r[1] for r in first}&{r[1] for r in second}
    database.init_db();assert len([r for r in search(database,'газовая',500) if r[0]=='materials'])==215


def test_delete_client_detaches_and_erases_personal_copies(tmp_path):
    from smetagaz.database import DatabaseManager
    from smetagaz.data_services import client_links, delete_client
    d = DatabaseManager(tmp_path / 'smetagaz.db')
    d.init_db()
    cid = d.execute("INSERT INTO crm.clients(name,phone,passport) VALUES('Иванов Иван','123','МР 1')").lastrowid
    keep = d.execute("INSERT INTO crm.clients(name) VALUES('Другой')").lastrowid
    est = d.execute("INSERT INTO estimates(title,total,client_id,client_name,client_phone) VALUES('Смета',10,?,'Иванов Иван','123')", (cid,)).lastrowid
    con = d.execute("INSERT INTO contracts(contract_number,client_id,client_name,client_phone,passport_series_number,client_address) VALUES('1',?,'Иванов Иван','123','МР 1','Минск')", (cid,)).lastrowid
    d.execute("INSERT INTO gsv_projects(pd_number,client_id,client_name,phone,passport) VALUES('01-26 ГСВ',?,'Иванов Иван','123','МР 1')", (cid,))
    d.execute("INSERT INTO kanban_tasks(title,client_id) VALUES('Позвонить',?)", (cid,))
    assert client_links(d, cid) == {'Сметы': 1, 'Монтаж ГСВ': 1, 'Проекты ГСВ': 1, 'Задачи': 1}
    delete_client(d, cid)
    assert d.fetchone('SELECT count(*) FROM crm.clients')[0] == 1 and d.fetchone('SELECT id FROM crm.clients')[0] == keep
    assert d.fetchone('SELECT client_id,client_name,client_phone FROM estimates WHERE id=?', (est,)) == (None, '', '')
    assert d.fetchone('SELECT client_id,client_name,passport_series_number,client_address FROM contracts WHERE id=?', (con,)) == (None, '', '', '')
    assert d.fetchone('SELECT client_id,client_name,phone,passport FROM gsv_projects') == (None, '', '', '')
    assert d.fetchone('SELECT title,client_id FROM kanban_tasks') == ('Позвонить', None)       # сами записи остались
    d.close()
    again = DatabaseManager(tmp_path / 'smetagaz.db')
    again.init_db()
    assert again.fetchone('SELECT count(*) FROM crm.clients')[0] == 1                          # после перезапуска клиент не воскресает
    with __import__('pytest').raises(ValueError):
        delete_client(again, cid)
    again.close()
