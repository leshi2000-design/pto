from pathlib import Path
from zipfile import ZipFile
import pytest
from docx import Document
from openpyxl import Workbook,load_workbook
from smetagaz.database import DatabaseManager
from smetagaz import report_templates as reports,template_domain
from smetagaz.table_templates import render
from smetagaz.backup import create_backup,restore_backup
@pytest.fixture
def db(tmp_path):
    d=DatabaseManager(tmp_path/'smetagaz.db');d.init_db();yield d;d.close()

def test_gsv_separated_preserves_user_configuration(db):
    # Pre-2.4 databases could carry an erroneous "ГСВ" document list; migrate() must archive it
    # under a distinct module name rather than mixing it with the real ГСВ (contracts) list below.
    # (The db fixture already ran migrate() once, seeding the real ГСН/ГСВ rows; resetting the
    # flag and adding one more row simulates an older database being opened again.)
    db.set_setting('gsv_templates_separated','0');db.execute("INSERT INTO executive_templates(module,slot,label,file_path) VALUES('ГСВ','legacy_slot','Старый','old.docx')")
    template_domain.migrate(db)
    assert len(template_domain.templates(db,'contracts'))==8;assert template_domain.templates(db,'gsv_projects')==[];assert len(template_domain.templates(db,'gsn_projects'))==12
    assert db.fetchone("SELECT file_path FROM executive_templates WHERE module LIKE 'ГСВ — архив%' AND slot='legacy_slot'")[0]=='old.docx'
    count_after_archival=db.fetchone('SELECT count(*) FROM executive_templates')[0]
    template_domain.migrate(db);assert db.fetchone('SELECT count(*) FROM executive_templates')[0]==count_after_archival

def test_word_repeated_rows_styles_and_empty(tmp_path):
    src=tmp_path/'t.docx';out=tmp_path/'out.docx';doc=Document();doc.add_paragraph('{{title}}');t=doc.add_table(rows=2,cols=3);t.rows[0].cells[0].text='Позиции';t.rows[1].cells[0].paragraphs[0].add_run('{{items.na').bold=True;t.rows[1].cells[0].paragraphs[0].add_run('me}}');t.rows[1].cells[1].text='{{items.quantity}}';t.rows[1].cells[2].text='{{items.amount}}';doc.save(src)
    render(src,{'title':'Смета'}, {'items':[{'name':'Труба','quantity':2,'amount':20},{'name':'Муфта','quantity':3,'amount':15}]},out)
    result=Document(out);assert len(result.tables[0].rows)==3;assert result.tables[0].cell(1,0).text=='Труба';assert result.tables[0].cell(2,0).text=='Муфта';assert result.tables[0].cell(1,0).paragraphs[0].runs[0].bold
    render(src,{'title':'Пустая'},{'items':[]},out);assert len(Document(out).tables[0].rows)==1

def test_excel_repeat_shift_merge_printarea_numbers(tmp_path):
    src=tmp_path/'t.xlsx';out=tmp_path/'out.xlsx';w=Workbook();s=w.active;s['A1']='{{title}}';s['A2']='{{items.name}}';s.merge_cells('A2:B2');s['C2']='{{items.quantity}}';s['D2']='{{items.amount}}';s['A3']='Итого';s['D3']='{{total}}';s.print_area='A1:D3';w.save(src)
    render(src,{'title':'Смета','total':35},{'items':[{'name':'Труба','quantity':2,'amount':20},{'name':'Муфта','quantity':3,'amount':15}]},out)
    s=load_workbook(out).active;assert s['A2'].value=='Труба';assert s['A3'].value=='Муфта';assert s['A4'].value=='Итого';assert s['D4'].value==35;assert s['C3'].value==3;assert 'A3:B3' in s.merged_cells;assert '$D$4' in str(s.print_area)

def test_repeat_formula_rejected_and_output_untouched(tmp_path):
    src=tmp_path/'bad.xlsx';out=tmp_path/'out.xlsx';out.write_bytes(b'original');w=Workbook();w.active['A1']='{{items.name}}';w.active['A2']='=1+1';w.save(src)
    with pytest.raises(ValueError,match='формул'):render(src,{}, {'items':[{'name':'A'}]},out)
    assert out.read_bytes()==b'original'

def test_estimate_context_totals_and_export(db,tmp_path):
    rid=db.execute("INSERT INTO estimates(title,date,total,paid,mat_adj_pct) VALUES('Смета','2026-09-01',132,32,10)").lastrowid
    db.execute("INSERT INTO estimate_items(estimate_id,name,item_type,quantity,price,sum) VALUES(?, 'Труба','Материал',10,10,100)",(rid,));ctx,tables=reports.context(db,'estimates',rid);assert ctx['calculated_total']==110 and ctx['vat_total']==0;assert ctx['debt']==100;assert tables['items'][0]['quantity']==10
    src=tmp_path/'t.docx';doc=Document();doc.add_paragraph('{{title}} {{total}}');doc.save(src);tid=db.execute('INSERT INTO report_templates(kind,name,file_path) VALUES(?,?,?)',('estimates','Смета',str(src))).lastrowid
    db.execute('UPDATE estimates SET prepared_by=? WHERE id=?',('Иванов И.И.',rid))
    result=tmp_path/'result.docx';reports.export(db,'estimates',rid,tid,result);assert Document(result).paragraphs[0].text=='Смета 132.0'
    with pytest.raises(ValueError):reports.export(db,'estimates',rid,tid,src)

def test_statistics_period_and_balances_filter(db):
    db.execute("INSERT INTO estimates(title,date,total,paid) VALUES('В периоде','2026-09-01',100,40)");db.execute("INSERT INTO estimates(title,date,total,paid) VALUES('Вне периода','2025-09-01',900,800)")
    ctx,tables=reports.context(db,'statistics',filters={'start':'2026-09-01','end':'2026-09-30'});assert ctx['total']==100;assert ctx['debt']==60;assert len(tables['estimates'])==1
    ctx,tables=reports.context(db,'balances',filters={'category':'Категория отсутствует'});assert ctx['item_count']==0
    with pytest.raises(ValueError):reports.context(db,'statistics',filters={'start':'2026-10-01','end':'2026-09-01'})

def test_report_template_backup(db,tmp_path):
    src=tmp_path/'t.docx';doc=Document();doc.add_paragraph('{{name}}');doc.save(src);db.execute('INSERT INTO report_templates(kind,name,file_path) VALUES(?,?,?)',('crm.clients','Карточка',str(src)))
    archive=tmp_path/'backup.zip';create_backup(db,archive);dest=tmp_path/'restored';restore_backup(archive,dest);other=DatabaseManager(dest/'smetagaz.db');other.init_db();file=other.fetchone('SELECT file_path FROM report_templates')[0];assert (dest/file).is_file();other.close()
