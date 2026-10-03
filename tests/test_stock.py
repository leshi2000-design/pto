import json
from pathlib import Path
import pytest
from smetagaz.database import DatabaseManager
from smetagaz import stock_domain as s
from smetagaz.filters import compile_rules
from smetagaz.stock_exports import export_document
from smetagaz.backup import create_backup,restore_backup

@pytest.fixture
def d(tmp_path):
    d=DatabaseManager(tmp_path/'smetagaz.db');d.init_db();yield d;d.close()

def setup_norm(d):
    mids=[r[0] for r in d.fetchall('SELECT id FROM stock_materials ORDER BY id')]
    pid=1
    s.save_profile(d,dict(name='Стык Ду15',diameter='15',thickness='2.8',basis='стык'),list(zip(mids,['0.1','0.2','0.03','0','0.05'])),pid)
    return pid,mids

def header(**kw):return dict(number='СП-1',act_date='2026-09-17',object_name='Произвольный объект',owner_type=None,owner_id=None,**kw)

def test_missing_norms_explicit_zero_and_decimal(d):
    with pytest.raises(ValueError,match='не задана норма'):s.calculate(d,[(1,'5')])
    pid,mids=setup_norm(d);lines,snapshot=s.calculate(d,[(pid,'10'),(pid,'5')])
    assert {r['material_id']:r['actual_qty'] for r in lines}=={mids[0]:1500000,mids[1]:3000000,mids[2]:450000,mids[4]:750000}
    assert s.fmt(10000000)=='10' and s.fmt(-1200000)=='-1.2'
    with pytest.raises(ValueError,match='целым'):s.calculate(d,[(pid,'1.5')])
    with pytest.raises(ValueError):s.units('nan')
    assert s.units('0,000001')==1
    # Paint can have its own norm per square metre.
    paint=s.save_profile(d,dict(name='Окраска',basis='м²'),[(mids[3],'0.15')]);assert s.calculate(d,[(paint,'2.5')])[0][0]['actual_qty']==375000

def test_posting_shortage_atomic_idempotent_and_reversal(d):
    pid,mids=setup_norm(d);lines,snapshot=s.calculate(d,[(pid,'10')]);aid=s.save_act(d,header(calculation=json.dumps(snapshot)),lines)
    s.receive(d,mids[0],'10','2026-09-01')
    with pytest.raises(ValueError,match='Недостаточно'):s.post_act(d,aid)
    assert s.balance(d,mids[0])==s.units(10);assert d.fetchone('SELECT count(*) FROM stock_moves WHERE act_id=?',(aid,))[0]==0
    for mid in mids[1:]:s.receive(d,mid,'10','2026-09-01')
    s.post_act(d,aid);s.post_act(d,aid);assert s.balance(d,mids[0])==s.units(9)
    with pytest.raises(ValueError,match='черновик'):s.save_act(d,header(),lines,aid)
    s.reverse_act(d,aid);s.reverse_act(d,aid);assert s.balance(d,mids[0])==s.units(10)
    copied=s.copy_act(d,aid);assert copied!=aid;s.post_act(d,copied);assert s.balance(d,mids[0])==s.units(9)

def test_dated_stock_prevents_future_receipts_and_bad_adjustment(d):
    mid=1;s.receive(d,mid,5,'2026-09-20')
    aid=s.save_act(d,header(),[dict(material_id=mid,norm_qty=None,actual_qty=s.units(2))])
    with pytest.raises(ValueError,match='2026-09-17'):s.post_act(d,aid)
    assert s.balance(d,mid)==s.units(5)
    with pytest.raises(ValueError):s.receive(d,mid,-6,'2026-09-21','Инвентаризация',True)
    assert s.balance(d,mid)==s.units(5)
    s.receive(d,mid,2,'2026-09-01');s.post_act(d,aid);assert s.balance(d,mid)==s.units(5)

def test_snapshot_and_concurrent_draft_edits(d):
    pid,mids=setup_norm(d);lines,snapshot=s.calculate(d,[(pid,'2')]);aid=s.save_act(d,header(calculation=json.dumps(snapshot)),lines)
    s.save_act(d,header(calculation=json.dumps(snapshot)),lines,aid,1)
    with pytest.raises(ValueError,match='другом окне'):s.save_act(d,header(),[],aid,1)
    d.execute('UPDATE norm_items SET rate=? WHERE profile_id=?',('99',pid))
    assert json.loads(d.fetchone('SELECT calculation FROM stock_acts WHERE id=?',(aid,))[0])[0]['items'][0]['rate']=='0.1'
    assert d.fetchone('SELECT actual_qty FROM stock_act_lines WHERE act_id=? AND material_id=?',(aid,mids[0]))[0]==200000
    with pytest.raises(ValueError,match='Единица'):s.save_material(d,'Газ','Газ','кг',mid=mids[0])

def test_defect_link_free_work_and_no_duplicate_writeoff(d):
    pid,mids=setup_norm(d)
    did=s.save_defect(d,header(),[dict(defect='Течь',work='Переварить участок',unit='стык',quantity='3',profile_id=pid),dict(defect='Коррозия',work='Покрасить',unit='м²',quantity='2')])
    aid=s.from_defect(d,did);assert s.from_defect(d,did)==aid
    assert 'Покрасить' in d.fetchone('SELECT notes FROM stock_acts WHERE id=?',(aid,))[0]
    assert d.fetchone('SELECT norm_qty FROM stock_act_lines WHERE act_id=? AND material_id=?',(aid,mids[0]))[0]==300000
    with pytest.raises(ValueError,match='списание'):s.save_defect(d,header(),[],did)
    with pytest.raises(ValueError,match='совпадать'):s.save_defect(d,header(),[dict(work='Окраска',defect='',unit='м²',quantity='1',profile_id=pid)])

def test_sql_filters_escape_dates_and_all_pages(d):
    d.executemany('INSERT INTO stock_materials(name,category,unit) VALUES(?,?,?)',[(f'Материал {i}','Тест','шт') for i in range(410)])
    clause,params=compile_rules([('name','Содержит','материал 409')],['name'])
    assert d.fetchone('SELECT name FROM stock_materials WHERE '+clause,params)[0]=='Материал 409'
    clause,params=compile_rules([('act_date','Дата с','17.09.2026'),('act_date','Дата по','2026-09-18')],['act_date']);assert params==['2026-09-17','2026-09-18']
    with pytest.raises(ValueError):compile_rules([('act_date','Дата с','bad')],['act_date'])
    assert compile_rules([('name','Содержит',"%' OR 1=1 --")],['name'])[1]==["%\\%' or 1=1 --%"]

def test_migrate_multiday_journal_once(d):
    jid=d.execute("INSERT INTO welding_jobs(title) VALUES('Старые работы')").lastrowid
    d.executemany('INSERT INTO welding_days(job_id,work_date) VALUES(?,?)',[(jid,'2026-08-01'),(jid,'2026-08-02')]);d.set_setting('stock_schema','0')
    # Simulate an upgrade without re-seeding catalogs in a real user database.
    s.migrate(d);rows=d.fetchall('SELECT job_id,work_date FROM welding_days ORDER BY work_date');assert len(rows)==2 and rows[0][0]!=rows[1][0]
    s.migrate(d);assert d.fetchone('SELECT count(*) FROM welding_days')[0]==2

def test_inventory_backup_restore_and_search(d,tmp_path):
    from smetagaz.data_services import search
    s.receive(d,1,5,'2026-09-01');aid=s.save_act(d,header(),[dict(material_id=1,norm_qty=None,actual_qty=s.units(2))]);s.post_act(d,aid)
    assert any(r[0]=='stock_acts' for r in search(d,'Произвольный'))
    path=tmp_path/'backup.zip';create_backup(d,path);restore_backup(path,tmp_path/'restore');new=DatabaseManager(tmp_path/'restore/smetagaz.db');new.init_db();assert s.balance(new,1)==s.units(3);assert new.fetchone('SELECT status FROM stock_acts')[0]=='posted';new.close()

def test_exports_all_formats_include_lines_and_no_formulas(d,tmp_path):
    pid,mids=setup_norm(d);lines,snapshot=s.calculate(d,[(pid,3)]);aid=s.save_act(d,header(calculation=json.dumps(snapshot)),lines)
    did=s.save_defect(d,header(),[dict(defect='Коррозия',work='=1+1',unit='м²',quantity='2.5')])
    for source,rid in [('stock_acts',aid),('defect_acts',did)]:
        for ext in ('docx','xlsx','pdf'):
            p=tmp_path/f'{source}.{ext}';export_document(d,source,rid,p,dict(mode='c29'));assert p.stat().st_size>1000
    from openpyxl import load_workbook
    wb=load_workbook(tmp_path/'defect_acts.xlsx');assert any(c.value=='=1+1' and c.data_type=='s' for row in wb.active for c in row)
    from docx import Document
    assert any('Ацетилен' in c.text for t in Document(tmp_path/'stock_acts.docx').tables for row in t.rows for c in row.cells)
