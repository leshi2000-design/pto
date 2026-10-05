import os
from datetime import date
import docx
import openpyxl
import pytest
from smetagaz.database import DatabaseManager
from smetagaz import gsvm_domain as m, gsvm_docs as t, gsv_project_domain as gd, payments_domain


@pytest.fixture
def d(tmp_path, monkeypatch):
    monkeypatch.setattr(gd, 'PROJECTS_DIR', tmp_path / 'projects')
    monkeypatch.setattr(gd, 'TEMPLATES_DIR', tmp_path / 'templates')
    monkeypatch.setattr(t, 'TEMPLATES_DIR', tmp_path / 'templates' / 'gsvm')
    d = DatabaseManager(tmp_path / 'smetagaz.db')
    d.init_db()
    yield d
    d.close()


def contract(d, **extra):
    cid = d.execute("INSERT INTO crm.clients(name,phone,address,passport,passport_issuer,passport_date) VALUES('Иванов Иван Иванович','375291112233 (основной)','Минск, Лесная 5','МР 1234567','РУВД','2015-03-04')").lastrowid
    cols = dict(contract_number='01-02/26', contract_date='2026-11-12', object_name='Жилой дом', object_address='д. Ключи, 12', client_name='Иванов Иван Иванович',
                client_id=cid, work_start_date='2026-11-12', work_end_date='2027-01-11', acceptance_act_date='2026-12-20', contract_amount=1250.5,
                designer_code='ГСВ-77', designer='ООО Проект', project_month='2026-10', seq_num=1, year_num=26)
    cols.update(extra)
    return d.execute(f"INSERT INTO contracts({','.join(cols)}) VALUES({','.join('?' for _ in cols)})", tuple(cols.values())).lastrowid


def pipeline(d, name, cert_number=None):
    cert = d.execute('INSERT INTO certificates(name,cert_number,file_path) VALUES(?,?,?)', ('Сертификат ' + name, cert_number, '/tmp/x.pdf')).lastrowid if cert_number else None
    return d.execute('INSERT INTO gsv_pipelines(name,unit,certificate_id) VALUES(?,?,?)', (name, 'м', cert)).lastrowid


def test_number_is_sequential_per_year_and_section(d):
    assert m.next_number(d, '2026-11-12') == (1, 26, '01-02/26')
    contract(d)
    assert m.next_number(d, '2026-11-12') == (2, 26, '02-02/26')
    assert m.next_number(d, '2027-01-05') == (1, 27, '01-02/27')
    d.execute("INSERT INTO contracts(contract_number) VALUES('05-02/26')")          # номер, введённый вручную, тоже учитывается
    assert m.next_number(d, '2026-02-01')[2] == '06-02/26'


def test_folder_name_format():
    assert m.folder_name('01-02/26', 'д. Ключи, 12', 'Иванов Иван Иванович') == '01-02.26, д. Ключи, 12 (Иванов И.И.)'
    assert m.folder_name('01-02/26', '', 'Иванов') == '01-02.26 (Иванов)'


def test_estimate_state_history_and_cost(d):
    cid = contract(d)
    assert m.estimate_state(d, cid) == 'none'
    eid = d.execute("INSERT INTO estimates(title,total) VALUES('Смета',0)").lastrowid
    m.link_estimate(d, cid, eid)
    assert m.estimate_state(d, cid) == 'unsynced'
    d.execute("INSERT INTO estimate_items(estimate_id,item_type,name,unit,quantity,price,sum) VALUES(?,?,?,?,?,?,?)", (eid, 'Материал', 'Труба 25', 'м', 10, 5, 50))
    d.execute("INSERT INTO estimate_items(estimate_id,item_type,name,unit,quantity,price,sum) VALUES(?,?,?,?,?,?,?)", (eid, 'Работа', 'Монтаж', 'шт', 1, 30, 30))
    d.execute('UPDATE estimates SET total=80 WHERE id=?', (eid,))
    s = m.sync_estimate(d, cid)
    assert (s['materials'], s['works'], s['total']) == (50, 30, 80) and m.estimate_state(d, cid) == 'fresh'
    assert d.fetchone('SELECT contract_amount,est_materials,est_works FROM contracts WHERE id=?', (cid,)) == (80, 50, 30)
    d.execute("UPDATE estimate_items SET quantity=12,sum=60 WHERE name='Труба 25'")
    assert m.estimate_state(d, cid) == 'stale'
    history = m.estimate_history(d, eid)
    assert any(h[1] == 'Изменено' and '10.0 → 12.0' in h[3] for h in history) and any(h[1] == 'Добавлено' for h in history)
    d.execute('UPDATE estimate_items SET sort_order=5')                                          # смена порядка строк историю не засоряет
    assert len(m.estimate_history(d, eid)) == len(history)
    m.sync_estimate(d, cid)
    assert m.estimate_state(d, cid) == 'fresh'
    other = contract(d, contract_number='02-02/26')
    with pytest.raises(ValueError, match='уже привязана'):
        m.link_estimate(d, other, eid)


def test_pipelines_joints_and_estimate(d):
    cid = contract(d)
    p25, p32 = pipeline(d, 'Труба 25', 'A-25'), pipeline(d, 'Труба ПЭ 32x3', 'A-32')
    m.save_pipelines(d, cid, [dict(pipeline_id=p25, quantity=10.456, note='ввод'), dict(pipeline_id=p32, quantity=3)])
    assert [r['quantity'] for r in m.pipelines(d, cid)] == [10.46, 3.0]
    assert [(j['name'], j['count']) for j in m.joints(d, cid)] == [('Труба 25', 0), ('Труба ПЭ 32x3', 0)]
    j = m.joints(d, cid)
    m.save_joints(d, cid, [dict(id=j[0]['id'], name='Труба 25', count=4, note=''), dict(id=j[1]['id'], name='Труба ПЭ 32x3', count=2, note=''),
                           dict(id=None, name='Сталь 57x3', count=3, note='произвольный', custom=1)])
    assert m.joints_total(d, cid) == 9 and len(m.joints(d, cid)) == 3
    m.save_pipelines(d, cid, [dict(pipeline_id=p25, quantity=10)])                   # убрали вторую трубу из верхней таблицы
    assert [j['name'] for j in m.joints(d, cid)] == ['Труба 25', 'Сталь 57x3']        # её строка исчезла из нижней, произвольная осталась
    eid = d.execute("INSERT INTO estimates(title,total) VALUES('Смета',0)").lastrowid
    m.link_estimate(d, cid, eid)
    d.execute("INSERT INTO estimate_items(estimate_id,item_type,name,unit,quantity,price,sum) VALUES(?,?,?,?,?,?,?)", (eid, 'Материал', ' труба  25 ', 'м', 22.5, 1, 22.5))
    d.execute("INSERT INTO estimate_items(estimate_id,item_type,name,unit,quantity,price,sum) VALUES(?,?,?,?,?,?,?)", (eid, 'Материал', 'Труба ПЭ 32x3', 'м', 7, 1, 7))
    assert m.pipelines_from_estimate(d, cid) == 2
    assert {r['name']: r['quantity'] for r in m.pipelines(d, cid)} == {'Труба 25': 22.5, 'Труба ПЭ 32x3': 7.0}
    d.execute("UPDATE estimate_items SET quantity=30 WHERE item_type='Материал' AND name LIKE '%25%'")
    assert m.refresh_estimate_pipelines(d, cid) == 1                                  # смета изменилась — строки из сметы обновились сами
    assert {r['name']: r['quantity'] for r in m.pipelines(d, cid)}['Труба 25'] == 30.0
    with pytest.raises(ValueError):
        m.save_pipelines(d, cid, [dict(pipeline_id=p25, quantity=-1)])


def test_certificates_groups_rules_and_exclusion(d):
    cid = contract(d)
    p32 = pipeline(d, 'Труба ПЭ 32x3', 'A-32')
    m.save_pipelines(d, cid, [dict(pipeline_id=p32, quantity=3)])
    m.save_equipment(d, cid, [dict(kind='Сигнализатор', model='СГ-1', serial='123', file_path='/tmp/sg.pdf'), dict(kind='Котел', model='Bosch', serial='9')])
    c_signal = d.execute("INSERT INTO certificates(name,cert_number,file_path) VALUES('Сертификат производителя сигнализаторов','S-1','/tmp/s1.pdf')").lastrowid
    c_fit = d.execute("INSERT INTO certificates(name,cert_number,file_path) VALUES('Сертификат фитингов 32','F-32','/tmp/f32.pdf')").lastrowid
    c_all = d.execute("INSERT INTO certificates(name,cert_number,file_path) VALUES('Сертификат на работы','W-1','/tmp/w1.pdf')").lastrowid
    c_unused = d.execute("INSERT INTO certificates(name,cert_number,file_path) VALUES('Для котлов другого типа','K-1','/tmp/k1.pdf')").lastrowid
    m.add_cert_rule(d, 'equipment_kind', 'Сигнализатор', c_signal)
    m.add_cert_rule(d, 'pipeline_text', '32', c_fit)
    m.add_cert_rule(d, 'equipment_kind', 'Колонка', c_unused)
    d.execute('INSERT INTO gsvm_default_certs(certificate_id) VALUES(?)', (c_all,))
    certs = m.collect_certs(d, cid)
    assert [(c['group'], c['number']) for c in certs] == [('equipment', '123'), ('pipeline', 'A-32'), ('dependent', 'S-1'), ('dependent', 'F-32'), ('default', 'W-1')]
    m.exclude_cert(d, cid, 'default:%d' % c_all)
    assert 'W-1' not in [c['number'] for c in m.collect_certs(d, cid)]
    m.add_custom_cert(d, cid, 'Свой документ', '/tmp/own.pdf')
    assert m.collect_certs(d, cid)[-1]['title'] == 'Свой документ'
    m.restore_certs(d, cid)
    assert 'W-1' in [c['number'] for c in m.collect_certs(d, cid)]


def test_tags_formats_and_empty_equipment(d):
    cid = contract(d)
    p25, p32 = pipeline(d, 'Труба 25', 'A-25'), pipeline(d, 'Труба ПЭ 32x3', 'A-32')
    m.save_pipelines(d, cid, [dict(pipeline_id=p25, quantity=10.5)])
    j = m.joints(d, cid)
    m.save_joints(d, cid, [dict(id=j[0]['id'], name='Труба 25', count=4)])
    m.save_equipment(d, cid, [dict(kind='Котел', model='Bosch 24', serial='SN-9')])
    payments_domain.add(d, 'contracts', cid, 500, '2026-11-15', 'аванс')
    tags = t.tag_map(d, cid)
    assert tags['НОМЕР_ДОГОВОРА'] == '01-02/26' and tags['ДАТА_ДОГОВОРА'] == '12 ноября 2026' and tags['КЛИЕНТ_СОКР'] == 'И.И. Иванов'
    assert tags['НАЧАЛО_РАБОТ'] == '12 ноября 2026' and tags['ОКОНЧАНИЕ_РАБОТ'] == '11 января 2027' and tags['ДАТА_АКТА'] == '«20» декабря 2026'
    assert tags['МЕСЯЦ_РАБОТ'] == 'декабрь' and tags['ДАТА_ПРОЕКТА'] == 'октябрь 2026' and tags['НОМЕР_ПРОЕКТА'] == 'ГСВ-77' and tags['ПРОЕКТИРОВЩИК'] == 'ООО Проект'
    assert tags['СТОИМОСТЬ'] == '1 250,50' and tags['СТОИМОСТЬ_ПРОПИСЬЮ'].lower().startswith('одна тысяча двести пятьдесят') and tags['ДАТА_ВЫДАЧИ'] == '04.03.2015'
    assert tags['КОТЕЛ_МОДЕЛЬ'] == 'Bosch 24' and tags['КОТЕЛ_СЕРИЙНЫЙ_НОМЕР'] == 'SN-9'
    assert tags['СЧЁТЧИК_ГАЗА_МОДЕЛЬ'] == '' and tags['ПГ_МОДЕЛЬ'] == ''                  # такого оборудования нет — пусто
    assert tags['ТРУБОПРОВОД_ТРУБА_25'] == 'Труба 25, сертификат № A-25' and tags['ТРУБА_25'].startswith('Труба 25') and tags['ДЛИНА_25'] == '10,50' and tags['СТЫКИ_25'] == '4'
    assert tags['ТРУБОПРОВОД_ТРУБА_ПЭ_32X3'] == '' and tags['СТЫКИ_32'] == '' and tags['ТРУБА_32'] == ''           # трубопровод не используется — пусто
    assert tags['СТЫКОВ_ВСЕГО'] == '4' and tags['ОПЛАЧЕНО'] == '500,00' and tags['ОСТАТОК'] == '750,50' and tags['СУММА_ОПЛАТЫ'] == '500,00'


def test_tag_editor_rename_custom_and_dependencies(d):
    cid = contract(d)
    tags = {r[1]: r for r in t.list_tags(d)}
    t.save_tag(d, tags['КЛИЕНТ'][0], 'заказчик фио', 'client.name', 'upper', '', True)            # переименование + формат
    t.save_tag(d, None, 'Телефон_или_нет', 'client.phone', 'text', 'телефон не указан')        # собственный тег + «если пусто»
    t.save_tag(d, None, 'ПРОЕКТ_МЕСЯЦ', 'project.date', 'month_name')
    out = t.tag_map(d, cid)
    assert out['ЗАКАЗЧИК_ФИО'] == 'ИВАНОВ ИВАН ИВАНОВИЧ' and 'КЛИЕНТ' not in out and out['ПРОЕКТ_МЕСЯЦ'] == 'октябрь'
    d.execute('UPDATE crm.clients SET phone=\'\'')
    assert t.tag_map(d, cid)['ТЕЛЕФОН_ИЛИ_НЕТ'] == 'телефон не указан'
    assert [r[1] for r in t.list_tags(d)].count('ЗАКАЗЧИК_ФИО') == 1 and 'КЛИЕНТ' not in [r[1] for r in t.list_tags(d)]       # переименованный не создаётся заново
    with pytest.raises(ValueError):
        t.save_tag(d, None, 'ТЕЛЕФОН', 'client.phone', 'text')
    with pytest.raises(ValueError):
        t.save_tag(d, None, 'плохое имя!', 'client.phone', 'text')
    custom = d.fetchone("SELECT id FROM gsvm_tags WHERE name='ПРОЕКТ_МЕСЯЦ'")[0]
    assert t.delete_tag(d, custom) == 'deleted' and t.delete_tag(d, tags['ТЕЛЕФОН'][0]) == 'disabled'
    assert 'ТЕЛЕФОН' not in t.tag_map(d, cid)


def test_documents_generation_state_and_link(d, tmp_path):
    cid = contract(d)
    with pytest.raises(ValueError, match='папку'):
        t.generate(d, cid, 'contract')
    folder = m.link_folder(d, cid, tmp_path / 'моя папка', create=True)
    p25 = pipeline(d, 'Труба 25', 'A-25')
    m.save_pipelines(d, cid, [dict(pipeline_id=p25, quantity=10)])
    path = t.generate(d, cid, 'contract')
    text = '\n'.join(p.text for p in docx.Document(path).paragraphs)
    assert os.path.dirname(path) == folder and '01-02/26' in text and '12 ноября 2026' in text and '{' not in text
    assert t.doc_state(d, cid, 'contract') == 'fresh'
    m.save_pipelines(d, cid, [dict(pipeline_id=p25, quantity=10)])
    d.execute('UPDATE contracts SET contract_amount=2000 WHERE id=?', (cid,))
    assert t.doc_state(d, cid, 'contract') == 'stale'
    for kind in t.DOC_KINDS:
        assert os.path.isfile(t.generate(d, cid, kind))
    assert all(t.doc_state(d, cid, k) == 'fresh' for k in t.DOC_KINDS)
    payments_domain.add(d, 'contracts', cid, 300, '2026-11-15', 'аванс')
    payments_domain.add(d, 'contracts', cid, 100, '2026-12-01', '')
    ws = openpyxl.load_workbook(t.generate(d, cid, 'card')).active
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'Дата оплаты')
    assert [ws.cell(row + 1, c).value for c in range(1, 7)] == ['15.11.2026', 300.0, 'аванс', '01.12.2026', 100.0, None]
    ext = tmp_path / 'готовый.docx'
    docx.Document().save(ext)
    t.link_doc(d, cid, 'warranty', str(ext))
    assert t.doc_state(d, cid, 'warranty') == 'linked' and t.doc_path(d, cid, 'warranty') == str(ext)
    t.remove_doc(d, cid, 'warranty', delete_file=True)
    assert ext.exists() and t.doc_state(d, cid, 'warranty') == 'none'                 # привязанный чужой файл не удаляется
    contract_file = t.doc_path(d, cid, 'contract')
    t.remove_doc(d, cid, 'contract', delete_file=True)
    assert not os.path.exists(contract_file)


def test_consumption_by_norms_and_schedule(d):
    from smetagaz import stock_domain
    cid = contract(d)
    welder = d.execute("INSERT INTO welders(name) VALUES('Сидоров')").lastrowid
    mats = [r[0] for r in d.fetchall('SELECT id FROM stock_materials ORDER BY id')]
    pid = d.fetchone("SELECT id FROM norm_profiles WHERE diameter='25'")[0]
    stock_domain.save_profile(d, dict(name='Ду 25', diameter='25', thickness='3.2', basis='стык'), [(mats[0], '0.5'), (mats[1], '1.5')] + [(x, '0') for x in mats[2:]], pid)
    assert m.guess_profile(d, 'Труба 25') == pid and m.guess_profile(d, 'Труба без числа') is None
    rows = m.calc_consumption(d, [(pid, 4)])
    assert {r['name']: r['qty'] for r in rows} == {d.fetchone('SELECT name FROM stock_materials WHERE id=?', (mats[0],))[0]: '2', d.fetchone('SELECT name FROM stock_materials WHERE id=?', (mats[1],))[0]: '6'}
    jid = m.create_work_entry(d, cid, '2026-11-20', 'Жилой дом', welder, 'Сидоров', [('Труба 25', 4), ('Труба 32', 0)], rows)
    assert d.fetchone('SELECT owner_type,owner_id,object_text,welder_text FROM welding_jobs WHERE id=?', (jid,)) == ('contracts', cid, 'Жилой дом', 'Сидоров')
    assert d.fetchone('SELECT work_date FROM welding_days WHERE job_id=?', (jid,))[0] == '2026-11-20'
    assert sorted(x[2] for x in m.consumption_summary(d, cid)) == ['2', '6'] and len(m.consumption_rows(d, cid)) == 2
    from smetagaz import agenda_domain as a
    assert any(e['kind'] == 'work' for e in a.events_between(d, date(2026, 11, 20), date(2026, 11, 20)))
