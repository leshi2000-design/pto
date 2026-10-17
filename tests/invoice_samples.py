"""Синтетические белорусские накладные для проверки разбора: Excel (ТН-2 с объединёнными заголовками) и картинка с текстом таблицы."""
ITEMS = [
    ('Труба ПЭ100 SDR11 32х3,0', 'м', 120, 3.45, 20),
    ('Муфта электросварная ПЭ 32', 'шт', 15, 6.80, 20),
    ('Кран шаровой газовый DN25 Ду25 муфтовый', 'шт', 4, 18.50, 20),
    ('Лента сигнальная жёлтая 250 м', 'рул', 2, 12.00, 20),
]


def amounts():
    rows = []
    for name, unit, qty, price, vat in ITEMS:
        s = round(qty * price, 2)
        v = round(s * vat / 100, 2)
        rows.append((name, unit, qty, price, s, vat, v, round(s + v, 2)))
    return rows


def make_excel(path, header_two_rows=True):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'ТОВАРНАЯ НАКЛАДНАЯ ТН-2 Серия АБ № 0456789'
    ws['A2'] = 'от 12.10.2026'
    ws['A3'] = 'Грузоотправитель: ООО "ГазТехСнаб", УНП 191234567, адрес: 220030, г. Минск, ул. Мясникова, 12, оф. 5'
    ws['A4'] = 'Расчётный счёт BY20 AKBB 3012 0000 0000 0000 0000, ОАО "АСБ Беларусбанк", БИК AKBBBY2X'
    ws['A5'] = 'Грузополучатель: ЧУП "Монтаж-Газ", УНП 590876543, адрес: г. Гомель, ул. Речная, 3'
    ws['A7'], ws['B7'], ws['C7'], ws['D7'], ws['E7'], ws['F7'], ws['G7'], ws['H7'] = 'Наименование товара', 'Единица измерения', 'Количество', 'Цена, руб.', 'Стоимость, руб.', 'Ставка НДС, %', 'Сумма НДС, руб.', 'Стоимость с НДС, руб.'
    first = 8
    if header_two_rows:
        ws.merge_cells('F7:G7')
        ws['A7'], ws['B7'], ws['C7'], ws['D7'], ws['E7'] = 'Наименование товара', 'Единица измерения', 'Количество', 'Цена', 'Стоимость без НДС'
        ws['F7'], ws['H7'] = 'НДС', 'Стоимость с НДС'
        ws['F8'], ws['G8'] = 'ставка, %', 'сумма'
        first = 9
    for i, col in enumerate('ABCDEFGH', 1):
        ws[f'{col}{first}'] = i
    r = first + 1
    for name, unit, qty, price, s, vat, v, total in amounts():
        for col, val in zip('ABCDEFGH', (name, unit, qty, price, s, vat, v, total)):
            ws[f'{col}{r}'] = val
        r += 1
    a = amounts()
    ws[f'A{r}'] = 'Итого'
    ws[f'E{r}'] = round(sum(x[4] for x in a), 2)
    ws[f'G{r}'] = round(sum(x[6] for x in a), 2)
    ws[f'H{r}'] = round(sum(x[7] for x in a), 2)
    ws[f'A{r + 2}'] = 'Всего стоимость товаров с НДС: ' + str(round(sum(x[7] for x in a), 2))
    wb.save(path)
    return path


def text_lines():
    a = amounts()
    lines = ['ТОВАРНАЯ НАКЛАДНАЯ ТН-2  Серия АБ № 0456789  от 12 октября 2026 г.', '',
             'Грузоотправитель: ООО "ГазТехСнаб", УНП 191234567',
             'Адрес: 220030, г. Минск, ул. Мясникова, 12, оф. 5',
             'Р/с BY20AKBB30120000000000000000 в ОАО "АСБ Беларусбанк", БИК AKBBBY2X',
             'Грузополучатель: ЧУП "Монтаж-Газ", УНП 590876543',
             'Адрес: г. Гомель, ул. Речная, 3', '',
             '№  Наименование товара                          Ед.  Кол-во   Цена     Стоимость  НДС%  Сумма НДС  Всего с НДС']
    for n, (name, unit, qty, price, s, vat, v, total) in enumerate(a, 1):
        lines.append(f'{n:<3}{name:<44}{unit:<5}{qty:<8}{price:<9.2f}{s:<11.2f}{vat:<6}{v:<11.2f}{total:.2f}'.replace('.', ','))
    lines.append('')
    lines.append(f'Итого                                                       {sum(x[4] for x in a):.2f}      {sum(x[6] for x in a):.2f}      {sum(x[7] for x in a):.2f}'.replace('.', ','))
    lines.append('Отпуск разрешил ____________    Груз принял ____________')
    return lines


def make_image(path, noise=False, rotate=0.0, jpeg=False):
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    import random
    font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf', 24)
    lines = text_lines()
    img = Image.new('RGB', (1900, 130 + 40 * len(lines)), 'white')
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((40, 60 + 40 * i), line, font=font, fill=(20, 20, 20))
    if rotate:
        img = img.rotate(rotate, expand=True, fillcolor='white')
    if noise:
        random.seed(3)
        px = img.load()
        for _ in range(60000):
            x, y = random.randrange(img.size[0]), random.randrange(img.size[1])
            g = random.randrange(120, 255)
            px[x, y] = (g, g, g)
        img = img.filter(ImageFilter.GaussianBlur(0.8))
    img.save(path, 'JPEG', quality=70) if jpeg else img.save(path)
    return path
