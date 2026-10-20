import pytest
from smetagaz import main_window, orders_view, docs_view, counterparty_ui, contract_card, gsn_card, volumes_ui, legal_entities_view  # noqa: F401  (до подмены базы в тестах)
from smetagaz.database import DatabaseManager
from smetagaz import orders_domain as od, staff_domain as sd, volumes_domain as vd, contracts_core as cc


@pytest.fixture
def db(tmp_path):
    d = DatabaseManager(tmp_path / 'o.db')
    d.init_db()
    yield d
    d.close()


def make_staff(db, *names):
    return [sd.save_staff(db, dict(name=n, role='Монтажник')) for n in names]


def order_values(**kw):
    v = dict(order_date='2026-10-01', start_date='2026-10-02', end_date='2026-10-05', object_text='Дом на Садовой')
    v.update(kw)
    return v


def test_order_numbers_workers_schedule_and_conflicts(db):
    a, b = make_staff(db, 'Иванов', 'Петров')
    crew = sd.save_crew(db, 'Бригада 1', leader_id=a, members=[a, b])
    o1 = od.save_order(db, order_values(crew_id=crew), [], [dict(kind='Работа', name='Монтаж', unit='м', plan_qty='10')])
    o2 = od.save_order(db, order_values(order_date='2026-10-03', object_text='Другой'), [a], [])
    assert od.get_order(db, o1)['number'] == '01/26' and od.get_order(db, o2)['number'] == '02/26'
    assert od.get_order(db, o1)['workers'] and len(od.get_order(db, o1)['workers']) == 2 and od.get_order(db, o1)['leader_id'] == a
    assert db.fetchone('SELECT count(*) FROM staff_assign WHERE order_id=?', (o1,))[0] == 2
    w = od.warnings(db, [a], '2026-10-02', '2026-10-05', o2)
    assert any('Иванов' in x and 'уже назначен' in x for x in w)
    with pytest.raises(ValueError):
        od.save_order(db, order_values(number='01/26'), [b], [])        # дубль номера
    with pytest.raises(ValueError):
        od.save_order(db, order_values(end_date='2026-09-01'), [b], [])
    with pytest.raises(ValueError):
        od.save_order(db, order_values(), [], [])                        # нет работников
    od.delete_order(db, o1)
    assert db.fetchone('SELECT count(*) FROM staff_assign WHERE order_id=?', (o1,))[0] == 0


def test_expired_permit_warning(db):
    (a,) = make_staff(db, 'Сидоров')
    sd.save_doc(db, a, dict(kind='Удостоверение', title='Аттестат', valid_to='2026-01-01'))
    assert any('просрочен допуск' in x for x in od.warnings(db, [a], '2026-10-02', '2026-10-05'))


def test_close_and_apply_facts_to_volumes(db):
    (a,) = make_staff(db, 'Иванов')
    lid = cc.save_legal(db, dict(name='ООО Ромашка', unp='190123456'))
    cid = cc.save_contract(db, 'smr', dict(client_id=lid, party_type='legal', legal_id=lid, contract_date='2026-09-01', amount=0))
    vd.save_items(db, 'smr_contracts', cid, [dict(kind='Работа', name='Монтаж', unit='м', quantity='10')])
    planned = od.lines_from_volumes(db, 'smr_contracts', cid)
    assert planned[0]['plan_qty'] == 10
    oid = od.save_order(db, order_values(owner_type='smr_contracts', owner_id=cid, object_text=''), [a], planned)
    with pytest.raises(ValueError):
        od.apply_facts_to_volumes(db, oid)                               # наряд ещё не закрыт
    line = od.get_order(db, oid)['lines'][0]
    od.close_order(db, oid, {line['id']: '4'})
    assert od.get_order(db, oid)['status'] == 'Закрыт' and db.fetchone('SELECT count(*) FROM staff_assign WHERE order_id=?', (oid,))[0] == 0
    assert od.apply_facts_to_volumes(db, oid) == 1
    assert vd.items(db, 'smr_contracts', cid)[0]['quantity'] == 14
    with pytest.raises(ValueError):
        od.apply_facts_to_volumes(db, oid)


def test_generate_document_word_and_excel(db, tmp_path):
    import docx, openpyxl
    (a,) = make_staff(db, 'Иванов')
    oid = od.save_order(db, order_values(task='Смонтировать газопровод'), [a], [dict(kind='Работа', name='Монтаж', unit='м', plan_qty='10')])
    out = od.generate(db, oid, str(tmp_path / 'out'))
    text = '\n'.join(p.text for p in docx.Document(out).paragraphs)
    assert '01/26' in text and 'Иванов' in text and 'Монтаж — 10 м' in text and 'Дом на Садовой' in text
    wb = openpyxl.Workbook()
    wb.active['A1'] = 'Наряд {НОМЕР_НАРЯДА} {ОБЪЕКТ}'
    tpl = tmp_path / 'tpl.xlsx'
    wb.save(tpl)
    db.set_setting(od.TEMPLATE_SETTING, str(tpl))
    out = od.generate(db, oid, str(tmp_path / 'out2'))
    assert out.endswith('.xlsx') and openpyxl.load_workbook(out).active['A1'].value == 'Наряд 01/26 Дом на Садовой'


def _patch_db(monkeypatch, db):
    import sys
    for name, mod in list(sys.modules.items()):
        if name.startswith('smetagaz') and getattr(mod, 'db', None) is not None and mod.db.__class__.__name__ == 'DatabaseManager':
            monkeypatch.setattr(mod, 'db', db, raising=False)


def test_orders_ui_flow(db, monkeypatch, tmp_path):
    _patch_db(monkeypatch, db)
    from PyQt6.QtWidgets import QMessageBox
    from PyQt6.QtCore import Qt
    from smetagaz import orders_view as ov, gsv_project_domain as gd
    monkeypatch.setattr(QMessageBox, 'question', lambda *a, **k: QMessageBox.StandardButton.Yes)
    monkeypatch.setattr(QMessageBox, 'warning', lambda *a, **k: 0)
    monkeypatch.setattr(QMessageBox, 'information', lambda *a, **k: 0)
    monkeypatch.setattr(gd, 'DATA_DIR', tmp_path)
    monkeypatch.setattr('smetagaz.gsvm_tabs.open_file', lambda *a, **k: None)
    a, b = make_staff(db, 'Иванов', 'Петров')
    crew = sd.save_crew(db, 'Бригада 1', leader_id=a, members=[a, b])
    lid = cc.save_legal(db, dict(name='ООО Ромашка', unp='190123456'))
    cid = cc.save_contract(db, 'smr', dict(client_id=lid, party_type='legal', legal_id=lid, contract_date='2026-09-01', amount=0))
    vd.save_items(db, 'smr_contracts', cid, [dict(kind='Работа', name='Монтаж', unit='м', quantity='10')])
    d = ov.OrderDialog()
    i = d.find_object('smr_contracts', cid)
    assert i >= 0
    d.object.setCurrentIndex(i)                       # предложит взять объёмы договора
    assert d.table.rowCount() == 1
    d.crew.setCurrentIndex(d.crew.findData(crew))
    assert len(d.checked_workers()) == 2 and d.leader.currentData() == a
    assert d.save()
    oid = d.order_id
    assert od.get_order(db, oid)['number'] == '01/26' or od.get_order(db, oid)['number'].endswith('/26')
    d.table.item(0, 4).setText('7')
    d.close_order()
    o = od.get_order(db, oid)
    assert o['status'] == 'Закрыт' and o['facts_applied'] == 1 and o['lines'][0]['fact_qty'] == 7
    assert vd.items(db, 'smr_contracts', cid)[0]['quantity'] == 17
    d.make_document()
    assert db.fetchone('SELECT doc_path FROM work_orders WHERE id=?', (oid,))[0].endswith('.docx')
    v = ov.OrdersView()
    v.load_data()
    assert v.table.rowCount() == 1 and v.table.item(0, 8).text() == 'Закрыт'


def test_menu_section_workers(db, monkeypatch):
    _patch_db(monkeypatch, db)
    from smetagaz import main_window as mw
    sections = dict(mw.NAV_SECTIONS)
    assert sections['Работники'] == ['orders', 'welders', 'staff']
    assert not {'welders', 'staff'} & set(sections['Дополнительно'])
    w = mw.MainWindow()
    names = [b.text() for b in w.tabs_buttons]
    assert 'Наряды' in names and 'Сварщики' in names
    w.switch_tab(names.index('Наряды'), w.tabs_buttons[names.index('Наряды')])
    w.close()


def test_permit_requires_fields_and_has_own_numbering(db, tmp_path):
    import docx
    a, b = make_staff(db, 'Иванов', 'Петров')
    base = order_values(kind=od.KIND_PERMIT)
    with pytest.raises(ValueError, match='вид работ'):
        od.save_order(db, base, [a, b], [])
    base.update(hazard='Газоопасные работы', safety='Проветривание, контроль загазованности', performer_id=b, supervisor_id=a, admitter_id=a, time_from='09:00', time_to='17:00',
                ppe='Противогаз', briefing=True, place='Колодец ГРП')
    with pytest.raises(ValueError, match='входить в состав'):
        od.save_order(db, dict(base, performer_id=None or 9999), [a], [])
    t1 = od.save_order(db, order_values(), [a], [])
    p1 = od.save_order(db, base, [a, b], [])
    p2 = od.save_order(db, dict(base, order_date='2026-10-03'), [a, b], [])
    assert od.get_order(db, t1)['number'] == '01/26' and od.get_order(db, p1)['number'] == 'Д-01/26' and od.get_order(db, p2)['number'] == 'Д-02/26'
    assert [o['kind'] for o in od.list_orders(db, kind=od.KIND_PERMIT)] == [od.KIND_PERMIT] * 2
    out = od.generate(db, p1, str(tmp_path / 'p'))
    text = '\n'.join(p.text for p in docx.Document(out).paragraphs)
    assert 'НАРЯД-ДОПУСК № Д-01/26' in text and 'Газоопасные работы' in text and 'Петров' in text and 'Колодец ГРП' in text and 'с 02.10.2026 09:00' in text


def test_document_registers(db):
    from smetagaz import docs_register_domain as dr
    from datetime import date
    out = dict(register='out', reg_date='2026-10-01', correspondent='Горгаз', subject='Запрос технических условий', due_date='2026-10-10')
    with pytest.raises(ValueError, match='номер'):
        dr.save_doc(db, out)                                          # номер исходящего вводится вручную
    o1 = dr.save_doc(db, dict(out, reg_number='Исх-01/26'))
    o2 = dr.save_doc(db, dict(out, reg_number='125-А', subject='Второе письмо'))
    p1 = dr.save_doc(db, dict(register='order', reg_date='2026-10-02', subject='О назначении ответственного за газовое хозяйство', executor='Иванов', due_date='2026-10-05'))
    d1 = dr.save_doc(db, dict(register='other', reg_date='2026-10-02', correspondent='Госпромнадзор', subject='Лицензия на монтаж', kind='Лицензия / разрешение', due_date='2026-11-01'))
    d2 = dr.save_doc(db, dict(register='other', reg_date='2026-10-02', subject='Страховой полис', kind='Страховой полис', due_date='2026-09-01'))
    nums = [dr.get_doc(db, i)['reg_number'] for i in (o1, o2, p1, d1, d2)]
    assert nums == ['Исх-01/26', '125-А', 'Пр-01/26', 'Д-01/26', 'Д-02/26']
    with pytest.raises(ValueError):
        dr.save_doc(db, dict(out, reg_number='Исх-01/26'))
    with pytest.raises(ValueError, match='кому'):
        dr.save_doc(db, dict(out, reg_number='5', correspondent=''))
    with pytest.raises(ValueError):
        dr.save_doc(db, dict(out, reg_number='6', due_date='2026-09-01'))
    with pytest.raises(ValueError):
        dr.save_doc(db, dict(register='order', reg_date='2026-10-02', subject='x', status='Отправлен'))        # статус чужого журнала
    today = date(2026, 10, 20)
    assert dr.listing(db, 'out', today=today)[0] == 2 and dr.listing(db, 'order')[0] == 1 and dr.listing(db, 'other')[0] == 2
    assert dr.problems(db, 'out', today) == 2 and dr.problems(db, 'order', today) == 1 and dr.problems(db, 'other', today) == 1       # полис истёк
    rows = {r['id']: r for r in dr.listing(db, 'other', today=today)[1]}
    assert rows[d1]['state'] == 'soon' and rows[d2]['state'] == 'overdue'                                                             # лицензия истекает в течение 30 дней
    dr.set_status(db, o1, 'Получен ответ')
    assert dr.problems(db, 'out', today) == 1
    with pytest.raises(ValueError):
        dr.set_status(db, p1, 'Отправлен')
    assert dr.listing(db, 'out', search='второе')[0] == 1
    # прежний журнал исходящих юрлиц попадает в «Регистрацию исходящих»
    db.execute("INSERT INTO le_outgoing(reg_number, reg_date, recipient, subject) VALUES('7','2026-02-01','ООО Ромашка','Письмо о сроках')")
    assert dr.import_outgoing(db) == 1 and dr.import_outgoing(db) == 0
    assert dr.listing(db, 'out', search='ромашка')[1][0]['reg_number'] == '7'                 # прежний номер сохраняется как есть


def test_migrate_from_4_7_journal(db):
    """Записи журнала версии 4.7.0 (без колонки register) раскладываются по новым журналам."""
    db.execute('DROP TABLE doc_register')
    db.execute("""CREATE TABLE doc_register(id INTEGER PRIMARY KEY, direction TEXT NOT NULL DEFAULT 'Входящий', seq_num INTEGER, year_num INTEGER, reg_number TEXT NOT NULL, reg_date TEXT NOT NULL,
        kind TEXT DEFAULT 'Письмо', doc_number TEXT DEFAULT '', doc_date TEXT DEFAULT '', correspondent TEXT NOT NULL DEFAULT '', client_id INTEGER, subject TEXT DEFAULT '',
        executor TEXT DEFAULT '', due_date TEXT DEFAULT '', status TEXT NOT NULL DEFAULT 'Зарегистрирован', answer_ref TEXT DEFAULT '', file_path TEXT DEFAULT '', note TEXT DEFAULT '',
        legacy_outgoing_id INTEGER UNIQUE)""")
    db.execute("INSERT INTO doc_register(direction,seq_num,year_num,reg_number,reg_date,correspondent,subject) VALUES('Исходящий',1,2026,'Исх-01/26','2026-10-01','А','x')")
    db.execute("INSERT INTO doc_register(direction,seq_num,year_num,reg_number,reg_date,correspondent,subject,status) VALUES('Входящий',1,2026,'Вх-01/26','2026-10-01','Б','y','На исполнении')")
    from smetagaz import docs_register_domain as dr
    dr.migrate(db)
    assert db.fetchone("SELECT register FROM doc_register WHERE reg_number='Исх-01/26'")[0] == 'out'
    assert db.fetchone("SELECT register,status FROM doc_register WHERE reg_number='Вх-01/26'") == ('other', 'Действует')


def test_documents_view_and_menu(db, monkeypatch):
    _patch_db(monkeypatch, db)
    from smetagaz import docs_view as dv, main_window as mw, counterparty_ui as cu
    v = dv.DocumentsView()
    names = [v.tabs.tabText(i) for i in range(v.tabs.count())]
    assert names == ['Регистрация исходящих', 'Регистрация приказов', 'Прочие документы по организации']
    for i in range(v.tabs.count()):
        v.tabs.setCurrentIndex(i)
        v.load_data()
    from smetagaz import docs_register_domain as dr
    for reg in (dr.OUT, dr.ORDER, dr.OTHER):
        d = dv.DocRegisterDialog(reg)
        d.subject.setPlainText('Тест')
        d.number.setText('77')
        d.correspondent.setText('Кто-то')
        assert d.values()['register'] == reg
        d.save()
        assert d.doc_id
        dv.DocRegisterDialog(reg, d.doc_id)
    assert 'Исходящая документация' not in [cu.LegalDirectoryView().tabs.tabText(i) for i in range(cu.LegalDirectoryView().tabs.count())]
    assert dict(mw.NAV_SECTIONS)['Документы'] == ['docs']
