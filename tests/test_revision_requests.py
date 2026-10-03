from pathlib import Path
from docx import Document
import pytest
from smetagaz.database import DatabaseManager
from smetagaz import payments_domain as pay,dossier_domain as dossier,template_domain as tpl
from smetagaz.backup import create_backup,restore_backup
@pytest.fixture
def db(tmp_path):
    d=DatabaseManager(tmp_path/'smetagaz.db');d.init_db();yield d;d.close()

def test_payment_shared_no_double_count_and_sections(db):
    e=db.execute("INSERT INTO estimates(title,date,total,paid) VALUES('Смета','2026-08-01',1000,0)").lastrowid
    c=db.execute("INSERT INTO gsn_projects(title,contract_number,contract_amount) VALUES('Дом','ГСН-1',1000)").lastrowid
    pay.add(db,'gsn_projects',c,'100.25','2026-09-01');pay.link(db,'gsn_projects',c,e);pay.add(db,'estimates',e,'200.35','2026-09-02')
    assert pay.summary(db,'gsn_projects',c)['paid']==pay.summary(db,'estimates',e)['paid']==__import__('decimal').Decimal('300.60')
    assert db.fetchone('SELECT paid FROM estimates WHERE id=?',(e,))[0]==300.6
    rows=pay.report(db,'2026-09-01','2026-09-30','gsn_projects');assert len(rows)==2;assert sum(r['amount'] for r in rows)==300.6
    assert not pay.report(db,'2026-09-01','2026-09-30','estimates')
    with pytest.raises(ValueError):pay.link(db,'gsn_projects',c,None)
    g=db.execute("INSERT INTO gsv_projects(object_name,cost) VALUES('Проект',500)").lastrowid
    with pytest.raises(ValueError):pay.link(db,'gsv_projects',g,e)
    with pytest.raises(ValueError):pay.add(db,'gsn_projects',c,'0.001','2026-09-03')
    with pytest.raises(Exception):db.execute('DELETE FROM gsn_projects WHERE id=?',(c,))

def test_payment_opening_kept_but_not_in_period(db):
    e=db.execute("INSERT INTO estimates(title,total,paid) VALUES('Старая оплата',1000,400)").lastrowid
    pay.add(db,'estimates',e,50,'2026-09-20');assert pay.summary(db,'estimates',e)['paid']==450
    assert sum(r['amount'] for r in pay.report(db,'2026-09-01','2026-09-30'))==50

def setup_dossier(db,tmp_path):
    c=db.execute("INSERT INTO gsn_projects(title,contract_number,contract_date,phone) VALUES('Дом','ГСН-2','2026-09-20','111')").lastrowid;folder=tmp_path/'object';folder.mkdir();tpl.save_details(db,'gsn_projects',c,str(folder),{},[])
    for slot,tag in [('title','номер_договора'),('contract','phone')]:
        source=tmp_path/(slot+'.docx');doc=Document();doc.add_paragraph('{{'+tag+'}}');doc.save(source);tpl.save_template(db,'gsn_projects',slot,str(source),slot)
    did=dossier.create(db,'ГСН','gsn_projects',c);return c,did,folder

def test_dossier_incremental_and_manual_protection(db,tmp_path):
    c,did,folder=setup_dossier(db,tmp_path);r=dossier.generate(db,did);assert r['changed']==2;paths={slot:Path(file) for slot,file in db.fetchall('SELECT slot,file_path FROM dossier_files WHERE dossier_id=?',(did,))};times={s:p.stat().st_mtime_ns for s,p in paths.items()}
    assert dossier.generate(db,did)['unchanged']==2;assert all(p.stat().st_mtime_ns==times[s] for s,p in paths.items())
    db.execute("UPDATE gsn_projects SET phone='222' WHERE id=?",(c,));r=dossier.generate(db,did);assert r['changed']==1 and r['unchanged']==1;assert paths['title'].stat().st_mtime_ns==times['title'];assert Document(paths['contract']).paragraphs[0].text=='222'
    paths['contract'].write_bytes(b'manually modified');db.execute("UPDATE gsn_projects SET phone='333' WHERE id=?",(c,))
    with pytest.raises(ValueError,match='вручную'):dossier.generate(db,did)
    assert paths['contract'].read_bytes()==b'manually modified'

def test_equipment_certificate_and_manual_extra(db,tmp_path):
    c,did,folder=setup_dossier(db,tmp_path);source=tmp_path/'equipment.pdf';source.write_bytes(b'passport');cert=db.execute('INSERT INTO certificates(name,cert_number,file_path) VALUES(?,?,?)',('Сигнализатор','P-1',str(source))).lastrowid
    from smetagaz.gsv_domain import save_equipment
    save_equipment(db,'gsn_projects',c,[dict(equipment_kind='Сигнализатор',equipment_model='A1',certificate_number='P-1',linked_cert_id=cert)])
    db.execute('INSERT INTO dossier_certificates VALUES(?,?)',(did,cert));assert len(dossier.certificates(db,did))==1
    result=dossier.generate(db,did);assert result['changed']==3;assert any(p.suffix=='.pdf' for p in Path(result['folder']).iterdir())

def test_dossier_backup_paths_and_regeneration(db,tmp_path):
    c,did,folder=setup_dossier(db,tmp_path);dossier.generate(db,did);backup=tmp_path/'backup.zip';create_backup(db,backup);restored=tmp_path/'restored';restore_backup(backup,restored);other=DatabaseManager(restored/'smetagaz.db');other.init_db()
    assert dossier.generate(other,did)['unchanged']==2
    other.execute("UPDATE gsn_projects SET phone='444' WHERE id=?",(c,));assert dossier.generate(other,did)['changed']==1;other.close()

def test_free_dossier_and_gsv_templates(db,tmp_path):
    # "Свободный" (owner-less) dossier still lists the real ГСВ document set; none of the 8
    # slots has a template file attached yet, so generation has nothing to produce.
    folder=tmp_path/'free';folder.mkdir();did=dossier.create(db,'ГСВ',title='Свободный объект',folder=str(folder));assert len(dossier.templates(db,did))==8
    with pytest.raises(ValueError):dossier.generate(db,did)

def test_signature_estimate_bottom(db,tmp_path):
    from smetagaz.report_templates import export
    rid=db.execute("INSERT INTO estimates(title,total,prepared_by) VALUES('Смета',123,'Иванов И.И.')").lastrowid;source=tmp_path/'signature.docx';doc=Document();doc.add_paragraph('{{title}}');doc.save(source);tid=db.execute('INSERT INTO report_templates(kind,name,file_path) VALUES(?,?,?)',('estimates','Смета',str(source))).lastrowid;out=tmp_path/'out.docx';export(db,'estimates',rid,tid,out);assert 'Составил: Иванов И.И.' in Document(out).paragraphs[-1].text
