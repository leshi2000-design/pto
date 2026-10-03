from pathlib import Path
from zipfile import ZipFile
import pytest
from docx import Document
from openpyxl import Workbook,load_workbook
from smetagaz.database import DatabaseManager
from smetagaz import template_domain as d
from smetagaz.template_engine import transform,generate
from smetagaz.gsv_domain import save_client,dossier
from smetagaz.backup import create_backup,restore_backup

@pytest.fixture
def base(tmp_path):
    data=tmp_path/'data';data.mkdir();db=DatabaseManager(data/'smetagaz.db');db.init_db()
    cid=save_client(db,dict(name='Иванов Иван Иванович',phone='123',address='Минск',passport='AB123',passport_issuer='РОВД',passport_date='2010-01-02'))
    rid=db.execute('INSERT INTO gsn_projects(title,contract_number,contract_date,client_id,client_name) VALUES(?,?,?,?,?)',('Газоснабжение дома','42/26','2026-09-19',cid,'Иванов Иван Иванович')).lastrowid
    folder=tmp_path/'object';folder.mkdir();d.save_details(db,'gsn_projects',rid,str(folder),{'project_number':'ПД-42','project_date':'2026-05-01'},[])
    yield db,rid,folder
    db.close()

def test_word_split_tags_headers_table_and_preserve(tmp_path):
    source=tmp_path/'template.docx';out=tmp_path/'out.docx';doc=Document();p=doc.add_paragraph();p.add_run('До {{фи').bold=True;p.add_run('о}} после {{номер_договора}}');doc.sections[0].header.paragraphs[0].text='{{фио}}';doc.add_table(rows=1,cols=1).cell(0,0).text='{{materials}}';doc.save(source);original=source.read_bytes()
    assert transform(source)=={'фио','номер_договора','materials'}
    transform(source,{'фио':'Иванов & Петров','номер_договора':'42','materials':'Труба\nМуфта'},out)
    result=Document(out);assert result.paragraphs[0].text=='До Иванов & Петров после 42';assert result.paragraphs[0].runs[0].bold;assert result.sections[0].header.paragraphs[0].text=='Иванов & Петров';assert result.tables[0].cell(0,0).text=='Труба\nМуфта';assert source.read_bytes()==original

def test_excel_formulas_styles_and_no_injection(tmp_path):
    source=tmp_path/'t.xlsx';out=tmp_path/'out.xlsx';w=Workbook();s=w.active;s['A1']='{{фио}}';s['A2']='=SUM(B1:B2)';s['B1']=2;s['B2']=3;s.merge_cells('C1:D1');s['C1']='№ {{номер_договора}}';s['C1'].number_format='@';w.save(source)
    transform(source,{'фио':'=HYPERLINK("bad")','номер_договора':'17'},out)
    result=load_workbook(out);assert result.active['A1'].data_type=='s';assert result.active['A2'].value=='=SUM(B1:B2)';assert result.active['C1'].value=='№ 17';assert 'C1:D1' in result.active.merged_cells;assert result.active['C1'].number_format=='@'

def test_context_client_quantities_certificates(base,tmp_path):
    db,rid,folder=base;certfile=tmp_path/'certificate.pdf';certfile.write_bytes(b'certificate');cid=db.execute('INSERT INTO certificates(name,cert_number,file_path) VALUES(?,?,?)',('Труба','C42',str(certfile))).lastrowid
    rows=[dict(kind='pe_pipe',name='Труба 25',unit='м',quantity='10.25',certificate_id=cid),dict(kind='pe_pipe',name='Труба 32',unit='м',quantity='5',certificate_id=cid)]
    d.save_details(db,'gsn_projects',rid,str(folder),{},rows);ctx,mats=d.context(db,'gsn_projects',rid)
    assert ctx['фио']=='Иванов Иван Иванович';assert ctx['ио_фамилия']=='И.И.Иванов';assert ctx['фамилия_ио']=='Иванов И.И.';assert ctx['дата_договора_кратко']=='19.09.26';assert ctx['pe_pipe_quantity']=='15.25';assert ctx['pe_pipe_certificate_number']=='C42';assert 'РОВД' in ctx['паспорт']
    client=db.fetchone('SELECT client_id FROM gsn_projects WHERE id=?',(rid,))[0];db.execute('UPDATE crm.clients SET name=? WHERE id=?',('Петров Пётр Петрович',client));assert db.fetchone('SELECT client_name FROM gsn_projects WHERE id=?',(rid,))[0]=='Петров Пётр Петрович'
    with pytest.raises(ValueError):d.save_details(db,'gsn_projects',rid,str(folder),{},[dict(kind='coupling',name='Муфта',unit='шт',quantity='1.5')])
    assert len(d.material_rows(db,'gsn_projects',rid))==2

def test_generation_atomic_unique_and_registry(base,tmp_path):
    db,rid,folder=base;src=tmp_path/'title.docx';doc=Document();doc.add_paragraph('{{фио}} / {{номер_договора}}');doc.save(src);d.save_template(db,'gsn_projects','title',str(src),'Титул {{номер_договора}}')
    files=generate(db,'gsn_projects',rid,['title']);second=generate(db,'gsn_projects',rid,['title']);assert set(files).isdisjoint(second);assert len(list(folder.iterdir()))==2
    bad=tmp_path/'bad.docx';doc=Document();doc.add_paragraph('{{missing}}');doc.save(bad);d.save_template(db,'gsn_projects','contract',str(bad),'Договор')
    with pytest.raises(ValueError,match='missing'):generate(db,'gsn_projects',rid,['title','contract'])
    assert len(list(folder.iterdir()))==2;assert db.fetchone('SELECT count(*) FROM executive_generated')[0]==2
    reg=tmp_path/'reg.xlsx';w=Workbook();w.active['A1']='{{реестр}}';w.save(reg);d.save_template(db,'gsn_projects','register',str(reg),'Реестр');outputs=generate(db,'gsn_projects',rid,['title','register']);assert Path(outputs[0]).name in load_workbook(outputs[1]).active['A1'].value

def test_overrides_and_backup(base,tmp_path):
    db,rid,folder=base;src=tmp_path/'title.docx';doc=Document();doc.add_paragraph('{{фио}}');doc.save(src);d.save_template(db,'gsn_projects','title',str(src),'Титул');d.save_template(db,'gsn_projects','title',str(src),'Личный',rid);generate(db,'gsn_projects',rid,['title']);assert d.templates(db,'gsn_projects',rid)[0]['override']
    backup=tmp_path/'backup.zip';manifest=create_backup(db,backup);assert not manifest['missing'];restored=tmp_path/'restore';restore_backup(backup,restored);other=DatabaseManager(restored/'smetagaz.db');other.init_db()
    linked,_=d.load_details(other,'gsn_projects',rid);assert d.path(other,linked).is_dir();assert d.path(other,d.templates(other,'gsn_projects',rid)[0]['file_path']).is_file();assert generate(other,'gsn_projects',rid,['title']);other.close()
    with pytest.raises(ValueError,match='внутри папки'):create_backup(db,folder/'bad.zip')

def test_gsv_certificate_dossier(base,tmp_path):
    db,_,folder=base;rid=db.execute('INSERT INTO contracts(object_name) VALUES(?)',('ГСВ',)).lastrowid;cid=db.execute('INSERT INTO certificates(name,cert_number) VALUES(?,?)',('Муфта','C1')).lastrowid;d.save_details(db,'contracts',rid,str(folder),{},[dict(kind='coupling',name='Муфта 25',unit='шт',quantity='3',certificate_id=cid)]);assert dossier(db,'contracts',rid)[0]['number']=='C1';assert d.context(db,'contracts',rid)[0]['coupling_certificate_number']=='C1'


def test_excel_numeric_quantity_and_invalid_tag(tmp_path):
    src=tmp_path/'q.xlsx';out=tmp_path/'q-out.xlsx';w=Workbook();w.active['A1']='{{труба_пэ_количество}}';w.active['A2']='=A1*2';w.save(src)
    transform(src,{'труба_пэ_количество':'12.5'},out);result=load_workbook(out);assert result.active['A1'].value==12.5;assert result.active['A1'].data_type=='n';assert result.active['A2'].value=='=A1*2'
    doc=Document();doc.add_paragraph('{{неверный тег}}');bad=tmp_path/'bad.docx';doc.save(bad)
    with pytest.raises(ValueError,match='Некорректный тег'):transform(bad)
