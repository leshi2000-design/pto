from datetime import datetime
import openpyxl
import pytest
from smetagaz.database import DatabaseManager
from smetagaz import gsv_project_import as imp
from smetagaz import gsv_project_domain as g

HEADERS = ['№ ПД', 'Наименование объекта строительства', 'Адрес', 'Заказчик', 'Телефон', 'Паспортные данные', '№ договора',
           'Дата заключения', 'Дата Акта', 'Стоимость', 'Срок выполнения', 'ТУ']


@pytest.fixture
def d(tmp_path, monkeypatch):
    monkeypatch.setattr(g, 'PROJECTS_DIR', tmp_path / 'projects')
    d = DatabaseManager(tmp_path / 'smetagaz.db')
    d.init_db()
    yield d
    d.close()


def make_xlsx(path, rows, title_rows=0):
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(title_rows):
        ws.append(['Реестр договоров на проектирование'])
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    wb.save(path)


ROWS = [
    ['01-26 ГСВ', 'Жилой дом', 'г. Минск, ул. Лесная 5', 'Иванов Иван Иванович', '375291112233', 'МР 1234567, выдан Фрунзенским РУВД 04.03.2015', '01-03/26',
     datetime(2026, 1, 12), datetime(2026, 2, 10), 250, datetime(2026, 2, 11), '№ 15 от 05.01.2026'],
    ['2-26', 'Баня', 'д. Ключи', 'Петров Пётр', '375447778899', 'КН 7654321', '', '15.02.2026', '', '320,50 руб.', '', ''],
    ['03-26 ГСВ', 'Гараж', 'Брест', 'иванов  иван иванович', '', '', '03-03/26', '12 марта 2026г.', '', '1 250,00', '12 апреля 2026 г.', 'ТУ 99'],
    ['04-26 ГСВ', 'Сарай', '', '', '', '', '', '', '', '', '', ''],
    ['01-26 ГСВ', 'Дубль в таблице', '', 'Сидоров', '', '', '', '', '', '', '', ''],
    ['', 'Без номера ПД', '', 'Козлов', '', '', '', '31.02.2026', '', 'abc', '', ''],
]


def test_auto_map_by_user_headers():
    m = imp.auto_map(HEADERS)
    assert [m[i] for i in range(12)] == ['pd_number', 'object_name', 'client_address', 'client_name', 'phone', 'passport', 'contract_number',
                                         'contract_date', 'act_date', 'cost', 'due_date', 'tu_text']


def test_parsers():
    assert imp.parse_date('12 ноября 2026г.') == ('2026-11-12', False)
    assert imp.parse_date('12.11.26') == ('2026-11-12', False)
    assert imp.parse_date(datetime(2026, 1, 2)) == ('2026-01-02', False)
    assert imp.parse_date(46000)[1] is False
    assert imp.parse_date('31.02.2026') == ('', True) and imp.parse_date('') == ('', False)
    assert imp.parse_cost('1 250,50 руб.') == (1250.5, False) and imp.parse_cost('abc') == (None, True) and imp.parse_cost(250) == (250.0, False)
    assert imp.parse_pd('2-26') == (2, 26, '02-26 ГСВ') and imp.parse_pd('01-26 ГСВ')[2] == '01-26 ГСВ' and imp.parse_pd('мусор')[0] is None


def test_read_table_finds_header_below_title(tmp_path):
    path = tmp_path / 'in.xlsx'
    make_xlsx(path, ROWS[:1], title_rows=2)
    headers, rows, header_row = imp.read_table(path)
    assert headers == HEADERS and header_row == 3 and len(rows) == 1


def test_full_import(d, tmp_path):
    path = tmp_path / 'in.xlsx'
    make_xlsx(path, ROWS)
    headers, rows, _ = imp.read_table(path)
    prepared = imp.prepare(d, rows, imp.auto_map(headers))
    assert [p['status'] for p in prepared] == ['new', 'new', 'new', 'error', 'error', 'new']
    assert 'нет заказчика' in prepared[3]['messages'][0] and 'повторяется' in prepared[4]['messages'][-1]
    assert len(prepared[5]['messages']) == 2 and prepared[5]['data']['contract_date'] == ''      # неверные дата и стоимость — замечания, строка импортируется
    report = imp.import_rows(d, prepared)
    assert (report['created'], report['updated'], len(report['errors'])) == (4, 0, 2)
    row = d.execute("SELECT * FROM gsv_projects WHERE pd_number='01-26 ГСВ'").fetchone()
    cols = [c[0] for c in d.execute('SELECT * FROM gsv_projects LIMIT 0').description]
    p = dict(zip(cols, row))
    assert p['object_name'] == 'Жилой дом' and p['client_address'] == 'г. Минск, ул. Лесная 5' and p['contract_number'] == '01-03/26'
    assert (p['contract_date'], p['act_date'], p['due_date'], p['cost'], p['tu_text']) == ('2026-01-12', '2026-02-10', '2026-02-11', 250, '№ 15 от 05.01.2026')
    assert (p['seq_num'], p['year_num'], p['contract_signed'], p['act_signed']) == (1, 26, 1, 1) and not p['project_folder']      # папку пользователь привязывает кнопкой
    assert g.project_statuses(d, p['id'], 'work') == ['Сделано'] and g.project_statuses(d, p['id'], 'client') == ['Акт подписан']
    client = d.fetchone('SELECT name,phone,address,passport FROM crm.clients WHERE id=?', (p['client_id'],))
    assert client == ('Иванов Иван Иванович', '375291112233', 'г. Минск, ул. Лесная 5', 'МР 1234567, выдан Фрунзенским РУВД 04.03.2015')
    second = d.fetchone("SELECT pd_number,contract_number,cost,due_date,contract_date FROM gsv_projects WHERE object_name='Баня'")
    assert second == ('02-26 ГСВ', '02-03/26', 320.5, '2026-03-17', '2026-02-15')            # срок = дата заключения + 30 дней
    third = d.fetchone("SELECT client_id,cost,due_date FROM gsv_projects WHERE object_name='Гараж'")
    assert third[0] == p['client_id'] and third[1:] == (1250.0, '2026-04-12')                 # тот же клиент по ФИО, второй проект
    assert d.fetchone('SELECT count(*) FROM crm.clients')[0] == 3          # Иванов, Петров, Козлов
    assert d.fetchone("SELECT pd_number,cost FROM gsv_projects WHERE object_name='Без номера ПД'")[1] == 250.0
    # календарь видит импортированные договоры
    from smetagaz import agenda_domain as a
    from datetime import date
    assert any(e['title'].startswith('Заключение договора Иванов И.И.') for e in a.events_between(d, date(2026, 1, 12), date(2026, 1, 12)))


def test_repeat_import_skips_or_updates(d, tmp_path):
    path = tmp_path / 'in.xlsx'
    make_xlsx(path, ROWS[:1])
    headers, rows, _ = imp.read_table(path)
    imp.import_rows(d, imp.prepare(d, rows, imp.auto_map(headers)))
    again = imp.prepare(d, rows, imp.auto_map(headers))
    assert again[0]['status'] == 'exists'
    assert imp.import_rows(d, again)['skipped'] == 1
    rows[0][9] = 400
    report = imp.import_rows(d, imp.prepare(d, rows, imp.auto_map(headers)), on_exists='update')
    assert report['updated'] == 1 and d.fetchone('SELECT count(*), max(cost) FROM gsv_projects') == (1, 400.0)
    assert d.fetchone('SELECT count(*) FROM crm.clients')[0] == 1


def test_manual_mapping_address_as_object_address(d, tmp_path):
    path = tmp_path / 'in.xlsx'
    make_xlsx(path, ROWS[:1])
    headers, rows, _ = imp.read_table(path)
    mapping = imp.auto_map(headers)
    mapping[2] = 'object_address'
    imp.import_rows(d, imp.prepare(d, rows, mapping))
    assert d.fetchone('SELECT address,client_address FROM gsv_projects') == ('г. Минск, ул. Лесная 5', '')


def test_tu_tag_and_export(d, tmp_path):
    path = tmp_path / 'in.xlsx'
    make_xlsx(path, ROWS[:1])
    headers, rows, _ = imp.read_table(path)
    imp.import_rows(d, imp.prepare(d, rows, imp.auto_map(headers)))
    pid = d.fetchone('SELECT id FROM gsv_projects')[0]
    g.link_folder(d, pid, tmp_path / 'папка проекта', create=True)
    assert g.project_tags(d, pid)['ТУ'] == '№ 15 от 05.01.2026'
    ws = openpyxl.load_workbook(g.generate(d, pid, 'card')).active
    assert any(ws.cell(r, 2).value == '№ 15 от 05.01.2026' for r in range(1, ws.max_row + 1))
