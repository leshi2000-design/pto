from pathlib import Path
from decimal import Decimal
import json
import hashlib
import sqlite3
import zipfile
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from shop.core import Store, scaled
from shop import backup, documents, sources


@pytest.fixture
def store(tmp_path):
    store = Store(tmp_path / 'data')
    yield store
    store.close()


def setup(store):
    product = store.product('A-1', 'Труба <25>', 'м')
    party = store.party('company', 'Покупатель', '123456789', 'Минск')
    return product, party


def doc(store, kind, pid, party, quantity='2', price='10.01', tax='20'):
    return store.document(kind, party, [{'product_id': pid, 'quantity': quantity, 'price': price, 'tax': tax}])


def test_ledger_and_shortage_atomic(store):
    pid, party = setup(store)
    receipt = doc(store, 'receipt', pid, party, '5')
    assert store.stock(pid) == 0
    store.post(receipt)
    assert store.stock(pid) == 5000
    invoice = doc(store, 'invoice', pid, party)
    store.post(invoice)
    assert store.stock(pid) == 5000
    shipment = store.shipment_from_invoice(invoice)
    store.post(shipment)
    assert store.stock(pid) == 3000
    with pytest.raises(ValueError):
        store.post(shipment)
    bad = doc(store, 'shipment', pid, party, '4')
    with pytest.raises(ValueError, match='Недостаточно'):
        store.post(bad)
    assert store.stock(pid) == 3000
    assert store.rows('SELECT status FROM documents WHERE id=?', (bad,))[0]['status'] == 'draft'
    with pytest.raises(ValueError, match='уже отгружен'):
        store.cancel(receipt)
    store.cancel(shipment)
    store.cancel(receipt)
    assert store.stock(pid) == 0


def test_duplicate_lines_and_transaction_rollback(store):
    pid, party = setup(store)
    receipt = doc(store, 'receipt', pid, party, '3')
    store.post(receipt)
    shipment = store.document('shipment', party, [dict(product_id=pid, quantity='2', price='1')]*2)
    with pytest.raises(ValueError):
        store.post(shipment)
    assert store.stock(pid) == 3000
    before = store.rows('SELECT * FROM documents')
    with pytest.raises(ValueError):
        store.document('receipt', party, [dict(product_id=pid, quantity='1', price='1'), dict(product_id=99999, quantity='1', price='1')])
    assert store.rows('SELECT * FROM documents') == before


@pytest.mark.parametrize('value', ['-1', 'NaN', 'Infinity', '0.0001', 'bad'])
def test_invalid_quantities(value):
    with pytest.raises(ValueError):
        scaled(value, 1000)


def test_snapshot_and_tax(store):
    pid, party = setup(store)
    invoice = doc(store, 'invoice', pid, party, '0.125', '10.01')
    store.post(invoice)
    store.offer(pid, 'Поставщик', '99', 'BYN', 'test')
    store.conn.execute('UPDATE products SET name=? WHERE id=?', ('Новое название', pid))
    store.conn.commit()
    _, values, items = documents.context(store, invoice)
    assert items[0]['name'] == 'Труба <25>'
    assert items[0]['price'] == '10.01'
    assert values['subtotal'] == '1.25'
    assert values['tax_total'] == '0.25'
    assert values['total'] == '1.50'


def test_import_atomic_and_price_only(store, tmp_path):
    path = tmp_path / 'prices.csv'
    path.write_text('Артикул;Название;Цена\nX;Кран;10,20\n', encoding='utf-8-sig')
    mapping = {'sku': 0, 'name': 1, 'price': 2}
    assert sources.import_price(store, path, mapping, 'Поставщик') == 1
    pid = store.catalog()[0]['id']
    assert store.stock(pid) == 0
    path.write_text('Артикул;Название;Цена\nX;Кран;11,20\n', encoding='utf-8')
    sources.import_price(store, path, mapping, 'Поставщик')
    assert store.rows('SELECT * FROM offers')[0]['price'] == 1120
    assert len(store.rows('SELECT * FROM offers')) == 1
    path.write_text('Артикул;Название;Цена\nY;Новый;1\nZ;Ошибка;нет\n', encoding='utf-8')
    with pytest.raises(ValueError):
        sources.import_price(store, path, mapping, 'Поставщик')
    assert len(store.catalog()) == 1


def test_web_parse_and_update(store):
    config = {'supplier': 'Сайт', 'currency': 'BYN'}
    html = '<script type="application/ld+json">' + json.dumps({'@type': 'Product', 'name': 'Кран', 'sku': 'K1', 'offers': {'price': '12.34', 'priceCurrency': 'BYN'}}) + '</script>'
    products, _ = sources.products_from_html(html, 'https://supplier.example/item', config)
    sources.apply_web_prices(store, config, products)
    products[0]['price'] = '14.50'
    sources.apply_web_prices(store, config, products)
    assert len(store.catalog()) == 1
    assert store.catalog()[0]['stock'] == 0
    assert store.rows('SELECT price FROM offers')[0]['price'] == 1450
    html = '<div class="card"><b class="name">Труба</b><span class="cost">15,25 руб.</span><a href="/p">Товар</a></div>'
    config.update(card_selector='.card', name_selector='.name', price_selector='.cost')
    parsed, _ = sources.products_from_html(html, 'https://supplier.example/category', config)
    assert parsed[0]['price'] == '15.25'
    assert parsed[0]['url'] == 'https://supplier.example/p'


def test_templates_html_docx_xlsx(store, tmp_path):
    pid, party = setup(store)
    invoice = doc(store, 'invoice', pid, party)
    store.post(invoice)
    documents.ensure_templates(store)
    output = tmp_path / 'invoice.html'
    documents.export(store, invoice, store.root / 'templates/invoice.html', output)
    assert 'Труба &lt;25&gt;' in output.read_text()
    assert '24.02' in output.read_text()
    word = Document()
    paragraph = word.add_paragraph()
    paragraph.add_run('{{num').bold = True
    paragraph.add_run('ber}}')
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '{{item.name}}'
    table.cell(0, 1).text = '{{item.total}}'
    template = tmp_path / 'template.docx'
    word.save(template)
    output = tmp_path / 'invoice.docx'
    documents.export(store, invoice, template, output)
    result = Document(output)
    assert result.paragraphs[0].text.startswith('СЧ-')
    assert result.paragraphs[0].runs[0].bold
    assert result.tables[0].cell(0, 0).text == 'Труба <25>'
    book = Workbook()
    book.active.append(['{{number}}'])
    book.active.append(['{{item.name}}', '{{item.quantity}}', '{{item.total}}'])
    template = tmp_path / 'template.xlsx'
    book.save(template)
    output = tmp_path / 'invoice.xlsx'
    documents.export(store, invoice, template, output)
    assert load_workbook(output).active['A2'].value == 'Труба <25>'


def test_backup_restore_integrity(store, tmp_path):
    pid, party = setup(store)
    receipt = doc(store, 'receipt', pid, party)
    store.post(receipt)
    documents.ensure_templates(store)
    (store.root / 'documents/example.html').write_text('test')
    archive = backup.create(store)
    root = backup.restore(archive, tmp_path / 'restored')
    restored = Store(root)
    assert restored.stock(pid) == 2000
    assert (root / 'documents/example.html').read_text() == 'test'
    restored.close()
    damaged = tmp_path / 'damaged.zip'
    with zipfile.ZipFile(archive) as source, zipfile.ZipFile(damaged, 'w') as target:
        for name in source.namelist():
            target.writestr(name, b'corrupted' if name == 'shop.db' else source.read(name))
    with pytest.raises(ValueError, match='Контрольная'):
        backup.restore(damaged, tmp_path / 'bad-restore')
    assert not (tmp_path / 'bad-restore').exists()


def test_restore_path_traversal(tmp_path):
    archive = tmp_path / 'evil.zip'
    with zipfile.ZipFile(archive, 'w') as z:
        empty_hash = hashlib.sha256(b'').hexdigest()
        z.writestr('manifest.json', json.dumps({'version': 1, 'files': {'shop.db': empty_hash, '../evil': empty_hash}}))
        z.writestr('shop.db', '')
        z.writestr('../evil', '')
    with pytest.raises(ValueError):
        backup.restore(archive, tmp_path / 'restored')
    assert not (tmp_path / 'evil').exists()


def test_crawl_category_pages_and_product_links(monkeypatch):
    pages = {
        'https://supplier.example/cat': '<a href="/cat?page=2">2</a><div class="product"><a href="/p/1">Товар</a></div><a href="/other">Другая категория</a>',
        'https://supplier.example/cat?page=2': '<script type="application/ld+json">{"@type":"Product","name":"Кран","sku":"K2","offers":{"price":"2.50"}}</script>',
        'https://supplier.example/p/1': '<h1>Труба</h1><span class="price">1,20 руб.</span>',
    }
    visited = []
    class Response:
        is_redirect = False
        def __init__(self, text):
            self.text = text
            self.content = text.encode()
        def raise_for_status(self):
            pass
    class Session:
        def __init__(self):
            self.headers = {}
        def get(self, url, **kwargs):
            visited.append(url)
            return Response(pages[url])
        def close(self):
            pass
    monkeypatch.setattr(sources.requests, 'Session', Session)
    monkeypatch.setattr(sources.time, 'sleep', lambda _: None)
    config = {'url': 'https://supplier.example', 'supplier': 'Поставщик', 'categories': 'https://supplier.example/cat', 'currency': 'BYN'}
    products, errors, limited = sources.crawl(config)
    assert {p['name'] for p in products} == {'Труба', 'Кран'}
    assert not errors and not limited
    assert 'https://supplier.example/other' not in visited


def test_multirow_templates_and_excel_formula_protection(store, tmp_path):
    pid, party = setup(store)
    second = store.product('B', '=HYPERLINK("https://example.org")')
    invoice = store.document('invoice', party, [dict(product_id=pid, quantity='1', price='1'),
                                              dict(product_id=second, quantity='2', price='2')])
    store.post(invoice)
    book = Workbook()
    book.active.append(['{{item.name}}', '{{item.total}}'])
    book.active.append(['Всего', '{{total}}'])
    template = tmp_path / 'template.xlsx'
    book.save(template)
    output = tmp_path / 'output.xlsx'
    documents.export(store, invoice, template, output)
    sheet = load_workbook(output).active
    assert sheet['A2'].data_type == 's'
    assert sheet['B3'].value == '5.00'
    word = Document()
    table = word.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '{{item.name}}'
    template = tmp_path / 'template.docx'
    word.save(template)
    output = tmp_path / 'output.docx'
    documents.export(store, invoice, template, output)
    assert len(Document(output).tables[0].rows) == 2
